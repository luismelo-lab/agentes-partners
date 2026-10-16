"""Planilha das empresas disponíveis para recorrência (Design System Save v0.5, modo claro). Lê recorrencia-todas.json (dados_proto.py)."""
import json, os, re, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image

sys.path.insert(0, os.path.expanduser("~/agentes-partners/lib")); sys.path.insert(0, os.path.expanduser("~/agentes-partners/.claude/worktrees/amazing-kirch-442f46/lib"))
import marca, caminhos, parecer_regras as PRG  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
C = marca.cores("claro") if hasattr(marca, "cores") else json.load(open(marca.PASTA / "marca.json"))["claro"]
hx = lambda c: c.lstrip("#").upper()
SANS = "IBM Plex Sans"; MONO = "IBM Plex Mono"
MASTERS = {m["parceiro"] for m in json.load(open(os.path.join(AQUI, "dados.json"), encoding="utf-8"))["masters"]} - {"Outros (fora do roster)"}
R = [r for r in json.load(open(os.path.join(AQUI, "recorrencia-todas.json"), encoding="utf-8")) if r["parceiro"] in MASTERS]  # só clientes dos Masters (Luis, 14/10)
HOJE = date(2026, 10, 14)

fmt_cnpj = lambda c: (lambda d: f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}" if len(d) == 14 else d)(re.sub(r"\D", "", str(c or "")).zfill(14)) if c else ""
dt = lambda s: date.fromisoformat(s[:10]) if s else None
sim = lambda c, item: next((x for x in c if x["item"].startswith(item)), {})

COLS = [("Empresa", 42, None), ("CNPJ", 20, "mono"), ("Parceiro", 18, None), ("Situação na Receita", 16, None), ("Data de abertura", 13, "data"), ("Regime (último parecer)", 17, None),
        ("Movimentação", 24, None), ("Valor zerado no parecer: motivo", 40, None), ("UF", 6, None), ("Cidade", 16, None), ("Segmento", 32, None),
        ("Último ganho", 13, "data"), ("Recorrências já feitas", 12, "int"), ("Período apurado", 19, None), ("Período conferido no parecer", 14, None),
        ("Competências novas", 13, "int"), ("Crédito anterior (R$)", 17, "rs"), ("Potencial em negócios (R$)", 19, "rs"), ("Honorário estimado (R$)", 17, "rs"),
        ("Repasse ao parceiro (R$)", 17, "rs"), ("Ação judicial própria (L12)", 22, None), ("Tese já aproveitada", 22, None), ("Teses do ganho anterior", 40, None),
        ("Sinais", 46, None), ("Negócio no HubSpot", 14, "link")]


def linha(r):
    c = r.get("conferencias", [])
    aj = sim(c, "Ação judicial"); ta = sim(c, "Tese já aproveitada")
    G = PRG.empresa(r.get("_deals") or [], r.get("_cnpj"), (r.get("_link") or "").rsplit("/", 1)[-1] or None)
    return [r["caso"], fmt_cnpj(r.get("_cnpj")), r["parceiro"], G["situacao_receita"], G["abertura"], G["regime_parecer"], G["movimentacao"], G["valor_zerado_motivo"] or "—",
            r.get("uf") or "", r.get("_cidade") or "", r.get("segmento") or "",
            dt(r.get("ultimo_ganho")), r["recorrencias_feitas"], r.get("periodo_apurado") or "estimado pela data do ganho", "Sim" if r.get("periodo_conferido") else "Não",
            r["meses_novos"], r["credito_anterior"], r["potencial_credito"], r["potencial_honorario"], r["potencial_repasse"],
            ("Sim: Jurídico confere" if aj.get("ok") is False else "Não" if aj.get("ok") else "Não informado") if aj else "Parecer não lido",
            (ta.get("resultado") if ta.get("ok") is False else "Não") if ta else "Parecer não lido",
            "; ".join(r.get("teses") or []) or "não registradas", "; ".join(r.get("sinais") or []), r.get("_link")]


borda = Border(bottom=Side(style="thin", color=hx(C["borda"])))
def aba(wb, titulo, sub, linhas, primeira=False, decisao=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo[:31]; ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 44
    try:
        img = Image(str(marca.logo("fundo_claro"))); img.height = 26; img.width = 26 / 0.1756 if img.width > img.height else 120; ws.add_image(img, "A1")
    except Exception:
        ws["A1"] = "Save Co."; ws["A1"].font = Font(name=SANS, bold=True, size=14, color=hx(C["texto"]))
    ws["C1"] = titulo; ws["C1"].font = Font(name=SANS, bold=True, size=16, color=hx(C["texto"])); ws["C1"].alignment = Alignment(vertical="center")
    ws["C2"] = sub; ws["C2"].font = Font(name=SANS, size=10, color=hx(C["texto_2"]))
    for col in range(1, len(COLS) + 1): ws.cell(3, col).fill = PatternFill("solid", fgColor=hx(C["ouro"]))
    ws.row_dimensions[3].height = 4
    H = 5
    for i, (nome, larg, _) in enumerate(COLS, 1):
        c = ws.cell(H, i, nome); c.font = Font(name=SANS, bold=True, size=10, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=hx(C["texto"]))
        c.alignment = Alignment(vertical="center", wrap_text=True); ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[H].height = 32
    for j, r in enumerate(linhas, H + 1):
        for i, (v, (nome, _, tipo)) in enumerate(zip(linha(r), COLS), 1):
            c = ws.cell(j, i, v); c.font = Font(name=MONO if tipo in ("mono", "rs", "int") else SANS, size=10, color=hx(C["texto"])); c.border = borda
            c.alignment = Alignment(vertical="top", wrap_text=nome in ("Sinais", "Teses do ganho anterior", "Segmento", "Valor zerado no parecer: motivo", "Movimentação"))
            if nome == "Situação na Receita" and v and v != "Ativa": c.font = Font(name=SANS, size=10, bold=True, color=hx(C["erro"]))
            if nome == "Movimentação" and v == "Empresa sem movimentação": c.font = Font(name=SANS, size=10, bold=True, color=hx(C["alerta"]))
            if tipo == "rs": c.number_format = '#,##0'
            if tipo == "data": c.number_format = "DD/MM/YYYY"
            if tipo == "link" and v: c.value = "Abrir"; c.hyperlink = v; c.font = Font(name=SANS, size=10, color=hx(C["info"]), underline="single")
            if nome == "Período conferido no parecer": c.font = Font(name=SANS, size=10, bold=True, color=hx(C["sucesso"] if v == "Sim" else C["alerta"]))
            if nome == "Ação judicial própria (L12)" and str(v).startswith("Sim"): c.font = Font(name=SANS, size=10, bold=True, color=hx(C["alerta"]))
    ws.freeze_panes = ws.cell(H + 1, 2)
    nc = len(COLS)
    if decisao:  # Geovana analisa mesmo assim (Luis, 16/10): decisão e observação na própria aba
        from openpyxl.worksheet.datavalidation import DataValidation
        for k, (nome, larg) in enumerate((("Decisão da Geovana", 24), ("Observação da Geovana", 40)), nc + 1):
            c = ws.cell(H, k, nome); c.font = Font(name=SANS, bold=True, size=10, color=hx(C["texto"])); c.fill = PatternFill("solid", fgColor=hx(C["ouro"]))
            c.alignment = Alignment(vertical="center", wrap_text=True); ws.column_dimensions[get_column_letter(k)].width = larg
            ws.cell(3, k).fill = PatternFill("solid", fgColor=hx(C["ouro"]))
        dv = DataValidation(type="list", formula1='"Gerar recorrência mesmo assim,Descartar: sem movimentação,Precisa de dado"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{get_column_letter(nc + 1)}{H + 1}:{get_column_letter(nc + 1)}{H + max(len(linhas), 1)}")
        for j in range(H + 1, H + len(linhas) + 1):
            for k in (nc + 1, nc + 2): ws.cell(j, k).border = borda; ws.cell(j, k).font = Font(name=SANS, size=10, color=hx(C["texto"]))
        nc += 2
    ws.auto_filter.ref = f"A{H}:{get_column_letter(nc)}{H + len(linhas)}"
    # totais
    t = H + len(linhas) + 2
    ws.cell(t, 1, f"Total: {len(linhas)} empresas").font = Font(name=SANS, bold=True, size=10, color=hx(C["texto"]))
    for i, (nome, _, tipo) in enumerate(COLS, 1):
        if tipo == "rs":
            c = ws.cell(t, i, f"=SUBTOTAL(9,{get_column_letter(i)}{H + 1}:{get_column_letter(i)}{H + len(linhas)})"); c.number_format = '#,##0'
            c.font = Font(name=MONO, bold=True, size=10, color=hx(C["texto"]))
    return ws


el = [r for r in R if r["elegivel"]]; ag = [r for r in R if not r["elegivel"]]
el.sort(key=lambda r: (not r["periodo_conferido"], -r["potencial_credito"]))
# sem movimentação (parecer da última venda com faturamento só até 2025): sai da lista de disponíveis e vai para a análise da Geovana
mov = lambda r: PRG.empresa(r.get("_deals") or [], r.get("_cnpj"), (r.get("_link") or "").rsplit("/", 1)[-1] or None)["movimentacao"]
sm = [r for r in el if mov(r) == "Empresa sem movimentação"]; el = [r for r in el if r not in sm]
ag.sort(key=lambda r: (r["faltam"], -r["credito_anterior"]))
wb = Workbook()
aba(wb, "Disponíveis para recorrência", f"{len(el)} empresas dos Masters com 6 competências novas ou mais · maior potencial primeiro, com período conferido no parecer no topo · gerado em {HOJE:%d/%m/%Y}", el, True)
aba(wb, "Sem movimentação · Geovana", f"{len(sm)} empresas cujo parecer da última venda traz faturamento só até 2025: fora do potencial, mas a Geovana analisa cada uma e registra a decisão nas duas últimas colunas", sm, decisao=True)
aba(wb, "Aguardando 6 meses", f"{len(ag)} empresas dos Masters que ainda não têm 6 competências novas · as mais próximas primeiro", ag)
ws = wb.create_sheet("Como ler"); ws.sheet_view.showGridLines = False; ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100
txt = [("Quem entra", "Só clientes dos Masters (fora do roster não entra). Clientes com negócio ganho, fora do Simples Nacional e do café, sem negócio aberto em nenhum pipeline (mesma raiz de CNPJ)."),
       ("Período apurado", "Lido no parecer do último fechamento. As competências novas contam do mês seguinte ao fim do período até 09/2026."),
       ("Período conferido = Não", "O parecer do último fechamento ainda não foi lido: os meses contam da data do ganho (estimativa conservadora)."),
       ("Potencial em negócios", "Crédito anterior ÷ 60 × competências novas (até 60). Ordem de grandeza para priorizar, não promessa."),
       ("Honorário e repasse", "Honorário da Save do negócio anterior (ou o típico do parceiro) e % oficial de repasse do cadastro."),
       ("Ação judicial (L12)", "Quando o parecer aponta ação judicial em alguma tese, o Jurídico confere antes de seguir."),
       ("Situação e abertura", f"Base dos Dados Abertos da Receita Federal, referência {PRG.referencia_receita()}. Empresa que não está Ativa aparece em vermelho."),
       ("Regime", "Só o que aparece no último parecer do técnico (regra da Geovana). Sem parecer lido: 'parecer não lido'."),
       ("Sem movimentação · Geovana", "Empresas tiradas da lista de disponíveis porque o parecer da última venda mostra faturamento só até 2025. A Geovana analisa cada uma e escolhe: gerar recorrência mesmo assim, descartar ou precisa de dado."),
       ("Movimentação", "Parecer da ÚLTIMA VENDA com a tabela de faturamento só até 2025 (e o período da venda contemplando esses anos) = Empresa sem movimentação. A análise considera o faturamento (Geovana e Luis, 16/10). Parecer de outro negócio: a regra não se aplica."),
       ("Valor zerado", "Quando o último parecer não cita valor, a coluna traz o motivo escrito no arquivo."),
       ("Uso", "Uso interno da Área Partners. Pré-análise: a decisão de recorrência é da Geovana no painel.")]
ws["A1"] = "Como ler esta planilha"; ws["A1"].font = Font(name=SANS, bold=True, size=14, color=hx(C["texto"]))
for i, (a, b) in enumerate(txt, 3):
    ws.cell(i, 1, a).font = Font(name=SANS, bold=True, size=10, color=hx(C["texto"])); c = ws.cell(i, 2, b); c.font = Font(name=SANS, size=10, color=hx(C["texto_2"])); c.alignment = Alignment(wrap_text=True, vertical="top")

dest = os.path.join(caminhos.onedrive(), caminhos.LUIS["areas"], f"Recorrência - empresas disponíveis - {HOJE:%Y-%m-%d}.xlsx")
os.makedirs(os.path.dirname(dest), exist_ok=True); wb.save(dest)
print(dest); print(caminhos.web(dest)); print(len(el), len(sm), len(ag), sum(r["potencial_credito"] for r in el), sum(r["potencial_credito"] for r in sm))
