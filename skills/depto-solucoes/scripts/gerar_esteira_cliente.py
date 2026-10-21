#!/usr/bin/env python3
"""Gera a Esteira do Cliente (Word, padrão Save): 12 passos em 3 movimentos, sem nome de produto, com os dados setoriais de dados-estudo/."""
import json, sys, glob, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parents[3]
SCR = Path(sys.argv[1])
OUT = Path(sys.argv[2])
M = json.load(open(REPO / "departamentos/marketing-marca/padrao/marca.json"))
C = M["claro"]; E = M["escuro"]
def rgb(h): h = h.lstrip("#"); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
SANS = "IBM Plex Sans"; MONO = "IBM Plex Mono"
LOGO_ESC = REPO / "departamentos/marketing-marca/padrao/logo-save-fundo-escuro.png"
LOGO_CLA = REPO / "departamentos/marketing-marca/padrao/logo-save-fundo-claro.png"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

def setfont(run, name=SANS, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = rgb(color)

st = doc.styles["Normal"]; st.font.name = SANS; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.15
for nome, tam, cor in (("Heading 1", 18, C["texto"]), ("Heading 2", 13, C["texto"]), ("Heading 3", 10.5, C["ouro_texto"])):
    s = doc.styles[nome]; s.font.name = SANS; s.font.size = Pt(tam); s.font.bold = True
    s.font.color.rgb = rgb(cor); s.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    s.element.rPr.rFonts.set(qn("w:ascii"), SANS); s.element.rPr.rFonts.set(qn("w:hAnsi"), SANS)
    s.paragraph_format.space_before = Pt(14 if nome != "Heading 3" else 8); s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True

def shade(cell, hexcor):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcor.lstrip("#"))
    tcPr.append(sh)

def borders(cell, cor, sz=4, lados=("top", "left", "bottom", "right")):
    tcPr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for l in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{l}")
        if l in lados: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), cor.lstrip("#"))
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def margens(cell, v=70, h=100):
    tcPr = cell._element.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for l, val in (("top", v), ("bottom", v), ("left", h), ("right", h)):
        e = OxmlElement(f"w:{l}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def para(texto="", size=10, bold=False, color=None, italic=False, align=None, after=None, before=None, style=None, mono=False):
    p = doc.add_paragraph(style=style)
    if texto:
        r = p.add_run(texto); setfont(r, MONO if mono else SANS, size, bold, color or C["texto"], italic)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p

def h(texto, n=1):
    p = doc.add_heading(level=n)
    r = p.add_run(texto)
    setfont(r, SANS, {1: 18, 2: 13, 3: 10.5}[n], True, C["texto"] if n < 3 else C["ouro_texto"])
    if n == 1:
        pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:color"), C["ouro"].lstrip("#")); b.set(qn("w:space"), "4")
        bd.append(b); pPr.append(bd)
    return p

def bullets(itens, size=10, cor=None):
    for i in itens:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(i); setfont(r, SANS, size, False, cor or C["texto"])

def tabela(cab, linhas, larguras, size=8.5, cab_escuro=True, zebra=True):
    t = doc.add_table(rows=1, cols=len(cab)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, c in enumerate(cab):
        cell = t.rows[0].cells[i]; cell.width = Cm(larguras[i]); shade(cell, E["fundo"] if cab_escuro else C["fundo_2"]); margens(cell)
        borders(cell, C["borda_2"]); cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = cell.paragraphs[0].add_run(c.upper()); setfont(r, MONO, 7.5, True, E["ouro"] if cab_escuro else C["texto"])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for k, ln in enumerate(linhas):
        row = t.add_row()
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        for i, v in enumerate(ln):
            cell = row.cells[i]; cell.width = Cm(larguras[i]); margens(cell); borders(cell, C["borda"])
            shade(cell, C["fundo"] if (zebra and k % 2) else C["cartao"])
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            neg = (i == 0)
            r = cell.paragraphs[0].add_run(str(v)); setfont(r, SANS, size, neg, C["texto"])
    for i,wd in enumerate(larguras): t.columns[i].width = Cm(wd)
    tblPr=t._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def caixa(titulo, texto, escuro=False):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; margens(c, 120, 180)
    shade(c, E["cartao"] if escuro else C["ouro_suave"]); borders(c, E["borda"] if escuro else C["ouro"], 6, ("left",))
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(titulo.upper()); setfont(r, MONO, 7.5, True, E["ouro"] if escuro else C["ouro_texto"])
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(texto); setfont(r, SANS, 10, False, E["texto"] if escuro else C["texto"])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def rodape():
    for s in doc.sections:
        s.different_first_page_header_footer = True
        f = s.footer; p = f.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run("Save Co. · Esteira completa · Pré-análise, sujeita a validação técnica. Sem promessa de valor, prazo ou êxito.   Pág. ")
        setfont(r, MONO, 7, False, C["texto_3"])
        for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            r2 = p.add_run(); setfont(r2, MONO, 7, False, C["texto_3"])
            if tipo:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r2._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r2._r.append(it)

def quebra(): doc.add_page_break()

# ───────────────────────── dados
setores = [json.load(open(f)) for f in sorted(glob.glob(str(SCR / "setores/*.json")))]
bench = json.load(open(SCR / "benchmark.json")) if (SCR / "benchmark.json").exists() else {}


# ───────────────────────── texto do cliente: sem nome de produto (regra do Luis, 21/10/2026)
SCRUB = [
    (r"Gestão de Contas \(([^)]*)\)", r"Cobrança organizada (\1)"), (r"Gestão de Contas", "Cobrança organizada"),
    (r"Gestão de Passivo", "Regularização de débitos"),
    (r"Antecipação de recebíveis, por instituição autorizada", "Receber antes, por instituição autorizada"),
    (r"Antecipação do recebível de frete", "Receber o frete antes"), (r"Antecipação de recebíveis de cartão e duplicatas", "Receber antes cartão e duplicatas"),
    (r"Antecipação de recebíveis", "Receber antes"), (r"antecipação de recebíveis", "recebimento antecipado"), (r"[Aa]ntecipação", "recebimento antecipado"),
    (r"Crédito internacional e FIDC", "Capital de longo prazo e mecanismo próprio de recebíveis"),
    (r"FIDC e modelo multicedente/multissacado", "Mecanismo próprio de recebíveis da cadeia"), (r"FIDC e multicedente", "Mecanismo próprio de recebíveis da cadeia"),
    (r"Mecanismo próprio de recebíveis \(FIDC\)", "Mecanismo próprio de recebíveis"), (r"Fundo de recebíveis \(FIDC\)", "Mecanismo próprio de recebíveis"),
    (r"Estrutura de recebíveis \(FIDC\)", "Mecanismo próprio de recebíveis"), (r"\bFIDC\b", "mecanismo próprio de recebíveis"),
    (r"Crédito internacional", "Capital de longo prazo, inclusive no exterior"), (r"crédito internacional", "capital de longo prazo, inclusive no exterior"),
    (r"Estruturação patrimonial", "Organização do patrimônio e sucessão"), (r"estruturação patrimonial", "organização do patrimônio e sucessão"),
    (r"Maquila: conversa de descoberta", "Produção no Paraguai: estudo de viabilidade"), (r"Maquila \([^)]*\)", "produção no Paraguai"), (r"Maquila", "produção no Paraguai"), (r"maquila", "produção no Paraguai"),
    (r"conta escrow", "conta garantia de recebíveis"), (r"escrow", "conta garantia de recebíveis"),
    (r"Captação de capital: sócio, investidores ou equity", "Capital: sócio, investidores ou equity"),
]
def limpa(t):
    for a, b in SCRUB: t = re.sub(a, b, t)
    t = re.sub(r"\b(em avaliação conjunta|Em avaliação conjunta)\b[^.]*\.?", "", t).strip()
    return t
import re  # noqa

PASSOS = [
 # n, movimento, nome, problema, fazemos, recebe, abre, quem, nota
 (1, "Destravar caixa", "Enxergar a empresa por inteiro",
  "Não sei ao certo onde está o meu dinheiro, o que devo e o que fazer primeiro.",
  "Lemos a documentação fiscal e financeira (notas, obrigações, DRE, balanço, endividamento) e montamos um mapa único: onde há crédito, onde há risco e qual é o melhor ponto de partida. Registramos a linha de base: caixa, carga tributária, prazo de recebimento e custo da dívida. Se você ainda não quer abrir todos os dados, começamos pelo que já pode compartilhar.",
  "Mapa da empresa, com prioridades, e a linha de base assinada por você.",
  "Define o primeiro passo certo para o seu caso (2, 3 ou 5) e deixa a medição pronta para provar o resultado de cada etapa seguinte.",
  "Gestor do projeto, equipe tributária e equipe de contas e soluções", ""),
 (2, "Destravar caixa", "Recuperar o que já é seu",
  "Pago imposto que talvez não devesse e não sei quanto volta.",
  "Identificamos créditos tributários dos últimos cinco anos, conforme a atividade, o regime e as compras da sua empresa. Preparamos e protocolamos os pedidos e acompanhamos até o retorno. Depois, mês a mês, capturamos as novas competências.",
  "Crédito apurado e acompanhado, com relatório do valor recuperado.",
  "O crédito vira caixa e pode abater débitos (passo 5). Os dados fiscais organizados alimentam o planejamento (passo 6) e o painel de controle (passo 7).",
  "Equipe tributária, revisão técnica e equipe jurídica", "Estimativa preliminar sujeita a validação técnica; sem promessa de valor ou prazo."),
 (3, "Destravar caixa", "Receber o que lhe devem, e mais cedo",
  "Vendo a prazo, o cliente atrasa e o meu caixa aperta.",
  "Organizamos a cobrança com mensagens automáticas (lembrete antes do vencimento, link de pagamento) sem mudar a sua rotina de emissão. É pago pelo resultado, sem mensalidade. Com a carteira em ordem, quem precisa de caixa antes pode receber antecipadamente, por instituição autorizada, com limite que acompanha a carteira registrada.",
  "Carteira organizada, previsibilidade de caixa e dados confiáveis de contas a receber.",
  "A carteira organizada é a base de qualquer crédito (passo 9) e do mecanismo próprio de recebíveis (passo 10), e gera os números que o passo 1 pedia.",
  "Equipe financeira da Save; a antecipação é realizada por instituição autorizada",
  "A Save organiza a carteira; a decisão de antecipar e as condições são da instituição."),
 (4, "Destravar caixa", "Aliviar a dívida que pesa",
  "Pago muita parcela de banco e sobra pouco para a operação.",
  "Com a consulta ao histórico de crédito e a curva de endividamento, estudamos alongar prazos, ganhar carência e trocar dívida cara por uma que cabe no fluxo. O crédito recuperado no passo 2 pode abater o que for mais caro.",
  "Plano de dívida com parcelas compatíveis e janela para respirar.",
  "O fôlego permite regularizar o que estiver pendente (passo 5) e investir (passos 9 a 11).",
  "Equipe financeira e equipe jurídica", "Avaliado caso a caso; sem condição definida antes da análise."),
 (5, "Proteger e controlar", "Deixar a empresa em dia",
  "Tenho débito ou auto de infração e não consigo certidão.",
  "Mapeamos débitos e autos, escolhemos o caminho adequado (pagamento, parcelamento, transação ou defesa), compensamos com o crédito do passo 2 quando cabe e controlamos todos os prazos com dupla checagem.",
  "Empresa regular, certidão em dia e prazos sob controle.",
  "A certidão negativa é pré-requisito para crédito de longo prazo, grandes clientes e licitações (passos 9 e 10) e libera o planejamento do passo 6 sem risco escondido.",
  "Equipe jurídica e equipe tributária", ""),
 (6, "Proteger e controlar", "Pagar o imposto certo e se preparar para 2027",
  "Não sei como a reforma tributária muda a minha margem.",
  "Revisamos regime e mix de produtos, simulamos a nova tributação (CBS e IBS) produto a produto até 2033 e usamos os créditos antigos antes da extinção do PIS e da COFINS.",
  "Plano tributário com cenários e menor carga dentro da lei.",
  "A margem e a previsibilidade alimentam o plano de crescimento (passos 9 a 11) e a organização do patrimônio (passo 8).",
  "Equipe tributária", "Reduz a carga quando há espaço legal; não promete percentual."),
 (7, "Proteger e controlar", "Dar governança ao dinheiro",
  "Falta controle: não confio nos meus números e o banco ou o investidor pede relatório.",
  "Montamos um painel mensal (resultado, endividamento, contas a receber e fluxo), regras claras para a conta onde entram os recebimentos e uma rotina de conferência acompanhada pelo gestor do projeto.",
  "Indicadores confiáveis e o nível de controle que quem financia exige.",
  "Sem esse controle, os passos de capital (9 e 10) não andam. O painel também abastece a revisão anual (passo 12).",
  "Gestor do projeto e equipe financeira", "A conta garantia de recebíveis dá governança; não é instrumento para evitar ordens judiciais."),
 (8, "Proteger e controlar", "Proteger o patrimônio e a família",
  "Se algo acontecer comigo, a empresa e o patrimônio aguentam? Pago mais imposto para distribuir lucro?",
  "Organizamos a estrutura patrimonial e empresarial, o acordo entre sócios e a sucessão, considerando a retenção sobre dividendos em vigor desde 2026.",
  "Estrutura documentada, protegida e preparada para a próxima geração.",
  "Ativos organizados e livres são a base das garantias que o capital pede (passo 9) e preparam a entrada de sócio ou investidor.",
  "Equipe jurídica", ""),
 (9, "Crescer", "Capital para crescer",
  "Tenho demanda ou projeto, mas o crédito do banco é caro e curto.",
  "Preparamos o projeto (plano de negócio, uso dos recursos, demonstrações, garantias) e o levamos a instituições e investidores adequados: crédito de longo prazo, inclusive em moeda forte e no exterior, caso a caso, ou sócio e investidor. A Save estrutura; a decisão e o crédito são da instituição.",
  "Capital com prazo compatível e disciplina de uso; projeto pronto para análise.",
  "Permite investir na capacidade (passo 10), nos mercados externos (passo 11) e escalar com controle (passo 7).",
  "Equipe de crédito e capital, com apoio jurídico e tributário", "Avaliado caso a caso; sem taxa, prazo ou condição antes da análise da instituição."),
 (10, "Crescer", "Financiar a sua cadeia com o seu próprio mecanismo",
  "Compro ou vendo para muita gente e os bancos ficam com a receita financeira da minha cadeia.",
  "Estruturamos um mecanismo próprio de recebíveis para antecipar fornecedores, produtores ou clientes dentro da cadeia, retendo a receita financeira e dando crédito mais barato a quem trabalha com você.",
  "Capital de giro da cadeia a custo menor e receita financeira que fica com a empresa.",
  "Fideliza a cadeia e financia o crescimento. Exige a carteira organizada (passo 3) e o controle (passo 7).",
  "Equipe financeira, com estrutura regulada de fundo", "Em empresas de maior porte, em torno de dezenas de milhões de reais de faturamento por mês, e sujeito a análise."),
 (11, "Crescer", "Vender e produzir além da fronteira",
  "Quero exportar ou produzir fora, mas os tributos e a burocracia assustam.",
  "Revisamos a tributação e a conformidade da importação e da exportação e, onde fizer sentido, fazemos o estudo de viabilidade para produzir no Paraguai pelo regime de maquila, com os números reais da sua empresa.",
  "Plano de mercado externo, com custos, riscos e caminho regulatório.",
  "Novos mercados voltam ao passo 6 (tributos) e ao passo 9 (capital).",
  "Equipe internacional e equipe tributária", "Estudo de viabilidade; não é planejamento para pagar menos imposto de renda e não promete economia."),
 (12, "Crescer", "Manter a esteira rodando",
  "Resolvi uma vez, mas os problemas voltam.",
  "Fazemos a revisão anual, capturamos as novas competências, realizamos o encontro trimestral com o gestor do projeto, atualizamos o painel de valor e lemos os novos sinais da empresa.",
  "Relatório de valor e plano dos próximos 90 dias.",
  "Cada sinal novo volta ao passo certo da esteira (2, 3, 6 ou 9): a esteira é contínua, não um projeto com fim.",
  "Gestor do projeto", ""),
]
NOME = {p[0]: p[2] for p in PASSOS}

def passo_de(nome):
    n = nome.lower()
    rules = [(r"maquila|paraguai|exporta|comércio exterior", 11), (r"mecanismo próprio|fidc|multicedente|fundo de recebíveis|estrutura de recebíveis", 10),
             (r"capital|crédito internacional|equity|sócio|funding", 9), (r"patrimonial|patrimônio|sucessão", 8), (r"planejamento|reforma", 6),
             (r"passivo|débitos|defesa|glosa|jurídic", 5), (r"cobrança|gestão de contas|antecipa|receber antes|recebimento antecipado", 3)]
    for rx, k in rules:
        if re.search(rx, n): return k
    return 2

# ───────────────────────── CAPA
t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
c = t.rows[0].cells[0]; shade(c, E["fundo"]); margens(c, 500, 500); borders(c, E["fundo"], 4, ())
p = c.paragraphs[0]; p.add_run().add_picture(str(LOGO_ESC), height=Cm(1.1)); p.paragraph_format.space_after = Pt(150)
p = c.add_paragraph(); r = p.add_run("ESTEIRA COMPLETA"); setfont(r, MONO, 9, True, E["ouro"])
p = c.add_paragraph(); r = p.add_run("Do aperto de caixa ao crescimento: um passo abre o próximo"); setfont(r, SANS, 26, True, E["texto"])
p.paragraph_format.space_after = Pt(10)
p = c.add_paragraph(); r = p.add_run("Como a Save resolve, em ordem, os problemas da sua empresa, com um gestor de projeto que conduz tudo"); setfont(r, SANS, 11, False, E["texto_2"])
p.paragraph_format.space_after = Pt(150)
p = c.add_paragraph(); r = p.add_run("<EMPRESA> · Setor: ____________ · Data-base: 21/10/2026"); setfont(r, MONO, 8.5, False, E["texto_2"])
p = c.add_paragraph(); r = p.add_run("PRÉ-ANÁLISE · ESTIMATIVA PRELIMINAR SUJEITA A VALIDAÇÃO TÉCNICA"); setfont(r, MONO, 8, True, E["ouro"])
for _ in range(2): c.add_paragraph()
quebra()

# ───────────────────────── 1. EM UMA PÁGINA
h("1. Em uma página")
para("Empresas costumam resolver um problema de cada vez: um crédito aqui, uma renegociação ali, um empréstimo quando o caixa aperta. Cada decisão ignora as outras, e o resultado é retrabalho e custo. A esteira da Save faz o contrário: resolve o seu problema mais urgente primeiro e usa o resultado como ponto de partida do passo seguinte, até a empresa estar organizada e pronta para crescer.", 10.5)
caixa("A ideia", "O fim de um passo é o começo do próximo. O crédito recuperado paga o passo seguinte; a carteira organizada abre o crédito; a empresa em dia abre o capital; o controle mantém tudo de pé.")
tabela(["Movimento", "A pergunta do empresário", "Os passos"],
       [["Destravar caixa", "Onde está o dinheiro que já é meu, e como faço o caixa respirar?", "1 · Enxergar a empresa por inteiro   2 · Recuperar o que já é seu   3 · Receber o que lhe devem, e mais cedo   4 · Aliviar a dívida que pesa"],
        ["Proteger e controlar", "Como organizo a casa para o problema não voltar?", "5 · Deixar a empresa em dia   6 · Pagar o imposto certo e se preparar para 2027   7 · Dar governança ao dinheiro   8 · Proteger o patrimônio e a família"],
        ["Crescer", "Como financio e expando com segurança?", "9 · Capital para crescer   10 · Financiar a sua cadeia   11 · Vender e produzir além da fronteira   12 · Manter a esteira rodando"]],
       [3.2, 5.0, 8.4], 9)
h("O que muda para você", 3)
bullets(["Um gestor de projeto conduz a esteira e conecta as equipes da Save. Você não precisa coordenar sete áreas.",
         "Cada passo tem entrega clara e termina com a pergunta: o que isto libera para o próximo?",
         "Você paga, sempre que possível, pelo resultado. Sem mensalidade na cobrança e sem cobrança antecipada em operações de crédito.",
         "A esteira se adapta ao seu ponto de partida e ao seu setor (capítulos 3 e 5)."])
quebra()

# ───────────────────────── 2. OS PASSOS
h("2. Os 12 passos, um a um")
para("Cada passo mostra o problema como o empresário o descreve, o que fazemos, o que você recebe e, principalmente, como ele abre o seguinte.", 10)
mov_atual = None
for n, mov, nome, prob, faz, recebe, abre, quem, nota in PASSOS:
    if mov != mov_atual:
        h(mov, 2); mov_atual = mov
    h(f"Passo {n} · {nome}", 3)
    tabela(["", ""], [["O problema", "“" + prob + "”"], ["O que fazemos", faz], ["O que você recebe", recebe], ["Como abre o próximo passo", abre], ["Quem atua", quem]] + ([["Cuidado", nota]] if nota else []), [4.0, 12.6], 9, cab_escuro=False, zebra=False)
    # remove a linha de cabeçalho vazia
    tb = doc.tables[-1]; tb._tbl.remove(tb.rows[0]._tr)
quebra()

# ───────────────────────── 3. PONTO DE ENTRADA E REGRAS DE PASSAGEM
h("3. Por onde começar e como um passo abre o outro")
para("Não é preciso começar pelo passo 1. Você entra pelo problema que mais dói e o gestor de projeto faz uma checagem rápida do que estiver antes na esteira, para não construir sobre areia.", 10)
tabela(["Se o seu problema é…", "Comece por", "Antes de seguir, conferimos"],
       [["Caixa apertado", "Passo 3 ou 2, o que liberar caixa mais rápido", "Se há débito (passo 5) que trave a certidão"],
        ["Débito, autuação ou falta de certidão", "Passo 5", "Se há crédito (passo 2) que possa abater o débito"],
        ["Dívida bancária pesada", "Passo 4", "Carteira (passo 3) e números (passo 7) para negociar com dados"],
        ["Pagar menos imposto e a reforma de 2027", "Passos 2 e 6", "Se o regime e o mix permitem; documentação em ordem"],
        ["Quero crescer e preciso de capital", "Passos 1, 7 e 8", "Certidão (passo 5) e ativos livres (passo 8)"],
        ["Quero exportar ou produzir fora", "Passos 6 e 11", "Capital (passo 9) e controle (passo 7)"]],
       [5.2, 5.4, 6.0], 9)
h("A regra de passagem: o que precisa estar pronto para o próximo passo", 3)
tabela(["De", "Para", "Condição para passar"],
       [["2 · Recuperar", "5 · Deixar em dia", "Crédito reconhecido para abater débito; senão, pagamento ou parcelamento"],
        ["5 · Deixar em dia", "9 · Capital", "Certidão negativa e prazos sob controle"],
        ["3 · Receber", "9 e 10 · Capital e mecanismo próprio", "Carteira organizada, com histórico de recebimento"],
        ["7 · Governança", "9 e 10 · Capital e mecanismo próprio", "Painel mensal, demonstrações em dia e conta de recebimentos com regras"],
        ["8 · Patrimônio", "9 · Capital", "Ativos livres e documentados, estrutura societária clara"],
        ["6 · Imposto certo", "11 · Mercados externos", "Carga conhecida e regime definido"],
        ["12 · Esteira contínua", "2, 3, 6 ou 9", "Sinal novo na empresa (nova competência, nova dívida, novo projeto)"]],
       [4.0, 5.2, 7.4], 9)
quebra()

# ───────────────────────── 4. GESTOR DE PROJETO
h("4. O seu gestor de projeto")
para("A esteira só funciona se alguém a conduz. Cada empresa tem um gestor de projeto, da equipe de atendimento da Save, que conhece o seu negócio e responde por todo o caminho.", 10.5)
tabela(["O gestor de projeto…", "Na prática"],
       [["Entende o seu negócio", "Faz a conversa inicial com o dono, monta o mapa da empresa e a linha de base."],
        ["Conecta as equipes", "Aciona tributário, jurídico, financeiro e internacional na hora certa, sem você repassar documentos de uma área a outra."],
        ["Garante a passagem entre passos", "Confere a condição de cada passagem e só avança quando o passo anterior entregou."],
        ["Presta contas", "Encontro trimestral de 45 minutos com pauta, relatório de valor e combinados com dono e data."],
        ["Fala a sua língua", "Traduz a parte técnica para decisão de negócio e avisa o que depende de você."]],
       [5.0, 11.6], 9)
quebra()

# ───────────────────────── 5. SETORES
h("5. A esteira para o seu setor")
para("A ordem dos passos é a mesma; o que muda é onde está o dinheiro e o risco em cada setor. Cada página traz o retrato do setor em 2026, as dores mais comuns, o seu caminho pela esteira e como medimos o resultado.", 10.5)
tabela(["Setor", "Perfis", "Caminho mais comum (passos)"],
       [[s["titulo"], s["subtitulo"], " → ".join(str(k) for k in sorted({passo_de(limpa(x["nome"])) for o in s["jornada"] for x in o["solucoes"]}))] for s in setores],
       [5.0, 6.4, 5.2], 9)
quebra()
for n, s in enumerate(setores, 1):
    h(f"5.{n} {s['titulo']}")
    para(s["subtitulo"].capitalize() + " · " + limpa(s["base_clientes"]), 9, color=C["texto_2"], italic=True)
    h("Retrato do setor em 2026", 3); bullets([limpa(x) for x in s["retrato"]], 9.5)
    h("O que mais pesa para o dono", 3); bullets([limpa(x) for x in s["dores"]], 9.5)
    h("O seu caminho pela esteira", 3)
    por_passo = {}
    for o in s["jornada"]:
        for x in o["solucoes"]:
            nm = limpa(x["nome"]); por_passo.setdefault(passo_de(nm), []).append((nm, limpa(x["porque"])))
    linhas = []
    for k in sorted(por_passo):
        for nm, pq in por_passo[k]: linhas.append([f"{k} · {NOME[k]}", nm, pq])
    tabela(["Passo", "No seu setor", "Por que se aplica"], linhas, [4.0, 4.4, 8.2], 8.5)
    h("Como medimos o resultado", 3)
    tabela(["Indicador", "Como medir"], [[limpa(k["indicador"]), limpa(k["como_medir"])] for k in s["kpis"]], [6.0, 10.6], 8.5)
    caixa("Em uma frase ao dono", limpa(s["frase_valor"]))
    h("Cuidados e limites", 3); bullets([limpa(x) for x in s["cuidados"]], 9.5)
    quebra()

# ───────────────────────── 6. VALOR
h("6. Como você vê o resultado")
para("Desde o primeiro dia registramos a linha de base (caixa, carga tributária, prazo de recebimento, custo da dívida). A cada trimestre o gestor apresenta o Mapa de Valor: o que foi entregue, o que significou para o caixa e o que a esteira libera a seguir.", 10.5)
tabela(["Passo", "Indicador principal"],
       [["2 · Recuperar", "Caixa liberado e carga tributária efetiva"], ["3 · Receber", "Prazo médio de recebimento e inadimplência recuperada"],
        ["4 · Aliviar a dívida", "Parcela mensal e custo médio da dívida"], ["5 · Deixar em dia", "Débito resolvido e certidão em dia"],
        ["6 · Imposto certo", "Variação da carga, em pontos percentuais, até 2033"], ["7 · Governança", "Painel mensal entregue e números conferidos"],
        ["8 · Patrimônio", "Estrutura documentada e sucessão definida"], ["9 e 10 · Capital", "Custo e prazo do capital contra a alternativa que você tinha"],
        ["11 · Mercados externos", "Estudo concluído e decisão tomada com números reais"], ["12 · Esteira contínua", "Novos sinais tratados dentro do trimestre"]],
       [5.0, 11.6], 9)
quebra()

# ───────────────────────── 7. PRIMEIROS 90 DIAS
h("7. Os primeiros 90 dias")
tabela(["Quando", "O que acontece", "O que você entrega", "O que você recebe"],
       [["Semanas 1–2", "Conversa inicial com o dono, acordo de confidencialidade, definição do ponto de entrada", "Documentos básicos (notas, obrigações, DRE, balanço, extrato de dívidas)", "Plano de trabalho e cronograma"],
        ["Semanas 3–4", "Mapa da empresa e linha de base (passo 1)", "Validação do mapa", "Mapa da empresa com prioridades"],
        ["Semanas 5–8", "Primeiro passo de execução (2, 3 ou 5, conforme o seu ponto de entrada)", "Acesso e respostas rápidas às dúvidas", "Pedidos protocolados ou carteira em operação"],
        ["Semanas 9–10", "Primeiro Mapa de Valor e desenho do passo seguinte", "Decisão sobre o próximo passo", "Mapa de Valor e plano do próximo passo"],
        ["Semanas 11–13", "Encontro trimestral com o dono", "Presença do dono", "Ata com combinados, dono e data"]],
       [2.4, 5.8, 4.4, 4.0], 8.5)
caixa("Próximo passo", "Escolher o ponto de entrada na seção 3, enviar os documentos básicos e agendar a conversa inicial com o seu gestor de projeto.")
h("Regras que protegem você", 3)
bullets(["Toda análise tributária e jurídica é pré-análise, validada por profissional habilitado, com fonte oficial citada e datada.",
         "Valores aparecem em faixa, com premissas. Não prometemos resultado, valor ou prazo.",
         "Crédito, antecipação e fundos são realizados por instituições autorizadas e reguladas. A Save estrutura e conduz o projeto.",
         "Honorários informados antes da contratação, sem cobrança antecipada em operações de crédito.",
         "Seus dados são tratados sob a LGPD, com acesso restrito e finalidade definida."])
para("Aviso: este documento é uma pré-análise de caráter geral. Não constitui oferta de crédito, de valores mobiliários nem parecer jurídico ou tributário. Estimativas dependem de documentos e de validação técnica. Estatísticas e dados de mercado são referências, com fonte e data indicadas.", 8, italic=True, color=C["texto_3"])
quebra()

# ───────────────────────── 8. FONTES
h("8. Fontes dos dados setoriais")
for s in setores:
    h(s["titulo"], 3)
    for f in s["fontes"]: para(f, 7.5, color=C["texto_2"], after=1)
rodape()
doc.save(OUT)
print("ok", OUT)
