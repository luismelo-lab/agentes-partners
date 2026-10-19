#!/usr/bin/env python3
"""Gera o Estudo de Jornada do Cliente (Word, padrão Save). Lê os JSON de setores e o JSON de benchmark."""
import json, sys, glob
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path("/home/user/agentes-partners")
SCR = Path(sys.argv[1])
OUT = Path(sys.argv[2])
M = json.load(open(REPO / "departamentos/marketing-marca/padrao/marca.json"))
C = M["claro"]; E = M["escuro"]
def rgb(h): h = h.lstrip("#"); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
SANS = "IBM Plex Sans"; MONO = "IBM Plex Mono"
LOGO_ESC = REPO / "departamentos/marketing-marca/padrao/logo-save-fundo-escuro.png"
LOGO_CLA = REPO / "departamentos/marketing-marca/padrao/logo-save-fundo-claro.png"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

def setfont(run, name=SANS, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = rgb(color)

st = doc.styles["Normal"]; st.font.name = SANS; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.15
for nome, tam, cor in (("Heading 1", 18, C["texto"]), ("Heading 2", 13, C["texto"]), ("Heading 3", 10.5, C["ouro_texto"])):
    s = doc.styles[nome]; s.font.name = SANS; s.font.size = Pt(tam); s.font.bold = True
    s.font.color.rgb = rgb(cor); s.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    s.element.rPr.rFonts.set(qn("w:ascii"), SANS); s.element.rPr.rFonts.set(qn("w:hAnsi"), SANS)
    s.paragraph_format.space_before = Pt(14 if nome != "Heading 3" else 8); s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.keep_with_next = True

def shade(cell, hexcor):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcor.lstrip("#"))
    tcPr.append(sh)

def borders(cell, cor, sz=4, lados=("top", "left", "bottom", "right")):
    tcPr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for l in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{l}")
        if l in lados: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), cor.lstrip("#"))
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def margens(cell, v=70, h=100):
    tcPr = cell._element.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for l, val in (("top", v), ("bottom", v), ("left", h), ("right", h)):
        e = OxmlElement(f"w:{l}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def para(texto="", size=10, bold=False, color=None, italic=False, align=None, after=None, before=None, style=None, mono=False):
    p = doc.add_paragraph(style=style)
    if texto:
        r = p.add_run(texto); setfont(r, MONO if mono else SANS, size, bold, color or C["texto"], italic)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p

def h(texto, n=1):
    p = doc.add_heading(level=n)
    r = p.add_run(texto)
    setfont(r, SANS, {1: 18, 2: 13, 3: 10.5}[n], True, C["texto"] if n < 3 else C["ouro_texto"])
    if n == 1:
        pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:color"), C["ouro"].lstrip("#")); b.set(qn("w:space"), "4")
        bd.append(b); pPr.append(bd)
    return p

def bullets(itens, size=10, cor=None):
    for i in itens:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(i); setfont(r, SANS, size, False, cor or C["texto"])

def tabela(cab, linhas, larguras, size=8.5, cab_escuro=True, zebra=True):
    t = doc.add_table(rows=1, cols=len(cab)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, c in enumerate(cab):
        cell = t.rows[0].cells[i]; cell.width = Cm(larguras[i]); shade(cell, E["fundo"] if cab_escuro else C["fundo_2"]); margens(cell)
        borders(cell, C["borda_2"]); cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        r = cell.paragraphs[0].add_run(c.upper()); setfont(r, MONO, 7.5, True, E["ouro"] if cab_escuro else C["texto"])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for k, ln in enumerate(linhas):
        row = t.add_row()
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        for i, v in enumerate(ln):
            cell = row.cells[i]; cell.width = Cm(larguras[i]); margens(cell); borders(cell, C["borda"])
            shade(cell, C["fundo"] if (zebra and k % 2) else C["cartao"])
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            neg = (i == 0)
            r = cell.paragraphs[0].add_run(str(v)); setfont(r, SANS, size, neg, C["texto"])
    for i,wd in enumerate(larguras): t.columns[i].width = Cm(wd)
    tblPr=t._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def caixa(titulo, texto, escuro=False):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; margens(c, 120, 180)
    shade(c, E["cartao"] if escuro else C["ouro_suave"]); borders(c, E["borda"] if escuro else C["ouro"], 6, ("left",))
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(titulo.upper()); setfont(r, MONO, 7.5, True, E["ouro"] if escuro else C["ouro_texto"])
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(texto); setfont(r, SANS, 10, False, E["texto"] if escuro else C["texto"])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def rodape():
    for s in doc.sections:
        s.different_first_page_header_footer = True
        f = s.footer; p = f.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run("Save Co. · Estudo de jornada do cliente · Pré-análise, sujeita a validação técnica. Sem promessa de valor, prazo ou êxito.   Pág. ")
        setfont(r, MONO, 7, False, C["texto_3"])
        for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            r2 = p.add_run(); setfont(r2, MONO, 7, False, C["texto_3"])
            if tipo:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); r2._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r2._r.append(it)

def quebra(): doc.add_page_break()

# ───────────────────────── dados
setores = [json.load(open(f)) for f in sorted(glob.glob(str(SCR / "setores/*.json")))]
bench = json.load(open(SCR / "benchmark.json")) if (SCR / "benchmark.json").exists() else {}

# ───────────────────────── CAPA
t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
c = t.rows[0].cells[0]; shade(c, E["fundo"]); margens(c, 500, 500); borders(c, E["fundo"], 4, ())
p = c.paragraphs[0]; p.add_run().add_picture(str(LOGO_ESC), height=Cm(1.1)); p.paragraph_format.space_after = Pt(150)
p = c.add_paragraph(); r = p.add_run("ESTUDO DE JORNADA DO CLIENTE"); setfont(r, MONO, 9, True, E["ouro"])
p = c.add_paragraph(); r = p.add_run("Do caixa à expansão: como a Save acompanha a sua empresa em cada etapa"); setfont(r, SANS, 26, True, E["texto"])
p.paragraph_format.space_after = Pt(10)
p = c.add_paragraph(); r = p.add_run("Modelo de estudo por setor · Recuperação tributária, gestão de contas, passivo, planejamento, estruturação patrimonial, crédito e capital"); setfont(r, SANS, 11, False, E["texto_2"])
p.paragraph_format.space_after = Pt(150)
p = c.add_paragraph(); r = p.add_run("<EMPRESA> · Setor: ____________ · Data-base: 19/10/2026"); setfont(r, MONO, 8.5, False, E["texto_2"])
p = c.add_paragraph(); r = p.add_run("PRÉ-ANÁLISE · ESTIMATIVA PRELIMINAR SUJEITA A VALIDAÇÃO TÉCNICA"); setfont(r, MONO, 8, True, E["ouro"])
for _ in range(2): c.add_paragraph()
quebra()

# ───────────────────────── 1. SUMÁRIO EXECUTIVO
h("1. Sumário executivo")
para("A maioria das empresas médias compra serviços um a um: um crédito tributário aqui, uma renegociação ali, um empréstimo quando o caixa aperta. O resultado é uma colcha de retalhos, em que cada decisão ignora as outras. Este estudo mostra a alternativa: uma jornada única, conduzida por um parceiro de negócio, na ordem em que o dinheiro e o risco pedem.", 10.5)
caixa("A tese em uma frase", "Caixa primeiro, estrutura depois, expansão por último. O caixa recuperado paga a etapa seguinte e abre os dados (DRE, balanço, débitos) de que a próxima etapa precisa.")
tabela(["Etapa", "Pergunta do dono", "O que a Save faz", "Janela"],
       [["1. Oxigenar", "Onde está o dinheiro que já é meu?", "Recuperação tributária, recorrência, crédito presumido, cobrança e gestão de contas", "0 a 6 meses"],
        ["2. Estruturar", "Como organizo o passivo e a carga para crescer?", "Gestão de passivo, defesa, planejamento e reforma tributária, antecipação, estruturação patrimonial", "3 a 12 meses"],
        ["3. Expandir", "Como financio o próximo salto?", "Crédito internacional, FIDC, captação de capital, novos mercados (em avaliação conjunta)", "6 a 24 meses"]],
       [2.6, 4.2, 7.0, 2.8], 9)
para("O que este documento entrega", 10.5, True, before=6)
bullets(["Um método de diagnóstico: os sinais que mostram qual serviço faz sentido, com o dado que o sustenta.",
         "Uma jornada em três ondas para cada um dos sete setores em que a Save mais atua, com aderência por serviço, indicadores de valor e cuidados.",
         "A forma como o valor é medido e apresentado ao longo do tempo, com revisão trimestral.",
         "As regras de conduta que protegem o cliente: independência, validação técnica e ausência de promessa de resultado.",
         "Um plano de 90 dias para colocar a jornada em marcha."])
para("Como ler: cada capítulo setorial é autônomo. O dono do negócio pode ir direto ao seu setor (capítulo 5) e voltar ao método (capítulos 2 a 4) para entender o porquê.", 9, italic=True, color=C["texto_2"])
quebra()

# ───────────────────────── 2. A JORNADA
h("2. A jornada em três ondas")
para("A ordem não é arbitrária. Ela segue cinco regras que valem para qualquer setor.", 10.5)
tabela(["#", "Regra", "Por que importa"],
       [["1", "O passivo vem antes do crédito novo", "Débito em aberto bloqueia a certidão, e sem certidão não há crédito mais barato."],
        ["2", "O crédito recuperado é o primeiro dinheiro novo", "A conversa da segunda onda começa mostrando o que a primeira trouxe."],
        ["3", "A reforma tributária cria urgência", "O PIS/COFINS acaba em 2027. A revisão das competências abertas deve ser fechada antes da transição."],
        ["4", "No máximo dois serviços novos por vez", "Para não sobrecarregar a empresa nem a equipe que entrega."],
        ["5", "Serviço em avaliação entra como conversa de descoberta", "Sem proposta, taxa ou prazo até a validação técnica e jurídica."]],
       [0.8, 6.0, 9.8], 9)
h("Onda 1 · Oxigenar", 2)
para("Objetivo: liberar caixa com risco baixo e abrir o acesso aos dados. Quem entrega: Tributário, com revisão técnica, e Contas e Soluções.")
bullets(["Recuperação tributária das teses aderentes ao perfil (CNAE × regime), nos últimos cinco anos.",
         "Recorrência: novas competências de teses já executadas.",
         "Crédito presumido de compras de pessoa física e Funrural, quando houver agroindústria.",
         "Gestão de Contas: cobrança automatizada por inteligência artificial (WhatsApp, Pix e boleto), paga no êxito, sem mensalidade."])
h("Onda 2 · Estruturar", 2)
para("Objetivo: organizar o passivo, a carga e a governança. Quem entrega: Jurídico, Tributário e Produtos Financeiros.")
bullets(["Gestão de passivo: transação, parcelamento e compensação com o próprio crédito, até a certidão negativa.",
         "Defesa e acompanhamento de débitos, intimações e ações.",
         "Planejamento tributário e reforma (CBS e IBS): regime, mix de produtos, créditos antes da extinção.",
         "Antecipação de recebíveis, por instituição autorizada, e alongamento de dívida bancária com carência.",
         "Estruturação patrimonial (proteção do patrimônio e preparação de garantias para a etapa seguinte)."])
h("Onda 3 · Expandir", 2)
para("Objetivo: financiar crescimento. Em avaliação conjunta, caso a caso, depois de passivo organizado e demonstrações em dia.")
bullets(["Crédito internacional em moeda forte, estruturado com projeto de uso dos recursos e acompanhamento. A Save estrutura; não concede crédito.",
         "FIDC próprio, para empresas com escala de recebíveis, capturando a receita financeira que hoje fica com bancos e fundos.",
         "Captação de capital: parceiros e investidores, nacionais e estrangeiros, com dossiê preparado.",
         "Novos mercados e internacionalização, em segunda rodada, depois de resolvida a capacidade de produção."])
quebra()

# ───────────────────────── 3. DIAGNÓSTICO DE ENTRADA
h("3. O diagnóstico de entrada: onde a jornada começa")
para("A jornada nasce de evidência, não de oferta. Cada serviço só entra na proposta se houver um sinal que o sustente e se estiver claro o que falta confirmar.", 10.5)
tabela(["Sinal no diagnóstico", "O que indica", "Serviço a considerar"],
       [["Contas a receber acima do faturamento mensal; vencidos altos", "Inadimplência e caixa apertado", "Gestão de Contas (cobrança), antecipação"],
        ["Despesa financeira alta e resultado baixo", "Dívida cara", "Alongamento e carência; crédito internacional, se houver ativos"],
        ["Dívida de curto prazo de 3 a 4 faturamentos", "Risco de insolvência", "Alongamento e carência, depois gestão de passivo"],
        ["Débito fiscal em aberto, sem certidão", "Bloqueio de crédito e de licitação", "Gestão de passivo e defesa"],
        ["Saldo credor, compras com alíquota zero, alimentos no Lucro Real", "Crédito a recuperar", "Recuperação tributária, recorrência, crédito presumido"],
        ["Faturamento acima de cerca de R$ 10 mi por mês e antecipação a banco", "Receita financeira capturada por terceiros", "FIDC próprio (em avaliação)"],
        ["Cadeia grande de fornecedores", "Oportunidade multicedente e multissacado", "FIDC (em avaliação)"],
        ["Ativos relevantes (terra, silo, galpão, lojas quitadas)", "Base para garantir crédito de longo prazo", "Estruturação patrimonial; crédito internacional"],
        ["Crescimento acima da capacidade; investimento antes de gerar caixa", "Necessidade de capital de expansão", "Captação de capital (dívida, equity, sócio estratégico)"]],
       [6.0, 4.4, 6.2], 8.5)
h("Documentos que abrem o diagnóstico", 3)
bullets(["Fiscal: EFD-Contribuições, EFD ICMS/IPI, eSocial/DCTFWeb e consulta de débitos.",
         "Financeiro: DRE e balanço dos últimos dois anos, curva ABC de endividamento e consulta ao SCR do Banco Central.",
         "Comercial: carteira de recebíveis com aging e relação dos maiores devedores.",
         "Patrimônio: matrículas, contratos de arrendamento e seguros, quando houver garantia em discussão."])
para("Se o dono ainda não se sente à vontade para abrir esses dados, a conversa começa pelo caixa: a cobrança e a gestão de contas resolvem uma dor imediata e constroem a confiança que libera o restante. Uma pergunta ajuda a achar a janela certa: em que mês do ano o seu caixa aperta mais?", 10)
quebra()

# ───────────────────────── 4. CATÁLOGO
h("4. O catálogo de serviços")
para("Seis frentes, cada uma com dono técnico dentro da Save e com a validação humana de profissional habilitado.", 10.5)
tabela(["Frente", "Serviços", "Estado", "Remuneração"],
       [["Tributário", "Recuperação por tese, recorrência, crédito presumido, planejamento e reforma, análise de créditos antes da extinção do PIS/COFINS", "Disponível", "Honorário de êxito, no resultado"],
        ["Jurídico", "Gestão de passivo (transação, parcelamento, compensação), defesa administrativa, contencioso e acompanhamento", "Disponível", "Conforme escopo"],
        ["Contas e caixa", "Cobrança por IA, antecipação por instituição autorizada, limite automático sobre a carteira registrada, integração gratuita ao sistema do cliente", "Em implantação e avaliação conjunta", "Cobrança paga no êxito; demais itens definidos na proposta"],
        ["Estrutura e patrimônio", "Estruturação patrimonial, certificado digital, apoio documental", "Em avaliação conjunta", "Conforme escopo"],
        ["Crédito e capital", "Crédito internacional, FIDC, captação de capital, home equity", "Em avaliação conjunta", "Êxito, só após a liberação; sem cobrança antecipada de análise"],
        ["Mercados externos", "Internacionalização, produção no Paraguai pelo regime de maquila (estudo de viabilidade)", "Conversa de descoberta", "A definir após o estudo"]],
       [3.0, 7.2, 3.0, 3.4], 8.5)
h("Como a Save trabalha em cada serviço", 3)
bullets(["Estrutura, nunca concede crédito: crédito internacional e antecipação são realizados por instituições autorizadas.",
         "Não distribui cotas de fundos nem faz oferta de valores mobiliários.",
         "A conta garantia (escrow) reduz o risco de bloqueio automático, mas não impede bloqueio por ordem judicial.",
         "A cobrança por IA tem pessoas nas exceções, e cada exceção melhora o sistema. Estatísticas de inadimplência citadas são referências de mercado, não resultados da Save."])
quebra()

# ───────────────────────── 5. SETORES
h("5. A jornada por setor")
para("Sete perfis, construídos sobre a base de empresas atendidas. Cada capítulo traz o retrato do setor em 2026, as dores, o que faz a empresa comprar agora, a jornada em três ondas, a aderência de cada serviço, os indicadores de valor e os cuidados.", 10.5)
tabela(["#", "Setor", "Perfis atendidos", "1º serviço além da recuperação tributária"],
       [[str(i + 1), s["titulo"], s["subtitulo"], ""] for i, s in enumerate(setores)],
       [0.8, 4.6, 6.0, 5.2], 9)
RESUMO = {  # espelha o quadro "3 produtos de maior aderência" do playbook
    "proteina-animal": "Antecipação de recebíveis (por instituição autorizada)", "cafe-graos": "Antecipação para o giro da safra",
    "varejo-alimentar": "Planejamento e reforma (mix por NCM)", "combustiveis-glp": "Gestão de passivo",
    "confeccao-textil": "Gestão de Contas", "transporte-carga": "Gestão de passivo", "panificacao-alimentos": "Planejamento e reforma"}
# preenche a coluna 4 da última tabela
tb = doc.tables[-1]
for i, s in enumerate(setores):
    cell = tb.rows[i + 1].cells[3]; cell.paragraphs[0].clear() if False else None
    r = cell.paragraphs[0].add_run(RESUMO.get(s["id"], "—")); setfont(r, SANS, 9, False, C["texto"])
quebra()

for n, s in enumerate(setores, 1):
    h(f"5.{n} {s['titulo']}")
    para(s["subtitulo"].capitalize() + " · " + s["base_clientes"], 9, color=C["texto_2"], italic=True)
    h("Retrato do setor em 2026", 3); bullets(s["retrato"], 9.5)
    h("Dores do dono e do CFO", 3); bullets(s["dores"], 9.5)
    h("O que faz a empresa decidir agora", 3); bullets(s["gatilhos"], 9.5)
    h("A jornada em três ondas", 2)
    for o in s["jornada"]:
        para(f"{o['onda']} · {o['janela']}", 10, True, C["ouro_texto"], after=1, before=4)
        para(o["objetivo"], 9.5, italic=True, after=2)
        tabela(["Serviço", "Por que se aplica", "Estado"], [[x["nome"], x["porque"], x["status"]] for x in o["solucoes"]], [4.2, 9.2, 3.2], 8.5)
        para("Quem entrega: " + o["entrega"], 8.5, color=C["texto_2"], after=4)
    h("Aderência por serviço", 3)
    tabela(["Serviço", "Aderência", "Observação"], [[a["solucao"], a["nivel"], a["observacao"]] for a in s["aderencia"]], [5.0, 2.2, 9.4], 8.5)
    h("Como medimos o valor", 3)
    tabela(["Indicador", "Como medir"], [[k["indicador"], k["como_medir"]] for k in s["kpis"]], [6.0, 10.6], 8.5)
    caixa("Em uma frase ao dono", s["frase_valor"])
    h("Como uma grande consultoria aborda este setor", 3); bullets(s["big_four"], 9.5)
    h("Cuidados e limites", 3); bullets(s["cuidados"], 9.5)
    quebra()

# ───────────────────────── 6. VALOR
h("6. Como o valor é medido e apresentado")
para("O cliente compra a etapa seguinte quando enxerga o ganho da anterior. Por isso, cada conta tem um Mapa de Valor, revisado a cada trimestre.", 10.5)
tabela(["Bloco do Mapa de Valor", "O que mostra"],
       [["O que já entregamos", "Teses executadas, valor do negócio e valor recebido. Sem dado confirmado, consta “em apuração”."],
        ["O que isso significou para o caixa", "Caixa liberado, compensado ou restituído e a variação da carga efetiva, em faixa."],
        ["O que ainda está disponível", "Para cada serviço: a evidência que o sustenta, a onda, a faixa de valor quando há base e as premissas."],
        ["Os próximos 90 dias", "O que entra agora, quem entrega e qual documento é necessário."],
        ["Riscos e cuidados", "Prazos (decadência de cinco anos, reforma em 2027), validações técnicas pendentes e limites de cada serviço."]],
       [5.0, 11.6], 9)
tabela(["Serviço", "Indicador principal", "Fonte do dado"],
       [["Recuperação tributária", "Caixa liberado e carga efetiva", "Resultado dos pedidos e tributo pago no período"],
        ["Recorrência", "Caixa novo por competência", "Média trimestral do diagnóstico"],
        ["Gestão de passivo", "Débito resolvido, desconto obtido, certidão", "Parecer e acordo firmado"],
        ["Planejamento e reforma", "Variação da carga em pontos percentuais, 2027 a 2033", "Simulação por NCM e regime"],
        ["Gestão de Contas", "Prazo médio de recebimento e inadimplência recuperada", "Carteira e vencidos do cliente"],
        ["Antecipação e FIDC", "Custo do capital de giro contra a alternativa bancária", "Taxa de mercado e proposta"],
        ["Crédito internacional", "Custo e prazo da dívida contra a atual", "Dívida atual e proposta aprovada"]],
       [4.2, 6.4, 6.0], 8.5)
para("Pauta do encontro trimestral (45 minutos): resultados do trimestre; os clientes com mais aderência; a onda que entra em cada um; pendências; combinados com dono e data.", 9.5)
quebra()

# ───────────────────────── 7. BENCHMARK
h("7. Como as grandes consultorias conduzem o relacionamento, e o que replicamos")
if bench:
    para(bench.get("intro", ""), 10)
    tabela(["Prática de mercado", "O que é", "Como a Save replica"], [[x["pratica"], x["descricao"], x["replica"]] for x in bench["praticas"]], [4.2, 6.0, 6.4], 8.5)
    if bench.get("referencias"):
        h("Referências de mercado usadas neste estudo", 3)
        tabela(["Dado", "Valor e data", "Fonte"], [[x["dado"], x["valor"], x["fonte"]] for x in bench["referencias"]], [5.0, 4.8, 6.8], 8)
else:
    para("[seção a preencher com o benchmark]", 10)
quebra()

# ───────────────────────── 8. GOVERNANÇA
h("8. Governança: independência, validação e transparência")
para("Quem atua em tributário, crédito e capital precisa mostrar com clareza onde termina um papel e começa o outro. As regras abaixo valem em toda a jornada.", 10.5)
tabela(["Regra", "Na prática"],
       [["Validação humana", "Toda análise tributária e jurídica é pré-análise, validada por técnico ou advogado habilitado, com fonte oficial citada e datada."],
        ["Sem promessa de resultado", "Valor, prazo e êxito nunca são garantidos. Valores aparecem em faixa, com premissas e a marca de estimativa preliminar."],
        ["Papéis separados", "A Save estrutura e facilita. Crédito, antecipação e fundos são realizados por instituições autorizadas e reguladas."],
        ["Transparência de remuneração", "Honorário de êxito informado antes da contratação; sem cobrança antecipada em crédito; custos de terceiros apresentados à parte."],
        ["Conflito de interesse", "Quando houver mais de um serviço para o mesmo cliente, a proposta apresenta a alternativa de mercado ao lado e quem se beneficia de cada decisão."],
        ["Proteção de dados", "Dados do cliente tratados sob a LGPD, com acesso restrito e finalidade definida; nada sai sem autorização."],
        ["Linguagem", "Termos precisos: estruturação patrimonial, crédito internacional, estimativa preliminar. Sem “garantido” e sem “dinheiro barato”."]],
       [4.2, 12.4], 9)
h("Quem faz o quê dentro da Save", 3)
tabela(["Papel", "Responsabilidade"],
       [["Consultor de negócio", "Fala de negócio com o dono, identifica a dor, aciona especialistas e mantém carteira pequena para garantir profundidade."],
        ["Responsável pelo projeto", "Traduz a demanda para a linguagem técnica e coordena financeiro, jurídico e tributário."],
        ["Equipes técnicas", "Tributário, Jurídico, Contas e Soluções, Produtos Financeiros e Internacionais, Contratos e Pós-venda."],
        ["Revisor de promessas e qualidade", "Confere linguagem, limites de cada serviço e aderência às regras antes de qualquer material sair."]],
       [4.6, 12.0], 9)
quebra()

# ───────────────────────── 9. ROADMAP
h("9. Plano de 90 dias")
tabela(["Quando", "O que acontece", "Responsável", "Entregável"],
       [["Semanas 1–2", "Conversa inicial, acordo de confidencialidade e coleta dos documentos de entrada", "Consultor + dono da empresa", "Lista de documentos e cronograma"],
        ["Semanas 3–4", "Diagnóstico fiscal e financeiro; mapa de sinais; escolha dos até dois primeiros serviços", "Tributário, Contas e Soluções", "Mapa de aderência da conta"],
        ["Semanas 5–8", "Execução da onda 1; início da cobrança, se aplicável; consulta de débitos", "Tributário, Jurídico, Contas", "Pedidos protocolados; régua de cobrança no ar"],
        ["Semanas 9–10", "Primeiro Mapa de Valor; desenho da onda 2", "Contas e Soluções", "Mapa de Valor e plano da onda 2"],
        ["Semanas 11–13", "Encontro trimestral com o dono; decisão sobre a próxima onda e conversa de descoberta de crédito e capital, se couber", "Consultor + dono", "Ata com combinados, dono e data"]],
       [2.4, 7.0, 3.6, 3.6], 8.5)
caixa("Próximo passo", "Escolher o setor do seu negócio no capítulo 5, enviar os documentos da seção 3 e agendar a conversa inicial. A partir dela, a Save propõe os primeiros serviços e os indicadores que serão acompanhados.")
para("Aviso: este documento é uma pré-análise de caráter geral. Não constitui oferta de crédito, de valores mobiliários nem parecer jurídico ou tributário. Estimativas dependem de documentos e de validação técnica. Estatísticas e dados de mercado são referências, com fonte e data indicadas, e não resultados da Save.", 8, italic=True, color=C["texto_3"])
quebra()

# ───────────────────────── 10. FONTES
h("10. Fontes")
for s in setores:
    h(s["titulo"], 3)
    for f in s["fontes"]: para(f, 7.5, color=C["texto_2"], after=1, mono=False)
rodape()
doc.save(OUT)
print("ok", OUT)
