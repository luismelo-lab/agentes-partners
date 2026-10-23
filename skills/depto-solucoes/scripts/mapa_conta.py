#!/usr/bin/env python3
"""Mapa de aderência da conta: o que mais cada cliente pode consumir da Save, em que ordem e com que evidência.

Somente leitura e local. Cruza o download do HubSpot do dia (todos os funis, pela raiz do CNPJ), os pareceres
lidos (celulas/geovana/pareceres/resumos.json) e a base local da Receita (cnpj.sqlite, read-only).
Nenhum CNPJ sai da máquina; não lê .env; não grava no HubSpot.

As regras estão na tabela PRODUTOS, uma linha por produto, explícitas para o time revisar. O playbook do setor
(references/playbooks-setoriais.md) explica o porquê; o script só aponta o sinal que existe nos dados.

Uso:
  python3 mapa_conta.py --negocio <id>            mapa de um cliente (JSON na tela)
  python3 mapa_conta.py --cnpj <cnpj>
  python3 mapa_conta.py --carteira [--master X]   todos os clientes ganhos → JSON agregado + Excel por Master
"""
import argparse, collections, glob, json, os, re, sqlite3, sys
from datetime import date
sys.path.insert(0, os.path.expanduser("~/agentes-partners/lib"))
from cnpj import normalizar_cnpj  # noqa: E402  CNPJ alfanumérico (IN RFB 2.229/2024)

AP = os.path.expanduser("~/agentes-partners")
OUT = os.path.join(AP, "departamentos", "solucoes")
HOJE = date.today()

# ── setores (cluster pelo CNAE de 4 dígitos) ─────────────────────────────────────────────────────────────
SETORES = [
    ("Proteína animal", ("1011", "1012", "1013", "4634", "4722")),
    ("Café e grãos", ("4621", "4622", "4623", "1061", "1081")),
    ("Varejo alimentar", ("4711", "4712", "4721", "4724", "4729", "4633", "4637", "4639", "4691")),
    ("Combustíveis e GLP", ("4731", "4732", "4681", "4682", "4784")),
    ("Confecção e têxtil", ("1311", "1321", "1330", "1340", "1351", "1352", "1353", "1354", "1359", "1411", "1412", "1413", "1414", "1421", "1422", "1531", "1532", "1533", "1539", "1540")),
    ("Transporte de carga", ("4930", "5211", "5212", "5250")),
    ("Panificação e alimentos", ("1091", "1092", "1093", "1094", "1095", "1096", "1099", "1052", "1053")),
]
INDUSTRIA = lambda c: c[:2].isdigit() and 10 <= int(c[:2]) <= 33
ATACADO = lambda c: c[:2] == "46"
VAREJO = lambda c: c[:2] == "47"


def setor(cnae):
    c = (cnae or "")[:4]
    for nome, cs in SETORES:
        if c in cs: return nome
    if INDUSTRIA(c): return "Indústria (outros)"
    if ATACADO(c): return "Atacado (outros)"
    if VAREJO(c): return "Varejo (outros)"
    return "Outros"


# faturamento anual: faixa do HubSpot (limite inferior) ou valor informado; None = desconhecido
FAIXA = {"R$81K - R$360K": 81e3, "R$360K - R$1M": 360e3, "R$1M - R$4.8M": 1e6, "R$4.8M - R$10M": 4.8e6, "R$10M - R$30M": 10e6,
         "R$30M - R$100M": 30e6, "R$100M - R$300M": 100e6, "R$300M+": 300e6}


def fat(ex):
    try:
        v = float(ex.get("faturamento") or 0)
        if v > 0: return v
    except Exception: pass
    return FAIXA.get(str(ex.get("faixa_faturamento") or ""))


SEM_DEB = re.compile(r"n[ãa]o (possui|consta|tem)|sem d[ée]bitos|n[ãa]o foi necess", re.I)
COM_DEB = re.compile(r"possui d[ée]bitos|d[ée]bitos? (em aberto|com exigibilidade|inscrit)|parcelamento|PGFN|d[íi]vida ativa|transa[çc][ãa]o", re.I)

# ── produtos: (id, nome, onda da esteira, status, área/skill, regra(s) → (aderência, motivo) ou None) ─────────
# onda: 1 Destravar caixa (caixa rápido, risco baixo) · 2 Proteger e controlar (passivo, carga, governança) · 3 Crescer (funding, novos mercados)
# status: pronto · em validação (trava do G2: sem taxa/prazo/fee) · em estruturação · em definição
# 19/10/2026 (Luis): GC, ANT e FIDC são prontos, ofertados pela Save Invest; EP é pronto, conduzido pelo jurídico; CI e MQ conduzidos pelo João Mattos


def r_recorrencia(s):
    n = s["competencias_novas"]
    if s["rt_ganho"] and n and n >= 12: return ("alta", f"{n} competências novas desde a última análise")
    if s["rt_ganho"] and n and n >= 6: return ("média", f"{n} competências novas; reavaliar quando chegar a 12")


def r_aliquota_zero(s):
    if s["setor"] in ("Varejo alimentar",) and s["cnae"][:4] in ("4711", "4712"): return None  # catálogo: fora supermercado
    if s["cnae"][:4] == "4621": return None  # catálogo: fora café
    if s["setor"] in ("Proteína animal", "Varejo alimentar", "Café e grãos") or ATACADO(s["cnae"]):
        if not s["teses"]: return ("média", "perfil do catálogo (revenda com alíquota zero); parecer não lido: confirmar se a tese já foi feita")
        if "Teses de aquisições" not in s["teses"]: return ("alta", "perfil do catálogo (revenda com alíquota zero) e tese de aquisições não aparece nos pareceres. Risco: EREsp 2.215.482 pendente na 1ª Seção do STJ e LC 224/2025 (radar tributário, 16/10)")
        if s["teses"]["Teses de aquisições"] in ("não analisada", "None", ""): return ("média", "tese de aquisições sem desfecho no parecer")


def r_previdenciaria(s):
    if not s["teses"]: return None
    if "Teses previdenciárias" not in s["teses"] and s["regime"] in ("Lucro Real", "Lucro Presumido"):
        if s["setor"] in ("Confecção e têxtil", "Transporte de carga"):
            return ("alta", "setor intensivo em folha, com reoneração da CPP até 2027 (playbook): CPP sobre verbas indenizatórias e Sistema S não aparecem nos pareceres")
        return ("média", "CPP sobre verbas indenizatórias e Sistema S não aparecem nos pareceres (confirmar folha)")


def r_agro_pf(s):
    c4 = s["cnae"][:4]
    if s["regime"] != "Lucro Real" or c4 not in PRESUMIDO_CADEIA: return None
    if c4 in ("1011", "1013"): return ("baixa", "crédito presumido de bovinos só na parcela exportada (Lei 12.058, art. 33); mercado interno só na via judicial, risco alto")
    return ("média", PRESUMIDO_CADEIA[c4])


def r_icms_st(s):
    if s["setor"] in ("Combustíveis e GLP",) or s["cnae"][:4] in ("4635", "4723", "4644", "4771", "4530"):
        return ("média", "revenda com ICMS-ST: exclusão do ICMS-ST da base (Tema 1.125 STJ); combustível é monofásico, conferir")


def r_passivo(s):
    if not s["passivo_debitos"]: return None
    vals = [float(v.replace(".", "").replace(",", ".")) for v in re.findall(r"R\$\s*([\d.]+,\d{2})", s.get("passivo_txt") or "")]
    if vals and max(vals) < 100e3:
        return ("baixa", "débito pequeno no parecer (R$ " + f"{max(vals):,.0f}".replace(",", ".") + "): resolver no próprio fluxo, sem produto à parte")
    return ("alta" if (not vals or max(vals) >= 500e3) else "média", "o parecer aponta débitos (Receita/PGFN): transação, parcelamento ou pagamento em DARF (na alíquota zero o crédito sai por PER, não por DCOMP)")


def r_juridico(s):
    if s["passivo_suspenso"]: return ("média", "débito com exigibilidade suspensa: acompanhar a defesa e os prazos (Produtos Jurídicos)")
    if s["acao_judicial"]:
        return ("baixa", "ação judicial em alguma tese: conferir se há acompanhamento contratado")


def r_planejamento(s):
    f = s["faturamento"] or 0
    if s["regime"] in ("Lucro Real", "Lucro Presumido") and f >= 4.8e6:
        a = "alta" if (f >= 30e6 or s["setor"] in ("Proteína animal", "Varejo alimentar", "Panificação e alimentos")) else "média"
        return (a, "Real/Presumido com faturamento ≥ R$ 4,8 mi: regime, reforma (CBS/IBS desde 2027; cesta básica no alimento) e créditos antigos antes da transição")


def r_gestao_contas(s):
    f = s["faturamento"] or 0
    b2b = ATACADO(s["cnae"]) or INDUSTRIA(s["cnae"]) or s["setor"] == "Transporte de carga"
    if b2b and f >= 4.8e6: return ("alta" if f >= 10e6 else "média", "vende a prazo para empresas (B2B): cobrança organizada e porta para DRE e balanço (oferta da Save Invest)")
    if VAREJO(s["cnae"]) and f >= 4.8e6: return ("baixa", "varejo recebe à vista ou no cartão: só se houver venda a prazo relevante")


def r_antecipacao(s):
    f = s["faturamento"] or 0
    if (ATACADO(s["cnae"]) or INDUSTRIA(s["cnae"])) and f >= 4.8e6: return ("média", "carteira a receber B2B: antecipação por instituição autorizada (a Save não antecipa)")


def r_fidc(s):
    f = s["faturamento"] or 0
    if f >= 120e6 and (ATACADO(s["cnae"]) or INDUSTRIA(s["cnae"])): return ("média", "faturamento ≥ ~R$ 10 mi por mês com recebíveis B2B (critério de 13/10)")


def r_ci(s):
    f = s["faturamento"] or 0
    if f >= 50e6 and (s["setor"] in ("Proteína animal", "Café e grãos") or INDUSTRIA(s["cnae"]) or ATACADO(s["cnae"])):
        return ("média", "faturamento ≥ R$ 50 mi/ano em setor do apetite; falta confirmar garantia livre (~2x) e necessidade ≥ US$ 5 mi")
    if f >= 30e6 and s["setor"] in ("Proteína animal", "Café e grãos"): return ("baixa", "agro entre R$ 30 e 50 mi: só com ativos > R$ 60 mi (a confirmar com o dono)")


def r_maquila(s):
    if s["setor"] == "Confecção e têxtil" or s["cnae"][:2] in ("22", "15"):
        return ("média" if s["regime"] == "Lucro Real" else "baixa", "indústria intensiva em mão de obra (corte do MQ)")
    if s["cnae"][:4] in ("1011", "1012", "1013"): return ("baixa", "carne em aberto no MQ (pergunta MQ-Q6, até 13/11); o simulador não roda para frigorífico (D59); a D&S é a proposta de Master piloto (MQ-D8)")


def r_patrimonial(s):
    f = s["faturamento"] or 0; cap = s["capital"] or 0
    if f >= 20e6 or cap >= 5e6: return ("média", "porte e patrimônio indicam estruturação patrimonial e sucessória")


ST_CNAES = ("4635", "4723", "4644", "4771", "4530", "4731", "4732", "4681", "4682", "4784")
SERVICO = lambda c: c[:2].isdigit() and int(c[:2]) >= 49 and c[:2] not in ("49",)


def teses_segmento(s):
    """Teses do catálogo (departamentos/tributario/catalogo-teses-save.md) que cabem no segmento e no regime.
    É o perfil do segmento, não o diagnóstico da empresa: o técnico confirma com os documentos.
    Revisado pelo Gerente Tributário em 16/10/2026 (ajustes C1–C14). Fora da lista (radar tributário):
    Sistema S com teto de 20 salários (STJ Tema 1.390), PIS/COFINS fora do IRPJ/CSLL no Presumido (Tema 1.312),
    PIS/COFINS de combustível monofásico (Tema 1.339), IPI não recuperável depois de 20/12/2022 (Tema 1.373),
    terço de férias (Tema 985 STF), DIFAL de 2022 sem ação até 29/11/2023 (Tema 1.266) e DCOMP na tese de alíquota zero."""
    c, c4, reg, f = s["cnae"], s["cnae"][:4], s["regime"], s["faturamento"] or 0
    real, rp = reg == "Lucro Real", reg in ("Lucro Real", "Lucro Presumido")
    saida_zero = s["setor"] == "Proteína animal"  # receita a alíquota zero: exclusões sem base
    mono = c4 in ("4731", "4732", "4771", "4530", "4723")
    nota = " (só sobre a receita fora do monofásico)" if mono else ""
    out = []
    az = (s["setor"] in ("Proteína animal", "Café e grãos") or c4 in ("4724", "4633") or ATACADO(c)) and c4 not in ("4621", "1081", "4711", "4712")
    if az and real: out.append("PIS/COFINS nas aquisições (alíquota zero; pedido de ressarcimento (PER); EREsp 2.215.482 pendente; desde 04/2026, recortar NCM fora da cesta)")
    if real and (INDUSTRIA(c) or ATACADO(c) or c[:2] == "49") and c4 not in ("4731", "4732"): out.append("PIS/COFINS sobre insumos")
    if rp and not saida_zero and not SERVICO(c) and f >= 3.6e6: out.append("Exclusão do ICMS da base do PIS/COFINS" + nota)
    if rp and not saida_zero and c4 in ST_CNAES: out.append("Exclusão do ICMS-ST da base" + nota)
    if rp and not saida_zero: out.append("Exclusão do PIS/COFINS da própria base" + nota)
    if real and (INDUSTRIA(c) or ATACADO(c) or VAREJO(c)): out.append("ICMS nos créditos de PIS/COFINS (entrada)")
    if rp and SERVICO(c): out.append("Exclusão do ISS da base do PIS/COFINS (JUD; Tema 118 STF empatado: não somar)")
    if real and (az or saida_zero): out.append("Crédito de PIS/COFINS em conta gráfica (conferir o saldo; escriturar, art. 378 da LC 214)")
    if real and INDUSTRIA(c) and c4 not in ("1011", "1012"):
        out.append("Crédito de IPI nas aquisições de atacadista")
        out.append("Crédito de IPI em conta gráfica" + (" (só se o produto não for NT na TIPI)" if c4 == "1013" else ""))
    if rp: out.append("CPP sobre verbas indenizatórias (sem terço de férias, Tema 985 STF; sem teto do Sistema S, Tema 1.390)")
    if real and (INDUSTRIA(c) or ATACADO(c)): out.append("Crédito presumido de ICMS fora do IRPJ/CSLL e do PIS/COFINS (se houver benefício; Tema 1.416 STJ suspenso: só quantificar)")
    if real: out.append("Creditamento ACT/CCT (se a convenção obrigar itens essenciais)")
    if real and c4 in PRESUMIDO_CADEIA:
        out += [PRESUMIDO_CADEIA[c4] + "; 90% desde 04/2026 (LC 224)", "Funrural (ADI 4.395 pendente: monitorar, não somar)"]
    return out


BOV = "Crédito presumido de bovinos (Lei 12.058, art. 33: só a parcela exportada; mercado interno = JUD, risco alto)"
AGRO = "Crédito presumido da agroindústria (Lei 10.925, art. 8º: só abate débito; ressarcimento = JUD)"
PRESUMIDO_CADEIA = {"1011": BOV, "1013": BOV, "1012": "Crédito presumido de suínos e aves (Lei 12.350, arts. 55, 56-A e 56-B; conferir)",
                    "1052": "Crédito presumido do leite (Lei 10.925, arts. 8º e 9º-A; ressarcível)",
                    "1081": "Crédito presumido do café (Lei 12.599, arts. 5º e 6º; só exportador)", "1061": AGRO, "1031": AGRO, "1033": AGRO}


def r_rt(s):
    """Recuperação Tributária como um produto só: junta os sinais das teses e da recorrência."""
    sub = [x for x in (r_recorrencia(s), r_aliquota_zero(s), r_agro_pf(s), r_previdenciaria(s), r_icms_st(s)) if x]
    if not sub and not teses_segmento(s): return None
    if not sub: return ("média", "teses do catálogo cabem no segmento; conferir o que o parecer ainda não cobriu")
    melhor = max(sub, key=lambda x: PESO.get(x[0], 0))
    return (melhor[0], "; ".join(m for _, m in sub))


PRODUTOS = [
    ("RT", "Recuperação Tributária", 1, "pronto", "depto-tributario + recuperador + revisor-diagnostico", r_rt),
    ("GC", "Gestão de Contas (cobrança com IA)", 1, "pronto", "Save Invest", r_gestao_contas),
    ("PAS", "Gestão de Passivo (transação e parcelamento)", 1, "pronto", "depto-juridico + depto-tributario", r_passivo),
    ("JUR", "Produtos Jurídicos (defesa e acompanhamento)", 2, "pronto", "depto-juridico", r_juridico),
    ("PLAN", "Planejamento tributário e reforma (CBS/IBS)", 2, "pronto", "depto-tributario", r_planejamento),
    ("EP", "Estruturação patrimonial", 2, "pronto", "depto-juridico", r_patrimonial),
    ("ANT", "Antecipação de recebíveis (instituição autorizada)", 1, "pronto", "Save Invest", r_antecipacao),
    ("FIDC", "FIDC", 3, "pronto", "Save Invest", r_fidc),
    ("CI", "Crédito Internacional (funding)", 3, "em validação", "João Mattos (depto-internacional)", r_ci),
    ("MQ", "Maquila no Paraguai", 3, "em estruturação", "João Mattos e Paulo (depto-internacional)", r_maquila),
]
ONDA = {1: "1 · Destravar caixa", 2: "2 · Proteger e controlar", 3: "3 · Crescer"}
PESO = {"alta": 3, "média": 2, "baixa": 1, "bloqueado": 0}
FUNIL_PRODUTO = {"Gestão de Passivo": "PAS", "Produtos Jurídicos": "JUR", "Planejamento Tributário": "PLAN"}


def carregar():
    f = sorted(glob.glob(os.path.join(AP, "departamentos/resultados/dados/hubspot-*.json")))[-1]
    neg = json.load(open(f))["negocios"]
    try: res = json.load(open(os.path.join(AP, "celulas/geovana/pareceres/resumos.json")))
    except Exception: res = {}
    db = sqlite3.connect("file:" + os.path.join(AP, "dados-publicos/receita/cnpj.sqlite") + "?mode=ro", uri=True)
    return os.path.basename(f), neg, res, db


def cnpj_de(n): return normalizar_cnpj((n.get("extras") or {}).get("cnpj")).zfill(14) if (n.get("extras") or {}).get("cnpj") else ""


def sinais(cnpj, negs, res, db):
    """Junta tudo o que a Save sabe da empresa (todos os negócios da mesma raiz de CNPJ)."""
    raiz = cnpj[:8]
    r = db.execute("select cnae, cnaes_sec, situacao, uf, inicio from estab where cnpj=?", (cnpj,)).fetchone() or ("", "", "", "", "")
    e = db.execute("select capital, porte, razao from empresa where basico=?", (raiz,)).fetchone() or (None, None, None)
    ganhos = [n for n in negs if n["status"] == "ganho" and n["pipeline"].startswith("RT")]
    # duplicidade suspeita: dois ganhos fechados no mesmo dia, um deles criado depois do fechamento (importação nov–dez/2025)
    dup = set(); por_dia = collections.defaultdict(list)
    for n in ganhos: por_dia[n.get("fechado")].append(n)
    for d, g in por_dia.items():
        if d and len(g) > 1:
            susp = [n for n in g if (n.get("criado") or "") > d]
            dup |= {str(n["id"]) for n in (susp if len(susp) < len(g) else susp[1:])}
    validos = [n for n in ganhos if str(n["id"]) not in dup]
    # extras: o valor mais recente que existir em qualquer negócio da raiz (o último card pode estar vazio)
    ex = {}
    for n in sorted(negs, key=lambda n: n.get("criado") or "", reverse=True):
        for k, v in (n.get("extras") or {}).items():
            if k not in ex and v not in (None, "", "None"): ex[k] = v
    rost = collections.Counter(n.get("parceiro") for n in negs if n.get("parceiro") and not str(n.get("parceiro")).startswith("Outros"))
    rec = [float(n["recebido"]) for n in validos if str(n.get("recebido")) not in ("None", "", "0", "0.0")]
    abertos = sorted([n for n in negs if n["status"] == "aberto"], key=lambda n: n.get("criado") or "")
    link = (abertos or sorted(validos, key=lambda n: n.get("fechado") or "") or negs)[-1]["id"]
    teses, passivo, acao, comp = {}, "", False, None
    for n in sorted(negs, key=lambda n: n.get("criado") or ""):
        p = res.get(str(n["id"])) or {}
        t = p.get("teses")
        if isinstance(t, dict):
            for k, v in t.items(): teses[k] = (v or {}).get("desfecho") if isinstance(v, dict) else str(v)
        if p.get("passivo"): passivo = str(p["passivo"])
        if "há ação judicial" in str(p.get("acoes_judiciais") or ""): acao = True
        if p.get("competencias_novas_desde_analise") not in (None, "None"):
            try: comp = int(p["competencias_novas_desde_analise"])
            except Exception: pass
    ja = {FUNIL_PRODUTO[n["pipeline"]]: n["status"] for n in negs if n["pipeline"] in FUNIL_PRODUTO}
    nome = e[2] or re.sub(r"^\s*(recorr[êe]ncia|recorrencia)\s*[-–:]?\s*(\d\s*tri\s*\d{4}\s*[-–:]?)?\s*", "", negs[0]["nome"], flags=re.I)
    return {"cnpj": cnpj, "nome": nome, "cnae": r[0] or "", "setor": setor(r[0]), "uf": r[3] or ex.get("uf"), "situacao": r[2],
            "regime": ex.get("regime"), "faturamento": fat(ex), "faixa": ex.get("faixa_faturamento"), "capital": e[0], "porte_receita": e[1],
            "master": (rost.most_common(1)[0][0] if rost else "Fora do roster"), "masters": sorted(rost),
            "rt_ganho": bool(ganhos), "rt_ganhos": len(validos), "rt_valor": round(sum(float(n.get("valor") or 0) for n in validos), 2),
            "rt_recebido": round(sum(rec), 2) if rec else None, "duplicidade_suspeita": sorted(dup), "link_negocio": str(link),
            "honorario_contratado": sorted({n.get("hContr") for n in validos if str(n.get("hContr")) not in ("None", "")}),
            "ultimo_ganho": max((n.get("fechado") or "" for n in validos), default=None),
            "teses": teses, "passivo_debitos": bool(passivo and COM_DEB.search(passivo) and not SEM_DEB.search(passivo)),
            "passivo_suspenso": "exigibilidade suspensa" in passivo.lower(), "passivo_txt": passivo, "acao_judicial": acao, "competencias_novas": comp,
            "ja_tem": ja, "negocios": [str(n["id"]) for n in negs]}


def mapa(s):
    ops = []
    for pid, nome, onda, status, area, regra in PRODUTOS:
        if pid in s["ja_tem"]:
            ops.append({"id": pid, "produto": nome, "aderencia": "já contratado" if s["ja_tem"][pid] == "ganho" else "em andamento",
                        "motivo": f"negócio no funil ({s['ja_tem'][pid]})", "onda": ONDA[onda], "status": status, "area": area}); continue
        try: x = regra(s)
        except Exception: x = None
        if x: ops.append({"id": pid, "produto": nome, "aderencia": x[0], "motivo": x[1], "onda": ONDA[onda], "status": status, "area": area})
    ops.sort(key=lambda o: (o["onda"], -PESO.get(o["aderencia"], -1)))
    falta = [k for k, v in (("faturamento real (só há faixa)", s["faturamento"] in (None, 3.6e6, 4.8e6)), ("parecer lido", not s["teses"]),
                            ("regime", not s["regime"]), ("CNAE na Receita", not s["cnae"])) if v]
    pontos = sum(PESO.get(o["aderencia"], 0) for o in ops if o["aderencia"] in PESO)
    alertas = ([f"duplicidade suspeita no HubSpot: {', '.join(s['duplicidade_suspeita'])} (fora da soma)"] if s["duplicidade_suspeita"] else []) + \
              ([f"mais de um parceiro na mesma raiz de CNPJ: {', '.join(s['masters'])} (conflito de canal)"] if len(s["masters"]) > 1 else [])
    return {**{k: s[k] for k in ("nome", "cnpj", "setor", "cnae", "uf", "regime", "faixa", "master", "rt_ganhos", "rt_valor", "rt_recebido", "ultimo_ganho", "honorario_contratado", "link_negocio")}, "alertas": alertas,
            "oportunidades": ops, "teses_segmento": teses_segmento(s) if any(o["id"] == "RT" for o in ops) else [], "indice_aderencia": pontos, "falta": falta, "negocios": s["negocios"]}


def por_cnpj(neg):
    g = collections.defaultdict(list)
    for n in neg:
        c = cnpj_de(n)
        if c: g[c].append(n)
    return g


def excel(linhas, destino):
    sys.path.insert(0, os.path.join(AP, "lib"))
    import openpyxl, marca  # noqa
    from openpyxl.styles import Font, PatternFill, Alignment
    M = marca.MARCA if hasattr(marca, "MARCA") else json.load(open(os.path.join(AP, "departamentos/marketing-marca/padrao/marca.json")))
    fonte, ouro = M["fontes"]["sans"], M["claro"]["ouro"].lstrip("#")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Carteira"
    cab = ["Master", "Cliente", "CNPJ", "Setor", "UF", "Regime", "Faixa de faturamento", "RT ganha (R$, sem duplicidade)", "Índice de aderência", "Recuperação Tributária", "Teses que cabem no segmento", "Onda 1 · Destravar caixa", "Onda 2 · Proteger e controlar", "Onda 3 · Crescer", "Falta", "Alertas", "Link"]
    ws.append(cab)
    for c in ws[1]: c.font = Font(name=fonte, bold=True, color="000000"); c.fill = PatternFill("solid", fgColor=ouro); c.alignment = Alignment(wrap_text=True, vertical="center")
    for m in linhas:
        onda = lambda o: "; ".join(f"{x['produto']} ({x['aderencia']}" + ("" if x["status"] == "pronto" else ", só descoberta") + ")" for x in m["oportunidades"] if x["onda"].startswith(o) and x["id"] != "RT" and x["aderencia"] in ("alta", "média"))
        rt = next((x for x in m["oportunidades"] if x["id"] == "RT"), None)
        c = m["cnpj"]; ws.append([m["master"] or "Fora do roster", m["nome"], f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}", m["setor"], m["uf"], m["regime"], m["faixa"], m["rt_valor"], m["indice_aderencia"],
                   (f"{rt['aderencia']}: {rt['motivo']}" if rt else "—"), ("\n".join(m["teses_segmento"]) if rt else "—"), onda("1"), onda("2"), onda("3"),
                   ", ".join(m["falta"]) or "nada", "; ".join(m["alertas"]) or "—", f"https://app.hubspot.com/contacts/47865021/deal/{m['link_negocio']}"])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if c.value in (None, "", "None"): c.value = "—"
            c.font = Font(name=fonte, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (14, 32, 19, 18, 5, 13, 14, 15, 10, 40, 46, 40, 42, 34, 24, 34, 16)): ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in r:
            if isinstance(c.value, (int, float)): c.number_format = '#,##0.00'
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    l = wb.create_sheet("Como ler")
    for t in ["Mapa de aderência da carteira (Desenvolvimento de Contas e Soluções). Uso interno.",
              "Onda 1 Destravar caixa: caixa rápido e risco baixo. Onda 2 Proteger e controlar: passivo, carga e governança. Onda 3 Crescer: funding e novos mercados.",
              "Recuperação Tributária aparece como um produto só; a coluna ao lado lista as teses do catálogo que cabem no segmento e no regime (perfil, não diagnóstico: o técnico confirma). Aderência alta ou média aparece nas ondas; o motivo de cada uma está no mapa individual (mapa_conta.py --negocio <id>).",
              "Crédito Internacional e Maquila (conduzidos pelo João Mattos) seguem em validação ou estruturação: só descoberta, sem taxa, prazo ou fee. Gestão de Contas, antecipação e FIDC (Save Invest) e estruturação patrimonial (jurídico) são produtos prontos.",
              "Índice de aderência: alta = 3, média = 2, baixa = 1. Serve para priorizar a conversa com o Master, não é valor em R$.",
              "RT ganha: soma do valor dos ganhos de Recuperação Tributária (crédito do cliente, não receita da Save), sem os cards com duplicidade suspeita. O campo recebido ainda não é preenchido no HubSpot.",
              "Alertas: duplicidade suspeita (ganho importado repetindo outro do mesmo dia) e conflito de canal (mais de um parceiro na mesma raiz de CNPJ). Tratar antes de falar com o Master."]:
        l.append([t]); l.cell(l.max_row, 1).font = Font(name=fonte, size=11)
    l.column_dimensions["A"].width = 120
    wb.save(destino)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--negocio"); ap.add_argument("--cnpj"); ap.add_argument("--carteira", action="store_true"); ap.add_argument("--master")
    ap.add_argument("--excel", help="caminho do Excel da carteira (padrão: departamentos/solucoes/carteira-<data>.xlsx)")
    a = ap.parse_args()
    base, neg, res, db = carregar(); G = por_cnpj(neg)
    if a.negocio or a.cnpj:
        c = normalizar_cnpj(a.cnpj).zfill(14) if a.cnpj else next((cnpj_de(n) for n in neg if str(n["id"]) == a.negocio), "")
        if not c or c not in G: sys.exit("Negócio ou CNPJ sem CNPJ no HubSpot.")
        raiz = [n for k, v in G.items() if k[:8] == c[:8] for n in v]
        print(json.dumps(mapa(sinais(c, raiz, res, db)), ensure_ascii=False, indent=1, default=str)); sys.exit()
    clientes = {c: v for c, v in G.items() if any(n["status"] == "ganho" and n["pipeline"].startswith("RT") for n in v)}
    M = [mapa(sinais(c, v, res, db)) for c, v in clientes.items()]
    if a.master: M = [m for m in M if (m["master"] or "").lower() == a.master.lower()]
    M.sort(key=lambda m: (m["master"] or "~", -m["indice_aderencia"]))
    os.makedirs(OUT, exist_ok=True)
    cont = collections.Counter(o["id"] for m in M for o in m["oportunidades"] if o["aderencia"] in ("alta", "média"))
    agg = {"gerado": HOJE.isoformat(), "fonte": base, "clientes": len(M), "por_produto": dict(cont),
           "por_setor": dict(collections.Counter(m["setor"] for m in M)), "mapas": M}
    json.dump(agg, open(os.path.join(OUT, f"carteira-{HOJE}.json"), "w"), ensure_ascii=False, indent=1, default=str)
    x = a.excel or os.path.join(OUT, f"carteira-{HOJE}.xlsx"); excel(M, x)
    print(f"{len(M)} clientes · oportunidades (alta/média) por produto: {dict(cont.most_common())}\n{x}")
