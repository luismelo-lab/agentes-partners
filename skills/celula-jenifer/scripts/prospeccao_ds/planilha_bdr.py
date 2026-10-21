"""Planilha de ligações da BDR (D&S): uma linha por empresa, pronta para usar durante a ligação. Design System Save (modo claro)."""
import collections, os, pickle, re, sys, urllib.parse
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

S = os.environ.get("PROSPECCAO_DS", os.path.expanduser("~/agentes-partners/departamentos/intel-comercial/prospeccao-ds")); AP = os.path.expanduser("~/agentes-partners")
sys.path.insert(0, os.path.expanduser("~/.claude/skills/depto-solucoes/scripts"))
import mapa_conta as MC  # noqa: E402  (teses do catálogo por segmento, revisadas pelo Tributário em 16/10)

M = pickle.load(open(S + "/modelo.pkl", "rb")); pros = M["pros"]
UFS = set((os.environ.get("UFS") or "MG").split(","))  # decisão do Luis (21/10): só Minas Gerais
pros = [p for p in pros if p["uf"] in UFS]
SO_CNPJS = os.environ.get("SO_CNPJS")  # Lote 2: só as empresas que não estavam no Lote 1
if SO_CNPJS: _ok = set(open(SO_CNPJS).read().split()); pros = [p for p in pros if p["cnpj"] in _ok]
import json
try: CONT = json.load(open(os.path.expanduser("~/agentes-partners/dados-publicos/receita/prospects-contatos.json")))
except Exception: CONT = {}
ORDEM_Q = ("Sócio-Administrador", "Administrador", "Titular", "Presidente", "Diretor", "Sócio")
def socios(c):
    ss = [x for x in (CONT.get(c) or {}).get("socios", []) if x.get("tipo") == "PF"]
    ss.sort(key=lambda x: next((i for i, q in enumerate(ORDEM_Q) if str(x["qualificacao"]).startswith(q)), 9))
    return "\n".join(f"{x['nome']} · {x['qualificacao']}" for x in ss[:3]) or None
def fones(c): t = (CONT.get(c) or {}).get("telefones", []); return (t + [None, None])[:2]
DESTINO = sys.argv[1]
LOTE = os.environ.get("LOTE", "Lote 1")
FONTE = "IBM Plex Sans"; OURO = "F3CD60"; OURO_T = "B58A1F"; OURO_S = "FBF1D7"; TX = "0F1115"; TX2 = "4A4F58"; BORDA = "E6E8EC"; FUNDO2 = "F1F2F5"


def nome(p): return (p["fantasia"] or p["razao"] or "").strip()


def anos(ab):
    try: return int((date.today() - date(int(ab[:4]), int(ab[4:6]), int(ab[6:8]))).days / 365.25)
    except Exception: return None


GANCHO = {
    "Proteína animal": "Açougues, frigoríficos e distribuidores de carne no Lucro Real costumam ter crédito de PIS/COFINS nas compras que não aparece na rotina da contabilidade.",
    "Café e grãos": "Empresas de grãos no Lucro Real costumam ter créditos de PIS/COFINS sobre compras e insumos que ficam fora da rotina.",
    "Varejo alimentar": "Hortifrútis, mercados e distribuidores de alimentos vendem muitos itens com tributação diferente; isso costuma deixar créditos de PIS/COFINS para trás.",
    "Combustíveis e GLP": "Revendas de combustível e gás têm tributação concentrada; vale conferir se há créditos de PIS/COFINS e ICMS-ST que ficaram de fora.",
    "Transporte de carga": "Transportadoras no Lucro Real costumam ter créditos de PIS/COFINS sobre combustível, peças e manutenção que não entram na rotina.",
    "Panificação e alimentos": "Indústrias de alimentos e padarias no Lucro Real costumam ter créditos de PIS/COFINS sobre insumos que ficam para trás.",
}
GANCHO_PADRAO = "Empresas do seu segmento no Lucro Real costumam ter créditos de PIS/COFINS dos últimos 5 anos que não aparecem na rotina da contabilidade."
URGENCIA = "Com a reforma, o PIS/COFINS acaba em 2027: vale revisar os últimos 5 anos enquanto dá tempo."


def setor_de(p): return MC.setor(p["cnae"])


def teses(p):
    s = {"cnae": p["cnae"], "regime": "Lucro Real", "faturamento": 3.6e6, "setor": setor_de(p)}
    t = MC.teses_segmento(s)
    return "\n".join("• " + x.split(" (")[0].strip() for x in t[:6])


JEN = "/Users/luisfernandomelo/Library/CloudStorage/OneDrive - SAVE COMPANY/Área Partners · Agentes/Central BDR · Jenifer/"
ANTIGAS = set()
for f in ("Listas/Lista-D-S-2026-09-29.xlsx", "Controle-Followup-Jenifer.xlsx"):
    try:
        for w in openpyxl.load_workbook(JEN + f, read_only=True, data_only=True).worksheets:
            for r in w.iter_rows(values_only=True):
                ANTIGAS |= {re.sub(r"\D", "", str(c))[:8] for c in r if c and len(re.sub(r"\D", "", str(c))) == 14}
    except Exception: pass
for p in pros:
    s = p["similaridade"]; p["tier"] = "A" if s >= 90 else "B" if s >= 85 else "C"
pros.sort(key=lambda p: (p["tier"], p["nivel_geo"], -p["potencial"]))

wb = openpyxl.Workbook()
fino = Side(style="thin", color=BORDA)

def fmt(c, bold=False, cor=TX, tam=10, fundo=None, wrap=True, h="left"):
    c.font = Font(name=FONTE, bold=bold, color=cor, size=tam)
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=h)
    if fundo: c.fill = PatternFill("solid", fgColor=fundo)

# ── Como usar ──
w0 = wb.active; w0.title = "Como usar"; w0.sheet_view.showGridLines = False; w0.column_dimensions["A"].width = 4; w0.column_dimensions["B"].width = 120
linhas = [
    ("Lista de ligações · Parceiro D&S", "titulo"),
    (f"{LOTE} · gerada em {date.today():%d/%m/%Y} · {len(pros)} empresas · Área Partners (Save Co.)", "sub"),
    ("", None),
    ("Quem está nesta lista", "h"),
    ("Empresas novas para a D&S em Minas Gerais: ativas na Receita, no Lucro Real (último ano declarado), dos segmentos e regiões em que a D&S já ganha.", None),
    ("Fora da lista: clientes da D&S (qualquer fase), qualquer empresa com negócio no HubSpot (qualquer parceiro), Simples, MEI e empresas em recuperação judicial ou falência.", None),
    ("", None),
    ("Como trabalhar", "h"),
    ("1. Abra a aba Ligações. Comece pela Prioridade A e, dentro dela, pelas empresas do mesmo bairro de um cliente da D&S (já estão no topo).", None),
    ("2. Antes de ligar, leia as colunas da faixa dourada (Para a conversa): o gancho e as teses do segmento.", None),
    ("3. Ligue para o Telefone 1 ou 2 (cadastro da empresa na Receita) e peça para falar com o sócio pelo nome (coluna Sócios). Se o número não funcionar, use 'Buscar contato' e anote o certo em Telefone confirmado.", None),
    ("4. Depois da ligação, preencha só as colunas de controle (Status, Data do contato, Próximo passo, Data de retorno e Anotações). Use os filtros do cabeçalho.", None),
    ("5. Reunião marcada ou interesse: registre no HubSpot (funil Save Partners - BDR/SDR), com o Master D&S como parceiro, e avise a Daniela.", None),
    ("6. Pediu para não contatar: marque o status e não ligue de novo.", None),
    ("", None),
    ("Regras da conversa (obrigatórias)", "h"),
    ("• Nunca prometer valor, percentual, prazo ou êxito. Falamos em 'verificar se há' e em 'diagnóstico'.", None),
    ("• Não citar clientes da D&S nem da Save. As colunas cinza (Uso interno) são só para você priorizar.", None),
    ("• O crédito de referência e o potencial são estimativas internas a partir de clientes parecidos. Não informe ao prospect.", None),
    ("• Confirme o regime na conversa: se não for Lucro Real, marque 'Não é Lucro Real'.", None),
    ("• Os dados de sócio e de contato são do cadastro público da empresa (CNPJ). Use só para esta prospecção, não repasse e respeite quem pedir para não ser contatado.", None),
    ("• Sempre ofereça a saída: 'se não fizer sentido, é só me avisar que não volto a ligar'.", None),
    ("", None),
    ("Prioridade", "h"),
    ("A: similaridade de 90% ou mais com um cliente ganho da D&S · B: de 85% a 90% · C: abaixo de 85%. Dentro de cada letra, a ordem é proximidade (bairro, região da cidade, cidade, cidades vizinhas) e potencial.", None),
]
for i, (t, k) in enumerate(linhas, 1):
    c = w0.cell(i, 2, t)
    if k == "titulo": fmt(c, True, TX, 18)
    elif k == "sub": fmt(c, False, TX2, 11)
    elif k == "h": fmt(c, True, OURO_T, 12)
    else: fmt(c, False, TX, 11)
w0.cell(3, 2).border = Border(top=Side(style="thick", color=OURO))

# ── Ligações ──
ws = wb.create_sheet("Ligações"); ws.sheet_view.showGridLines = False
COLS = [  # (título, largura, grupo)
    ("Nº", 5, "c"), ("Prioridade", 9, "c"), ("Status", 22, "c"), ("Data do contato", 12, "c"), ("Próximo passo", 26, "c"), ("Data de retorno", 12, "c"), ("Anotações da ligação", 40, "c"),
    ("Empresa", 30, "e"), ("Razão social", 30, "e"), ("CNPJ", 19, "e"), ("Segmento", 26, "e"), ("Cidade", 16, "e"), ("Bairro", 18, "e"), ("Endereço", 30, "e"), ("CEP", 10, "e"),
    ("Sócios para pedir na ligação (nome · cargo)", 30, "e"), ("Telefone 1 (Receita)", 15, "e"), ("Telefone 2 (Receita)", 15, "e"), ("E-mail da empresa (Receita)", 24, "e"), ("Telefone confirmado (anotar)", 15, "e"), ("Buscar contato", 12, "e"), ("Anos de empresa", 9, "e"), ("Porte", 14, "e"), ("Capital social (R$)", 14, "e"), ("Regime (ano)", 16, "e"),
    ("Gancho para abrir a ligação", 48, "p"), ("Teses que cabem no segmento (para o diagnóstico)", 44, "p"),
    ("Perto de cliente D&S?", 20, "i"), ("Cliente D&S parecido (não citar)", 26, "i"), ("Similaridade", 11, "i"), ("Crédito de referência (interno)", 15, "i"), ("Potencial esperado (interno)", 15, "i"),
]
GRUPO = {"c": ("Controle da ligação · preencha aqui", "0F1115", "FFFFFF"), "e": ("Empresa", "4A4F58", "FFFFFF"),
         "p": ("Para a conversa", OURO, TX), "i": ("Uso interno · não citar ao prospect", "C9CCD3", TX)}
# linha 1: faixas de grupo
ini = 1
for g in ("c", "e", "p", "i"):
    n = sum(1 for x in COLS if x[2] == g); fim = ini + n - 1
    ws.merge_cells(start_row=1, start_column=ini, end_row=1, end_column=fim)
    c = ws.cell(1, ini, GRUPO[g][0]); fmt(c, True, GRUPO[g][2], 10, GRUPO[g][1], False)
    for j in range(ini, fim + 1): ws.cell(1, j).fill = PatternFill("solid", fgColor=GRUPO[g][1])
    ini = fim + 1
for j, (t, w, g) in enumerate(COLS, 1):
    c = ws.cell(2, j, t); fmt(c, True, TX, 10, OURO_S if g != "i" else FUNDO2); c.border = Border(bottom=Side(style="medium", color=OURO))
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[2].height = 32

for i, p in enumerate(pros, 1):
    c14 = p["cnpj"]; cnpj = f"{c14[:2]}.{c14[2:5]}.{c14[5:8]}/{c14[8:12]}-{c14[12:]}"
    busca = "https://www.google.com/search?q=" + urllib.parse.quote(f"{nome(p)} {p['municipio'].title()} {p['uf']} telefone")
    ex = p["exemplo"]; ag = p["ancora_geo"]
    perto = p["nivel_geo"][4:] + (f" ({(ag['rf']['fantasia'] or ag['rf']['razao'] or '').strip()[:24]})" if not p["nivel_geo"].startswith("4") else "")
    cep = p["cep"] or ""
    v = [i, p["tier"], "A ligar", None, None, None, ("Já estava na Lista D&S de 29/09: conferir no Controle de follow-up antes de ligar." if c14[:8] in ANTIGAS else None),
         nome(p), p["razao"], cnpj, MC.setor(p["cnae"]) + " · " + (p["cnae_desc"] or "")[:60], (p["municipio"] or "").title(), (p["bairro"] or "").title(),
         (p["logradouro"] or "").title(), f"{cep[:5]}-{cep[5:]}" if len(cep) == 8 else cep, socios(c14), *fones(c14), (CONT.get(c14) or {}).get("email"), None, "Buscar", anos(p["abertura"]), p["porte"], p["capital"],
         (p["regime"] or "").title().replace("Lucro Real", "Lucro Real"),
         GANCHO.get(setor_de(p), GANCHO_PADRAO) + " " + URGENCIA, teses(p),
         perto, (ex["rf"]["fantasia"] or ex["rf"]["razao"] or "").strip(), p["similaridade"] / 100, round(p["ref"], -3), round(p["potencial"], -3)]
    r = i + 2
    for j, x in enumerate(v, 1):
        c = ws.cell(r, j, x); g = COLS[j - 1][2]
        fmt(c, False, TX2 if g == "i" else TX, 10, FUNDO2 if g == "i" else None, j not in (1, 2, 10, 15, 17, 18, 19, 21, 22, 30, 31, 32), "center" if j in (1, 2, 22, 30) else "left")
        c.border = Border(bottom=fino)
    ws.cell(r, 21).hyperlink = busca; ws.cell(r, 21).font = Font(name=FONTE, color=OURO_T, underline="single", size=10)
    ws.cell(r, 2).font = Font(name=FONTE, bold=True, color=OURO_T if p["tier"] == "A" else TX, size=11)
    ws.cell(r, 4).number_format = ws.cell(r, 6).number_format = "dd/mm/yyyy"
    ws.cell(r, 24).number_format = '#,##0'; ws.cell(r, 30).number_format = "0%"
    ws.cell(r, 31).number_format = ws.cell(r, 32).number_format = '"R$" #,##0'
    ws.row_dimensions[r].height = 96
fim = len(pros) + 2
STATUS = ["A ligar", "Sem resposta", "Retornar depois", "Contato errado", "Respondeu: interesse", "Reunião marcada", "Pediu material",
          "Não é o decisor (indicou outro)", "Respondeu: sem interesse agora", "Já trabalha com outra consultoria", "Não é Lucro Real", "Pediu para não contatar"]
dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"C3:C{fim}")
PASSO = ["Ligar de novo", "Enviar e-mail de apresentação", "Pedir apresentação à D&S", "Marcar reunião", "Reunião marcada: registrar no HubSpot", "Encerrar"]
dv2 = DataValidation(type="list", formula1='"' + ",".join(PASSO) + '"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f"E3:E{fim}")
dvd = DataValidation(type="date", allow_blank=True); ws.add_data_validation(dvd); dvd.add(f"D3:D{fim}"); dvd.add(f"F3:F{fim}")
# cor da linha pelo status
for st, cor in (("Reunião marcada", "DCEFE0"), ("Respondeu: interesse", "E8F5E9"), ("Pediu para não contatar", "FBE3E2"), ("Não é Lucro Real", "F1F2F5")):
    ws.conditional_formatting.add(f"A3:G{fim}", FormulaRule(formula=[f'$C3="{st}"'], fill=PatternFill("solid", fgColor=cor)))
ws.freeze_panes = "I3"; ws.auto_filter.ref = f"A2:{get_column_letter(len(COLS))}{fim}"

# ── Roteiro ──
wr = wb.create_sheet("Roteiro da ligação"); wr.sheet_view.showGridLines = False; wr.column_dimensions["A"].width = 30; wr.column_dimensions["B"].width = 110
rot = [
    ("Roteiro da ligação (3 a 5 minutos)", None, "titulo"),
    ("1. Abertura", "Bom dia, [NOME]. Sou a Jenifer, da Save, parceira da D&S em recuperação de créditos tributários. Falo rápido: dois minutos, pode ser?", None),
    ("2. Gancho", "Use a coluna 'Gancho para abrir a ligação' da empresa. Termine com uma pergunta: 'Isso já foi olhado aí na empresa?'", None),
    ("3. Qualificação (R-D-A)", "R · Regime: 'Vocês estão no Lucro Real hoje? E nos últimos 5 anos?'\nD · Decisor: 'Quem cuida dessa decisão: o sócio ou o financeiro?'\nA · Arquivos: 'Vocês têm as EFD-Contribuições dos últimos 5 anos com o contador?'", None),
    ("4. Contexto (C-H-U)", "C · Contador: 'Quem é o contador de vocês?' (ele acompanha o diagnóstico)\nH · Histórico: 'Já fizeram alguma revisão de créditos? De qual período?'\nU · Urgência: 'Com a reforma, o PIS/COFINS acaba em 2027: faz sentido revisar agora?'", None),
    ("5. Fechamento", "'Faz sentido uma conversa de 30 minutos com o nosso especialista para ver se vale fazer o diagnóstico? Tenho [DIA/HORA] ou [DIA/HORA].'\nSe não fizer sentido: 'Sem problema. Se mudar, é só me chamar; não volto a ligar.'", None),
    ("", None, None),
    ("Objeções frequentes", None, "h"),
    ("'Meu contador já faz isso.'", "Ótimo, o diagnóstico é complementar e o contador acompanha. Olhamos pontos específicos do seu segmento que costumam ficar fora da rotina.", None),
    ("'Isso é arriscado?'", "Cada tese tem um grau de risco, e o técnico explica no diagnóstico antes de qualquer pedido. Trabalhamos por pedido de ressarcimento.", None),
    ("'Quanto vou recuperar?'", "Só dá para dizer depois do diagnóstico sobre os arquivos. Não trabalhamos com estimativa antes disso.", None),
    ("'Não tenho tempo.'", "Entendo. São 30 minutos; depois o trabalho é feito sobre os arquivos, com pouca demanda do seu time. Prefere [DIA] ou [DIA]?", None),
    ("'Já fizemos com outra empresa.'", "Faz sentido. Qual período e quais teses foram analisadas? Às vezes ficam períodos ou teses de fora.", None),
    ("", None, None),
    ("Nunca dizer", None, "h"),
    ("", "'Garantido', 'sem risco', 'você tem R$ X a recuperar', percentual de êxito, prazo de recebimento, nome de cliente da D&S ou da Save.", None),
]
for i, (a, b, k) in enumerate(rot, 1):
    ca = wr.cell(i, 1, a); cb = wr.cell(i, 2, b)
    if k == "titulo": fmt(ca, True, TX, 16, wrap=False)
    elif k == "h": fmt(ca, True, OURO_T, 12, wrap=False)
    else: fmt(ca, True, TX, 11); fmt(cb, False, TX, 11)
    if b and "\n" in b: wr.row_dimensions[i].height = 18 * (b.count("\n") + 1) + 6

# ── Resumo ──
wz = wb.create_sheet("Resumo por cidade"); wz.sheet_view.showGridLines = False
cab = ["Cidade", "Empresas", "Prioridade A", "No mesmo bairro de cliente D&S", "Na mesma região da cidade", "Potencial esperado (interno, R$)"]
for j, t in enumerate(cab, 1):
    c = wz.cell(1, j, t); fmt(c, True, TX, 10, OURO_S); c.border = Border(bottom=Side(style="medium", color=OURO)); wz.column_dimensions[get_column_letter(j)].width = [26, 11, 12, 18, 18, 22][j - 1]
ag = collections.defaultdict(lambda: [0, 0, 0, 0, 0.0])
for p in pros:
    a = ag[(p["municipio"] or "").title()]; a[0] += 1; a[1] += p["tier"] == "A"; a[2] += p["nivel_geo"].startswith("1"); a[3] += p["nivel_geo"].startswith("2"); a[4] += p["potencial"]
for i, (cid, a) in enumerate(sorted(ag.items(), key=lambda x: -x[1][1]), 2):
    for j, x in enumerate([cid] + a[:4] + [round(a[4], -3)], 1):
        c = wz.cell(i, j, x); fmt(c, False, TX, 10, wrap=False); c.border = Border(bottom=fino)
    wz.cell(i, 6).number_format = '"R$" #,##0'
wz.freeze_panes = "A2"; wz.auto_filter.ref = f"A1:F{len(ag) + 1}"

wb.active = 1
wb.save(DESTINO); print("ok", DESTINO, len(pros), collections.Counter(p["tier"] for p in pros))
