#!/usr/bin/env python3
"""Central BDR (célula da Jenifer) — lista de contas-alvo priorizadas para um parceiro Master.

Determinístico: nada de IA nos números. Fontes (todas locais):
- HubSpot: último download em ~/agentes-partners/departamentos/resultados/dados/hubspot-*.json (não chama a API)
- Receita Federal (Dados Abertos do CNPJ): ~/agentes-partners/dados-publicos/receita/cnpj.sqlite
- ICP / lead scoring v1: ~/agentes-partners/departamentos/intel-comercial/icp-*.json
- Matriz tese × perfil: ~/.claude/skills/recuperador/references/matriz.json
- Parâmetros: ~/agentes-partners/celulas/jenifer/config-jenifer.json (prioridades do Q4, pesos, cadastro dos Masters)

Regras embutidas:
- Só dados de EMPRESA (CNPJ, CNAE, porte, UF, capital, regime). Nenhum contato pessoal (LGPD, até validação do DPO).
- Conflito de canal: a raiz do CNPJ é checada no HubSpot (ganho, aberto, perdido, desqualificado pelo BDR) e em
  listas anteriores de outros Masters. Conflito vai para a aba "Conflitos de canal", nunca para a lista principal.
- Minha Receita (regime) só com --permitir-externo e só para prospects (CNPJ cuja raiz não está no HubSpot).
- Teses são indicativas (matriz), sem valores. Café fora da meta (regra do Luis, abr/2026).

Uso:
  python3 lista_master.py --master "D&S" [--uf MG,SP] [--cnaes 1011,4634] [--n 50] [--permitir-externo] [--sem-q4]
  python3 lista_master.py --master "Cardoso" --parceiros [--n 50]     # escritórios contábeis (CNAE 6920-6/01) na região
"""
import argparse, glob, json, os, re, sqlite3, subprocess, sys, time, unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

H = os.path.expanduser("~")
AP = os.path.join(H, "agentes-partners")
HS_DIR = os.path.join(AP, "departamentos/resultados/dados")
ENR = os.path.join(AP, "departamentos/intel-comercial/enriquecimento-receita.json")
ICP_DIR = os.path.join(AP, "departamentos/intel-comercial")
MAP = os.path.join(AP, "departamentos/resultados/mapeamento-hubspot.json")
DB = os.path.join(AP, "dados-publicos/receita/cnpj.sqlite")
RCFG = os.path.join(AP, "departamentos/intel-comercial/receita-config.json")
CACHE = os.path.join(AP, "dados-publicos/receita/regime-cache.json")
MATRIZ = os.path.join(H, ".claude/skills/recuperador/references/matriz.json")
CEL = os.path.join(AP, "celulas/jenifer")
CFG = os.path.join(CEL, "config-jenifer.json")
OUT = os.path.join(CEL, "listas")
HOJE = date.today()


# ---------------------------------------------------------------- utilidades
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()


SUFIXOS = {"LTDA", "LTD", "ME", "EPP", "EIRELI", "SA", "S A", "S", "A", "CIA", "COMERCIO", "E", "DE", "DA", "DO", "DOS", "DAS", "EM", "ALIMENTOS", "COM", "IND", "INDUSTRIA"}


def nome_chave(s):
    """Nome normalizado sem sufixos societários e palavras genéricas (para achar conflito quando falta CNPJ)."""
    t = [w for w in norm(s).split() if w not in SUFIXOS and not w.isdigit()]
    return " ".join(t)


def raiz(cnpj):
    d = re.sub(r"\D", "", str(cnpj or ""))
    return d.zfill(14)[:8] if d else ""


def dt(s):
    try: return datetime.fromisoformat(str(s)[:10]).date()
    except (ValueError, TypeError): return None


def fmt_cnpj(c):
    c = str(c).zfill(14); return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


def ultimo(pat):
    fs = sorted(glob.glob(pat))
    if not fs: sys.exit(f"Arquivo não encontrado: {pat}")
    return fs[-1]


def cfg():
    if not os.path.exists(CFG): sys.exit(f"Falta o arquivo de parâmetros {CFG}")
    return json.load(open(CFG))


# ---------------------------------------------------------------- matriz tese × perfil
M = json.load(open(MATRIZ))["matriz"]


def linha_matriz(cnae, regime_curto):
    cands = [(len(pf), m) for m in M if m["regime"] == regime_curto for pf in m["prefixos"] if str(cnae).startswith(pf)]
    return max(cands, key=lambda x: x[0])[1] if cands else None


def regime_curto(reg):
    r = norm(reg)
    if r.startswith("LUCRO REAL"): return "Real"
    if r.startswith("LUCRO PRESUMIDO"): return "Presumido"
    if r.startswith("SIMPLES"): return "Simples"
    return None


# ---------------------------------------------------------------- HubSpot
def carregar_hubspot():
    f = ultimo(os.path.join(HS_DIR, "hubspot-*.json")); D = json.load(open(f))
    E = json.load(open(ENR))["negocios"] if os.path.exists(ENR) else {}
    por_raiz, por_nome = defaultdict(list), defaultdict(list)
    for d in D["negocios"]:
        e = E.get(d["id"], {})
        d["_raiz"] = raiz(e.get("cnpj") or d["extras"].get("cnpj"))
        d["_cnae"] = str(e.get("cnae_oficial") or d.get("cnae") or "")
        d["_cnae_desc"] = e.get("cnae_desc_oficial") or d.get("segmento") or ""
        d["_uf"] = (e.get("uf_oficial") or d["extras"].get("uf") or "").strip().upper()
        d["_cidade"] = norm(e.get("municipio_oficial") or d["extras"].get("cidade"))
        if d["_raiz"]: por_raiz[d["_raiz"]].append(d)
        k = nome_chave(d["nome"])
        if len(k) >= 8: por_nome[k].append(d)
    return f, D, por_raiz, por_nome


def resolver_master(txt):
    ros = json.load(open(MAP))["escopo"]["parceiros_roster"]; t = norm(txt)
    for m, aliases in ros.items():
        if t == norm(m) or any(t == norm(a) for a in aliases): return m, ros
    hits = [m for m, aliases in ros.items() if any(t in norm(a) or norm(a) in t for a in [m] + aliases if len(norm(a)) >= 3)]
    if len(hits) == 1: return hits[0], ros
    sys.exit(f"Master '{txt}' não reconhecido (ou ambíguo: {hits}). Opções: {', '.join(ros)}")


def eh_cafe(cnae, nome, cafe_cnaes, cafe_pal):
    c = str(cnae or "")
    return any(c.startswith(x) or c == x.zfill(7) for x in cafe_cnaes) or any(p in (nome or "").lower() for p in cafe_pal)


def perfil_master(master, D, C, cafe_cnaes, cafe_pal):
    P = C["perfil_master"]; desde = HOJE - timedelta(days=30 * P["janela_meses"])
    G = [d for d in D["negocios"] if d["parceiro"] == master and d["status"] == "ganho" and (dt(d["fechado"]) or HOJE) >= desde]
    cafe = [d for d in G if eh_cafe(d["_cnae"], d["nome"] + " " + d["_cnae_desc"], cafe_cnaes, cafe_pal)] if C.get("cafe_fora_da_meta", True) else []
    ids_cafe = {d["id"] for d in cafe}; Gm = [d for d in G if d["id"] not in ids_cafe]
    # Território (UF e municípios) = onde o Master atua, com TODOS os ganhos (café inclusive: a regra do café é de segmento, não de região).
    # Segmentos (classes CNAE) = só ganhos fora do café.
    uf = Counter(d["_uf"] for d in G if d["_uf"]); tot = sum(uf.values()) or 1
    ufs = [u for u, n in uf.most_common() if n / tot * 100 >= P["uf_participacao_minima_pct"]] or [u for u, _ in uf.most_common(1)]
    classes = Counter(d["_cnae"][:4] for d in Gm if len(d["_cnae"]) >= 4)
    cls = [c for c, n in classes.most_common(P["max_classes_cnae"]) if n >= P["classe_cnae_minimo_ganhos"]]
    cid = Counter(d["_cidade"] for d in G if d["_cidade"])
    return {"ganhos_janela": len(G), "ganhos_cafe_excluidos": len(cafe), "ganhos_considerados": len(Gm),
            "ufs": ufs, "uf_contagem": uf.most_common(), "uf_share": {u: n / tot for u, n in uf.items()}, "classes": cls, "classe_contagem": classes.most_common(12),
            "municipios": [c for c, _ in cid.most_common(15)], "municipio_contagem": cid.most_common(15),
            "regimes": Counter((d["extras"].get("regime") or "não informado") for d in Gm).most_common(), "fonte_perfil": "ganhos no HubSpot"}


# ---------------------------------------------------------------- conflito de canal
def conflito(r, nome_razao, por_raiz, por_nome, C, master):
    """Retorna (bloqueia: bool, texto). Olha a raiz do CNPJ e, se não houver, o nome normalizado."""
    Lc = C["conflito"]; deals = por_raiz.get(r, []); via = "CNPJ (raiz)"
    if not deals:
        k = nome_chave(nome_razao)
        deals = por_nome.get(k, []) if len(k) >= 8 else []; via = "nome (sem CNPJ no card: conferir)"
    if not deals: return False, ""
    txt = []; bloq = False
    sev = lambda x: (0 if x["status"] == "ganho" else 1 if x["status"] == "aberto" else 2, -(dt(x.get("fechado") or x.get("criado")) or HOJE).toordinal())
    for d in sorted(deals, key=sev):
        p = d["parceiro"] if not d["parceiro"].startswith("Outros") else (d.get("parceiro_bruto") or "fora do roster")
        quando = d.get("fechado") or d.get("criado") or "?"
        bdr = d["pipeline"].startswith("Save Partners")
        if d["status"] == "ganho":
            bloq = True; txt.append(f"JÁ É CLIENTE ({p}, {d['pipeline']}, fechado {quando})")
        elif d["status"] == "aberto" and not (bdr and "desqualif" in (d["etapa"] or "").lower()):
            bloq = True; txt.append(f"NEGÓCIO ABERTO com {p} ({d['pipeline']} · {d['etapa']}, criado {d.get('criado')})")
        elif bdr and "desqualif" in (d["etapa"] or "").lower():
            dias = (HOJE - (dt(d.get("criado")) or HOJE)).days
            if dias <= Lc["desqualificado_bdr_bloqueia_dias"]: bloq = True
            txt.append(f"Desqualificado no BDR/SDR ({p}, criado {d.get('criado')}, há {dias} dias)")
        else:
            dias = (HOJE - (dt(quando) or HOJE)).days
            if dias <= Lc["perdido_recente_dias"]: bloq = True
            txt.append(f"Perdido com {p} ({d['pipeline']}, {quando}{', recente' if dias <= Lc['perdido_recente_dias'] else ''})")
    return bloq, f"[{via}] " + " | ".join(txt[:3]) + (f" (+{len(txt) - 3})" if len(txt) > 3 else "")


RESERVAS = os.path.join(OUT, "reservas.json")


def reservas_carregar(C):
    """Registro interno de quem recebeu cada raiz primeiro (evita sugerir a mesma empresa a dois Masters).
    Vale por C['conflito']['reserva_dias'] dias; é proposta, quem decide conflito é o Luis."""
    R = json.load(open(RESERVAS)) if os.path.exists(RESERVAS) else {}
    lim = (HOJE - timedelta(days=C["conflito"].get("reserva_dias", 45))).isoformat()
    return {k: v for k, v in R.items() if v["data"] >= lim}


def reservas_gravar(R, master, raizes):
    for r in raizes: R.setdefault(r, {"master": master, "data": HOJE.isoformat()})
    json.dump(R, open(RESERVAS, "w"), ensure_ascii=False, indent=0)


def listas_anteriores(master):
    """CNPJs já sugeridos em listas anteriores: {raiz: [(master, data)]}."""
    from openpyxl import load_workbook
    R = defaultdict(list)
    for f in sorted(glob.glob(os.path.join(OUT, "Lista-*.xlsx"))):
        if "Indicadores" in f: continue
        try:
            wb = load_workbook(f, read_only=True); ws = wb["Prospects"]; rows = ws.iter_rows(values_only=True); h = next(rows)
            iC, iM = h.index("CNPJ"), h.index("Master")
            data = re.search(r"(\d{4}-\d{2}-\d{2})", f).group(1)
            for l in rows:
                if l[iC]: R[raiz(l[iC])].append((l[iM], data))
        except Exception: continue
    return R


# ---------------------------------------------------------------- regime (Minha Receita, opcional)
def regime_cache():
    return json.load(open(CACHE)) if os.path.exists(CACHE) else {}


def consultar_regime(cnpjs, cache, hs_raizes):
    feitos = 0
    for cn in cnpjs:
        if cn in cache or raiz(cn) in hs_raizes: continue  # proteção: nunca consultar cliente/negócio existente
        r = subprocess.run(["curl", "-sS", "-m", "20", f"https://minhareceita.org/{cn}"], capture_output=True, text=True)
        try:
            rt = sorted(json.loads(r.stdout).get("regime_tributario") or [], key=lambda x: x.get("ano", 0))
            cache[cn] = f"{rt[-1]['forma_de_tributacao']} ({rt[-1]['ano']})" if rt else "sem registro"
        except Exception: cache[cn] = "erro na consulta"
        feitos += 1; time.sleep(1.0)
    json.dump(cache, open(CACHE, "w"))
    return feitos


# ---------------------------------------------------------------- ICP
def icp():
    f = ultimo(os.path.join(ICP_DIR, "icp-*.json")); I = json.load(open(f))
    pts = {(x["dimensao"], x["valor"]): x["pontos_sugeridos"] for x in I["lead_scoring_v1"]}
    import importlib.util
    spec = importlib.util.spec_from_file_location("icp", os.path.join(H, ".claude/skills/depto-intel-comercial/scripts/icp.py"))
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    return f, pts, mod.grupo


# ---------------------------------------------------------------- Excel
def planilha(dest, abas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); wb.remove(wb.active)
    HF, HFont = PatternFill("solid", fgColor="1F1F1F"), Font(bold=True, color="F5C518")
    for nome, cab, linhas, larg, dv in abas:
        ws = wb.create_sheet(nome[:31]); ws.append(cab)
        for j, c in enumerate(cab, 1):
            x = ws.cell(row=1, column=j); x.fill = HF; x.font = HFont; x.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[x.column_letter].width = (larg or {}).get(c, 16)
        for l in linhas: ws.append(l)
        ws.freeze_panes = "A2" if len(cab) > 1 else None
        if len(cab) > 1 and linhas: ws.auto_filter.ref = ws.dimensions
        for col, opcoes in (dv or {}).items():
            if col in cab:
                v = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
                L = ws.cell(row=1, column=cab.index(col) + 1).column_letter; v.add(f"{L}2:{L}{max(len(linhas) + 50, 100)}"); ws.add_data_validation(v)
    wb.save(dest)


# ---------------------------------------------------------------- modo clientes
VIEW = """SELECT e.cnpj, e.basico, e.matriz, COALESCE(m.razao,'') razao, e.fantasia, e.situacao, e.inicio, e.cnae, c.descricao cnae_desc,
          e.uf, mu.descricao municipio, m.capital, m.porte, s.simples, s.mei, e.no_hubspot
          FROM estab e LEFT JOIN empresa m ON m.basico=e.basico LEFT JOIN simples s ON s.basico=e.basico
          LEFT JOIN ref c ON c.tipo='cnae' AND c.codigo=e.cnae LEFT JOIN ref mu ON mu.tipo='municipio' AND mu.codigo=e.municipio"""


def q(sql, args=()):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    return [dict(r) for r in con.execute(sql, args)]


def persona(r, C):
    if (r["capital"] or 0) >= C["persona"]["capital_cfo_a_partir_de"] and r["porte"] == "Demais": return "CFO / Controller (contador como influenciador)"
    return "Sócio-administrador (contador como influenciador)"


def teses(cnae, reg_c, q4_item):
    """Texto indicativo, sem valores. Se o regime é desconhecido, mostra as duas leituras."""
    def um(rc):
        m = linha_matriz(cnae, rc)
        if not m: return None, "sem linha na matriz: pedir pré-análise ao Tributário"
        t = f"{m['variante']} · {m['tese']}"
        if m.get("tese_save"): t = "Tese Save: aquisições com alíquota zero (via PER; risco registrado) + alternativa " + t
        return m, t
    if reg_c in ("Real", "Presumido"):
        m, t = um(reg_c); ader = m["aderencia"] if m else "a verificar"
    else:
        mr, tr = um("Real"); mp, tp = um("Presumido")
        t = f"SE LUCRO REAL: {tr} || SE PRESUMIDO: {tp}"; ader = "a verificar"; m = mr
    if q4_item and q4_item["id"] == "P2": t += " · Q4 (se D7-b aprovada): só variante B"
    docs = m["documentos"] if m else "EFD-Contribuições (5 anos), balancete"
    return ader, t, docs


def modo_clientes(a, C, master, ros):
    fhs, D, por_raiz, por_nome = carregar_hubspot()
    MP = json.load(open(MAP))["regras"]; cafe_cnaes, cafe_pal = MP["cafe_cnae"], MP["cafe_palavras"]
    rc = json.load(open(RCFG))
    P = perfil_master(master, D, C, cafe_cnaes, cafe_pal)
    cad = C.get("masters_cadastro", {}).get(master)
    if cad: P.update({k: v for k, v in cad.items() if k in ("ufs", "classes", "municipios")}); P["fonte_perfil"] = "cadastro (config-jenifer.json)"
    if a.uf: P["ufs"] = [u.strip().upper() for u in a.uf.split(",")]; P["fonte_perfil"] += " + --uf"
    q4 = C["prioridades_q4"] if (C["prioridades_q4"]["ativo"] and not a.sem_q4) else None
    q4_pref = [(it, pf) for it in (q4["itens"] if q4 else []) for pf in it["prefixos"]]
    if a.cnaes: alvo = [c.strip() for c in a.cnaes.split(",")]; origem_alvo = "--cnaes"
    else:
        alvo = list(P["classes"]); origem_alvo = "classes CNAE dos ganhos do Master"
        if q4 and q4.get("incluir_prefixos_na_busca"): alvo += [pf for _, pf in q4_pref if pf not in alvo]; origem_alvo += " + prioridades Q4 (parâmetro)"
    if C.get("cafe_fora_da_meta", True): alvo = [c for c in alvo if not any(c[:4] == x.zfill(7)[:4] for x in cafe_cnaes)]
    fora_base = [c for c in alvo if not any(c.startswith(p) or p.startswith(c) for p in rc["cnaes_alvo"])]
    avisos = []
    if P["ganhos_considerados"] < 15 and not cad and not a.cnaes:
        avisos.append(f"Perfil de segmento fraco: só {P['ganhos_considerados']} ganhos fora do café em {C['perfil_master']['janela_meses']} meses. A lista usa o território do Master (UF/municípios) + prioridades do Q4; a Daniela confirma com o {master} se ele quer atuar nesses segmentos (ou preencher masters_cadastro / usar --cnaes).")
    if fora_base: avisos.append(f"CNAEs sem cobertura na base local da Receita (incluir em receita-config.json 'cnaes_alvo' e recarregar): {', '.join(fora_base)}")
    ufs_fora = [u for u in P["ufs"] if u not in rc["ufs_alvo"]]
    if ufs_fora: avisos.append(f"UFs do Master fora da base local (receita-config.json 'ufs_alvo'): {', '.join(ufs_fora)}")
    alvo_ok = [c for c in alvo if c not in fora_base]
    if not alvo_ok or not P["ufs"]: sys.exit("Sem CNAE/UF utilizável para este Master. " + " ".join(avisos) + " Use --uf e --cnaes.")

    ph = ",".join("?" * len(P["ufs"])); portes = rc["portes_prospeccao"]; pp = ",".join("?" * len(portes))
    cond_cnae = " OR ".join("e.cnae LIKE ?" for _ in alvo_ok)
    sql = VIEW + f" WHERE e.situacao='Ativa' AND e.uf IN ({ph}) AND ({cond_cnae})"
    rows_all = q(sql, (*P["ufs"], *[c + "%" for c in alvo_ok]))
    # sinais: filiais na base (mesmo recorte) e abertura recente
    fil = defaultdict(list)
    for r in rows_all:
        if r["matriz"] == "filial": fil[r["basico"]].append(r["inicio"])
    rows = [r for r in rows_all if r["matriz"] == "matriz" and r["porte"] in portes and (r["capital"] or 0) >= rc["capital_social_minimo_prospeccao"]
            and not (rc.get("excluir_simples_na_prospeccao", True) and (r["simples"] == "S" or r["mei"] == "S"))
            and not (C.get("cafe_fora_da_meta", True) and eh_cafe(r["cnae"], (r["razao"] or "") + " " + (r["fantasia"] or ""), cafe_cnaes, cafe_pal))]
    teto = rc.get("capital_social_maximo_prospeccao")
    fic, pts, grupo = icp(); cache = regime_cache(); hs_raizes = set(por_raiz)
    antes = listas_anteriores(master); RES = reservas_carregar(C)
    S = C["score"]; mun_m = set(P["municipios"]); cls_m = set(P["classes"])
    seg_max = max([v for (d, _), v in pts.items() if d == "Segmento"] or [30])
    reg_max = max([v for (d, _), v in pts.items() if d == "Regime"] or [17])
    MAXP = seg_max + reg_max + max(S["aderencia_matriz"].values()) + S["classe_cnae_do_master"] + S["municipio_do_master"] + S.get("uf_do_master_max", 10) + S["sinal_filial_12m"] + max([it["pontos"] for it in (q4["itens"] if q4 else [])] or [0])

    def pontuar(r):
        cache_reg = cache.get(r["cnpj"], "")
        rcurto = regime_curto(cache_reg) if cache_reg and not cache_reg.startswith(("sem", "erro")) else None
        q4_item = next((it for it, pf in q4_pref if r["cnae"].startswith(pf)), None)
        ader, tese, docs = teses(r["cnae"], rcurto, q4_item)
        seg = grupo(r["cnae_desc"])
        p_reg = next((v for (d, k), v in pts.items() if d == "Regime" and rcurto and norm(k).startswith("LUCRO " + norm(rcurto))), 0)
        f12 = [x for x in fil.get(r["basico"], []) if dt(f"{x[:4]}-{x[4:6]}-{x[6:8]}") and dt(f"{x[:4]}-{x[4:6]}-{x[6:8]}") >= HOJE - timedelta(days=365)]
        comp = {"segmento": pts.get(("Segmento", seg), 0), "regime": p_reg,
                "aderencia": S["aderencia_matriz"].get(ader, 0), "classe_master": S["classe_cnae_do_master"] if r["cnae"][:4] in cls_m else 0,
                "municipio_master": S["municipio_do_master"] if norm(r["municipio"]) in mun_m else 0,
                "uf_peso_master": round(S.get("uf_do_master_max", 10) * P.get("uf_share", {}).get(r["uf"], 0)),
                "filial_12m": S["sinal_filial_12m"] if f12 else 0, "q4": q4_item["pontos"] if q4_item else 0}
        alerta = "Recuperação judicial/falência no nome: avaliar risco com o Jurídico antes de abordar" if re.search(r"RECUPERACAO JUDICIAL|FALIDA|MASSA FALIDA|EM LIQUIDACAO", norm(r["razao"])) else ""
        comp["alerta"] = -15 if alerta else 0
        r.update({"alerta": alerta, "segmento": seg, "regime": cache_reg or "a confirmar", "regime_c": rcurto, "aderencia": ader, "tese": tese, "docs": docs,
                  "q4": f"{q4_item['id']} {q4_item['nome']}" if q4_item else "", "componentes": comp,
                  "score": max(0, round(sum(comp.values()) / MAXP * 100)), "filiais": len(fil.get(r["basico"], [])), "filial_12m": max(f12) if f12 else ""})

    for r in rows: pontuar(r)
    rows.sort(key=lambda r: (-r["score"], -(r["capital"] or 0)))
    # regime externo para os melhores ainda sem regime (só prospects)
    consultados = 0
    if a.permitir_externo:
        topo = [r["cnpj"] for r in rows if r["basico"] not in hs_raizes and not r["regime_c"]][: a.n * 2]
        consultados = consultar_regime(topo, cache, hs_raizes)
        for r in rows: pontuar(r)
        rows.sort(key=lambda r: (-r["score"], -(r["capital"] or 0)))

    principal, conflitos, estrategicas, vistos = [], [], [], set()
    for r in rows:
        if r["basico"] in vistos: continue
        vistos.add(r["basico"])
        tem, txt = conflito(r["basico"], r["razao"], por_raiz, por_nome, C, master)
        res = RES.get(r["basico"]); outros = [(res["master"], res["data"])] if res and res["master"] != master else []
        if outros: tem, txt = True, (txt + " | " if txt else "") + "Já sugerido a outro Master: " + ", ".join(f"{m} em {d}" for m, d in outros[:3])
        r["conflito"] = txt; r["bloqueia"] = tem
        r["mesmo_master_antes"] = ", ".join(d for m, d in antes.get(r["basico"], []) if m == master)
        if tem: conflitos.append(r)
        elif teto and (r["capital"] or 0) > teto: estrategicas.append(r)
        else: principal.append(r)
    top, reserva = principal[: a.n], principal[a.n: a.n + 200]
    reservas_gravar(RES, master, [r["basico"] for r in top])

    def prox(r):
        if r["regime_c"] == "Simples": return "Descartar: Simples Nacional (sem crédito não cumulativo)"
        if not r["regime_c"]: return "Confirmar regime (Minha Receita com autorização ou pergunta ao contador) → depois pedir apresentação ao Master"
        if r["aderencia"] == "baixa": return "Baixa aderência no regime atual: só se o Master pedir; rascunhar abordagem genérica"
        return f"Pedir apresentação ao Master {master} (toque 1) → rascunho de abordagem por persona"

    cab = ["Prioridade", "Score", "Master", "CNPJ", "Razão social", "Nome fantasia", "CNAE", "Descrição CNAE", "Segmento (ICP)", "Prioridade Q4",
           "UF", "Município", "Porte", "Capital social (R$)", "Início atividade", "Regime", "Aderência (matriz)", "Teses prováveis (indicativo, sem valores)",
           "Documentos para a pré-análise", "Filiais na base", "Filial aberta 12m", "Persona sugerida", "Alerta", "Conflito de canal", "Já listado p/ este Master",
           "Próximo passo", "Pontos (composição)", "Decisão Jenifer", "Motivo / observação"]

    def linha(i, r):
        ini = r["inicio"] or ""; ini = f"{ini[6:8]}/{ini[4:6]}/{ini[:4]}" if len(ini) == 8 else ini
        f12 = r["filial_12m"]; f12 = f"{f12[6:8]}/{f12[4:6]}/{f12[:4]}" if f12 else ""
        return [i, r["score"], master, fmt_cnpj(r["cnpj"]), r["razao"], r["fantasia"], r["cnae"], r["cnae_desc"], r["segmento"], r["q4"],
                r["uf"], r["municipio"], r["porte"], r["capital"], ini, r["regime"], r["aderencia"], r["tese"], r["docs"], r["filiais"], f12,
                persona(r, C), r.get("alerta", ""), r.get("conflito", ""), r.get("mesmo_master_antes", ""), "Não abordar sem decisão do Luis (conflito de canal)" if r.get("bloqueia") else (("Histórico antigo no HubSpot: avisar o Master e a Daniela antes · " if r.get("conflito") else "") + prox(r)),
                json.dumps(r["componentes"], ensure_ascii=False), "", ""]

    larg = {"Razão social": 38, "Descrição CNAE": 34, "Teses prováveis (indicativo, sem valores)": 60, "Documentos para a pré-análise": 40, "Conflito de canal": 50,
            "Próximo passo": 45, "Pontos (composição)": 40, "Persona sugerida": 30, "Município": 20, "Motivo / observação": 30}
    dv = {"Decisão Jenifer": ["Abordar", "Descartar", "Precisa de dado", "Reservar"]}
    meta = {x["chave"]: x["valor"] for x in q("SELECT * FROM meta")}
    perfil_l = [["Master", master], ["Aliases (roster)", ", ".join(ros[master])], ["Fonte do perfil", P["fonte_perfil"]],
                ["Ganhos na janela (" + str(C["perfil_master"]["janela_meses"]) + " meses)", P["ganhos_janela"]],
                ["Ganhos de café excluídos (regra do café)", P["ganhos_cafe_excluidos"]], ["Ganhos considerados", P["ganhos_considerados"]],
                ["UFs-alvo", ", ".join(P["ufs"])], ["UFs (contagem)", str(P["uf_contagem"])],
                ["Classes CNAE do Master", ", ".join(P["classes"]) or "nenhuma com amostra mínima"], ["Classes (contagem)", str(P["classe_contagem"])],
                ["Municípios onde já ganha", ", ".join(P["municipios"])], ["Regimes dos ganhos", str(P["regimes"])],
                ["CNAEs buscados", ", ".join(alvo_ok) + f"  ({origem_alvo})"],
                ["Prioridades Q4", ("ATIVAS como parâmetro — " + C["prioridades_q4"]["_status"]) if q4 else "desligadas (--sem-q4 ou config)"]] + [["Aviso", t] for t in avisos]
    resumo = Counter((r["segmento"], r["uf"]) for r in top)
    fonte = [[meta.get("fonte", "Receita Federal")], [f"HubSpot: {os.path.basename(fhs)} (download local; a API não foi chamada)"], [f"ICP/lead scoring v1: {os.path.basename(fic)} (vale por padrão; ajuste da Inteligência Comercial com a Jenifer)"],
             ["Matriz tese × perfil: recuperador/references/matriz.json (Depto Tributário, 29/09/2026 — pré-análise, validar com o técnico)"],
             [f"Filtros: ativas, matriz, porte {portes}, capital ≥ R$ {rc['capital_social_minimo_prospeccao']:,.0f}, sem Simples/MEI, café fora. Acima de R$ {teto:,.0f} → aba 'Contas estratégicas'."],
             [f"Score 0–100 = (segmento + regime do lead scoring v1 + aderência da matriz + peso da UF no território do Master + classe do Master + município do Master + filial nova 12m + prioridade Q4 − alerta) ÷ {MAXP} pontos."],
             ["Conflito de canal: raiz do CNPJ no HubSpot (ganho, aberto, perdido < 12 meses, desqualificado BDR < 180 dias) ou nome igual sem CNPJ, e listas anteriores de outros Masters. Limite: o download cobre só a Área Master desde 2024 (Daniela + roster); negócios de outros parceiros fora desse escopo não aparecem."],
             [f"Regime: cache da Minha Receita (consultas desta execução: {consultados}; externas só com --permitir-externo e só para prospects)." ],
             ["LGPD: sem sócios, e-mails ou telefones. Abordagem preferencial via apresentação do Master. Nada foi enviado nem gravado no HubSpot."],
             ["Teses: indicativas, sem valores e sem promessa de êxito. A pré-análise formal é do Depto Tributário / técnico humano."]]
    stamp = HOJE.isoformat(); slug = re.sub(r"[^A-Za-z0-9]+", "-", master).strip("-")
    dest = os.path.join(OUT, f"Lista-{slug}-{stamp}.xlsx"); os.makedirs(OUT, exist_ok=True)
    planilha(dest, [("Prospects", cab, [linha(i, r) for i, r in enumerate(top, 1)], larg, dv),
                    ("Conflitos de canal", cab, [linha(i, r) for i, r in enumerate(conflitos[:300], 1)], larg, {}),
                    ("Contas estratégicas", cab, [linha(i, r) for i, r in enumerate(estrategicas[:100], 1)], larg, dv),
                    ("Reserva", cab, [linha(i, r) for i, r in enumerate(reserva, a.n + 1)], larg, dv),
                    ("Perfil do Master", ["Item", "Valor"], perfil_l, {"Item": 34, "Valor": 120}, {}),
                    ("Resumo", ["Segmento", "UF", "Prospects na lista"], [[s, u, n] for (s, u), n in resumo.most_common()], {"Segmento": 30}, {}),
                    ("Fonte e regras", ["Fonte e regras"], fonte, {"Fonte e regras": 160}, {})])
    dup = len(top) - len({r["basico"] for r in top})
    print(json.dumps({"arquivo": dest, "master": master, "ufs": P["ufs"], "cnaes_buscados": alvo_ok, "candidatos_filtrados": len(rows),
                      "lista": len(top), "conflitos": len(conflitos), "conflitos_bloqueio_cliente": sum(1 for r in conflitos if "JÁ É CLIENTE" in r["conflito"]),
                      "conflitos_aberto": sum(1 for r in conflitos if "NEGÓCIO ABERTO" in r["conflito"]), "conflitos_por_nome": sum(1 for r in conflitos if "[nome" in r["conflito"]),
                      "estrategicas": len(estrategicas), "reserva": len(reserva), "duplicados_na_lista": dup, "regime_conhecido_na_lista": sum(1 for r in top if r["regime_c"]),
                      "consultas_externas": consultados, "segmentos_top": Counter(r["segmento"] for r in top).most_common(6), "avisos": avisos}, ensure_ascii=False, indent=1))


# ---------------------------------------------------------------- modo parceiros indicadores (escritórios contábeis)
def modo_parceiros(a, C, master, ros):
    fhs, D, por_raiz, por_nome = carregar_hubspot()
    MP = json.load(open(MAP))["regras"]
    P = perfil_master(master, D, C, MP["cafe_cnae"], MP["cafe_palavras"])
    cad = C.get("masters_cadastro", {}).get(master)
    if cad: P.update({k: v for k, v in cad.items() if k in ("ufs", "classes", "municipios")})
    if a.uf: P["ufs"] = [u.strip().upper() for u in a.uf.split(",")]
    ph = ",".join("?" * len(P["ufs"]))
    rows = q(VIEW + f" WHERE e.situacao='Ativa' AND e.matriz='matriz' AND e.cnae='6920601' AND e.uf IN ({ph})", tuple(P["ufs"]))
    # carteira potencial local: empresas-alvo (Q4 + classes do Master) ativas por município
    alvo = list(P["classes"]) + [pf for it in C["prioridades_q4"]["itens"] for pf in it["prefixos"]] if C["prioridades_q4"]["ativo"] else list(P["classes"])
    cond = " OR ".join("e.cnae LIKE ?" for _ in alvo) or "0"
    cart = Counter(r["municipio"] for r in q(VIEW + f" WHERE e.situacao='Ativa' AND e.matriz='matriz' AND e.uf IN ({ph}) AND ({cond}) AND COALESCE(s.simples,'N')<>'S'",
                                                (*P["ufs"], *[c + "%" for c in alvo])))
    parceiros_hs = {nome_chave(d.get("parceiro_bruto")) for d in D["negocios"] if d.get("parceiro_bruto")} | {nome_chave(x) for v in ros.values() for x in v}
    parceiros_hs = {p for p in parceiros_hs if len(p) >= 5}
    mun_m = set(P["municipios"]); L = []
    for r in rows:
        k = nome_chave(r["razao"]); kf = nome_chave(r["fantasia"])
        conf = next((p for p in parceiros_hs if p and (p == k or p == kf or (len(p) >= 8 and (p in k or p in kf)))), "")
        anos = (HOJE.year - int(r["inicio"][:4])) if r["inicio"] else 0
        score = (40 if norm(r["municipio"]) in mun_m else 0) + min(30, cart.get(r["municipio"], 0) // 5) + min(15, anos) + (15 if r["porte"] in ("Demais", "Empresa de Pequeno Porte") else 5)
        L.append((score, r, anos, conf))
    L.sort(key=lambda x: (-x[0], -(x[1]["capital"] or 0)))
    cab = ["Prioridade", "Score", "Master", "CNPJ", "Razão social", "Nome fantasia", "UF", "Município", "Município onde o Master já ganha", "Empresas-alvo no município",
           "Anos de atividade", "Porte", "Capital social (R$)", "Simples", "Possível conflito (já parceiro/indicador)", "Próximo passo", "Decisão Jenifer", "Motivo / observação"]
    lin = []
    teto_mun = max(3, a.n // 6); por_mun = Counter(); escolhidos = []
    for x in L:  # espalha a lista entre municípios (no máx. teto_mun por município)
        if x[3] or por_mun[x[1]["municipio"]] >= teto_mun: continue
        por_mun[x[1]["municipio"]] += 1; escolhidos.append(x)
        if len(escolhidos) >= a.n: break
    for i, (s, r, anos, conf) in enumerate(escolhidos, 1):
        lin.append([i, s, master, fmt_cnpj(r["cnpj"]), r["razao"], r["fantasia"], r["uf"], r["municipio"], "sim" if norm(r["municipio"]) in mun_m else "não",
                    cart.get(r["municipio"], 0), anos, r["porte"], r["capital"], r["simples"] or "", "",
                    f"A Daniela confirma com o Master {master} se ele quer sub-indicadores na região → rascunho de convite (persona contador)", "", ""])
    conf_l = [[i, s, master, fmt_cnpj(r["cnpj"]), r["razao"], r["fantasia"], r["uf"], r["municipio"], "", "", anos, r["porte"], r["capital"], r["simples"] or "",
               f"nome parecido com parceiro existente: {c}", "Não abordar sem checar com a Daniela", "", ""] for i, (s, r, anos, c) in enumerate([x for x in L if x[3]], 1)]
    dest = os.path.join(OUT, f"Lista-Indicadores-{re.sub(r'[^A-Za-z0-9]+', '-', master).strip('-')}-{HOJE.isoformat()}.xlsx")
    planilha(dest, [("Escritórios", cab, lin, {"Razão social": 40, "Próximo passo": 60, "Possível conflito (já parceiro/indicador)": 40}, {"Decisão Jenifer": ["Abordar", "Descartar", "Precisa de dado"]}),
                    ("Possível conflito", cab, conf_l, {"Razão social": 40}, {}),
                    ("Fonte e regras", ["Fonte e regras"], [["Receita Federal · Dados Abertos do CNPJ (base local). CNAE 6920-6/01 (atividades de contabilidade), ativas, matriz, nas UFs do Master: " + ", ".join(P["ufs"])],
                     ["Score = 40 se o município é onde o Master já ganha + até 30 pela quantidade de empresas-alvo no município (Q4 + classes do Master, fora Simples) + até 15 por anos de atividade + 15 se EPP/Demais."],
                     ["Conflito: nome parecido com parceiros/consultores já registrados no HubSpot (campo nome_do_parceiro) ou no roster dos Masters."],
                     ["LGPD: só dados da empresa. Recrutar sub-indicador na região de um Master é decisão do Luis (política de canal e remuneração pendentes com o Jurídico)."]], {"Fonte e regras": 160}, {})])
    print(json.dumps({"arquivo": dest, "master": master, "ufs": P["ufs"], "escritorios": len(L), "na_lista": len(lin), "possivel_conflito": len(conf_l),
                      "no_municipio_do_master": sum(1 for x in L if norm(x[1]["municipio"]) in mun_m)}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--master", required=True); ap.add_argument("--uf"); ap.add_argument("--cnaes"); ap.add_argument("--n", type=int, default=50)
    ap.add_argument("--permitir-externo", action="store_true", help="consulta o regime na Minha Receita (só prospects)")
    ap.add_argument("--sem-q4", action="store_true", help="ignora as prioridades do Q4 (D7-b pendente)")
    ap.add_argument("--parceiros", action="store_true", help="lista escritórios contábeis (recrutamento de indicadores)")
    a = ap.parse_args(); C = cfg(); master, ros = resolver_master(a.master)
    (modo_parceiros if a.parceiros else modo_clientes)(a, C, master, ros)
