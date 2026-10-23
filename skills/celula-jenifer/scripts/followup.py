#!/usr/bin/env python3
"""Central BDR (célula da Jenifer) — planilha de controle de follow-up.

Monta/atualiza ~/agentes-partners/celulas/jenifer/Controle-Followup-Jenifer.xlsx a partir das listas do lista_master.py.
Determinístico e idempotente: preserva o que a Jenifer preencheu (colunas editáveis) e só recalcula as colunas automáticas.

- Importa da lista as empresas marcadas "Abordar" na coluna "Decisão Jenifer" (ou as N primeiras com --top N).
- Próxima ação e data vêm da cadência em config-jenifer.json (dias úteis a partir da "Data de início"), conforme os "Toques feitos".
- Reconfere o conflito de canal no último download do HubSpot (se apareceu negócio para a raiz do CNPJ depois da lista, avisa).
- Abas: Controle · Agenda (vencidos + próximos 5 dias úteis) · Propostas HubSpot (RASCUNHOS para a Jenifer aprovar) · Painel · Cadência · Como usar.
- Nada é enviado nem gravado no HubSpot. Sem contatos pessoais (LGPD).

Uso:
  python3 followup.py                                   # só recalcula o controle existente
  python3 followup.py --lista <Lista-...xlsx>           # importa as marcadas "Abordar"
  python3 followup.py --lista <Lista-...xlsx> --top 10  # importa as 10 primeiras (teste / sem marcação)
  python3 followup.py --lista <...> --inicio 2026-10-01
"""
import argparse, glob, json, os, re, sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
sys.path.insert(0, os.path.expanduser("~/agentes-partners/lib"))
from cnpj import normalizar_cnpj  # noqa: E402

H = os.path.expanduser("~")
CEL = os.path.join(H, "agentes-partners/celulas/jenifer")
CFG = os.path.join(CEL, "config-jenifer.json")
CTRL = os.path.join(CEL, "Controle-Followup-Jenifer.xlsx")
HS_DIR = os.path.join(H, "agentes-partners/departamentos/resultados/dados")
ENR = os.path.join(H, "agentes-partners/departamentos/intel-comercial/enriquecimento-receita.json")
HOJE = date.today()

ETAPAS = ["Conta-alvo", "Aguardando apresentação do Master", "Em cadência", "Respondeu", "Reunião agendada", "Reunião realizada",
          "Qualificado (SQL)", "Enviado ao Master/Técnico", "Nutrição", "Descartado"]
RESULTADOS = ["Sem resposta", "Respondeu: interesse", "Respondeu: sem interesse agora", "Pediu material", "Reunião marcada",
              "Não é o decisor (indicou outro)", "Contato errado", "Já trabalha com outra consultoria", "Pediu para não contatar"]
MOTIVOS = ["Simples/regime sem tese", "Já é cliente/conflito de canal", "Sem interesse", "Já recuperou com outra empresa", "Porte/valor baixo",
           "Sem documentos (EFD)", "Pediu para não contatar (opt-out)", "Master não aceitou", "Outro"]
# (coluna, largura, editável pela Jenifer?)
COLS = [("CNPJ", 20, False), ("Empresa", 36, False), ("Master", 14, True), ("UF", 5, False), ("Município", 18, False), ("Segmento", 20, False), ("Regime", 16, False),
        ("Score", 7, False), ("Persona", 26, True), ("Teses prováveis (indicativo)", 45, False),
        ("Etapa", 24, True), ("Data de início", 12, True), ("Toques feitos", 8, True), ("Data último contato", 12, True),
        ("Resultado último toque", 26, True), ("Próximo toque (nº)", 8, False), ("Próxima ação (cadência)", 45, False), ("Data próxima ação", 12, False),
        ("Situação", 14, False), ("Data da reunião", 12, True), ("Participantes (cargos)", 22, True), ("Próximo passo acordado", 36, True),
        ("Motivo descarte", 24, True), ("Aceite do Master (data)", 12, True), ("Conflito (reconferido)", 45, False), ("Origem (lista)", 26, False),
        ("Observações", 36, True)]
NOMES = [c for c, _, _ in COLS]


def raiz(c):
    d = normalizar_cnpj(c); return d.zfill(14)[:8] if d else ""


def dt(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    s = str(v or "").strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try: return datetime.strptime(s[:10], f).date()
        except ValueError: pass
    return None


def n_toques(r):
    try: return int(float(str(r.get("Toques feitos") or 0).replace(",", ".")))
    except (TypeError, ValueError): return 0


def mais_uteis(d, n):
    while d.weekday() >= 5: d += timedelta(days=1)
    k = 0
    while k < n:
        d += timedelta(days=1)
        if d.weekday() < 5: k += 1
    return d


def hubspot_raizes():
    f = sorted(glob.glob(os.path.join(HS_DIR, "hubspot-*.json")))
    if not f: return {}, "sem download do HubSpot"
    N = json.load(open(f[-1]))["negocios"]; E = json.load(open(ENR))["negocios"] if os.path.exists(ENR) else {}
    R = defaultdict(list)
    for d in N:
        r = raiz(E.get(d["id"], {}).get("cnpj") or d["extras"].get("cnpj"))
        if r: R[r].append(d)
    return R, os.path.basename(f[-1])


def ler_controle():
    if not os.path.exists(CTRL): return []
    from openpyxl import load_workbook
    ws = load_workbook(CTRL)["Controle"]; it = ws.iter_rows(values_only=True); h = list(next(it))
    return [{k: (l[h.index(k)] if k in h else None) for k in NOMES} for l in it if any(l)]


def importar(lista, top, inicio, atuais):
    from openpyxl import load_workbook
    ws = load_workbook(lista, read_only=True)["Prospects"]; it = ws.iter_rows(values_only=True); h = list(next(it))
    g = lambda l, k: l[h.index(k)] if k in h else None
    linhas = [l for l in it if g(l, "CNPJ")]
    marcadas = [l for l in linhas if str(g(l, "Decisão Jenifer") or "").strip().lower() == "abordar"]
    esc = marcadas or (linhas[:top] if top else [])
    if not esc: print("Nenhuma empresa marcada 'Abordar' na lista. Marque na coluna 'Decisão Jenifer' ou use --top N."); return 0
    ja = {raiz(r["CNPJ"]) for r in atuais}; novos = 0
    for l in esc:
        if raiz(g(l, "CNPJ")) in ja: continue
        atuais.append({"CNPJ": g(l, "CNPJ"), "Empresa": g(l, "Razão social"), "Master": g(l, "Master"), "UF": g(l, "UF"), "Município": g(l, "Município"),
                       "Segmento": g(l, "Segmento (ICP)"), "Regime": g(l, "Regime"), "Score": g(l, "Score"), "Persona": g(l, "Persona sugerida"),
                       "Teses prováveis (indicativo)": g(l, "Teses prováveis (indicativo, sem valores)"), "Etapa": "Aguardando apresentação do Master",
                       "Data de início": inicio.strftime("%d/%m/%Y"), "Toques feitos": 0, "Origem (lista)": os.path.basename(lista)})
        ja.add(raiz(g(l, "CNPJ"))); novos += 1
    return novos


def recalcular(rows, C, R):
    cad = C["cadencia"]; fim = {"Descartado", "Enviado ao Master/Técnico", "Nutrição"}
    for r in rows:
        ini = dt(r["Data de início"]) or HOJE
        feitos = n_toques(r)
        et = r["Etapa"] or "Conta-alvo"
        if et in fim:
            r["Próximo toque (nº)"] = ""; r["Data próxima ação"] = ""
            r["Próxima ação (cadência)"] = {"Descartado": "Encerrado (registrar motivo)", "Enviado ao Master/Técnico": "Acompanhar o aceite do Master / negócio no HubSpot",
                                           "Nutrição": f"Retomar em {C['sla']['nutricao_retomar_dias']} dias com conteúdo do setor"}[et]
            if et == "Nutrição":
                base = dt(r["Data último contato"]) or ini; r["Data próxima ação"] = (base + timedelta(days=C["sla"]["nutricao_retomar_dias"])).strftime("%d/%m/%Y")
        elif et in ("Reunião agendada",):
            r["Próximo toque (nº)"] = ""; r["Próxima ação (cadência)"] = "Preparar briefing da reunião (Agenda & Preparação) e confirmar na véspera"
            dr = dt(r["Data da reunião"]); r["Data próxima ação"] = mais_uteis(dr, 0).strftime("%d/%m/%Y") if dr else ""
        elif et in ("Reunião realizada", "Qualificado (SQL)"):
            r["Próximo toque (nº)"] = ""; dr = dt(r["Data da reunião"]) or dt(r["Data último contato"]) or HOJE
            r["Próxima ação (cadência)"] = ("Ata + follow-up em até 24 h (Pós-reunião) e proposta de registro no HubSpot" if et == "Reunião realizada"
                                            else f"Passar ao Master com a ficha; cobrar aceite em {C['sla']['master_aceite_dias_uteis']} dias úteis")
            r["Data próxima ação"] = mais_uteis(dr, 1).strftime("%d/%m/%Y")
        elif feitos >= len(cad):
            r["Próximo toque (nº)"] = ""; r["Próxima ação (cadência)"] = "Cadência concluída sem resposta → mover para Nutrição ou Descartado"; r["Data próxima ação"] = HOJE.strftime("%d/%m/%Y")
        else:
            t = cad[feitos]; r["Próximo toque (nº)"] = t["toque"]; r["Próxima ação (cadência)"] = f"{t['canal']}: {t['acao']}"
            r["Data próxima ação"] = mais_uteis(ini, t["dia_util"]).strftime("%d/%m/%Y")
        d = dt(r["Data próxima ação"])
        r["Situação"] = "" if not d else ("VENCIDA" if d < HOJE else "HOJE" if d == HOJE else "no prazo")
        # conflito reconferido
        ds = R.get(raiz(r["CNPJ"]), [])
        ab = [x for x in ds if x["status"] in ("ganho", "aberto") and not (x["pipeline"].startswith("Save Partners - BDR"))]
        r["Conflito (reconferido)"] = ("ATENÇÃO: " + " | ".join(f"{x['status']} com {x['parceiro']} ({x['pipeline']} · {x['etapa']}, {x.get('criado')})" for x in ab[:2])) if ab else "sem conflito no HubSpot"
    rows.sort(key=lambda r: (ETAPAS.index(r["Etapa"]) if r["Etapa"] in ETAPAS else 0, dt(r["Data próxima ação"]) or date.max))


def proposta_hubspot(r):
    """Rascunho de registro no HubSpot (nomes internos do mapeamento-hubspot.json). NÃO é gravado."""
    et = r["Etapa"]
    etapa_bdr = {"Respondeu": "Follow - Up", "Em cadência": "Follow - Up", "Reunião agendada": "[CONFIRMAR etapa 'Reunião agendada' no pipeline BDR/SDR]",
                 "Reunião realizada": "[CONFIRMAR etapa]", "Qualificado (SQL)": "[CONFIRMAR: passar ao pipeline 'RT - Real & Presumido' · Agendamento?]",
                 "Descartado": "Desqualificado"}.get(et)
    if not etapa_bdr: return None
    return ["RASCUNHO · aguardando aprovação da Jenifer", "Criar/atualizar negócio", r["Empresa"], "Save Partners - BDR/SDR", etapa_bdr,
            r["Master"], normalizar_cnpj(r["CNPJ"]), r["UF"], r["Município"], (r.get("Regime") or "[confirmar regime]").split(" (")[0].title(), r["Segmento"],
            "[ID da Jenifer no HubSpot: CONFIRMAR]", (r["Motivo descarte"] or "") if et == "Descartado" else (r["Próximo passo acordado"] or ""),
            r["Conflito (reconferido)"]]


def salvar(rows, C, fonte_hs):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    HF, HFont = PatternFill("solid", fgColor="1F1F1F"), Font(bold=True, color="F5C518")
    ED = PatternFill("solid", fgColor="FFF8DC")
    wb = Workbook(); ws = wb.active; ws.title = "Controle"; ws.append(NOMES)
    for j, (c, w, e) in enumerate(COLS, 1):
        x = ws.cell(row=1, column=j); x.fill = HF; x.font = HFont; x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[x.column_letter].width = w
    for r in rows: ws.append([r.get(c) for c in NOMES])
    last = max(len(rows) + 200, 300)
    for j, (c, w, e) in enumerate(COLS, 1):
        if e:
            for i in range(2, len(rows) + 2): ws.cell(row=i, column=j).fill = ED
    for col, ops in (("Etapa", ETAPAS), ("Resultado último toque", RESULTADOS), ("Motivo descarte", MOTIVOS)):
        L = ws.cell(row=1, column=NOMES.index(col) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(ops) + '"', allow_blank=True); dv.add(f"{L}2:{L}{last}"); ws.add_data_validation(dv)
    Ls = ws.cell(row=1, column=NOMES.index("Situação") + 1).column_letter
    ws.conditional_formatting.add(f"{Ls}2:{Ls}{last}", FormulaRule(formula=[f'{Ls}2="VENCIDA"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"{Ls}2:{Ls}{last}", FormulaRule(formula=[f'{Ls}2="HOJE"'], fill=PatternFill("solid", fgColor="FFE699")))
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(NOMES)).column_letter}{max(len(rows) + 1, 2)}"

    def aba(nome, cab, linhas, larg=18):
        w = wb.create_sheet(nome); w.append(cab)
        for j in range(1, len(cab) + 1):
            x = w.cell(row=1, column=j); x.fill = HF; x.font = HFont; w.column_dimensions[x.column_letter].width = larg if len(cab) > 1 else 150
        for l in linhas: w.append(l)
        if len(cab) > 1: w.freeze_panes = "A2"
        return w
    lim = mais_uteis(HOJE, 5)
    ag = [r for r in rows if dt(r["Data próxima ação"]) and dt(r["Data próxima ação"]) <= lim]
    aba("Agenda", ["Situação", "Data próxima ação", "Empresa", "Master", "Etapa", "Próxima ação (cadência)", "Persona", "Conflito (reconferido)"],
        [[r["Situação"], r["Data próxima ação"], r["Empresa"], r["Master"], r["Etapa"], r["Próxima ação (cadência)"], r["Persona"], r["Conflito (reconferido)"]]
         for r in sorted(ag, key=lambda r: dt(r["Data próxima ação"]))], 24)
    aba("Propostas HubSpot", ["Status", "Operação", "dealname", "pipeline", "dealstage (sugerida)", "nome_do_parceiro", "cnpj_", "estado", "cidade",
                              "qual_e_o_regime_tributario_da_empresa_", "descricao_do_cnae (segmento)", "hubspot_owner_id", "Nota / próximo passo", "Conflito"],
        [p for p in (proposta_hubspot(r) for r in rows) if p], 22)
    et = Counter(r["Etapa"] for r in rows); pm = defaultdict(Counter)
    for r in rows: pm[r["Master"]][r["Etapa"]] += 1
    toques = sum(n_toques(r) for r in rows)  # 2.0 (número do Excel) ficava fora da soma com o isdigit()
    reun = sum(1 for r in rows if r["Etapa"] in ("Reunião agendada", "Reunião realizada", "Qualificado (SQL)", "Enviado ao Master/Técnico"))
    painel = [[f"Atualizado em {HOJE.strftime('%d/%m/%Y')} · HubSpot: {fonte_hs}"], [""], ["Indicador", "Valor"], ["Empresas em controle", len(rows)], ["Toques feitos (total)", toques],
              ["Reuniões (agendadas ou além)", reun], ["Taxa empresa → reunião", f"{reun / len(rows) * 100:.0f}%" if rows else "—"],
              ["Ações vencidas", sum(1 for r in rows if r["Situação"] == "VENCIDA")], ["Com conflito novo no HubSpot", sum(1 for r in rows if r["Conflito (reconferido)"].startswith("ATENÇÃO"))], [""], ["Etapa", "Empresas"]] + \
             [[e, et.get(e, 0)] for e in ETAPAS] + [[""], ["Master"] + ETAPAS] + [[m] + [c.get(e, 0) for e in ETAPAS] for m, c in sorted(pm.items())]
    w = wb.create_sheet("Painel")
    for l in painel: w.append(l)
    w.column_dimensions["A"].width = 34
    aba("Cadência", ["Toque", "Dia útil (a partir do início)", "Canal", "Ação"], [[t["toque"], t["dia_util"], t["canal"], t["acao"]] for t in C["cadencia"]], 30)
    aba("Como usar", ["Como usar"], [[t] for t in [
        "1. Preencha só as colunas em amarelo (Etapa, Toques feitos, Data último contato, Resultado, Reunião, Próximo passo, Motivo, Aceite do Master, Observações).",
        "2. Rode de novo o followup.py (ou peça à Central BDR 'atualizar follow-up'): a próxima ação, a data e a situação são recalculadas; o que você escreveu é preservado.",
        "3. A aba Agenda mostra o que venceu e o que vence nos próximos 5 dias úteis.",
        "4. A aba Propostas HubSpot traz RASCUNHOS de registro/atualização. Nada é gravado: você confere e registra (ou aprova) no HubSpot.",
        "5. Mensagens da cadência: modelos em ~/.claude/skills/celula-jenifer/references/modelos-mensagem.md. Todas são rascunhos; nada é enviado automaticamente.",
        "6. LGPD: não registre aqui e-mail ou telefone pessoal. Use o canal que o Master apresentou.",
        "7. 'Pediu para não contatar' → Etapa Descartado + motivo opt-out. Não volta para cadência.",
        f"SLA propostos (valem por padrão; Jenifer e Daniela ajustam): aceite do Master em {C['sla']['master_aceite_dias_uteis']} dias úteis · follow-up pós-reunião em {C['sla']['followup_pos_reuniao_horas']} h · nutrição retoma em {C['sla']['nutricao_retomar_dias']} dias."]])
    wb.save(CTRL)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--lista"); ap.add_argument("--top", type=int, default=0); ap.add_argument("--inicio")
    a = ap.parse_args(); C = json.load(open(CFG))
    rows = ler_controle(); novos = 0
    if a.lista:
        if not os.path.exists(a.lista): sys.exit(f"Lista não encontrada: {a.lista}")
        novos = importar(a.lista, a.top, dt(a.inicio) or HOJE, rows)
    R, fonte = hubspot_raizes(); recalcular(rows, C, R); salvar(rows, C, fonte)
    print(json.dumps({"arquivo": CTRL, "empresas": len(rows), "importadas_agora": novos, "vencidas": sum(1 for r in rows if r["Situação"] == "VENCIDA"),
                      "conflito_novo": sum(1 for r in rows if str(r["Conflito (reconferido)"]).startswith("ATENÇÃO")),
                      "propostas_hubspot": sum(1 for r in rows if proposta_hubspot(r))}, ensure_ascii=False))
