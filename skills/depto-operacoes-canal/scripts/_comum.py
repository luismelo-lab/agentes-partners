"""Funções comuns do Departamento de Operações de Canal & Comissionamento (somente leitura).

Lê o download do HubSpot gerado por depto-resultados/scripts/hubspot_fonte.py. Não chama API,
não lê .env, não grava nada fora de ~/agentes-partners/departamentos/operacoes-canal/.
"""
import glob, json, os, re, unicodedata
from datetime import datetime

BASE = os.path.expanduser("~/agentes-partners/departamentos")
SRC = os.path.join(BASE, "resultados/dados")
OUT = os.path.join(BASE, "operacoes-canal")
MAPA = os.path.join(BASE, "resultados/mapeamento-hubspot.json")
CFG = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "references/config-canal.json")
LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"


def carregar(arquivo=None):
    f = arquivo or sorted(glob.glob(os.path.join(SRC, "hubspot-*.json")))[-1]
    D = json.load(open(f))
    cfg = json.load(open(CFG))
    roster = json.load(open(MAPA))["escopo"]["parceiros_roster"]
    return f, D, cfg, roster


def data_ref(D):
    return (D.get("baixado_em") or datetime.now().strftime("%Y-%m-%d"))[:10]


def sem_acento(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def cnpj_limpo(v):
    d = re.sub(r"\D", "", str(v or ""))
    if not d or len(d) > 14: return None
    return d.zfill(14)


def cnpj_valido(c):
    if not c or len(c) != 14 or len(set(c)) == 1: return False
    def dv(base, pesos):
        s = sum(int(a) * b for a, b in zip(base, pesos)) % 11
        return "0" if s < 2 else str(11 - s)
    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]; p2 = [6] + p1
    return c[12] == dv(c[:12], p1) and c[13] == dv(c[:12] + c[12], p2)


def canal(d, cfg, roster):
    """Retorna (canal_canônico, tipo, observação). tipo: Master | Fora do roster | Sem parceiro."""
    p = d.get("parceiro") or ""
    if p in roster:
        return p, "Master", ""
    bruto = (d.get("parceiro_bruto") or "").strip()
    nb = sem_acento(bruto)
    if not nb:
        return "Sem parceiro identificado", "Sem parceiro", "campo parceiro vazio"
    if nb in {sem_acento(x) for x in cfg.get("parceiro_invalido", [])}:
        return "Sem parceiro identificado", "Sem parceiro", f"campo parceiro com valor numérico/ inválido ('{bruto}')"
    for nome, trechos in cfg.get("aliases_fora_roster", {}).items():
        if nome.startswith("_"): continue
        if any(t in nb for t in trechos):
            return nome, "Fora do roster", (f"apelido '{bruto}' agrupado" if sem_acento(nome) != nb else "")
    obs = ""
    alvo = cfg.get("aliases_a_confirmar", {}).get(nb)
    if alvo: obs = f"possível mesmo parceiro que {alvo} (apelido a confirmar)"
    return bruto.title() if bruto.isupper() else bruto, "Fora do roster", obs


# ---------------- Excel ----------------
def planilha(caminho, abas):
    """abas: lista de (nome, cabeçalhos, linhas, formatos{coluna: formato}, larguras{coluna: n})."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    cab = PatternFill("solid", fgColor="1F2A36"); ouro = Font(bold=True, color="E9C46A")
    for nome, heads, rows, fmts, wid in abas:
        ws = wb.create_sheet(nome[:31])
        if heads:
            ws.append(heads)
            for c in ws[1]: c.fill = cab; c.font = ouro; c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows: ws.append(list(r))
        if heads and rows:
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(heads or [], 1):
            L = get_column_letter(i)
            ws.column_dimensions[L].width = (wid or {}).get(h, max(10, min(45, len(str(h)) + 4)))
            f = (fmts or {}).get(h)
            if f:
                for c in ws[L][1:]: c.number_format = f
            if h == "Link HubSpot":
                for c in ws[L][1:]:
                    if c.value: c.hyperlink = c.value; c.font = Font(color="0563C1", underline="single")
        if not heads:
            ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 110
            for row in ws.iter_rows():
                for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
                if row[0].value and (len(row) < 2 or row[1].value in (None, "")):
                    row[0].font = Font(bold=True, color="1F2A36", size=12)
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    wb.save(caminho)
    return caminho


RS = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'
