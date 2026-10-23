#!/usr/bin/env python3
"""Dados & BI — qualidade do CRM e scorecard de parceiros (cálculo determinístico).

Lê o download mais recente do HubSpot (gerado por depto-resultados/scripts/hubspot_fonte.py)
e gera ~/agentes-partners/departamentos/dados/Dados-BI-AAAA-MM-DD.xlsx + resumo .json.

Uso: python3 dados_bi.py [--baixar] [--parados 30]
"""
import argparse, glob, json, os, re, statistics, subprocess, sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
sys.path.insert(0, os.path.expanduser("~/agentes-partners/lib"))
from cnpj import normalizar_cnpj  # noqa: E402

BASE = os.path.expanduser("~/agentes-partners/departamentos")
SRC = os.path.join(BASE, "resultados/dados")
OUT = os.path.join(BASE, "dados")
HS = os.path.expanduser("~/.claude/skills/depto-resultados/scripts/hubspot_fonte.py")
LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"



def ef(d):  # data efetiva (dicionário metricas.md): a menor entre criação e fechamento; em aberto, a criação
    return min((x for x in (d["criado"], d["fechado"] if d["status"] != "aberto" else "") if x), default="")

def dt(s):
    try: return datetime.fromisoformat((s or "")[:10]).date()
    except ValueError: return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--baixar", action="store_true", help="baixa o HubSpot antes (≈1–2 min)")
    ap.add_argument("--parados", type=int, default=30, help="dias sem atividade para considerar negócio aberto parado")
    a = ap.parse_args()
    if a.baixar:
        subprocess.run([sys.executable, HS, "--baixar"], check=True)
    f = sorted(glob.glob(os.path.join(SRC, "hubspot-*.json")))[-1]
    D = json.load(open(f)); N = D["negocios"]; hoje = date.today()
    os.makedirs(OUT, exist_ok=True)
    R = {"fonte": f, "baixado_em": D["baixado_em"], "negocios": len(N)}

    # ---------- 1. Completude por campo e status ----------
    campos = [("valor", lambda d: d["valor"] > 0), ("parceiro", lambda d: not d["parceiro"].startswith("Outros")),
              ("% honorário", lambda d: d["hContr"] is not None), ("% parceiro", lambda d: d["hParc"] is not None),
              ("tese aplicada", lambda d: bool(d.get("variante"))), ("CNAE/segmento", lambda d: bool(d.get("segmento"))),
              ("CNPJ", lambda d: bool(d["extras"].get("cnpj"))), ("UF", lambda d: bool(d["extras"].get("uf"))),
              ("regime", lambda d: bool(d["extras"].get("regime"))), ("faturamento", lambda d: bool(d["extras"].get("faturamento"))),
              ("última atividade", lambda d: bool(d["extras"].get("ultima_atividade")))]
    comp = []
    for nome, fn in campos:
        linha = {"campo": nome}
        for st in ("ganho", "aberto", "perdido"):
            G = [d for d in N if d["status"] == st]
            linha[st] = round(sum(1 for d in G if fn(d)) / len(G) * 100, 1) if G else None
        comp.append(linha)
    R["completude"] = comp
    # índice de qualidade dos ganhos: média dos campos essenciais para margem e ROI
    ess = ["valor", "% honorário", "% parceiro", "tese aplicada", "CNAE/segmento", "CNPJ"]
    R["indice_qualidade_ganhos"] = round(statistics.mean(l["ganho"] for l in comp if l["campo"] in ess), 1)

    # ---------- 2. Listas de correção ----------
    def row(d, motivo):
        return {"motivo": motivo, "negócio": d["nome"], "parceiro": d["parceiro"], "owner": d["owner"], "status": d["status"],
                "valor": round(d["valor"]), "criado": d["criado"], "fechado": d["fechado"], "link": LINK.format(d["id"])}
    incons = []
    for d in N:
        if d["status"] == "ganho":
            if d["valor"] <= 0: incons.append(row(d, "Ganho sem valor"))
            if d["hContr"] is None: incons.append(row(d, "Ganho sem % de honorário"))
            if not d.get("variante"): incons.append(row(d, "Ganho sem tese registrada"))
        if d["hContr"] is not None and d["hContr"] > 0.5: incons.append(row(d, f"Honorário atípico ({d['hContr']*100:.0f}% do crédito)"))
        if d["hParc"] is not None and d["hParc"] > 1: incons.append(row(d, f"% do parceiro acima de 100% ({d['hParc']*100:.0f}%)"))
        if d["status"] == "aberto" and d["valor"] <= 0: incons.append(row(d, "Aberto sem valor"))
    R["inconsistencias"] = dict(Counter(r["motivo"].split(" (")[0] for r in incons))
    corte = hoje - timedelta(days=a.parados)
    parados = []
    for d in N:
        if d["status"] != "aberto": continue
        ua = dt(d["extras"].get("ultima_atividade")) or dt(d["extras"].get("ultima_modificacao"))
        if not ua or ua < corte:
            r = row(d, f"Aberto sem atividade há mais de {a.parados} dias"); r["última atividade"] = ua.isoformat() if ua else "nunca"
            r["dias parado"] = (hoje - ua).days if ua else None; parados.append(r)
    parados.sort(key=lambda r: -(r["valor"] or 0))
    R["abertos_parados"] = {"n": len(parados), "rs": sum(r["valor"] for r in parados)}
    # duplicidade: mesmo CNPJ com 2+ negócios abertos ao mesmo tempo (vários ganhos do mesmo CNPJ podem ser teses diferentes)
    porc = defaultdict(list)
    for d in N:
        c = normalizar_cnpj(d["extras"].get("cnpj"))
        if c and d["status"] == "aberto": porc[c].append(d)
    dup = [dict(row(d, "CNPJ com 2+ negócios abertos"), cnpj=c) for c, L in porc.items() if len(L) > 1 for d in L]
    R["duplicidade_abertos"] = {"cnpjs": sum(1 for L in porc.values() if len(L) > 1), "negocios": len(dup)}
    imp = Counter((d["criado"][:7]) for d in N if str(d["extras"].get("importado")).lower() == "true")
    R["importacoes_por_mes"] = dict(sorted(imp.items())[-18:])

    # ---------- 3. Scorecard de parceiros ----------
    d90, d180, d365 = hoje - timedelta(days=90), hoje - timedelta(days=180), hoje - timedelta(days=365)
    P = defaultdict(list)
    for d in N: P[d["parceiro"]].append(d)
    sc = []
    for p, L in P.items():
        cri90 = [d for d in L if (dt(ef(d)) or date.min) >= d90]
        g = lambda a_, b_: [d for d in L if d["status"] == "ganho" and a_ <= (dt(d["fechado"]) or date.min) < b_]
        g90, g90a, g12 = g(d90, hoje + timedelta(days=1)), g(d180, d90), g(d365, hoje + timedelta(days=1))
        p12 = [d for d in L if d["status"] == "perdido" and (dt(d["fechado"]) or date.min) >= d365]
        ab = [d for d in L if d["status"] == "aberto"]
        ciclos = [c for c in ((dt(d["fechado"]) - dt(d["criado"])).days for d in g12 if dt(d["fechado"]) and dt(d["criado"])) if c >= 0]  # ciclo negativo = criação regravada pela importação
        liq = [d["valor"] * d["hContr"] * (1 - min(d["hParc"] or 0, 1)) for d in g12 if d["hContr"] is not None]
        conv = len(g12) / (len(g12) + len(p12)) * 100 if (g12 or p12) else None
        v90, v90a = sum(d["valor"] for d in g90), sum(d["valor"] for d in g90a)
        tend = (v90 / v90a - 1) * 100 if v90a else (100.0 if v90 else None)
        hon_pct = len([d for d in g12 if d["hContr"] is not None]) / len(g12) * 100 if g12 else None
        par = sum(1 for d in ab if d["id"] in {r["link"].rsplit("/", 1)[1] for r in parados})
        # índice de saúde 0–100, comparando o parceiro com a média trimestral dele mesmo nos 12 meses
        # (pesos: ritmo de entrada 30 · conversão 25 · ritmo de valor fechado 25 · qualidade do dado 20)
        cri_ant = [d for d in L if d365 <= (dt(ef(d)) or date.min) < d90]
        base_cri = len(cri_ant) / 3 if cri_ant else None
        r_at = (len(cri90) / base_cri) if base_cri else (0.5 if cri90 else 0)
        base_val = sum(d["valor"] for d in g12) / 4
        r_val = (v90 / base_val) if base_val else 0
        s_at = min(r_at, 1.2) / 1.2 * 30
        s_cv = (min(conv / 35, 1) * 25) if conv is not None else 0
        s_td = min(r_val, 1.2) / 1.2 * 25
        s_q = (hon_pct or 0) / 100 * 20
        saude = round(s_at + s_cv + s_td + s_q)
        tend = round((r_val - 1) * 100, 1) if base_val else None
        sc.append({"parceiro": p, "saude": saude, "sinal": ("⚪ base pequena" if len(g12) < 5 else "🟢" if saude >= 70 else "🟡" if saude >= 45 else "🔴"),
                   "criados_90d": len(cri90), "ganhos_90d": len(g90), "fechado_90d_rs": round(v90), "ritmo_valor_vs_media_trimestral_pct": tend, "ritmo_entrada_vs_media_pct": round((r_at - 1) * 100, 1) if base_cri else None,
                   "ganhos_12m": len(g12), "fechado_12m_rs": round(sum(d["valor"] for d in g12)), "conversao_12m_pct": round(conv, 1) if conv is not None else None,
                   "ticket_12m_rs": round(sum(d["valor"] for d in g12) / len(g12)) if g12 else None, "ciclo_mediano_dias": round(statistics.median(ciclos)) if ciclos else None,
                   "honorario_liquido_12m_rs": round(sum(liq)), "honorario_preenchido_pct": round(hon_pct, 1) if hon_pct is not None else None,
                   "abertos": len(ab), "abertos_rs": round(sum(d["valor"] for d in ab)), "abertos_parados": par,
                   "ultimo_negocio_criado": max((d["criado"] for d in L), default="")})
    sc.sort(key=lambda x: -x["fechado_12m_rs"])
    R["scorecard"] = sc

    # ---------- Saída ----------
    stamp = hoje.isoformat()
    json.dump(R, open(os.path.join(OUT, f"dados-bi-{stamp}.json"), "w"), ensure_ascii=False, indent=1, default=str)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F1F1F")
    def aba(nome, linhas, primeira=False):
        ws = wb.active if primeira else wb.create_sheet(nome); ws.title = nome
        if not linhas: ws["A1"] = "Nada a listar."; return
        cols = list(linhas[0].keys())
        for j, c in enumerate(cols, 1):
            x = ws.cell(row=1, column=j, value=c); x.font = H; x.fill = HF
            ws.column_dimensions[x.column_letter].width = min(max(12, len(c) + 2), 48)
        for i, l in enumerate(linhas, 2):
            for j, c in enumerate(cols, 1):
                v = l[c]; x = ws.cell(row=i, column=j, value=v)
                if c == "link" and v: x.hyperlink = v; x.font = Font(color="1F5FBF", underline="single")
        ws.freeze_panes = "A2"
    resumo = [{"indicador": "Índice de qualidade dos ganhos (média dos campos essenciais)", "valor": f"{R['indice_qualidade_ganhos']}%"},
              {"indicador": f"Abertos parados (> {a.parados} dias sem atividade)", "valor": f"{R['abertos_parados']['n']} · R$ {R['abertos_parados']['rs']/1e6:.1f} mi"},
              {"indicador": "CNPJs com 2+ negócios abertos (possível duplicidade)", "valor": R["duplicidade_abertos"]["cnpjs"]}]
    resumo += [{"indicador": k, "valor": v} for k, v in R["inconsistencias"].items()]
    aba("Resumo", resumo, primeira=True); aba("Scorecard parceiros", sc); aba("Completude", comp)
    aba("Abertos parados", parados); aba("Inconsistências", incons); aba("Duplicidade", dup)
    x = os.path.join(OUT, f"Dados-BI-{stamp}.xlsx"); wb.save(x); print(x)


if __name__ == "__main__":
    main()
