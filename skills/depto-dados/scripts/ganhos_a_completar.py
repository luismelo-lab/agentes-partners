#!/usr/bin/env python3
"""Ganhos do mês com dado essencial faltando (% de honorário, % do parceiro, tese), para a Daniela completar no HubSpot (D62).

Lê o download mais recente do HubSpot (depto-resultados/scripts/hubspot_fonte.py). Só leitura: nada é gravado no HubSpot.
O % de honorário só é informado no ganho (D71 e D73, Luis 22/10/2026), por isso a lista olha só os ganhos.

  python3 ganhos_a_completar.py [--mes 2026-10] [--saida <arquivo.xlsx>]   padrão: pasta da Daniela no OneDrive
"""
import argparse, glob, json, os, sys
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

AP = os.path.expanduser("~/agentes-partners"); sys.path.insert(0, os.path.join(AP, "lib"))
import caminhos, marca
LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"
C = {k: v.lstrip("#") for k, v in marca.C.items() if not k.startswith("_")}; F = "IBM Plex Sans"; FM = "IBM Plex Mono"
COLS = [("Negócio", 38), ("Parceiro", 20), ("Responsável", 18), ("Valor (R$)", 15), ("Fechado em", 12), ("O que falta", 30),
        ("% honorário", 12), ("% parceiro", 11), ("Tese aplicada", 26), ("Feito no HubSpot?", 15), ("Observação", 30), ("Abrir no HubSpot", 18)]


def main():
    a = argparse.ArgumentParser(); a.add_argument("--mes", default=date.today().strftime("%Y-%m")); a.add_argument("--saida"); a = a.parse_args()
    fonte = sorted(glob.glob(os.path.join(AP, "departamentos/resultados/dados/hubspot-*.json")))[-1]
    N = json.load(open(fonte))["negocios"]
    L = [d for d in N if d["status"] == "ganho" and (d.get("fechado") or "")[:7] == a.mes]
    linhas = []
    for d in L:
        falta = [n for n, ok in (("% honorário", d["hContr"] is not None), ("% parceiro", d["hParc"] is not None), ("tese", bool(d.get("variante")))) if not ok]
        if falta: linhas.append((d, falta))
    linhas.sort(key=lambda x: (-len(x[1]), -(x[0]["valor"] or 0)))
    saida = a.saida or os.path.join(caminhos.pasta("daniela"), f"Ganhos-a-completar-{a.mes}.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Completar"; ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
    c = ws.cell(1, 1, f"Ganhos de {a.mes[5:]}/{a.mes[:4]} com dado faltando · {len(linhas)} de {len(L)} · complete no HubSpot e marque 'Feito'")
    c.font = Font(name=F, bold=True, size=13, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro"]); ws.row_dimensions[1].height = 26
    for j, (t, w) in enumerate(COLS, 1):
        c = ws.cell(2, j, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro_suave"])
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=Side(style="medium", color=C["ouro"]))
        ws.column_dimensions[get_column_letter(j)].width = w
    fino = Border(bottom=Side(style="thin", color=C["borda"]))
    for i, (d, falta) in enumerate(linhas, 3):
        vals = [d["nome"], d["parceiro"], d["owner"], round(d["valor"] or 0), d["fechado"][:10], ", ".join(falta),
                d["hContr"], d["hParc"], d.get("variante") or None, None, None, "Abrir"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.border = fino; c.alignment = Alignment(vertical="top", wrap_text=j in (1, 6, 11))
            c.font = Font(name=FM if j in (4, 7, 8) else F, size=10, color=C["texto"])
        ws.cell(i, 4).number_format = "#,##0"; ws.cell(i, 7).number_format = ws.cell(i, 8).number_format = "0.0%"
        for j in (7, 8, 9):
            if ws.cell(i, j).value is None: ws.cell(i, j).fill = PatternFill("solid", fgColor=C["ouro_suave"])
        ws.cell(i, 12).hyperlink = LINK.format(d["id"]); ws.cell(i, 12).font = Font(name=F, size=10, color=C["ouro_texto"], underline="single")
    fim = len(linhas) + 2
    dv = DataValidation(type="list", formula1='"Feito,Falta dado,Não se aplica"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"J3:J{max(3, fim)}")
    ws.freeze_panes = "B3"; ws.auto_filter.ref = f"A2:L{max(2, fim)}"
    u = wb.create_sheet("Como usar"); u.column_dimensions["A"].width = 110; u.sheet_view.showGridLines = False
    for i, t in enumerate(["Como usar",
            "1. Abra cada negócio pelo link e preencha no HubSpot o que está em ouro: % de honorário, % do parceiro e tese aplicada.",
            "2. Marque 'Feito no HubSpot?'. Se faltar o dado (contrato sem o %), marque 'Falta dado' e diga em 'Observação' quem tem.",
            "3. O % de honorário é informado só no ganho (decisão do Luis, 22/10/2026). O repasse segue o percentual negociado de cada parceiro.",
            f"4. Fonte: download do HubSpot de {os.path.basename(fonte)[8:18]}. O depto-dados confere no download seguinte; esta planilha não altera o HubSpot."], 1):
        c = u.cell(i, 1, t); c.font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11, color=C["texto"]); c.alignment = Alignment(wrap_text=True)
    wb.save(saida); print("ok", saida, f"{len(linhas)} de {len(L)} ganhos com dado faltando")


if __name__ == "__main__":
    main()
