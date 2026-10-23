#!/usr/bin/env python3
"""Lista de saneamento dos ganhos possivelmente duplicados no HubSpot (D87), para o time de CRM decidir. Nada é alterado no HubSpot.

Método (references/especialistas/qualidade-saneamento.md, seção 5):
  1. Foto congelada: o download mais recente do HubSpot (departamentos/resultados/dados/hubspot-AAAA-MM-DD.json).
  2. Chaves normalizadas: CNPJ em maiúsculas sem . / - e espaços (preserva as letras do CNPJ alfanumérico, lib/cnpj.py);
     raiz = 8 primeiras posições; tese = variantes do card sem acento e em minúsculas; valor arredondado ao real;
     competência = período de apuração lido no parecer do técnico (celulas/geovana/pareceres/resumos.json), quando há.
  3. Bloqueio pela raiz do CNPJ: só se comparam ganhos da mesma raiz.
  4. Regras por par:
       R1 mesmo id_bitrix ............................................................ alta
       R2 mesmo valor (diferença até R$ 1, valor > 0) ................................. alta
       R3 mesma data de fechamento + mesma tese (ou tese vazia em um) + selo de importado alta
       R3b mesma data de fechamento + mesma tese (ou vazia), nenhum importado ......... média
       R4 mesma tese + competências sobrepostas + valores diferentes até 10% ......... média
       R5 mesma data de fechamento + teses diferentes ................................ baixa (provável legítimo: aba à parte)
  5. Grupos: união dos pares de alta e média confiança. Grupo "alta" quando todos os cards se ligam só por pares de alta.
  6. Regra de sobrevivência (qual card fica): (a) mais campos essenciais preenchidos (valor, % honorário, % parceiro, tese, CNAE, CNPJ);
     (b) empate: tem atividade, anexo de diagnóstico ou parecer; (c) empate: criado no HubSpot (não importado); (d) empate: menor ID.
     O que o card que sai tem e o que fica não tem vai em "levar para o card que fica".
  Valor em duplicidade = soma do valor dos cards que saem. É valor de crédito estimado, não honorário nem caixa.
  Grupo com card acima de R$ 10 mi vai ao Luis antes da mescla. D76: a coluna "Escopo Partners" separa roster e Daniela do resto.

  python3 saneamento_duplicados.py [--saida <arquivo.xlsx>]
  padrão: Saneamento-Ganhos-Duplicados-<hoje>.xlsx em 01 · Comercial e Parceiros/Relatórios das áreas (caminhos.luis('areas'))
"""
import argparse, glob, json, os, re, sys, unicodedata
from collections import Counter, defaultdict
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

AP = os.path.expanduser("~/agentes-partners"); sys.path.insert(0, os.path.join(AP, "lib"))
import caminhos, marca  # noqa: E402
from cnpj import completar_cnpj, raiz_cnpj  # noqa: E402

LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"
HS = os.path.join(AP, "departamentos/resultados/dados")
RES = os.path.join(AP, "celulas/geovana/pareceres/resumos.json")
MAPA = os.path.join(AP, "departamentos/resultados/mapeamento-hubspot.json")
C = {k: v.lstrip("#") for k, v in marca.C.items() if not k.startswith("_")}; F = "IBM Plex Sans"; FM = "IBM Plex Mono"
NIVEL = {"alta": 3, "média": 2, "baixa": 1}
REGRAS = {"R1": ("mesmo id_bitrix", "alta"), "R2": ("mesmo valor", "alta"),
          "R3": ("mesma data de fechamento + mesma tese (ou vazia) + importado", "alta"),
          "R3b": ("mesma data de fechamento + mesma tese (ou vazia), sem selo de importado", "média"),
          "R4": ("mesma tese + competências sobrepostas + valor até 10% diferente", "média"),
          "R5": ("mesma data de fechamento + teses diferentes", "baixa")}


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def teses(v):
    return frozenset(re.sub(r"\s+", " ", sem_acento(x).lower()).strip() for x in str(v or "").split(";") if x.strip())


def competencia(txt):
    """'12/2020 a 10/2025' -> (202012, 202510)."""
    m = re.findall(r"(\d{2})/(\d{4})", str(txt or ""))
    if len(m) < 2: return None
    a, b = int(m[0][1]) * 100 + int(m[0][0]), int(m[-1][1]) * 100 + int(m[-1][0])
    return (min(a, b), max(a, b))


def essenciais(d):
    e = d["extras"]
    return sum(1 for ok in (d.get("valor"), d.get("hContr") is not None, d.get("hParc") is not None, d.get("variante"), d.get("cnae"), e.get("cnpj")) if ok)


def evidencia(d):
    e = d["extras"]
    return bool(e.get("ultima_atividade") or e.get("anexo_diagnostico") or e.get("parecer_rt") or e.get("parecer_juridico"))


def importado(d):
    return str(d["extras"].get("importado") or "").lower() == "true"


def par(a, b):
    """Regras que casam o par (lista de códigos)."""
    r = []
    ba, bb = a["extras"].get("id_bitrix"), b["extras"].get("id_bitrix")
    if ba and ba == bb: r.append("R1")
    va, vb = round(a.get("valor") or 0), round(b.get("valor") or 0)
    if va > 0 and abs(va - vb) <= 1: r.append("R2")
    ta, tb = a["_teses"], b["_teses"]
    compat = ta == tb or not ta or not tb
    fa, fb = (a.get("fechado") or "")[:10], (b.get("fechado") or "")[:10]
    if fa and fa == fb:
        if compat: r.append("R3" if importado(a) or importado(b) else "R3b")
        else: r.append("R5")
    ca, cb = a["_comp"], b["_comp"]
    if ta and tb and (ta & tb) and ca and cb and ca[0] <= cb[1] and cb[0] <= ca[1] and va > 0 and vb > 0 and va != vb \
            and abs(va - vb) / max(va, vb) <= 0.10:
        r.append("R4")
    return r


class UF:
    def __init__(s): s.p = {}
    def f(s, x):
        s.p.setdefault(x, x)
        while s.p[x] != x: s.p[x] = s.p[s.p[x]]; x = s.p[x]
        return x
    def u(s, a, b): s.p[s.f(a)] = s.f(b)


def sobrevivente(cards):
    chave = lambda d: (-essenciais(d), not evidencia(d), importado(d), int(d["id"]))
    ordem = sorted(cards, key=chave); fica = ordem[0]; seg = ordem[1] if len(ordem) > 1 else None
    if seg is None: motivo = "único"
    elif essenciais(fica) > essenciais(seg): motivo = f"(a) mais campos essenciais preenchidos ({essenciais(fica)} de 6 contra {essenciais(seg)})"
    elif evidencia(fica) and not evidencia(seg): motivo = "(b) empate em campos; tem atividade, anexo ou parecer"
    elif not importado(fica) and importado(seg): motivo = "(c) empate; criado no HubSpot (não importado)"
    else: motivo = "(d) empate em tudo; menor ID"
    return fica, motivo


def levar(fica, sai):
    nomes = [("% honorário", lambda d: d.get("hContr")), ("% parceiro", lambda d: d.get("hParc")), ("tese", lambda d: d.get("variante")),
             ("CNAE", lambda d: d.get("cnae")), ("parecer RT", lambda d: d["extras"].get("parecer_rt")),
             ("anexo do diagnóstico", lambda d: d["extras"].get("anexo_diagnostico")), ("parecer jurídico", lambda d: d["extras"].get("parecer_juridico"))]
    return ", ".join(n for n, g in nomes if g(sai) not in (None, "") and g(fica) in (None, "")) or None


def montar():
    fonte = sorted(glob.glob(os.path.join(HS, "hubspot-*.json")))[-1]; D = json.load(open(fonte))
    R = json.load(open(RES)) if os.path.exists(RES) else {}
    mapa = json.load(open(MAPA)); roster = set(mapa["escopo"]["parceiros_roster"]); owners = set(mapa["escopo"].get("owners") or [])
    G = [d for d in D["negocios"] if d["status"] == "ganho"]
    bloco = defaultdict(list); sem_cnpj = 0; alfa = 0
    for d in G:
        c = completar_cnpj(d["extras"].get("cnpj")); d["_cnpj"] = c; d["_raiz"] = raiz_cnpj(c)
        d["_teses"] = teses(d.get("variante")); d["_comp"] = competencia((R.get(d["id"]) or {}).get("periodo_apuracao"))
        d["_escopo"] = d["parceiro"] in roster or d.get("owner") in owners
        if re.search(r"[A-Z]", c): alfa += 1
        if d["_raiz"]: bloco[d["_raiz"]].append(d)
        else: sem_cnpj += 1
    uf_ah, uf_all = UF(), UF(); arestas = defaultdict(set); baixas = []
    for rz, ds in bloco.items():
        for i in range(len(ds)):
            for j in range(i + 1, len(ds)):
                a, b = ds[i], ds[j]; rs = par(a, b)
                if not rs: continue
                nivel = max(NIVEL[REGRAS[x][1]] for x in rs)
                if nivel == 1: baixas.append((a, b, rs)); continue
                uf_all.u(a["id"], b["id"]); arestas[frozenset((a["id"], b["id"]))] = set(rs)
                if nivel == 3: uf_ah.u(a["id"], b["id"])
    porid = {d["id"]: d for d in G}; comp = defaultdict(list)
    for k in list(uf_all.p): comp[uf_all.f(k)].append(porid[k])
    grupos = []
    for cards in comp.values():
        ids = {d["id"] for d in cards}
        alta = len({uf_ah.f(i) for i in ids}) == 1
        regras = sorted({r for k, rs in arestas.items() if k <= ids for r in rs})
        fica, motivo = sobrevivente(cards)
        saem = [d for d in cards if d is not fica]
        cnpjs = {d["_cnpj"] for d in cards}
        grupos.append({"cards": sorted(cards, key=lambda d: d is not fica), "fica": fica, "motivo": motivo, "alta": alta, "regras": regras,
                       "excedente": sum(d.get("valor") or 0 for d in saem),
                       "excedente_min": sum(d.get("valor") or 0 for d in cards) - max(d.get("valor") or 0 for d in cards),
                       "valores_dif": len({round(d.get("valor") or 0) for d in cards}) > 1, "n_saem": len(saem), "filial": len(cnpjs) > 1,
                       "grande": any((d.get("valor") or 0) > 10_000_000 for d in cards), "escopo": any(d["_escopo"] for d in cards),
                       "parceiro": fica["parceiro"]})
    grupos.sort(key=lambda g: (not g["alta"], not g["escopo"], -g["excedente"]))
    for i, g in enumerate(grupos, 1): g["id"] = f"G{i:03d}"
    meta = {"fonte": os.path.basename(fonte), "baixado": D.get("baixado_em", ""), "ganhos": len(G), "sem_cnpj": sem_cnpj, "alfa": alfa,
            "raizes": len(bloco), "pares_baixa": len(baixas)}
    return grupos, baixas, meta


COLS = [("Grupo", 7), ("Confiança", 9), ("Regras que casaram", 16), ("Fica ou sai", 9), ("Negócio", 36), ("CNPJ", 19), ("Parceiro", 14),
        ("Escopo Partners", 9), ("Funil / etapa", 22), ("Valor (R$)", 15), ("Fechado em", 11), ("Criado em", 11), ("Tese no card", 30),
        ("Competência (parecer)", 13), ("id_bitrix", 9), ("Importado", 9), ("Campos essenciais (de 6)", 9), ("Por que fica / o que levar", 40),
        ("Alerta", 26), ("Decisão do CRM", 14), ("Observação", 28), ("Abrir no HubSpot", 12)]


def cabecalho(ws, titulo, cols):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    c = ws.cell(1, 1, titulo); c.font = Font(name=F, bold=True, size=13, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro"])
    c.alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 28
    for j, (t, w) in enumerate(cols, 1):
        c = ws.cell(2, j, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro_suave"])
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=Side(style="medium", color=C["ouro"]))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[2].height = 32


def comp_txt(c):
    return f"{c[0] % 100:02d}/{c[0] // 100} a {c[1] % 100:02d}/{c[1] // 100}" if c else None


def linha_card(g, d):
    fica = d is g["fica"]
    alerta = []
    if g["grande"]: alerta.append("card acima de R$ 10 mi: Luis antes de mesclar")
    if re.search(r"recorr|rean[aá]lis|habilita|complementar", d["nome"], re.I): alerta.append("nome indica recorrência, reanálise ou habilitação: pode ser negócio legítimo")
    if g["valores_dif"]: alerta.append("valores diferentes no grupo: conferir qual valor está certo antes de mesclar")
    if g["filial"]: alerta.append("CNPJs diferentes da mesma raiz (matriz/filial): conferir se é um pedido só")
    return [g["id"], "Alta" if g["alta"] else "Média", ", ".join(g["regras"]), "Fica" if fica else "Sai", d["nome"],
            d["_cnpj"], d["parceiro"], "Sim" if d["_escopo"] else "Não", f"{d['pipeline']} / {d['etapa']}", round(d.get("valor") or 0),
            (d.get("fechado") or "")[:10] or None, (d.get("criado") or "")[:10] or None, d.get("variante"), comp_txt(d["_comp"]),
            d["extras"].get("id_bitrix"), "Sim" if importado(d) else "Não", essenciais(d),
            g["motivo"] if fica else (("Levar para o card que fica: " + levar(g["fica"], d)) if levar(g["fica"], d) else "Nada a levar"),
            "; ".join(alerta) or None, None, None, "Abrir"]


def main():
    a = argparse.ArgumentParser(); a.add_argument("--saida"); a = a.parse_args()
    hoje = date.today().isoformat()
    grupos, baixas, meta = montar()
    saida = a.saida or os.path.join(caminhos.luis("areas"), f"Saneamento-Ganhos-Duplicados-{hoje}.xlsx")
    wb = openpyxl.Workbook()
    # Resumo
    ga = [g for g in grupos if g["alta"]]; gm = [g for g in grupos if not g["alta"]]
    exc = lambda L: sum(g["excedente"] for g in L); exm = lambda L: sum(g["excedente_min"] for g in L); n = lambda L: sum(g["n_saem"] for g in L)
    brl = lambda v: f"R$ {v:,.0f}".replace(",", ".")
    reg = Counter(); regv = Counter()
    for g in grupos:
        for r in g["regras"]: reg[r] += 1; regv[r] += g["excedente"]
    porp = defaultdict(lambda: [0, 0, 0.0])
    for g in grupos: p = porp[g["parceiro"]]; p[0] += 1; p[1] += g["n_saem"]; p[2] += g["excedente"]
    ws = wb.active; ws.title = "Resumo"; ws.sheet_view.showGridLines = False; ws.column_dimensions["A"].width = 62; ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16; ws.column_dimensions["D"].width = 22
    linhas = [("Saneamento dos ganhos duplicados (D87) · " + hoje, None, None, None), (None,) * 4,
              ("Foto do HubSpot", meta["fonte"] + " · " + meta["baixado"], None, None),
              ("Ganhos analisados", meta["ganhos"], None, None), ("Ganhos sem CNPJ (fora do bloqueio)", meta["sem_cnpj"], None, None),
              ("Ganhos com CNPJ alfanumérico", meta["alfa"], None, None), (None,) * 4,
              ("", "Grupos", "Cards que saem", "Valor em duplicidade"),
              ("Alta confiança", len(ga), n(ga), exc(ga)), ("Média confiança", len(gm), n(gm), exc(gm)),
              ("Total", len(grupos), n(grupos), exc(grupos)),
              ("  dos quais no escopo Partners (roster ou Daniela, D76)", sum(g["escopo"] for g in grupos), n([g for g in grupos if g["escopo"]]), exc([g for g in grupos if g["escopo"]])),
              ("  com card acima de R$ 10 mi (Luis decide antes)", sum(g["grande"] for g in grupos), n([g for g in grupos if g["grande"]]), exc([g for g in grupos if g["grande"]])),
              ("  com CNPJs diferentes da mesma raiz (matriz/filial)", sum(g["filial"] for g in grupos), n([g for g in grupos if g["filial"]]), exc([g for g in grupos if g["filial"]])),
              ("Valor mínimo em duplicidade (soma do grupo menos o maior valor)", None, None, exm(grupos)),
              ("  grupos com valores diferentes entre os cards (conferir o valor certo)", sum(g["valores_dif"] for g in grupos), None, None),
              ("Pares de baixa confiança (aba à parte, provável legítimo)", meta["pares_baixa"], None, None), (None,) * 4,
              ("Por regra (um grupo pode ter mais de uma)", "Grupos", None, "Valor em duplicidade")]
    linhas += [(f"{r} · {REGRAS[r][0]} ({REGRAS[r][1]})", reg[r], None, regv[r]) for r in ("R1", "R2", "R3", "R3b", "R4")]
    linhas += [(None,) * 4, ("Por parceiro do card que fica", "Grupos", "Cards que saem", "Valor em duplicidade")]
    linhas += [(p, v[0], v[1], v[2]) for p, v in sorted(porp.items(), key=lambda x: -x[1][2])]
    linhas += [(None,) * 4, ("Valor em duplicidade = soma do valor dos cards que saem; o mínimo supõe que o card certo é o de maior valor. É crédito estimado no card, não honorário nem caixa.", None, None, None)]
    for i, l in enumerate(linhas, 1):
        for j, v in enumerate(l, 1):
            c = ws.cell(i, j, v); c.font = Font(name=FM if (j > 1 and isinstance(v, (int, float))) else F, size=13 if i == 1 else 11,
                                                bold=(i == 1 or (l[0] or "").startswith(("Por ", "Total")) or l[0] == ""), color=C["texto"])
            if j == 4 and isinstance(v, (int, float)): c.number_format = '"R$" #,##0'
        if i == 1:
            for j in range(1, 5): ws.cell(1, j).fill = PatternFill("solid", fgColor=C["ouro"])
    # Grupos
    g1 = wb.create_sheet("Grupos")
    cabecalho(g1, f"Ganhos possivelmente duplicados · {len(grupos)} grupos · {n(grupos)} cards que saem · {brl(exc(grupos))} em duplicidade · decida a coluna T", COLS)
    fino = Border(bottom=Side(style="thin", color=C["borda"])); grosso = Border(top=Side(style="medium", color=C["borda_3"]), bottom=Side(style="thin", color=C["borda"]))
    i = 3
    for g in grupos:
        for k, d in enumerate(g["cards"]):
            for j, v in enumerate(linha_card(g, d), 1):
                c = g1.cell(i, j, v); c.border = grosso if k == 0 else fino; c.alignment = Alignment(vertical="top", wrap_text=j in (5, 9, 13, 18, 19, 21))
                c.font = Font(name=FM if j in (6, 10, 11, 12, 14, 15, 17) else F, size=10, color=C["texto"], bold=(j == 4 and k == 0))
            g1.cell(i, 10).number_format = "#,##0"
            g1.cell(i, 4).font = Font(name=F, size=10, bold=True, color=C["sucesso"] if k == 0 else C["erro"])
            if k == 0: g1.cell(i, 20).value = None
            g1.cell(i, 20).fill = PatternFill("solid", fgColor=C["ouro_suave"])
            g1.cell(i, 22).hyperlink = LINK.format(d["id"]); g1.cell(i, 22).font = Font(name=F, size=10, color=C["ouro_texto"], underline="single")
            i += 1
    fim = max(3, i - 1)
    dv = DataValidation(type="list", formula1='"Mesclar,Não é duplicado,Verificar"', allow_blank=True); g1.add_data_validation(dv); dv.add(f"T3:T{fim}")
    g1.freeze_panes = "F3"; g1.auto_filter.ref = f"A2:{get_column_letter(len(COLS))}{fim}"
    # Pares de baixa confiança
    cb = [("Negócio A", 34), ("Negócio B", 34), ("CNPJ A", 19), ("CNPJ B", 19), ("Parceiro", 14), ("Fechado em", 11), ("Tese A", 30), ("Tese B", 30),
          ("Valor A", 14), ("Valor B", 14), ("Decisão do CRM", 14), ("Link A", 10), ("Link B", 10)]
    b1 = wb.create_sheet("Pares baixa confiança"); cabecalho(b1, f"Mesma raiz e mesma data de fechamento, teses diferentes · {len(baixas)} pares · provável legítimo (duas teses)", cb)
    for i, (x, y, rs) in enumerate(sorted(baixas, key=lambda t: t[0]["nome"]), 3):
        vals = [x["nome"], y["nome"], x["_cnpj"], y["_cnpj"], x["parceiro"], (x.get("fechado") or "")[:10], x.get("variante"), y.get("variante"),
                round(x.get("valor") or 0), round(y.get("valor") or 0), None, "Abrir", "Abrir"]
        for j, v in enumerate(vals, 1):
            c = b1.cell(i, j, v); c.border = fino; c.alignment = Alignment(vertical="top", wrap_text=j in (1, 2, 7, 8))
            c.font = Font(name=FM if j in (3, 4, 6, 9, 10) else F, size=10, color=C["texto"])
        for j in (9, 10): b1.cell(i, j).number_format = "#,##0"
        b1.cell(i, 11).fill = PatternFill("solid", fgColor=C["ouro_suave"])
        for j, d in ((12, x), (13, y)):
            b1.cell(i, j).hyperlink = LINK.format(d["id"]); b1.cell(i, j).font = Font(name=F, size=10, color=C["ouro_texto"], underline="single")
    dv2 = DataValidation(type="list", formula1='"Mesclar,Não é duplicado,Verificar"', allow_blank=True); b1.add_data_validation(dv2); dv2.add(f"K3:K{max(3, len(baixas) + 2)}")
    b1.freeze_panes = "C3"
    # Como usar e Regras
    u = wb.create_sheet("Como usar"); u.column_dimensions["A"].width = 115; u.sheet_view.showGridLines = False
    texto = ["Como usar",
             "Para quem: time de CRM da Save (quem mescla, D18). Nada nesta planilha altera o HubSpot; ela só propõe (D87, item 10b de 22/10/2026).",
             "1. Aba 'Grupos': cada grupo reúne ganhos da mesma empresa (raiz do CNPJ) que as regras apontam como o mesmo negócio. A primeira linha (Fica) é o card proposto para ficar; as demais (Sai) são as que seriam mescladas nele.",
             "2. Abra os cards pelo link e marque 'Decisão do CRM' em cada linha que sai: Mesclar, Não é duplicado (dois negócios legítimos, por exemplo outra tese ou outra competência) ou Verificar (precisa de quem vendeu).",
             "3. Antes de mesclar, copie para o card que fica o que está em 'o que levar' (% honorário, % parceiro, tese, pareceres).",
             "4. Comece pelos grupos de confiança Alta. Grupo com card acima de R$ 10 mi: o Luis decide antes. Grupo com CNPJs diferentes da mesma raiz (matriz e filial): confirme se foi um pedido só.",
             "5. Validação da regra: leia 30 grupos de Alta e 20 de Média e anote o acerto. A regra vale se acertar 95% ou mais na Alta e 80% ou mais na Média; abaixo disso, devolva ao depto-dados para ajuste.",
             "6. A aba 'Pares baixa confiança' traz ganhos da mesma empresa fechados no mesmo dia com teses diferentes: em geral são legítimos. Só marque se souber que é duplicado.",
             "7. O depto-dados confere no download seguinte do HubSpot e recalcula o índice de qualidade e a carteira.",
             "Cuidados: o valor em duplicidade é crédito estimado no card, não honorário nem caixa. Escopo Partners (roster e Daniela, D76) está na coluna H; o resto aparece para o CRM ter a base inteira.",
             "Script: skills/depto-dados/scripts/saneamento_duplicados.py (repositório agentes-partners).", "",
             "Regras (versão 1, 23/10/2026)"]
    texto += [f"{r}: {REGRAS[r][0]} → confiança {REGRAS[r][1]}" for r in REGRAS]
    texto += ["Bloqueio: raiz do CNPJ (8 primeiras posições). CNPJ normalizado em maiúsculas, sem . / - e espaços; letras do CNPJ alfanumérico preservadas.",
              "Sobrevivência: (a) mais campos essenciais preenchidos (valor, % honorário, % parceiro, tese, CNAE, CNPJ); (b) tem atividade, anexo ou parecer; (c) criado no HubSpot, não importado; (d) menor ID.",
              "Competência: período de apuração lido no parecer do técnico, quando existe; sem parecer, a regra R4 não se aplica."]
    for i, t in enumerate(texto, 1):
        c = u.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, bold=t in ("Como usar", "Regras (versão 1, 23/10/2026)"), size=13 if i == 1 else 11, color=C["texto"])
        if i == 1: c.fill = PatternFill("solid", fgColor=C["ouro"])
    wb.move_sheet("Como usar", offset=-(len(wb.sheetnames) - 2))
    wb.save(saida)
    print("ok", saida)
    print(f"ganhos {meta['ganhos']} · grupos {len(grupos)} (alta {len(ga)}, média {len(gm)}) · cards que saem {n(grupos)} (alta {n(ga)}, média {n(gm)})")
    print(f"valor mínimo em duplicidade {brl(exm(grupos))} · grupos com valores diferentes {sum(g['valores_dif'] for g in grupos)}")
    print(f"valor em duplicidade {brl(exc(grupos))} (alta {brl(exc(ga))}, média {brl(exc(gm))}) · escopo Partners {sum(g['escopo'] for g in grupos)} grupos")
    print(f"acima de R$ 10 mi: {sum(g['grande'] for g in grupos)} grupos · matriz/filial: {sum(g['filial'] for g in grupos)} · pares baixa: {meta['pares_baixa']} · sem CNPJ {meta['sem_cnpj']} · alfanumérico {meta['alfa']}")
    print("por regra:", {r: (reg[r], brl(regv[r])) for r in reg})


if __name__ == "__main__":
    main()
