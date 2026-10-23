#!/usr/bin/env python3
"""Revisor de Diagnóstico — 2ª frente da célula da Geovana (segunda opinião antes de o diagnóstico ir ao cliente).

Fluxo determinístico:
  1. seleção: negócios abertos na etapa "Apresentação de Diagnóstico" (pipeline RT - Real & Presumido;
     Simples, Gestão de Passivo, café e CNPJ baixado/inapto vão para a aba "Fora da revisão");
  2. perfil oficial: CNAE principal e secundários, situação, porte, capital, início e Simples da base local da Receita
     (depto-intel-comercial) + regime do HubSpot (com alerta de divergência);
  3. teses esperadas: linha da matriz tese × perfil (recuperador/references/matriz.json) → teses atômicas do catálogo
     (revisor-diagnostico/references/teses.json), com as regras de atividade do Tributário (C só na indústria; insumos fora do comércio);
  4. teses cobertas: leitura do texto já extraído do parecer/diagnóstico do card (celulas/geovana/pareceres/textos) —
     só leitura, nada é baixado aqui;
  5. lacunas: tese esperada e não analisada. Sem texto legível, a lacuna é "provável" quando o técnico costuma
     deixar a tese de fora neste perfil (cobertura típica nos pareceres da base) e "a conferir" nos demais casos;
  6. look-alike: o que os ganhos parecidos (mesma atividade, CNAE/classe e regime; nunca a própria raiz do CNPJ) tiveram;
  7. saída: Revisor-Diagnostico-AAAA-MM-DD.xlsx (Revisão, Lacunas, Fora da revisão, Cobertura típica, Como usar)
     + fichas/AAAA-MM-DD.md + revisor-AAAA-MM-DD.json;
  8. aprendizado: as decisões da Geovana na planilha anterior (Incluir tese / Não se aplica / Precisa de dado) são
     copiadas para a nova e acumuladas em decisoes.json; "Não se aplica" repetido rebaixa a lacuna no perfil.

N1: tudo é pré-análise para a Geovana validar. Sem valores prometidos: o look-alike mostra contagem de casos, não R$.
Uso:  python3 revisor.py [--etapa "Apresentação de Diagnóstico"] [--ids 123,456] [--fichas 30 | --sem-fichas]
"""
import argparse, glob, json, os, re, sqlite3, sys
from collections import Counter, defaultdict
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fila import normalizar_cnpj, perfil, parecido, raiz, linha_matriz, dt  # noqa: E402  (reuso: mesmo perfil e look-alike do Recuperador)

H = os.path.expanduser
SRC = H("~/agentes-partners/departamentos/resultados/dados")
ENR = H("~/agentes-partners/departamentos/intel-comercial/enriquecimento-receita.json")
DB = H("~/agentes-partners/dados-publicos/receita/cnpj.sqlite")
REGCACHE = H("~/agentes-partners/dados-publicos/receita/regime-cache.json")
TXT = H("~/agentes-partners/celulas/geovana/pareceres/textos")
OUT = os.environ.get("REVISOR_OUT") or H("~/agentes-partners/celulas/geovana/revisor")  # REVISOR_OUT só para teste
CAT = json.load(open(H("~/.claude/skills/revisor-diagnostico/references/teses.json")))
TESES = CAT["teses"]; PRI = CAT["prioridade"]
LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"
DECISOES = ("Incluir tese", "Não se aplica", "Precisa de dado")
CORTE = date.fromisoformat(PRI["corte_urgencia"])

# ---------------------------------------------------------------- classificação de tese em trecho de texto
def classificar(s):
    """Códigos de tese citados num trecho (linha do parecer, rótulo do HubSpot)."""
    s = s.lower(); out = set()
    if re.search(r"\birpj\b|\bcsll\b", s) and not re.search(r"saldo negativo", s):
        return {"IRPJ_BASE"}
    if re.search(r"aquisi[çc][õo]es", s): out.add("AQ_SAVE")
    if re.search(r"cr[ée]dito\s+presumido|presumido agro", s) and not re.search(r"exclus[ãa]o do cr[ée]dito presumido|presumido de icms", s): out.add("C")
    if re.search(r"manuten[çc][ãa]o (de|dos) cr[ée]ditos|art\.?\s*17 da lei|cr[ée]ditos? (acumulad|vinculad)|saldo credor|pis e cofins acumulado|ressarcimento de cr[ée]ditos", s): out.add("B")
    if re.search(r"\binsumos?\b", s): out.add("INSUMOS")
    if re.search(r"pr[óo]pria base", s): out.add("PROPRIA_BASE")
    if re.search(r"diesel|\bglp\b", s): out.add("DIESEL")
    if re.search(r"\biss\b", s) and re.search(r"exclus|base|bc\b", s): out.add("ISS")
    if re.search(r"icms", s):
        if re.search(r"difal", s): out.add("DIFAL")
        elif re.search(r"inclus[ãa]o", s) and re.search(r"cr[ée]dito", s): out.add("ICMS_CRED")
        elif re.search(r"st (pago )?a maior|restitui[çc][ãa]o (do |de )?icms|ressarcimento (do |de )?icms[- ]?st|complemento", s): out.add("ST_REST")
        elif re.search(r"ciap|icms acumulado|cr[ée]dito acumulado de icms", s): out.add("ICMS_CIAP")
        elif re.search(r"exclus|base|\bbc\b|tema 69", s):
            prop = re.search(r"pr[óo]prio", s); spec_st = re.search(r"\(icms[\s-]*st\)", s)
            if prop and not spec_st: out.add("EX_ICMS")
            elif spec_st: out.add("EX_ST")
            elif re.search(r"icms\s*/\s*st", s): out |= {"EX_ICMS", "EX_ST"}
            elif re.search(r"icms[\s-]*st\b|substitui[çc][ãa]o tribut", s): out.add("EX_ST")
            else: out.add("EX_ICMS")
    if re.search(r"tema 201", s): out.add("ST_REST")
    if re.search(r"erro de cst|revis[ãa]o d[eo] pis|\bcst\b|monof[áa]sic|segrega[çc][ãa]o de receitas", s): out.add("CST")
    if re.search(r"sistema s", s): out.add("MSC_S")
    elif re.search(r"verbas|previdenci|\binss\b|\brat\b|\bfap\b|aviso pr[ée]vio|maternidade|hora extra|descanso semanal|indenizat", s): out.add("PREV")
    return out


VAL = re.compile(r"R\$\s?([\d.]*\d(?:,\d{2})?)")


def cobertura_texto(t):
    """{código: status} — 'apurada' (valor > 0), 'analisada sem valor/não aplicada' ou 'citada'."""
    cov = {}; secao = None
    rank = {"citada": 0, "analisada sem valor/não aplicada": 1, "apurada": 2}
    L = [re.sub(r"\s+", " ", l).strip() for l in t.splitlines()]
    L = [l for l in L if l]
    for i, l in enumerate(L):
        ll = l.lower()
        if ll.startswith("teses aplicadas"): secao = "apl"
        elif ll.startswith("teses não aplicadas") or ll.startswith("teses nao aplicadas"): secao = "nao"
        elif ll.startswith(("previdenciário", "ms coletivo")) and secao: pass
        elif ll.startswith(("o que esperar", "litispend", "atividades exercidas", "pagamentos efetuados")): secao = None
        trecho = l if len(l) > 25 or i + 1 >= len(L) else l + " " + L[i + 1]  # rótulo quebrado em duas linhas
        cods = classificar(trecho)
        if not cods: continue
        vals = [float(v.replace(".", "").replace(",", ".")) for v in VAL.findall(trecho) if re.search(r"\d", v)]
        if not vals:  # valor na linha seguinte (PDF quebra o rótulo e o valor)
            for j in (i + 1, i + 2):
                if j < len(L) and re.fullmatch(r"R\$\s?[\d.,]+", L[j]): vals = [float(L[j][2:].strip().replace(".", "").replace(",", ".") or 0)]; break
        neg = re.search(r"insuficien|n[ãa]o (foram|foi) (identificad|encontrad)|inviabiliz|n[ãa]o se mostr|pouco[s]? expressiv|n[ãa]o h[áa]|sem oportunidade|n[ãa]o apresent", ll)
        if secao == "nao" or neg or (vals and max(vals) == 0): st = "analisada sem valor/não aplicada"
        elif (vals and max(vals) > 0) or secao == "apl": st = "apurada"
        else: st = "citada"
        for c in cods:
            if rank[st] >= rank.get(cov.get(c), -1): cov[c] = st
    return cov


def textos_por_negocio():
    T = defaultdict(list)
    if not os.path.isdir(TXT): return T
    for f in os.listdir(TXT):
        if not f.endswith(".txt"): continue
        did, kind = f.split("_", 1)[0], f.split("_", 1)[1].rsplit("_", 1)[0]
        if kind in ("parecer_rt", "anexo_diagnostico"): T[did].append(os.path.join(TXT, f))
    return T


def data_analise(t):
    m = re.search(r"data da an[áa]lise:?\s*(\d{2}/\d{2}/\d{4})", t, re.I)
    return m.group(1) if m else None

# ---------------------------------------------------------------- perfil e teses esperadas
def atividade(cnae):
    d = (cnae or "")[:2]
    if not d.isdigit(): return "?"
    n = int(d)
    return "agro/pesca" if n <= 3 else "indústria" if 5 <= n <= 33 else "atacado" if n == 46 else "varejo" if n == 47 else "comércio" if n == 45 else "serviço"


def esperadas(mz, cnae):
    """Teses atômicas que a matriz espera para a linha do perfil (regras de atividade do Tributário)."""
    if not mz: return []
    tese = mz["tese"].lower(); tok = set(re.findall(r"[A-Za-z]+", mz["variante"].upper())); obs = (mz.get("observacao") or "").lower()
    e = []
    if mz.get("tese_save"): e.append("AQ_SAVE")
    if "C" in tok or "crédito presumido" in tese: e.append("C")
    if "B" in tok or "ressarcimento" in tese: e.append("B")
    if "INSUMOS" in tok or "insumo" in tese: e.append("INSUMOS")
    if "EX" in tok or "exclus" in tese: e.append("EX_ICMS")
    if "EX_ICMS" in e and re.search(r"icms-st|icms e icms-st", tese): e.append("EX_ST")
    if "CST" in tok or "cst" in tese: e.append("CST")
    if "ST" in tok or "restitui" in tese or ("icms-st" in tese and "EX_ICMS" not in e): e.append("ST_REST")
    if "PREV" in tok or "previdenci" in tese: e.append("PREV")
    if "ciap" in tese: e.append("ICMS_CIAP")
    div = (cnae or "")[:2]
    if div in CAT["atividade"]["sem_C"]: e = [x for x in e if x != "C"]
    if div in CAT["atividade"]["sem_INSUMOS"]: e = [x for x in e if x != "INSUMOS"]
    if mz["regime"] == "Presumido" and re.search(r"n[ãa]o (h[áa]|tem o que) exclu|sem exclus|nada a recuperar|n[ãa]o h[áa] exclus", obs):
        e = [x for x in e if x not in ("EX_ICMS", "EX_ST")]  # V12: receita quase toda em alíquota zero
    return list(dict.fromkeys(e))


def secundarios(cnpj):
    if not cnpj or not os.path.exists(DB): return []
    try:
        con = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
        r = con.execute("SELECT cnaes_sec FROM estab WHERE cnpj=?", (cnpj,)).fetchone(); con.close()
        return [c for c in (r[0] or "").split(",") if c] if r else []
    except sqlite3.Error:
        return []


CAFE = CAT["cafe"]


def eh_cafe(d, p):
    nome = (d.get("nome") or "").lower()
    return p["cnae"] in CAFE["cnae"] or any(re.search(rf"\b{re.escape(w)}\b", nome) for w in CAFE["palavras_inteiras"]) or any(r in nome for r in CAFE["radicais"])


SEG_NOME = [("Hortifrúti", r"hortifr|sacol|frutas|verdur|legume"), ("Carnes", r"carne|frigor|a[çc]ougue|abate|bovin|su[íi]n|aves\b"), ("Pescados", r"pescad|peixe"),
            ("Supermercados", r"supermerc|mercearia|minimerc|mercado\b|atacarejo"), ("Laticínios", r"latic|leite|queijo"), ("Padaria", r"padaria|panific")]
SEG_CNAE = {"Hortifrúti": ("4633", "4724"), "Carnes": ("1011", "1012", "1013", "4634", "4722"), "Pescados": ("1020", "4634", "4722", "03"),
            "Supermercados": ("4711", "4712", "4729", "4691", "4639", "4637"), "Laticínios": ("1051", "1052", "4631"), "Padaria": ("1091", "4721")}


def incoerencia_nome(nome, cnae):
    n = (nome or "").lower()
    for s, rx in SEG_NOME:
        if re.search(rx, n) and not cnae.startswith(SEG_CNAE[s]):
            return f"nome sugere {s}, mas o CNAE principal é {cnae or 'vazio'}: conferir o cadastro/atividade real"
    return None

# ---------------------------------------------------------------- planilha anterior e aprendizado
def decisoes_anteriores():
    """(link, código) → decisão registrada na planilha mais recente."""
    fs = sorted(glob.glob(os.path.join(OUT, "Revisor-Diagnostico-*.xlsx")))
    if not fs: return {}, None
    from openpyxl import load_workbook
    try: ws = load_workbook(fs[-1], read_only=True, data_only=True)["Lacunas"]
    except (KeyError, Exception): return {}, fs[-1]
    rows = list(ws.iter_rows(values_only=True)); h = list(rows[0]); out = {}
    for r in rows[1:]:
        g = dict(zip(h, r))
        if g.get("decisão da Geovana") or g.get("motivo da decisão"):
            out[(g.get("link"), g.get("código"))] = {k: g.get(k) for k in ("decisão da Geovana", "motivo da decisão", "minutos no caso", "levado ao técnico em",
                                                                           "segmento (matriz)", "regime", "tese", "negócio")}
    return out, fs[-1]


def acumular(prev):
    f = os.path.join(OUT, "decisoes.json"); D = json.load(open(f)) if os.path.exists(f) else {}
    for (lk, cod), v in prev.items():
        if v.get("decisão da Geovana") in DECISOES:
            D[f"{lk}|{cod}"] = {"link": lk, "codigo": cod, **v, "registrado_em": D.get(f"{lk}|{cod}", {}).get("registrado_em") or date.today().isoformat()}
    if D: json.dump(D, open(f, "w"), ensure_ascii=False, indent=1, default=str)
    hist = defaultdict(Counter)
    for v in D.values(): hist[(v.get("segmento (matriz)"), v.get("regime"), v.get("codigo"))][v.get("decisão da Geovana")] += 1
    return hist

# ---------------------------------------------------------------- principal
def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--etapa", default="Apresentação de Diagnóstico"); ap.add_argument("--ids", help="ids separados por vírgula (teste)")
    ap.add_argument("--sem-fichas", action="store_true"); ap.add_argument("--fichas", type=int, default=30, help="quantos negócios ganham ficha em Markdown (os de mais lacunas altas)")
    a = ap.parse_args()
    fonte = sorted(glob.glob(os.path.join(SRC, "hubspot-*.json")))[-1]
    N = json.load(open(fonte))["negocios"]
    E = json.load(open(ENR))["negocios"] if os.path.exists(ENR) else {}
    RC = json.load(open(REGCACHE)) if os.path.exists(REGCACHE) else {}
    T = textos_por_negocio(); hoje = date.today(); d36 = (hoje - timedelta(days=1095)).isoformat(); d24 = (hoje - timedelta(days=730)).isoformat()
    COV = {}
    def cov(did):
        if did not in COV:
            c = {}
            for f in T.get(did, []):
                for k, v in cobertura_texto(open(f, errors="ignore").read()).items():
                    if {"apurada": 2, "analisada sem valor/não aplicada": 1, "citada": 0}[v] >= {"apurada": 2, "analisada sem valor/não aplicada": 1, "citada": 0}.get(c.get(k), -1): c[k] = v
            COV[did] = c
        return COV[did]

    # 1) cobertura típica: em quantos pareceres legíveis de cada perfil (linha da matriz) o técnico analisou a tese
    tip = defaultdict(Counter); tip_n = Counter()
    for d in N:
        if d["id"] not in T or (d["criado"] or "") < d24: continue  # prática atual do técnico: últimos 24 meses
        p = perfil(d, E); mz = linha_matriz(p["cnae"], p["regime"])
        if not mz: continue
        k = (mz["segmento"], mz["regime"]); c = cov(d["id"])
        if not c: continue
        tip_n[k] += 1
        for cod in c: tip[k][cod] += 1

    # 2) referência do look-alike: ganhos (e fechados) dos últimos 36 meses, fora café; teses de cada ganho = campo do HubSpot + texto
    ref = []
    for d in N:
        if d["status"] in ("ganho", "perdido") and (d["fechado"] or "") >= d36:
            p = perfil(d, E)
            if eh_cafe(d, p): continue
            tes = set()
            if d["status"] == "ganho":
                for lab in (d.get("variante") or "").split(";"): tes |= classificar(lab)
                tes |= {k for k, v in cov(d["id"]).items() if v == "apurada"}
            ref.append((d, p, tes))
    por_raiz = defaultdict(list)
    for d in N: por_raiz[raiz(d, E)].append(d)

    prev, prev_f = decisoes_anteriores(); hist = acumular(prev)

    # 3) seleção
    alvo = [d for d in N if d.get("etapa") == a.etapa and d["status"] == "aberto"]
    if a.ids: alvo = [d for d in N if d["id"] in set(a.ids.split(","))]
    revis, lacunas, fora = [], [], []
    vistos = {}
    for d in sorted(alvo, key=lambda x: x["criado"] or ""):
        p = perfil(d, E); e = E.get(d["id"], {}); cnpj = normalizar_cnpj(e.get("cnpj") or d["extras"].get("cnpj")).zfill(14)
        rz = cnpj[:8]
        base = {"negócio": d["nome"], "parceiro": d["parceiro"], "pipeline": d["pipeline"], "etapa": d.get("etapa"), "criado": d["criado"],
                "regime": p["regime"] or "não informado", "cnae_oficial": f"{p['cnae']} {p['cnae_desc'] or ''}".strip(), "uf": p["uf"], "link": LINK.format(d["id"])}
        motivo = None
        if "passivo" in (d.get("pipeline") or "").lower() or "(passivo" in (d.get("nome") or "").lower(): motivo = "V4 Gestão de Passivo: outro produto"
        elif p["regime"].startswith("Simples") or "simples" in (d.get("pipeline") or "").lower(): motivo = "V2 Simples Nacional: fora da matriz de crédito não cumulativo (se houver produto Simples, é outra fila)"
        elif p["situacao"] and p["situacao"] not in ("Ativa", "Suspensa"): motivo = f"V1 CNPJ {p['situacao'].lower()} na Receita"
        elif eh_cafe(d, p): motivo = "V3 Café: fora da meta desde abr/2026"
        if motivo: fora.append({"motivo": motivo, **base}); continue
        if rz in vistos and rz != "00000000":  # V6: mesma raiz já na revisão → um registro por grupo
            vistos[rz]["outros negócios do grupo em diagnóstico"] = (vistos[rz].get("outros negócios do grupo em diagnóstico") or "") + f"{d['nome'][:40]} ({LINK.format(d['id'])}) · "
            fora.append({"motivo": "V6 duplicidade: mesma raiz de CNPJ já está na revisão", **base}); continue
        # perfil oficial e alertas
        alertas = []
        reg = p["regime"]
        if p["simples"] == "S" and not reg.startswith("Simples"): alertas.append(f"Receita indica optante do Simples, HubSpot diz {reg}: conferir o regime antes de tudo")
        rc = RC.get(cnpj)
        if rc and reg and rc.split()[0:2] != reg.upper().split()[0:2]: alertas.append(f"regime declarado mais recente (base pública): {rc} × HubSpot {reg}")
        if not reg: alertas.append("regime não informado no HubSpot")
        if p["situacao"] == "Suspensa": alertas.append("CNPJ suspenso: pendência cadastral (precisa de dado)")
        inc = incoerencia_nome(d["nome"], p["cnae"])
        if inc: alertas.append(inc)
        ini = e.get("inicio_atividade")
        if ini and len(ini) == 8:
            meses = (hoje.year - int(ini[:4])) * 12 + hoje.month - int(ini[4:6])
            if meses < 12 and p["porte"] == "Microempresa": alertas.append("V15: empresa com menos de 12 meses e porte ME: base pequena, reavaliar em 6 meses")
        if not p["cnae"]: alertas.append("sem CNAE oficial (CNPJ não achado na base local da Receita)")
        mz = linha_matriz(p["cnae"], reg)
        esp = esperadas(mz, p["cnae"])
        sec = secundarios(cnpj); extra = []
        for sc in sec:  # CNAE secundário de abate/pescado/laticínio abre o crédito presumido (C) que o principal (comércio) não abre
            if sc.startswith(("1011", "1012", "1013", "1020", "1051", "1052")) and p["cnae"][:2] in ("45", "46", "47") and "C" not in esp:
                if "C" not in [x[0] for x in extra]: extra.append(("C", sc))
        # cobertura deste negócio
        c_neg = cov(d["id"]); legivel = bool(T.get(d["id"])) and bool(c_neg)
        datas = [data_analise(open(f, errors="ignore").read()) for f in T.get(d["id"], [])]
        # histórico do grupo: outras análises da mesma raiz com texto legível
        hist_g = {}; hist_ref = []
        for o in por_raiz.get(rz, []) if rz != "00000000" else []:
            if o["id"] == d["id"] or not T.get(o["id"]): continue
            for k, v in cov(o["id"]).items():
                if k not in hist_g or v == "apurada": hist_g[k] = v
            hist_ref.append(f"{o['nome'][:30]} ({o['status']}, {o['criado']})")
        ganhos_grupo = [o for o in por_raiz.get(rz, []) if o["status"] == "ganho" and rz != "00000000"]
        tes_grupo = set()
        for o in ganhos_grupo:
            for lab in (o.get("variante") or "").split(";"): tes_grupo |= classificar(lab)
        # look-alike (mesma atividade, nunca a própria raiz, uma empresa conta uma vez)
        viz, vr = [], set()
        for s_, dd, tes in sorted(((parecido(p, q), dd, tes) for dd, q, tes in ref if raiz(dd, E) != rz and atividade(q["cnae"]) == atividade(p["cnae"])), key=lambda x: -x[0]):
            r_ = raiz(dd, E)
            if s_ < 8 or (r_ is not None and r_ in vr): continue  # sem CNPJ: cada negócio conta por si
            vr.add(r_); viz.append((dd, tes))
            if len(viz) >= 40: break
        vg = [(dd, tes) for dd, tes in viz if dd["status"] == "ganho" and dd["valor"] > 0]
        teses_viz = Counter(c for _, tes in vg for c in tes)
        k_tip = (mz["segmento"], mz["regime"]) if mz else None
        ncas = len(vg)
        # lacunas
        linhas_neg = []
        for cod, origem in [(c, "matriz") for c in esp] + [(c, f"CNAE secundário {sc}") for c, sc in extra]:
            st = c_neg.get(cod)
            if legivel and st: continue  # coberta
            tip_pct = (tip[k_tip][cod] / tip_n[k_tip]) if (k_tip and tip_n[k_tip] >= 5) else None
            if legivel: status = "lacuna confirmada (não citada no parecer)"
            elif tip_pct is not None and tip_pct < PRI["cobertura_tipica_limite_provavel"]: status = "lacuna provável (o técnico costuma não analisar neste perfil)"
            else: status = "a conferir (parecer ainda não legível)"
            if not legivel and tip_pct is not None and tip_pct >= PRI["cobertura_tipica_limite_provavel"] and origem == "matriz":
                continue  # normalmente coberta pelo modelo padrão: fica só na lista de conferência da aba Revisão
            ader = {"alta": 3, "média": 2, "baixa": 1}.get(mz["aderencia"] if mz else "", 1)
            pts = ader; motivos = [f"aderência {mz['aderencia'] if mz else '?'} na matriz"]
            kv = teses_viz.get(cod, 0)
            if kv >= 1: pts += 1; motivos.append(f"{kv} de {ncas} ganhos parecidos tiveram")
            if ncas >= 5 and kv / ncas >= 0.2: pts += 1
            if tip_pct is not None and tip_pct < 0.3: pts += 1; motivos.append(f"coberta em só {tip_pct*100:.0f}% dos pareceres do perfil")
            if TESES[cod]["tipo"] == "saldo credor" and hoje <= CORTE: pts += 1; motivos.append("saldo credor: escriturar até 31/12/2026")
            if origem != "matriz": pts -= 1
            if hist_g.get(cod) == "analisada sem valor/não aplicada": pts -= 2; motivos.append("análise anterior do grupo já olhou esta tese e não achou valor: ver o motivo antes de pedir")
            elif hist_g.get(cod) == "apurada": motivos.append("análise anterior do grupo APUROU esta tese: conferir por que não está neste diagnóstico")
            if cod in tes_grupo: pts -= 1; motivos.append("tese já feita no ganho do grupo: ver se é recorrência")
            hh = hist.get((mz["segmento"] if mz else None, reg, cod), Counter())
            if hh.get("Não se aplica", 0) >= 3 and not hh.get("Incluir tese"): pts -= 3; motivos.append(f"Geovana marcou 'Não se aplica' {hh['Não se aplica']}x neste perfil")
            if hh.get("Incluir tese"): motivos.append(f"Geovana já incluiu esta tese {hh['Incluir tese']}x neste perfil")
            prio = "alta" if pts >= PRI["alta"] else "média" if pts >= PRI["media"] else "baixa"
            if origem != "matriz" and prio == "alta": prio = "média"  # CNAE secundário: confirmar se a atividade existe de fato
            ex = [f"{dd['nome'][:35]} ({dd['parceiro']})" for dd, tes in vg if cod in tes][:3]
            lk = LINK.format(d["id"]); pv = prev.get((lk, cod), {})
            linhas_neg.append({"prioridade": prio, "pontos": pts, "negócio": d["nome"], "parceiro": d["parceiro"], "regime": reg, "segmento (matriz)": mz["segmento"] if mz else "sem linha na matriz",
                               "código": cod, "tese": TESES[cod]["nome"], "origem": origem, "status": status,
                               "por que": " · ".join(motivos), "cobertura típica no perfil": f"{tip_pct*100:.0f}% de {tip_n[k_tip]}" if tip_pct is not None else "amostra pequena",
                               "ganhos parecidos com a tese": f"{kv} de {ncas}" if ncas else "sem ganhos parecidos", "exemplos (look-alike)": " · ".join(ex),
                               "analisada antes no grupo": hist_g.get(cod) or ("feita no ganho do grupo" if cod in tes_grupo else ""),
                               "base legal": TESES[cod]["base_legal"], "grau de risco": TESES[cod]["risco"], "procedimento": TESES[cod]["procedimento"],
                               "tipo": TESES[cod]["tipo"], "documento que decide": TESES[cod]["documentos"],
                               "decisão da Geovana": pv.get("decisão da Geovana") or "", "motivo da decisão": pv.get("motivo da decisão") or "",
                               "minutos no caso": pv.get("minutos no caso") or "", "levado ao técnico em": pv.get("levado ao técnico em") or "", "link": lk})
        ordem = {"alta": 0, "média": 1, "baixa": 2}
        linhas_neg.sort(key=lambda r: (ordem[r["prioridade"]], -r["pontos"]))
        lacunas += linhas_neg
        conferir = [TESES[c]["nome"].split(" (")[0] for c in esp if c not in [l["código"] for l in linhas_neg] and not (legivel and c_neg.get(c))]
        outras_viz = [f"{CAT['outras_no_catalogo_save'].get(c, TESES.get(c, {}).get('nome', c)).split(' (')[0]} ({n})" for c, n in teses_viz.most_common() if c not in esp][:4]
        r = {**base, "segmento (matriz)": mz["segmento"] if mz else "sem linha na matriz", "aderência": mz["aderencia"] if mz else "a verificar",
             "atividade": atividade(p["cnae"]), "cnaes secundários": ", ".join(sec[:8]), "porte_receita": p["porte"], "capital_social": p["capital"] or None,
             "situação": p["situacao"], "valor no HubSpot": round(d["valor"]) or None,
             "parecer legível": "sim" if legivel else ("não: arquivo no HubSpot (aguarda escopo files)" if any((d["extras"].get(k) or "").strip().isdigit() for k in ("parecer_rt", "anexo_diagnostico")) else "não: sem texto extraído"),
             "data da análise": next((x for x in datas if x), None),
             "teses esperadas (matriz)": ", ".join(esp) or "—",
             "teses cobertas no parecer": ", ".join(f"{k} ({v})" for k, v in c_neg.items() if k in TESES) if legivel else "não verificável",
             "lacunas": len(linhas_neg), "lacunas alta": sum(l["prioridade"] == "alta" for l in linhas_neg),
             "principais lacunas": " · ".join(f"{l['código']} [{l['prioridade']}]" for l in linhas_neg[:4]),
             "conferir no diagnóstico (normalmente coberta)": ", ".join(conferir),
             "ganhos parecidos": ncas, "o que os parecidos tiveram além da matriz": " · ".join(outras_viz),
             "grupo já é cliente": f"sim ({len(ganhos_grupo)} ganho(s); teses: {', '.join(sorted(tes_grupo)) or 'não registradas'})" if ganhos_grupo else "não",
             "análises anteriores do grupo": " · ".join(hist_ref[:3]), "alertas de perfil": " | ".join(alertas),
             "hipótese principal (matriz)": (f"Tese Save: {mz['tese_save']}. Alternativa conservadora: {mz['tese']} [{mz['variante']}]" if mz and mz.get("tese_save") else (f"{mz['tese']} [{mz['variante']}]" if mz else "regime/CNAE sem linha na matriz: conferir")),
             "observação da matriz": mz["observacao"] if mz else "", "outros negócios do grupo em diagnóstico": ""}
        vistos[rz] = r; revis.append(r)
    revis.sort(key=lambda r: (-r["lacunas alta"], -r["lacunas"]))
    for i, r in enumerate(revis, 1): r["ordem"] = i

    # ---------------------------------------------------------------- saída
    os.makedirs(OUT, exist_ok=True)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); H1 = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F1F1F")
    largos = {"negócio", "cnae_oficial", "tese", "por que", "base legal", "grau de risco", "procedimento", "documento que decide", "exemplos (look-alike)",
              "principais lacunas", "conferir no diagnóstico (normalmente coberta)", "o que os parecidos tiveram além da matriz", "alertas de perfil",
              "hipótese principal (matriz)", "observação da matriz", "teses cobertas no parecer", "análises anteriores do grupo", "motivo", "status", "motivo da decisão",
              "outros negócios do grupo em diagnóstico", "grupo já é cliente"}
    def aba(ws, L, cols):
        ws.append(cols)
        for j, c in enumerate(cols, 1):
            x = ws.cell(row=1, column=j); x.font = H1; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[x.column_letter].width = 44 if c in largos else 15
        for i, l in enumerate(L, 2):
            ws.append([l.get(c) for c in cols])
            if "link" in cols: x = ws.cell(row=i, column=cols.index("link") + 1); x.hyperlink = l["link"]; x.font = Font(color="1F5FBF", underline="single")
        ws.freeze_panes = "C2"
    c_rev = ["ordem", "negócio", "parceiro", "regime", "segmento (matriz)", "aderência", "atividade", "cnae_oficial", "cnaes secundários", "uf", "porte_receita", "capital_social",
             "situação", "criado", "valor no HubSpot", "parecer legível", "data da análise", "teses esperadas (matriz)", "teses cobertas no parecer", "lacunas", "lacunas alta",
             "principais lacunas", "conferir no diagnóstico (normalmente coberta)", "ganhos parecidos", "o que os parecidos tiveram além da matriz", "grupo já é cliente",
             "análises anteriores do grupo", "outros negócios do grupo em diagnóstico", "alertas de perfil", "hipótese principal (matriz)", "observação da matriz", "link"]
    ws = wb.active; ws.title = "Revisão"; aba(ws, revis, c_rev)
    c_lac = ["prioridade", "pontos", "negócio", "parceiro", "regime", "segmento (matriz)", "código", "tese", "origem", "status", "por que", "cobertura típica no perfil",
             "ganhos parecidos com a tese", "exemplos (look-alike)", "analisada antes no grupo", "base legal", "grau de risco", "procedimento", "tipo", "documento que decide",
             "decisão da Geovana", "motivo da decisão", "minutos no caso", "levado ao técnico em", "link"]
    wl = wb.create_sheet("Lacunas"); aba(wl, lacunas, c_lac)
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISOES) + '"', allow_blank=True); wl.add_data_validation(dv)
    col = wl.cell(row=1, column=c_lac.index("decisão da Geovana") + 1).column_letter; dv.add(f"{col}2:{col}{max(2, len(lacunas) + 1)}")
    # resumo por tese
    wt = wb.create_sheet("Lacunas por tese"); wt.append(["código", "tese", "lacunas", "alta", "média", "baixa", "confirmadas", "prováveis", "a conferir", "base legal", "grau de risco"])
    for j in range(1, 12): wt.cell(row=1, column=j).font = H1; wt.cell(row=1, column=j).fill = HF
    porT = defaultdict(Counter)
    for l in lacunas:
        porT[l["código"]][l["prioridade"]] += 1; porT[l["código"]]["n"] += 1; porT[l["código"]][l["status"].split(" (")[0]] += 1
    for cod, cn in sorted(porT.items(), key=lambda x: -x[1]["n"]):
        wt.append([cod, TESES[cod]["nome"], cn["n"], cn["alta"], cn["média"], cn["baixa"], cn["lacuna confirmada"], cn["lacuna provável"], cn["a conferir"], TESES[cod]["base_legal"], TESES[cod]["risco"]])
    wt.column_dimensions["B"].width = 60; wt.column_dimensions["J"].width = 60; wt.column_dimensions["K"].width = 40
    aba(wb.create_sheet("Fora da revisão"), fora, ["motivo", "negócio", "parceiro", "pipeline", "regime", "cnae_oficial", "uf", "criado", "link"])
    wc = wb.create_sheet("Cobertura típica"); cods = list(TESES)
    wc.append(["segmento (matriz)", "regime", "pareceres legíveis"] + cods)
    for j in range(1, len(cods) + 4): wc.cell(row=1, column=j).font = H1; wc.cell(row=1, column=j).fill = HF
    for k, n in sorted(tip_n.items(), key=lambda x: -x[1]):
        wc.append([k[0], k[1], n] + [f"{tip[k][c]/n*100:.0f}%" for c in cods])
    wc.column_dimensions["A"].width = 36
    como = wb.create_sheet("Como usar"); como.column_dimensions["A"].width = 150
    for l in [
        "REVISOR DE DIAGNÓSTICO — pré-análise para a Geovana (N1). Nada daqui vai ao parceiro, ao cliente ou ao técnico sem a validação dela.",
        f"Base: HubSpot {os.path.basename(fonte)} · Receita (base local) · pareceres já extraídos em textos/ · gerado em {hoje.isoformat()}.",
        "",
        "1. Aba Revisão: um negócio por linha (uma empresa por raiz de CNPJ), ordenado pelas lacunas de prioridade alta. Leia os 'alertas de perfil' antes de tudo (regime divergente, CNAE incoerente com o nome).",
        "2. 'Teses esperadas' = o que a matriz tese × perfil (Depto Tributário, 29/09/2026) espera para o CNAE e o regime. 'Teses cobertas' = o que o texto do parecer/diagnóstico cita.",
        "3. Aba Lacunas: tese esperada e não analisada. Status: 'confirmada' (parecer lido, tese não citada) · 'provável' (parecer ainda não legível e o técnico costuma não analisar essa tese neste perfil, ver aba Cobertura típica) · 'a conferir' (sem histórico suficiente).",
        "   Teses que o modelo padrão do técnico quase sempre cobre ficam só na coluna 'conferir no diagnóstico' da aba Revisão, para não gerar ruído.",
        "4. Para cada lacuna, marque a 'decisão da Geovana': Incluir tese (vale pedir ao técnico) · Não se aplica (explique o motivo: vira regra) · Precisa de dado (qual documento).",
        "   Preencha os minutos gastos. Na próxima rodada, o script copia as decisões (pelo link + código da tese) e acumula em decisoes.json; 'Não se aplica' 3x no mesmo perfil rebaixa a tese.",
        "5. Look-alike: ganhos parecidos = mesma atividade (indústria/atacado/varejo), mesma classe ou CNAE e mesmo regime, nunca a própria empresa (raiz do CNPJ). Mostra CONTAGEM de casos, nunca valor.",
        "6. Regras fixas: fora Simples (V2), Gestão de Passivo (V4), café (V3, fora da meta desde abr/2026), CNPJ baixado/inapto (V1); uma linha por grupo (V6). C só para indústria/produção; insumos nunca no comércio.",
        "   Tese de alíquota zero (entendimento Save) conta como hipótese válida nos perfis de alimentos em Lucro Real (D22), sempre via PER, nunca DCOMP. Vetos V7–V10 (D/A como principal, monofásicos, ICMS-ST pago, Temas 1008/985/1079) nunca aparecem como lacuna.",
        "7. Toda tese traz base legal e grau de risco (catálogo em ~/.claude/skills/revisor-diagnostico/references/teses.json). Itens marcados 'confirmar' foram citados de memória técnica pelo Tributário: validar antes de usar fora da célula.",
        "LIMITAÇÃO ATUAL: os pareceres dos negócios em apresentação de diagnóstico estão como arquivo do HubSpot (exigem o escopo 'files'); até lá a cobertura é inferida pela cobertura típica do perfil.",
    ]: como.append([l])
    f = os.path.join(OUT, f"Revisor-Diagnostico-{hoje.isoformat()}.xlsx"); wb.save(f)
    resumo = {"gerado_em": hoje.isoformat(), "fonte": os.path.basename(fonte), "etapa": a.etapa, "selecionados": len(alvo), "revisados": len(revis), "fora": Counter(x["motivo"].split(":")[0] for x in fora),
              "parecer_legivel": sum(r["parecer legível"] == "sim" for r in revis), "lacunas": len(lacunas), "por_tese": {k: dict(v) for k, v in porT.items()},
              "por_prioridade": Counter(l["prioridade"] for l in lacunas), "por_status": Counter(l["status"].split(" (")[0] for l in lacunas),
              "decisoes_copiadas": sum(1 for l in lacunas if l["decisão da Geovana"]), "planilha_anterior": prev_f}
    json.dump({"resumo": resumo, "revisao": revis, "lacunas": lacunas}, open(os.path.join(OUT, f"revisor-{hoje.isoformat()}.json"), "w"), ensure_ascii=False, indent=1, default=str)
    if not a.sem_fichas: fichas(revis[:a.fichas], lacunas, hoje)
    print(f, "|", json.dumps(resumo, ensure_ascii=False, default=str))


def fichas(revis, lacunas, hoje):
    porN = defaultdict(list)
    for l in lacunas: porN[l["link"]].append(l)
    L = [f"# Revisor de Diagnóstico — fichas de {hoje.strftime('%d/%m/%Y')}", "",
         "> PRÉ-ANÁLISE para a Geovana validar (N1). Nada vai ao parceiro, ao cliente ou ao técnico sem ela. Sem valores: o look-alike é contagem de casos.", ""]
    for r in [r for r in revis if r["lacunas"]]:
        L += [f"## {r['ordem']}. {r['negócio']} — {r['parceiro']}", "",
              f"- **Perfil oficial:** {r['cnae_oficial']} · {r['atividade']} · {r['regime']} · {r['uf']} · porte {r['porte_receita']} · situação {r['situação']}" + (f" · secundários {r['cnaes secundários']}" if r["cnaes secundários"] else ""),
              f"- **Matriz:** {r['segmento (matriz)']} (aderência {r['aderência']}) · esperadas: {r['teses esperadas (matriz)']}",
              f"- **Parecer:** {r['parecer legível']}" + (f" · análise de {r['data da análise']}" if r["data da análise"] else "") + f" · cobertas: {r['teses cobertas no parecer']}"]
        if r["alertas de perfil"]: L.append(f"- **Alertas:** {r['alertas de perfil']}")
        if r["grupo já é cliente"] != "não": L.append(f"- **Grupo já é cliente:** {r['grupo já é cliente']}")
        L += ["", "| Prioridade | Tese | Status | Por que | Base legal | Risco | Documento que decide |", "|---|---|---|---|---|---|---|"]
        for l in porN[r["link"]]:
            L.append(f"| {l['prioridade']} | {l['tese']} ({l['origem']}) | {l['status']} | {l['por que']}" + (f"; ex.: {l['exemplos (look-alike)']}" if l["exemplos (look-alike)"] else "") +
                     f" | {l['base legal']} | {l['grau de risco']} | {l['documento que decide']} |")
        if r["conferir no diagnóstico (normalmente coberta)"]: L.append(f"\nConferir também (normalmente coberta): {r['conferir no diagnóstico (normalmente coberta)']}.")
        L += ["", f"**Decisão da Geovana** (na planilha, aba Lacunas): Incluir tese / Não se aplica / Precisa de dado · [card]({r['link']})", ""]
    os.makedirs(os.path.join(OUT, "fichas"), exist_ok=True)
    open(os.path.join(OUT, "fichas", f"{hoje.isoformat()}.md"), "w").write("\n".join(L))


if __name__ == "__main__":
    main()
