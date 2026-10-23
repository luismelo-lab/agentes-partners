#!/usr/bin/env python3
"""Folha curta de decisão do lote do Recuperador: só os casos do lote, só as colunas que a Geovana precisa para decidir.

A fila completa tem ~300 linhas e 32 colunas; o lote 1 foi reenviado duas vezes sem nenhuma decisão (19/10 e 22/10/2026).
Esta folha vai para a pasta da Geovana no OneDrive com o nome Fila-Recuperador-Lote<N>-decidir-<data>.xlsx e a aba "Fila",
para que fila.py (decisoes_anteriores) leia as decisões marcadas aqui na próxima rodada, sem passo manual.

  python3 lote_decidir.py [--lote 1] [--fila <Fila-Recuperador-AAAA-MM-DD.xlsx>]
"""
import argparse, glob, os, sys
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

AP = os.path.expanduser("~/agentes-partners"); sys.path.insert(0, os.path.join(AP, "lib"))
import caminhos, marca
C = {k: v.lstrip("#") for k, v in marca.C.items() if not k.startswith("_")}; F = "IBM Plex Sans"; FM = "IBM Plex Mono"
COLS = [("prioridade", 9), ("negócio", 34), ("parceiro", 14), ("honorário esperado (look-alike)", 15), ("saldo credor (urgente até 30/11)", 14),
        ("sinal do parecer", 22), ("hipótese preliminar", 46), ("documentos a pedir", 34),
        ("decisão da Geovana", 20), ("motivo da decisão", 32), ("valor validado (R$)", 14), ("minutos no caso", 10), ("link", 14)]
EDITA = {"decisão da Geovana", "motivo da decisão", "valor validado (R$)", "minutos no caso"}


def main():
    a = argparse.ArgumentParser(); a.add_argument("--lote", type=int, default=1); a.add_argument("--fila"); a.add_argument("--saida"); a = a.parse_args()
    filas = sorted(glob.glob(os.path.join(os.environ.get("RECUPERADOR_OUT") or os.path.join(AP, "celulas/geovana/recuperador"), "Fila-Recuperador-????-??-??.xlsx")))
    fila = a.fila or (filas[-1] if filas else sys.exit("sem Fila-Recuperador-AAAA-MM-DD.xlsx: rode fila.py antes"))
    rows = list(openpyxl.load_workbook(fila, read_only=True, data_only=True)["Fila"].iter_rows(values_only=True)); h = list(rows[0])
    L = [dict(zip(h, r)) for r in rows[1:]]; L = [g for g in L if str(g.get("lote") or "") == str(a.lote)]
    L.sort(key=lambda g: g.get("prioridade") or 999)
    # decisões já marcadas em qualquer folha anterior (inclusive esta mesma, se for gerada de novo no dia) não se perdem:
    # antes, rodar de novo a partir de uma fila mais antiga sobrescrevia o que a Geovana já tinha marcado na folha curta.
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); from fila import decisoes_anteriores
    ant = decisoes_anteriores()
    for g in L:
        for k, v in (ant.get(g.get("link")) or {}).items():
            if g.get(k) in (None, ""): g[k] = v
    saida = a.saida or os.path.join(caminhos.pasta("geovana"), f"Fila-Recuperador-Lote{a.lote}-decidir-{date.today()}.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Fila"; ws.sheet_view.showGridLines = False
    for j, (t, w) in enumerate(COLS, 1):
        c = ws.cell(1, j, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor=C["ouro"] if t in EDITA else C["ouro_suave"]); c.border = Border(bottom=Side(style="medium", color=C["ouro"]))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 34; fino = Border(bottom=Side(style="thin", color=C["borda"]))
    for i, g in enumerate(L, 2):
        for j, (t, _) in enumerate(COLS, 1):
            v = g.get(t); c = ws.cell(i, j, v if t != "link" else "Abrir"); c.border = fino
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name=FM if "R$" in t or "honorário" in t else F, size=10, color=C["texto"])
            if t == "link" and v: c.value = v; c.hyperlink = v; c.font = Font(name=F, size=10, color=C["ouro_texto"], underline="single")
            if "honorário" in t or "R$" in t: c.number_format = "#,##0"
        ws.row_dimensions[i].height = 75
    fim = len(L) + 1
    dv = DataValidation(type="list", formula1='"Aprovar para reanálise,Vetar,Precisa de dado"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"I2:I{max(2, fim)}"); ws.freeze_panes = "C2"
    u = wb.create_sheet("Como usar"); u.column_dimensions["A"].width = 110; u.sheet_view.showGridLines = False
    for i, t in enumerate([f"Lote {a.lote} · {len(L)} casos · cerca de 10 minutos por caso",
            "1. Em 'decisão da Geovana', escolha: Aprovar para reanálise (vai ao técnico), Vetar (diga o motivo: vira regra) ou Precisa de dado (qual documento).",
            "2. Se aprovar, anote em 'minutos no caso' quanto tempo levou. É o indicador do piloto (horas por caso).",
            "3. Salve o arquivo nesta pasta. Na segunda, a fila nova lê as suas decisões sozinha; não precisa copiar para outra planilha.",
            "4. A ficha completa de cada caso está em 'Fichas do Recuperador' (Word do lote)."], 1):
        c = u.cell(i, 1, t); c.font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11, color=C["texto"]); c.alignment = Alignment(wrap_text=True)
    wb.save(saida); print("ok", saida, len(L), "casos")


if __name__ == "__main__":
    main()
