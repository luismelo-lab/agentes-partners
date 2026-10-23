#!/usr/bin/env python3
"""Candidatos à transação de pequeno valor da Receita Federal (Edital de Transação RFB nº 10/2026), para a Daniela confirmar com o Master.

O edital (DOU de 15/07/2026, seção 3, p. 77; página oficial gov.br/receitafederal › Transação Tributária › Pequeno Valor,
conferida em 23/10/2026):
  - quem adere: pessoa natural, MEI, empresário individual, ME e EPP (item 1.1);
  - o quê: débito em contencioso administrativo fiscal da RFB (impugnação, manifestação de inconformidade ou recurso com efeito
    suspensivo) ou ainda no prazo de impugnar, até 60 salários mínimos POR PROCESSO (R$ 97.260,00 em 2026) (itens 1.1, 2.1 e 2.3);
  - fora: débitos do Simples Nacional, salvo multa por atraso de obrigação acessória (item 2.2); dívida ativa da PGFN não é objeto
    deste edital (a PGFN tem o seu: Edital PGDAU 6/2026, pequeno valor);
  - exige desistência da discussão e confissão irrevogável (itens 3.1 e 3.2); todo o processo entra, sem inclusão parcial (3.6);
  - condições: redução de 50% em até 12 prestações, 40% em 24, 35% em 36, 30% em até 55; prestação mínima de R$ 200 (6.1 e 6.3);
  - prazo: 30/10/2026, 20h59min59s, horário de Brasília, no Portal de Serviços da RFB (4.1).

Fontes cruzadas (todas locais; só leitura, nada vai ao HubSpot):
  1. HubSpot (download mais recente, todos os funis, ganhos, abertos e perdidos): funil "Gestão de Passivo".
  2. Fila do Recuperador (a mais recente): aba "Gestão de Passivo" (veto V4).
  3. Pareceres do técnico já lidos pelo Recuperador (celulas/geovana/pareceres/resumos.json): campos "passivo" e "perdcomp".
  4. Porte, Simples e MEI oficiais da Receita (dados-publicos/receita/cnpj.sqlite), pela raiz do CNPJ.
Critério: porte ME/EPP (ou MEI) E pelo menos um sinal de débito ou processo. A lista é de HIPÓTESES: o HubSpot não diz se o débito
está em contencioso administrativo nem se o processo cabe em 60 salários mínimos. Quem confirma é a Daniela, com o Master.
D76: candidato sem Master do roster vai para uma aba à parte.

  python3 transacao_pequeno_valor.py [--saida <arquivo.xlsx>] [--texto <arquivo.txt>]
  padrão: Transacao-Pequeno-Valor-candidatos-<hoje>.xlsx na pasta da Daniela (OneDrive)
"""
import argparse, glob, json, os, re, sqlite3, sys
from collections import defaultdict
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

AP = os.path.expanduser("~/agentes-partners"); sys.path.insert(0, os.path.join(AP, "lib"))
import caminhos, marca  # noqa: E402
from cnpj import completar_cnpj, raiz_cnpj  # noqa: E402

LINK = "https://app.hubspot.com/contacts/47865021/deal/{}"
HS = os.path.join(AP, "departamentos/resultados/dados")
FILA = os.path.join(AP, "celulas/geovana/recuperador")
RES = os.path.join(AP, "celulas/geovana/pareceres/resumos.json")
DB = os.path.join(AP, "dados-publicos/receita/cnpj.sqlite")
MAPA = os.path.join(AP, "departamentos/resultados/mapeamento-hubspot.json")
LIMITE = 97260.00  # 60 salários mínimos de 2026, valor citado pela RFB na página do edital
PRAZO = "30/10/2026, 20h59 (horário de Brasília)"
EDITAL = "Edital de Transação RFB nº 10/2026 (DOU 15/07/2026); página oficial da RFB conferida em 23/10/2026"
C = {k: v.lstrip("#") for k, v in marca.C.items() if not k.startswith("_")}; F = "IBM Plex Sans"; FM = "IBM Plex Mono"
FORCA = {"Alta": 0, "Média": 1, "Baixa": 2}

TEXTO_MASTER = """Olá, {master}. Tudo bem?

A Receita Federal abriu uma negociação por adesão para débitos federais de pequeno valor que estão em discussão administrativa (impugnação, manifestação de inconformidade ou recurso) ou ainda no prazo de impugnar. Vale para microempresa, empresa de pequeno porte, MEI e pessoa física, com processo de até 60 salários mínimos (R$ 97.260,00 em 2026). É o Edital de Transação RFB nº 10/2026, e o prazo de adesão termina em 30/10/2026, às 20h59 (horário de Brasília).

No cadastro de {empresa} encontramos um sinal de débito federal que pode se enquadrar. Ainda não sabemos se cabe: depende do tipo do débito, de ele estar em discussão na Receita e do valor de cada processo. Pode haver redução prevista no edital, sujeita à análise do processo e ao deferimento da adesão pela Receita. A adesão também exige desistir da discussão e confessar o débito, por isso a decisão é do cliente, com o advogado dele.

Você consegue confirmar com o cliente, até 27/10, se existe processo administrativo aberto na Receita (o número do processo já ajuda) ou nos passar a procuração no e-CAC? Com isso o nosso time jurídico faz a pré-análise ainda dentro do prazo.

Obrigado!"""


def fmt_cnpj(c):
    c = completar_cnpj(c)
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else c


def sinal_passivo(txt):
    """Lê o campo 'Gestão de passivo' do parecer. Devolve (sinal, força, onde_cabe) ou None."""
    t = re.sub(r"\s+", " ", str(txt or "")).strip(); b = re.sub(r"(?<=\d)\.(?=\d)", "", t.lower())
    if not t: return None
    virada = re.search(r"por[ée]m|, mas |entretanto|contudo", b)
    neg = re.match(r"^(a )?(empresas? )?(n[ãa]o (possui|consta|foi|tem)|sem d[íi]vida)", b) and not virada
    if neg or re.search(r"^n[ãa]o, (a )?empresa n[ãa]o possui", b): return None
    RFB = r"(receita|rfb)"
    if re.search(r"exigibilidade suspensa[^.]{0,60}?" + RFB, b) or (re.search(r"exigibilidade suspensa", b) and not re.search(r"suspensa[^.]{0,20}pgfn", b)):
        return ("Parecer: débito com exigibilidade suspensa na Receita (pode ser impugnação, manifestação ou recurso; ou parcelamento)", "Alta", "RFB 10")
    if re.search(r"processo (fiscal|administrativo)[^.]{0,40}" + RFB + r"|pend[êe]ncia de processo", b):
        return ("Parecer: processo fiscal pendente na Receita", "Alta", "RFB 10")
    rfb = re.search(r"(d[ée]bitos?|d[íi]vidas?)[^.]{0,90}?(na|junto [àa]|junto a) " + RFB + r"|d[ée]bitos? rfb", b) and not re.search(r"n[ãa]o possui d[ée]bitos?[^.]{0,30}(na|em aberto na) " + RFB, b)
    pgfn = (re.search(r"(d[ée]bitos?|d[íi]vidas?|valor)[^.]{0,90}?(na|junto [àa]|junto a|registrados junto [àa]) (pgfn|procuradoria)|^n[ãa]o, possui d[íi]vida|pgfn possui", b)
            and not re.search(r"pgfn,? n[ãa]o (constam|possui)|n[ãa]o (possui|consta|tem)[^.,]{0,40}pgfn|negociad|extint", b))
    if rfb:
        return ("Parecer: débito em aberto na Receita", "Baixa", "só se for auto de infração ainda no prazo de impugnar (débito declarado em DCTF não entra)")
    if ("gestão de passivo" in b or "gestao de passivo" in b) and "envi" in b:
        return ("Parecer: enviada para Gestão de Passivo", "Média", "a confirmar")
    if pgfn:
        return ("Parecer: dívida na PGFN (dívida ativa)", "Baixa", "fora do RFB 10; ver Edital PGDAU 6/2026 pequeno valor")
    if re.search(r"possui d[ée]bitos?|d[ée]bitos? ativos", b):
        return ("Parecer: possui débitos (órgão não informado)", "Média", "a confirmar")
    return None


def sinal_perdcomp(txt):
    b = str(txt or "").lower()
    if re.search(r"n[ãa]o homologad|homologa[çc][ãa]o parcial|despacho decis|indeferid", b):
        return ("Parecer: PER/DCOMP não homologada ou indeferida", "Alta", "RFB 10 se houver manifestação de inconformidade")
    return None


def valores_citados(txt):
    v = []
    for m in re.findall(r"R\$\s?([\d.]+,\d{2})", str(txt or "")):
        try: v.append(float(m.replace(".", "").replace(",", ".")))
        except ValueError: pass
    return v


def carregar():
    fonte = sorted(glob.glob(os.path.join(HS, "hubspot-*.json")))[-1]
    D = json.load(open(fonte)); N = D["negocios"]
    fila = sorted(glob.glob(os.path.join(FILA, "Fila-Recuperador-*.xlsx")))[-1]
    v4 = {}
    wb = openpyxl.load_workbook(fila, read_only=True)
    if "Gestão de Passivo" in wb.sheetnames:
        rows = list(wb["Gestão de Passivo"].iter_rows(values_only=True)); cab = rows[0]
        for r in rows[1:]:
            d = dict(zip(cab, r)); m = re.search(r"deal/(\d+)", str(d.get("link") or ""))
            if m: v4[m.group(1)] = d.get("motivo") or "V4 Gestão de Passivo"
    R = json.load(open(RES)) if os.path.exists(RES) else {}
    mapa = json.load(open(MAPA)); roster = set(mapa["escopo"]["parceiros_roster"])
    return fonte, D.get("baixado_em", ""), N, fila, v4, R, roster


def receita(raizes):
    db = sqlite3.connect(DB); q = ",".join("?" * len(raizes)); rz = list(raizes)
    emp = {r[0]: {"razao": r[1], "porte": r[2], "capital": r[3]} for r in db.execute(f"select basico,razao,porte,capital from empresa where basico in ({q})", rz)}
    for b, s, m in db.execute(f"select basico,simples,mei from simples where basico in ({q})", rz):
        emp.setdefault(b, {})["simples"] = s; emp[b]["mei"] = m
    mes = dict(db.execute("select chave,valor from meta").fetchall()).get("mes_referencia", "?")
    return emp, mes


def montar():
    fonte, baixado, N, fila, v4, R, roster = carregar()
    por_raiz = defaultdict(list)
    for d in N:
        r = raiz_cnpj(d["extras"].get("cnpj"))
        if r: por_raiz[r].append(d)
    sinais = defaultdict(list)  # raiz -> [(sinal, força, fonte, onde_cabe, deal_id)]
    valores = defaultdict(list); datas = defaultdict(set); textos = defaultdict(set)
    for r, ds in por_raiz.items():
        for d in ds:
            if d["pipeline"] == "Gestão de Passivo":
                sinais[r].append((f"Funil Gestão de Passivo ({d['status']}, {d['etapa']})", "Média", "HubSpot · funil Gestão de Passivo", "a confirmar", d["id"]))
            if d["id"] in v4:
                sinais[r].append(("Veto V4 do Recuperador: caso de Gestão de Passivo", "Média", f"Fila do Recuperador · aba Gestão de Passivo ({os.path.basename(fila)[18:28]})", "a confirmar", d["id"]))
            p = R.get(d["id"]) or {}
            for s in (sinal_passivo(p.get("passivo")), sinal_perdcomp(p.get("perdcomp"))):
                if s:
                    sinais[r].append((s[0], s[1], f"Parecer do técnico no card (análise de {(p.get('data_analise') or 's/ data').strip('. ')})", s[2], d["id"]))
                    textos[r].add(re.sub(r"\s+", " ", str(p.get("passivo") or p.get("perdcomp")))[:220])
                    valores[r] += valores_citados(p.get("passivo")); datas[r].add((p.get("data_analise") or "").strip(". "))
    emp, mes_rf = receita(set(sinais))
    linhas, fora, descartados = [], [], defaultdict(int)
    for r, ss in sinais.items():
        e = emp.get(r, {}); porte = e.get("porte"); mei = e.get("mei") == "S"
        if porte == "Demais" and not mei: descartados["porte Demais (fora do edital)"] += 1; continue
        ds = por_raiz[r]
        # card de referência: o do funil de passivo; senão, o mais recente
        ref = sorted(ds, key=lambda d: (d["pipeline"] != "Gestão de Passivo", d["id"] not in {s[4] for s in ss}, -(int(d["id"]))))[0]
        masters = sorted({d["parceiro"] for d in ds if d["parceiro"] in roster})
        uniq = {}
        for s in ss: uniq.setdefault(s[0], s)
        ss = sorted(uniq.values(), key=lambda s: FORCA[s[1]])
        forca = ss[0][1]
        so_pgfn = all("PGFN" in s[0] for s in ss)
        so_aberto = all(s[0] == "Parecer: débito em aberto na Receita" for s in ss)
        porte_txt = (porte or "não encontrado na base da Receita") + (" · MEI" if mei else "")
        simples = e.get("simples") == "S"
        confirmar = ["se o débito está em discussão na Receita (impugnação, manifestação de inconformidade ou recurso) ou ainda no prazo de impugnar",
                     f"valor de cada processo até R$ {LIMITE:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
                     "se o cliente aceita desistir da discussão e confessar o débito"]
        if simples: confirmar.insert(0, "é optante do Simples: débito do Simples não entra (só multa por atraso de obrigação acessória)")
        if not porte: confirmar.insert(0, "porte ME/EPP no cadastro da Receita (empresa não achada na base local)")
        if so_pgfn: confirmar = ["dívida só na PGFN não entra no RFB 10; se houver débito na Receita em discussão, entra", "alternativa: Edital PGDAU 6/2026, pequeno valor, até 29/01/2027"]
        vmax = max(valores[r]) if valores[r] else None
        if vmax and vmax > LIMITE: confirmar.append(f"o parecer cita R$ {vmax:,.0f}".replace(",", ".") + " (pode ser a soma de processos; o limite é por processo)")
        if so_aberto: confirmar.insert(0, "débito em aberto costuma ser declarado (DCTF) e não entra; só entra se for lançamento de ofício no prazo de impugnar")
        sug = "Não se aplica (ver PGDAU 6)" if so_pgfn else ("Provável não" if so_aberto else "Confirmar")
        if not porte: sug = "Confirmar porte"
        linha = {
            "empresa": e.get("razao") or ref["nome"], "cnpj": fmt_cnpj(ref["extras"].get("cnpj")), "master": ", ".join(masters) or ref["parceiro"],
            "porte": porte_txt, "simples": "Sim" if simples else ("Não" if e.get("simples") == "N" else "não consta"),
            "forca": forca if not so_pgfn else "Baixa", "sinal": " · ".join(s[0] for s in ss), "fonte": " · ".join(sorted({s[2] for s in ss})),
            "onde": "; ".join(sorted({s[3] for s in ss})), "trecho": " | ".join(sorted(textos[r]))[:400] or None,
            "data_parecer": ", ".join(sorted(x for x in datas[r] if x)) or None,
            "funis": " · ".join(sorted({f"{d['pipeline']} / {d['etapa']} ({d['status']})" for d in ds})), "n_neg": len(ds),
            "link": LINK.format(ref["id"]), "uf": ref["extras"].get("uf"), "confirmar": "; ".join(confirmar), "sugestao": sug,
        }
        (linhas if masters else fora).append(linha)
    chave = lambda x: (x["master"], FORCA[x["forca"]], x["empresa"])
    linhas.sort(key=chave); fora.sort(key=chave)
    meta = {"fonte": os.path.basename(fonte), "baixado": baixado, "fila": os.path.basename(fila), "mes_rf": mes_rf,
            "com_sinal": len(sinais), "descartados": dict(descartados), "pareceres": len(R)}
    return linhas, fora, meta


COLS = [("Empresa", 34, "empresa"), ("CNPJ", 19, "cnpj"), ("Master", 16, "master"), ("Porte oficial (Receita)", 18, "porte"),
        ("Optante do Simples?", 10, "simples"), ("Força do sinal", 9, "forca"), ("Sinal encontrado", 44, "sinal"), ("De onde", 30, "fonte"),
        ("Trecho do parecer", 40, "trecho"), ("Data da análise", 11, "data_parecer"), ("Funil / etapa no HubSpot", 36, "funis"),
        ("Negócios da empresa", 9, "n_neg"), ("UF", 5, "uf"), ("O que confirmar com o Master", 50, "confirmar"),
        ("Leitura do agente", 16, "sugestao"), ("Decisão da Daniela", 14, None), ("Observação", 30, None), ("Abrir no HubSpot", 13, "link")]


def aba(wb, titulo, linhas, cabecalho, primeira=False):
    ws = wb.active if primeira else wb.create_sheet(); ws.title = titulo; ws.sheet_view.showGridLines = False
    n = len(COLS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(1, 1, cabecalho); c.font = Font(name=F, bold=True, size=13, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro"])
    ws.row_dimensions[1].height = 28; c.alignment = Alignment(vertical="center")
    for j, (t, w, _) in enumerate(COLS, 1):
        c = ws.cell(2, j, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro_suave"])
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=Side(style="medium", color=C["ouro"]))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[2].height = 32
    fino = Border(bottom=Side(style="thin", color=C["borda"]))
    cor = {"Alta": C["sucesso"], "Média": C["alerta"], "Baixa": C["texto_3"]}
    for i, l in enumerate(linhas, 3):
        for j, (_, _, k) in enumerate(COLS, 1):
            v = "Abrir" if k == "link" else (l.get(k) if k else None)
            c = ws.cell(i, j, v); c.border = fino; c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(name=FM if k in ("cnpj", "n_neg", "data_parecer") else F, size=10, color=C["texto"])
        ws.cell(i, 6).font = Font(name=F, size=10, bold=True, color=cor[l["forca"]])
        ws.cell(i, 16).fill = PatternFill("solid", fgColor=C["ouro_suave"])
        ws.cell(i, n).hyperlink = l["link"]; ws.cell(i, n).font = Font(name=F, size=10, color=C["ouro_texto"], underline="single")
    fim = max(3, len(linhas) + 2)
    dv = DataValidation(type="list", formula1='"Abordar,Não se aplica,Confirmar"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"P3:P{fim}")
    ws.freeze_panes = "B3"; ws.auto_filter.ref = f"A2:{get_column_letter(n)}{fim}"
    return ws


def texto_linhas(wb, titulo, linhas, largura=110):
    u = wb.create_sheet(titulo); u.column_dimensions["A"].width = largura; u.sheet_view.showGridLines = False
    for i, t in enumerate(linhas, 1):
        c = u.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, bold=(i == 1 or t.endswith(":")), size=13 if i == 1 else 11, color=C["texto"])
        if i == 1: c.fill = PatternFill("solid", fgColor=C["ouro"])
    return u


def main():
    a = argparse.ArgumentParser(); a.add_argument("--saida"); a.add_argument("--texto"); a = a.parse_args()
    hoje = date.today().isoformat()
    linhas, fora, meta = montar()
    saida = a.saida or os.path.join(caminhos.pasta("daniela"), f"Transacao-Pequeno-Valor-candidatos-{hoje}.xlsx")
    wb = openpyxl.Workbook()
    aba(wb, "Candidatos", linhas, f"Transação de pequeno valor da RFB · {len(linhas)} hipóteses com Master · prazo {PRAZO} · decida a coluna P", primeira=True)
    aba(wb, "Sem Master (D76)", fora, f"Sem Master do roster (D76: fora das listas da Área Partners) · {len(fora)} · só para conhecimento")
    por_master = defaultdict(lambda: defaultdict(int))
    for l in linhas: por_master[l["master"]][l["forca"]] += 1
    res = [f"Resumo · gerado em {hoje}", "",
           f"Candidatos com Master: {len(linhas)}  ·  sem Master (aba à parte): {len(fora)}",
           f"Empresas com algum sinal de débito ou processo (antes do filtro de porte): {meta['com_sinal']}",
           "Saíram pelo porte: " + (", ".join(f"{k}: {v}" for k, v in meta["descartados"].items()) or "nenhuma"), "", "Por Master (força do sinal):"]
    for m in sorted(por_master):
        d = por_master[m]; res.append(f"  {m}: {sum(d.values())}  (Alta {d['Alta']} · Média {d['Média']} · Baixa {d['Baixa']})")
    res += ["", f"Fontes: HubSpot {meta['fonte']} (baixado {meta['baixado']}); {meta['fila']} (aba Gestão de Passivo); "
            f"{meta['pareceres']} pareceres já lidos pelo Recuperador; porte, Simples e MEI da Receita (Dados Abertos do CNPJ, {meta['mes_rf']}).",
            f"Edital: {EDITAL}."]
    texto_linhas(wb, "Resumo", res)
    como = ["Como usar",
            "O que é: lista de HIPÓTESES de clientes e leads ME, EPP ou MEI com algum sinal de débito federal ou processo, para conferir se cabem na transação de pequeno valor da Receita (Edital RFB 10/2026). Nada aqui confirma que o débito cabe no edital.",
            "Regra do edital (texto oficial):",
            "• Quem: pessoa física, MEI, empresário individual, ME e EPP.",
            "• O quê: débito na Receita em discussão administrativa (impugnação, manifestação de inconformidade ou recurso) ou ainda no prazo de impugnar, de até 60 salários mínimos por processo (R$ 97.260,00 em 2026). Entra o processo inteiro.",
            "• Fora: débito do Simples Nacional (salvo multa por atraso de obrigação acessória) e dívida ativa da PGFN (essa tem o Edital PGDAU 6/2026, pequeno valor, até 29/01/2027).",
            "• Exige: desistir da discussão e confessar o débito de forma irrevogável. Adesão no Portal de Serviços da Receita, até 30/10/2026, 20h59 (Brasília).",
            "Passo a passo:",
            "1. Comece pelos de força Alta (parecer cita débito com exigibilidade suspensa na Receita: costuma ser impugnação ou manifestação em curso, mas também pode ser parcelamento).",
            "2. Mande ao Master o texto da aba 'Texto ao Master' (troque o nome da empresa). Peça o número do processo ou a procuração no e-CAC (D85: nunca senha). Peça retorno até 27/10 para dar tempo da pré-análise.",
            "3. Marque a coluna 'Decisão da Daniela': Abordar (cliente tem processo que parece caber), Não se aplica (sem processo na Receita, valor acima do limite, débito do Simples, porte Demais) ou Confirmar (aguardando o Master).",
            "4. Os casos 'Abordar' vão ao Jurídico (Gestão de Passivo) para a pré-análise do processo e validação do advogado antes de falar em adesão ao cliente.",
            "Cuidados:",
            "• Não fale em percentual de desconto ao cliente nem ao Master: diga 'pode haver redução prevista no edital, sujeita à análise do processo'.",
            "• A força do sinal mede só a evidência que temos no CRM, não a chance de o cliente aderir.",
            "• O porte vem do cadastro da Receita (declarado pela empresa; pode estar desatualizado). A Receita confere o porte na adesão.",
            "• Pareceres antigos (veja a data da análise) podem não refletir a situação de hoje.",
            "• Esta planilha não altera o HubSpot. Dados de cliente: fica só nesta pasta (LGPD).",
            "Script: skills/depto-juridico/scripts/transacao_pequeno_valor.py (repositório agentes-partners)."]
    texto_linhas(wb, "Como usar", como)
    texto_linhas(wb, "Texto ao Master", ["Texto curto de abordagem ao Master (rascunho; nada é enviado sem aprovação)", ""] +
                 TEXTO_MASTER.format(master="[nome do Master]", empresa="[empresa]").split("\n"))
    ordem = ["Como usar", "Resumo", "Candidatos", "Texto ao Master", "Sem Master (D76)"]
    wb._sheets.sort(key=lambda w: ordem.index(w.title)); wb.active = 2
    wb.save(saida)
    if a.texto:
        open(a.texto, "w").write(TEXTO_MASTER.format(master="[nome do Master]", empresa="[empresa]"))
    print("ok", saida); print("\n".join(res[2:]))


if __name__ == "__main__":
    main()
