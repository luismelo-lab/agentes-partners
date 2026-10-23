#!/usr/bin/env python3
"""Painel de janelas de receita (Marketing de canal · Estrategista de Campanhas & Calendário Fiscal).

Monta Janelas-de-Receita-<data>.xlsx a partir de docs/lancamento-2026/janelas.json: uma linha por janela, com prazo,
dias que faltam, produto ou tese, público, base legal (da skill da área), área de agentes, responsável, base da Save por
Master, campanha proposta (lista, kit e roteiro), status e próximo passo. Uso interno. Não envia nada.

Para atualizar: edite o janelas.json (nova janela, prazo, status) e rode de novo. O texto passa pelo auditor de promessas.

  python3 janelas_receita.py                         gera em 04 · Gestão/Central de Agentes (caminhos.luis('central'))
  python3 janelas_receita.py --saida arquivo.xlsx    gera em outro lugar
  python3 janelas_receita.py --auditar               só roda o auditor de promessas nos textos do JSON
  python3 janelas_receita.py --recalcular-base       recalcula a base por Master (HubSpot + Receita local; leva alguns minutos)

Regras: nada de promessa de valor, prazo ou êxito; teto de edital é limite, não regra; alíquota zero só por PER.
"""
import argparse, json, os, subprocess, sys, tempfile
from collections import Counter
from datetime import date
from pathlib import Path

AP = Path(os.path.expanduser("~/agentes-partners"))
RAIZ = Path(__file__).resolve().parents[3]  # repositório (worktree ou checkout principal)
JSON_PADRAO = next((p for p in (RAIZ / "docs/lancamento-2026/janelas.json", AP / "docs/lancamento-2026/janelas.json") if p.exists()),
                   AP / "docs/lancamento-2026/janelas.json")
sys.path.insert(0, str(AP / "lib"))
import caminhos, marca  # noqa: E402

C = {k: v.lstrip("#") for k, v in marca.C.items() if not k.startswith("_")}
E = {k: v.lstrip("#") for k, v in marca.E.items() if not k.startswith("_")}
F, FM = "IBM Plex Sans", "IBM Plex Mono"
ORDEM_STATUS = ["Urgente", "Preparar", "Travado (decisão do Luis)", "Baixa aderência", "Monitorar"]
COLS = [("Prazo", 13), ("Faltam (dias)", 9), ("Janela", 34), ("Produto / tese", 24), ("Público (perfil)", 40), ("Base legal", 44),
        ("Área de agentes", 24), ("Responsável", 24), ("Base da Save por Master", 30), ("Campanha · lista", 34), ("Campanha · kit", 32),
        ("Campanha · roteiro", 32), ("Status", 15), ("Próximo passo", 40), ("Fonte interna", 34)]


def ler(arq):
    return json.load(open(arq, encoding="utf-8"))


def base_txt(j, perfil, n=5):
    if not perfil: return "Sem recorte na carteira"
    d = (j.get("base_save") or {}).get(perfil) or {}
    if not d: return "Nenhuma raiz no perfil"
    xs = [(m, q) for m, q in d.items() if m != "Outros (fora do roster)"][:n]
    fora = d.get("Outros (fora do roster)")
    return " · ".join(f"{m} {q}" for m, q in xs) + (f" · fora do roster {fora}" if fora else "") + f" (total {sum(d.values())})"


def textos(j):
    for w in j["janelas"]:
        yield w["id"], "\n".join(str(w.get(k) or "") for k in ("janela", "prazo_texto", "produto_tese", "publico", "base_legal",
                                                               "campanha_lista", "campanha_kit", "campanha_roteiro", "proximo_passo"))


def auditar(j):
    aud = AP / "skills/depto-qualidade/scripts/auditar_texto.py"
    if (RAIZ / "skills/depto-qualidade/scripts/auditar_texto.py").exists(): aud = RAIZ / "skills/depto-qualidade/scripts/auditar_texto.py"
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "janelas-de-receita.md"
        f.write_text("\n\n".join(f"## {i}\n{t}" for i, t in textos(j)), encoding="utf-8")
        r = subprocess.run([sys.executable, str(aud), str(f)], capture_output=True, text=True)
    print(r.stdout.strip() or r.stderr.strip()); return r.returncode


def recalcular_base(arq):
    """Raízes de CNPJ com negócio ganho ou aberto, por Master, em perfis aproximados (CNAE, regime, porte). Só leitura."""
    import partners_dados as PD
    d = json.load(open(PD.ultimo_download()))["negocios"]
    base = {}
    for x in d:
        if x["status"] not in ("ganho", "aberto"): continue
        e = x.get("extras") or {}; r = PD.raiz(e.get("cnpj") or "")
        if not r or (x["parceiro"], r) in base: continue
        cn = str(x.get("cnae") or ""); rec = PD.receita(e.get("cnpj")) if e.get("cnpj") else None
        if not cn[:2].isdigit(): cn = str((rec or {}).get("cnae") or "")
        base[(x["parceiro"], r)] = dict(cnae=cn, regime=e.get("regime") or "", porte=str((rec or {}).get("porte") or ""))
    serv = {41, 42, 43, 52, 53} | set(range(58, 64)) | set(range(68, 76)) | set(range(77, 83)) | set(range(85, 89)) | set(range(90, 97))
    mono = ("4731", "4771", "4530", "4635", "4723", "4772")
    dv = lambda c: int(c[:2]) if c[:2].isdigit() else -1
    perf = {
        "iss": lambda b: dv(b["cnae"]) in serv and "Simples" not in b["regime"],
        "t843": lambda b: (10 <= dv(b["cnae"]) <= 33 or dv(b["cnae"]) == 46) and b["regime"] == "Lucro Real",
        "saldo": lambda b: b["regime"] == "Lucro Real" and dv(b["cnae"]) in (1, 2, 3, 10, 11, 46),
        "estoque381": lambda b: (b["regime"].startswith("Lucro Presumido") and 45 <= dv(b["cnae"]) <= 47) or b["cnae"][:4] in mono,
        "peqvalor": lambda b: b["porte"] in ("Microempresa", "Empresa de Pequeno Porte") and "Simples" not in b["regime"],
        "agro": lambda b: dv(b["cnae"]) in (1, 2, 3),
        "todos": lambda b: True,
    }
    j = ler(arq)
    j["base_save"] = {p: dict(Counter(m for (m, _), b in base.items() if f(b)).most_common()) for p, f in perf.items()}
    j["atualizado_em"] = date.today().isoformat()
    json.dump(j, open(arq, "w", encoding="utf-8"), ensure_ascii=False, indent=2); print("base recalculada:", len(base), "raízes")


def gerar(j, saida):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    hoje = date.today()
    ws_ = sorted(j["janelas"], key=lambda w: (w.get("prazo") is None, w.get("prazo") or "", ORDEM_STATUS.index(w["status"]) if w["status"] in ORDEM_STATUS else 9))
    wb = Workbook(); ws = wb.active; ws.title = "Janelas"; ws.sheet_view.showGridLines = False
    n = len(COLS); fino = Border(bottom=Side(style="thin", color=C["borda"]))
    # faixa de título (modo escuro, como capa de painel)
    for r in (1, 2):
        for c in range(1, n + 1): ws.cell(r, c).fill = PatternFill("solid", fgColor=E["fundo"])
    ws.cell(1, 1, "Janelas de receita · Área Partners").font = Font(name=F, bold=True, size=16, color=E["ouro"])
    ws.cell(2, 1, f"Uso interno · atualizado em {hoje.strftime('%d/%m/%Y')} · fonte: docs/lancamento-2026/janelas.json · nada aqui promete valor, prazo ou êxito").font = Font(name=F, size=10, color=E["texto_2"])
    ws.row_dimensions[1].height = 28
    h = 4
    for k, (t, w) in enumerate(COLS, 1):
        c = ws.cell(h, k, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro_suave"])
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=Side(style="medium", color=C["ouro"]))
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.row_dimensions[h].height = 30
    cor_status = {"Urgente": C["erro"], "Preparar": C["ouro_texto"], "Travado (decisão do Luis)": C["alerta"], "Baixa aderência": C["texto_3"], "Monitorar": C["texto_3"]}
    for i, w in enumerate(ws_, h + 1):
        prazo = date.fromisoformat(w["prazo"]) if w.get("prazo") else None
        falta = (prazo - hoje).days if prazo else None
        vals = [prazo, falta if falta is not None else "sem data", w["janela"], w["produto_tese"], w["publico"], w["base_legal"], w["area_agentes"],
                w["responsavel"], base_txt(j, w.get("base_perfil")), w["campanha_lista"], w["campanha_kit"], w["campanha_roteiro"],
                w["status"], w["proximo_passo"], w["fonte_interna"]]
        for k, v in enumerate(vals, 1):
            c = ws.cell(i, k, v); c.border = fino; c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name=FM if k in (1, 2) else F, size=10, color=C["texto"], bold=(k in (3, 13)))
        ws.cell(i, 1).number_format = "dd/mm/yyyy"
        if not prazo: ws.cell(i, 1, w["prazo_texto"].split(":")[0]); ws.cell(i, 1).font = Font(name=F, size=10, color=C["texto_2"])
        ws.cell(i, 13).font = Font(name=F, size=10, bold=True, color=cor_status.get(w["status"], C["texto"]))
        if falta is not None and falta <= 10: ws.cell(i, 2).font = Font(name=FM, size=10, bold=True, color=C["erro"])
        if i % 2 == 0:
            for k in range(1, n + 1): ws.cell(i, k).fill = PatternFill("solid", fgColor=C["fundo_2"])
        ws.row_dimensions[i].height = 96
    ws.freeze_panes = ws.cell(h + 1, 4); ws.auto_filter.ref = f"A{h}:{get_column_letter(n)}{h + len(ws_)}"

    # base por Master
    b = wb.create_sheet("Base por Master"); b.sheet_view.showGridLines = False
    perfis = {"todos": "Carteira (ganho ou aberto)", "peqvalor": "ME ou EPP fora do Simples", "saldo": "Saldo credor (Real)", "t843": "Indústria ou atacado (Real)",
              "estoque381": "Estoque (Presumido ou monofásico)", "iss": "Serviços com ISS", "agro": "Agro (CNAE 01 a 03)"}
    bs = j.get("base_save") or {}
    masters = list((bs.get("todos") or {}).keys())
    b.cell(1, 1, "Raízes de CNPJ com negócio ganho ou aberto, por Master e perfil").font = Font(name=F, bold=True, size=13, color=C["texto"])
    b.cell(2, 1, j.get("base_save_nota", "")).font = Font(name=F, size=9, color=C["texto_2"]); b.cell(2, 1).alignment = Alignment(wrap_text=True)
    b.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(perfis) + 1); b.row_dimensions[2].height = 70
    for k, t in enumerate(["Master"] + list(perfis.values()), 1):
        c = b.cell(4, k, t); c.font = Font(name=F, bold=True, size=10, color=C["texto"]); c.fill = PatternFill("solid", fgColor=C["ouro_suave"])
        c.alignment = Alignment(wrap_text=True, vertical="center"); b.column_dimensions[get_column_letter(k)].width = 22 if k == 1 else 16
    for i, m in enumerate(masters, 5):
        b.cell(i, 1, m).font = Font(name=F, size=10, color=C["texto"])
        for k, p in enumerate(perfis, 2):
            c = b.cell(i, k, (bs.get(p) or {}).get(m, 0)); c.font = Font(name=FM, size=10, color=C["texto"]); c.border = fino

    # como ler
    u = wb.create_sheet("Como ler"); u.sheet_view.showGridLines = False; u.column_dimensions["A"].width = 120
    for i, t in enumerate([
        "Como ler o painel de janelas de receita",
        "1. Uma linha por janela, em ordem de prazo. 'Faltam (dias)' conta a partir da data em que o arquivo foi gerado; em vermelho, 10 dias ou menos.",
        "2. Janela sem data (Temas 118 e 843) depende de ato do STF: agir antes da ata ou do acórdão, sem levar valor ao cliente antes disso.",
        "3. 'Base da Save por Master' é tamanho de carteira no perfil aproximado, não lista de leads. A lista real sai pela Jenifer (lista_master.py) e pelo depto-dados.",
        "4. Status: Urgente (prazo curto e produto pronto), Preparar (montar lista e kit), Travado (falta decisão do Luis), Baixa aderência, Monitorar (confirmar a fonte).",
        "5. Nenhum texto daqui vai ao parceiro ou ao cliente sem passar pela área dona, pelo validador humano e pelo auditor de promessas. Teto de edital é limite, não regra.",
        "6. Para atualizar: edite docs/lancamento-2026/janelas.json e rode skills/depto-marketing/scripts/janelas_receita.py."], 1):
        c = u.cell(i, 1, t); c.font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11, color=C["texto"]); c.alignment = Alignment(wrap_text=True)
    os.makedirs(os.path.dirname(os.path.abspath(saida)), exist_ok=True); wb.save(saida); return saida  # abspath: "--saida arquivo.xlsx" dava makedirs("")


def main():
    a = argparse.ArgumentParser(); a.add_argument("--json", default=str(JSON_PADRAO)); a.add_argument("--saida")
    a.add_argument("--auditar", action="store_true"); a.add_argument("--recalcular-base", action="store_true"); a = a.parse_args()
    if a.recalcular_base: recalcular_base(a.json)
    j = ler(a.json)
    if a.auditar: sys.exit(auditar(j))
    saida = a.saida or os.path.join(caminhos.luis("central"), f"Janelas-de-Receita-{date.today().isoformat()}.xlsx")
    print("ok", gerar(j, saida), len(j["janelas"]), "janelas")
    try:  # fora do OneDrive (--saida em pasta local) caminhos.web montava um link com "../.." que não abre
        if os.path.realpath(saida).startswith(os.path.realpath(caminhos.onedrive()) + os.sep): print("link da pasta:", caminhos.web(saida))
        else: print("fora do OneDrive: sem link")
    except Exception as e: print("link indisponível:", e)


if __name__ == "__main__":
    main()
