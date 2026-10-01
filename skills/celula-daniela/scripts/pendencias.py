#!/usr/bin/env python3
"""Controle de pendências dos Masters (célula da Daniela · subagente Acompanhamento & SLA).

Cria e mantém ~/agentes-partners/celulas/daniela/Controle-Pendencias-Masters.xlsx.
A planilha é da Daniela: ela pode editar à mão. O script só acrescenta linhas aprovadas e
recalcula as colunas automáticas (dias em aberto, situação do SLA) e a aba de resumo.
Prazos vêm de references/sla.json (PROPOSTA a validar). Não envia nada, não grava no HubSpot/GLPI.

Uso:
  python3 pendencias.py --criar
  python3 pendencias.py --adicionar --master "D&S" --tipo POS-EXTRATO --canal WhatsApp \
      --descricao "Extrato de saldo de créditos para reunião de 01/10" --cliente "Frigorífico X" --negocio 123 \
      --aprovado-por Daniela
  python3 pendencias.py --adicionar-json fichas.json --aprovado-por Daniela   # lista de fichas da Triagem
  python3 pendencias.py --atualizar [--master AGILE]     # recalcula e lista atrasos e vencimentos
  python3 pendencias.py --status PM-2026-0001 "Encaminhada" [--responsavel "Fulana (Pós-venda)"]
  python3 pendencias.py --cobranca PM-2026-0001          # registra que a Daniela cobrou hoje
  python3 pendencias.py --concluir PM-2026-0001
  python3 pendencias.py --tipos                          # lista os tipos de demanda e prazos
"""
import argparse, json, os, sys
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from comum import PENDENCIAS, SKILL, dias_uteis_depois, resolver_master, roster

SLA = json.load(open(os.path.join(SKILL, "references/sla.json")))
TIPOS = SLA["tipos"]
COLS = ["ID", "Data abertura", "Master", "Contato no parceiro", "Canal", "Nº GLPI", "Tipo de demanda", "Categoria",
        "Descrição", "Cliente", "Negócio HubSpot", "Área de destino", "Responsável na área", "Status",
        "Prazo 1ª resposta", "Prazo solução", "Dias em aberto", "Situação SLA", "Última cobrança", "Próxima ação",
        "Data conclusão", "Foi ponte? (S/N)", "Aprovado por", "Observações"]
STATUS = ["Aberta", "Aguardando dados do parceiro", "Encaminhada", "Em andamento na área", "Aguardando retorno ao parceiro", "Concluída", "Cancelada"]
CANAIS = ["WhatsApp", "GLPI", "E-mail", "Reunião", "Telefone", "Teams", "Outro"]
LARG = {"ID": 13, "Descrição": 50, "Próxima ação": 36, "Observações": 30, "Área de destino": 28, "Tipo de demanda": 30, "Cliente": 28}


def estilo():
    from openpyxl.styles import Alignment, Font, PatternFill
    return Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F"), Alignment(wrap_text=True, vertical="top")


def criar():
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    if os.path.exists(PENDENCIAS):
        print(f"Já existe: {PENDENCIAS} (nada alterado)."); return
    os.makedirs(os.path.dirname(PENDENCIAS), exist_ok=True)
    H, HF, WR = estilo(); wb = Workbook(); ws = wb.active; ws.title = "Pendências"
    for j, c in enumerate(COLS, 1):
        x = ws.cell(row=1, column=j, value=c); x.font = H; x.fill = HF; x.alignment = WR
        ws.column_dimensions[x.column_letter].width = LARG.get(c, 16)
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLS)).column_letter}1"
    col = lambda n: ws.cell(row=1, column=COLS.index(n) + 1).column_letter
    R, _ = roster()
    for nome, lista in (("Master", list(R.keys())), ("Canal", CANAIS), ("Tipo de demanda", list(TIPOS.keys())), ("Status", STATUS), ("Foi ponte? (S/N)", ["S", "N"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(lista) + '"', allow_blank=True); ws.add_data_validation(dv)
        dv.add(f"{col(nome)}2:{col(nome)}1000")
    t = wb.create_sheet("Tipos e SLA")
    cab = ["Código", "Tipo", "Categoria", "Área de destino", "Rota", "1ª resposta (dias úteis)", "Solução (dias úteis)", "Dados mínimos"]
    for j, c in enumerate(cab, 1):
        x = t.cell(row=1, column=j, value=c); x.font = H; x.fill = HF
    for i, (k, v) in enumerate(TIPOS.items(), 2):
        for j, val in enumerate([k, v["rotulo"], v["categoria"], v["area"], v["rota"], v["primeira_resposta_du"], v["solucao_du"] if v["solucao_du"] is not None else v.get("solucao_obs"), "; ".join(v["dados_minimos"])], 1):
            t.cell(row=i, column=j, value=val).alignment = WR
    for L, wdt in zip("ABCDEFGH", (18, 40, 14, 30, 40, 12, 12, 60)): t.column_dimensions[L].width = wdt
    t.cell(row=len(TIPOS) + 3, column=1, value=SLA["_status"])
    c = wb.create_sheet("Como usar")
    for i, l in enumerate([
        "Controle de pendências dos 9 Masters — dono: Daniela Maciel. Uso interno (LGPD: só o mínimo de dado de cliente; nada de CPF, telefone ou valores de crédito na descrição).",
        "1. Cada linha é UMA demanda de parceiro. A Triagem sugere a ficha; a linha só entra depois que a Daniela aprova (coluna 'Aprovado por').",
        "2. Atualize a coluna Status à mão ou com: pendencias.py --status <ID> \"Encaminhada\".",
        "3. Colunas automáticas (não editar): Prazo 1ª resposta, Prazo solução, Dias em aberto, Situação SLA. Rode pendencias.py --atualizar toda manhã.",
        "4. Situação SLA: 🟢 no prazo · 🟡 vence hoje ou amanhã · 🔴 atrasada · ✅ concluída · ⚪ sem prazo.",
        "5. 'Foi ponte?' = S quando a Daniela só recebeu, repassou e devolveu (sem trabalho próprio). É o que mais se ganha delegando.",
        "6. Prazos (aba 'Tipos e SLA') são PROPOSTA: validar com as áreas antes de cobrar alguém por eles.",
        "7. O agente nunca envia cobrança nem resposta: ele rascunha, a Daniela aprova e envia."], 1):
        c.cell(row=i, column=1, value=l)
    c.column_dimensions["A"].width = 140
    wb.create_sheet("Resumo por Master")
    wb.save(PENDENCIAS); os.chmod(PENDENCIAS, 0o600); print(PENDENCIAS)


def abrir():
    from openpyxl import load_workbook
    if not os.path.exists(PENDENCIAS):
        criar()
    wb = load_workbook(PENDENCIAS); ws = wb["Pendências"]
    idx = {c.value: i + 1 for i, c in enumerate(ws[1])}
    return wb, ws, idx


def linhas(ws, idx):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=idx["ID"]).value:
            yield r


def como_data(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    try: return datetime.strptime(str(v), "%d/%m/%Y").date()
    except (TypeError, ValueError):
        try: return date.fromisoformat(str(v)[:10])
        except ValueError: return None


def novo_id(ws, idx, hoje):
    pref = f"PM-{hoje.year}-"; n = 0
    for r in linhas(ws, idx):
        v = str(ws.cell(row=r, column=idx["ID"]).value)
        if v.startswith(pref): n = max(n, int(v[len(pref):] or 0))
    return f"{pref}{n + 1:04d}"


def adicionar(fichas, aprovado_por):
    if not aprovado_por:
        sys.exit("Informe --aprovado-por: a linha só entra no controle depois da aprovação da Daniela.")
    wb, ws, idx = abrir(); hoje = date.today(); _, _, WR = estilo(); ids = []
    for f in fichas:
        tipo = f.get("tipo")
        if tipo not in TIPOS: sys.exit(f"Tipo '{tipo}' desconhecido. Use --tipos para ver a lista.")
        T = TIPOS[tipo]; ab = como_data(f.get("data")) or hoje
        v = {"ID": novo_id(ws, idx, hoje), "Data abertura": ab, "Master": resolver_master(f["master"]), "Contato no parceiro": f.get("contato"),
             "Canal": f.get("canal") or "WhatsApp", "Nº GLPI": f.get("glpi"), "Tipo de demanda": tipo, "Categoria": T["categoria"],
             "Descrição": f.get("descricao"), "Cliente": f.get("cliente"), "Negócio HubSpot": f.get("negocio"),
             "Área de destino": f.get("area") or T["area"], "Responsável na área": f.get("responsavel"), "Status": f.get("status") or "Aberta",
             "Prazo 1ª resposta": dias_uteis_depois(ab, T["primeira_resposta_du"]), "Prazo solução": dias_uteis_depois(ab, T["solucao_du"]) if T["solucao_du"] is not None else (como_data(f.get("prazo")) or None),
             "Próxima ação": f.get("proxima_acao"), "Foi ponte? (S/N)": f.get("ponte"), "Aprovado por": aprovado_por, "Observações": f.get("obs")}
        r = ws.max_row + 1
        while r > 2 and not any(ws.cell(row=r - 1, column=j).value for j in range(1, len(COLS) + 1)): r -= 1
        for k, val in v.items():
            x = ws.cell(row=r, column=idx[k], value=val); x.alignment = WR
            if isinstance(val, date): x.number_format = "DD/MM/YYYY"
        ids.append(v["ID"])
    wb.save(PENDENCIAS); atualizar(silencioso=True)
    print("Registradas:", ", ".join(ids))


def achar(ws, idx, pid):
    for r in linhas(ws, idx):
        if ws.cell(row=r, column=idx["ID"]).value == pid: return r
    sys.exit(f"{pid} não encontrado.")


def mudar(pid, campo, valor, extra=None):
    wb, ws, idx = abrir(); r = achar(ws, idx, pid)
    x = ws.cell(row=r, column=idx[campo], value=valor)
    if isinstance(valor, date): x.number_format = "DD/MM/YYYY"
    for k, v in (extra or {}).items():
        if v is not None:
            y = ws.cell(row=r, column=idx[k], value=v)
            if isinstance(v, date): y.number_format = "DD/MM/YYYY"
    wb.save(PENDENCIAS); atualizar(silencioso=True); print(f"{pid}: {campo} = {valor}")


def atualizar(master=None, silencioso=False):
    from openpyxl.styles import PatternFill
    wb, ws, idx = abrir(); hoje = date.today(); H, HF, _ = estilo()
    cor = {"🔴": "F8D7DA", "🟡": "FFF3CD", "🟢": "D4EDDA", "✅": "E2E3E5", "⚪": "FFFFFF"}
    res = {}; alertas = []
    for r in linhas(ws, idx):
        g = lambda k: ws.cell(row=r, column=idx[k]).value
        st = g("Status") or "Aberta"; ab = como_data(g("Data abertura")); p1 = como_data(g("Prazo 1ª resposta")); ps = como_data(g("Prazo solução"))
        fim = como_data(g("Data conclusão")) if st in ("Concluída", "Cancelada") else None
        dias = ((fim or hoje) - ab).days if ab else None
        if st in ("Concluída", "Cancelada"):
            sit = "✅ " + ("no prazo" if (ps and fim and fim <= ps) else "fora do prazo" if (ps and fim) else st.lower())
        elif not ps: sit = "⚪ sem prazo"
        elif hoje > ps: sit = f"🔴 atrasada {(hoje - ps).days}d"
        elif (ps - hoje).days <= 1: sit = "🟡 vence " + ("hoje" if ps == hoje else "amanhã")
        else: sit = "🟢 no prazo"
        if st == "Aberta" and p1 and hoje > p1 and not sit.startswith("🔴"): sit += " · 1ª resposta atrasada"
        ws.cell(row=r, column=idx["Dias em aberto"], value=dias)
        x = ws.cell(row=r, column=idx["Situação SLA"], value=sit); x.fill = PatternFill("solid", fgColor=cor[sit[0]] if sit[0] in cor else "FFFFFF")
        m = g("Master") or "(sem Master)"
        R = res.setdefault(m, {"abertas": 0, "atrasadas": 0, "vencendo": 0, "concluidas": 0, "no_prazo_concl": 0, "dias": [], "ponte": 0})
        if st in ("Concluída", "Cancelada"):
            R["concluidas"] += 1; R["no_prazo_concl"] += 1 if "no prazo" in sit else 0
            if dias is not None: R["dias"].append(dias)
        else:
            R["abertas"] += 1; R["atrasadas"] += sit.startswith("🔴"); R["vencendo"] += sit.startswith("🟡")
            if (sit.startswith("🔴") or sit.startswith("🟡") or "1ª resposta atrasada" in sit) and (not master or m == master):
                alertas.append((sit, g("ID"), m, g("Tipo de demanda"), g("Área de destino"), (g("Descrição") or "")[:70], g("Última cobrança")))
        R["ponte"] += 1 if str(g("Foi ponte? (S/N)") or "").upper() == "S" else 0
    s = wb["Resumo por Master"]; s.delete_rows(1, s.max_row)
    cab = ["Master", "Abertas", "Atrasadas", "Vencem hoje/amanhã", "Concluídas", "% concluídas no prazo", "Dias médios até concluir", "Demandas de ponte", "Atualizado em"]
    for j, c in enumerate(cab, 1):
        x = s.cell(row=1, column=j, value=c); x.font = H; x.fill = HF; s.column_dimensions[x.column_letter].width = 18
    for i, (m, R) in enumerate(sorted(res.items()), 2):
        for j, v in enumerate([m, R["abertas"], R["atrasadas"], R["vencendo"], R["concluidas"],
                               round(R["no_prazo_concl"] / R["concluidas"] * 100) if R["concluidas"] else None,
                               round(sum(R["dias"]) / len(R["dias"]), 1) if R["dias"] else None, R["ponte"], hoje.strftime("%d/%m/%Y")], 1):
            s.cell(row=i, column=j, value=v)
    wb.save(PENDENCIAS)
    if silencioso: return
    tot = {k: sum(R[k] for R in res.values()) for k in ("abertas", "atrasadas", "vencendo")}
    print(f"{PENDENCIAS}\nAbertas: {tot['abertas']} · atrasadas: {tot['atrasadas']} · vencem hoje/amanhã: {tot['vencendo']}")
    for a in sorted(alertas):
        print(" | ".join((x.strftime("%d/%m/%Y") if hasattr(x, "strftime") else str(x)) if x is not None else "—" for x in a))


def main():
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--criar", action="store_true"); g.add_argument("--adicionar", action="store_true"); g.add_argument("--adicionar-json")
    g.add_argument("--atualizar", action="store_true"); g.add_argument("--status", nargs=2, metavar=("ID", "STATUS"))
    g.add_argument("--cobranca", metavar="ID"); g.add_argument("--concluir", metavar="ID"); g.add_argument("--tipos", action="store_true")
    for o in ("master", "tipo", "canal", "descricao", "cliente", "negocio", "glpi", "contato", "area", "responsavel", "data", "proxima-acao", "ponte", "obs", "aprovado-por", "prazo"):
        ap.add_argument("--" + o)
    a = ap.parse_args()
    if a.criar: criar()
    elif a.tipos:
        for k, v in TIPOS.items(): print(f"{k:17} {v['rotulo']} → {v['area']} · 1ª resp. {v['primeira_resposta_du']} du · solução " + (f"{v['solucao_du']} du" if v['solucao_du'] is not None else v.get('solucao_obs', 'a definir')))
    elif a.adicionar:
        if not (a.master and a.tipo and a.descricao): ap.error("--adicionar exige --master, --tipo e --descricao")
        adicionar([{"master": a.master, "tipo": a.tipo, "canal": a.canal, "descricao": a.descricao, "cliente": a.cliente, "negocio": a.negocio,
                    "glpi": a.glpi, "contato": a.contato, "area": a.area, "responsavel": a.responsavel, "data": a.data,
                    "proxima_acao": a.proxima_acao, "ponte": a.ponte, "obs": a.obs, "prazo": a.prazo}], a.aprovado_por)
    elif a.adicionar_json:
        F = json.load(open(os.path.expanduser(a.adicionar_json))); adicionar(F if isinstance(F, list) else [F], a.aprovado_por)
    elif a.atualizar: atualizar(resolver_master(a.master) if a.master else None)
    elif a.status:
        extra = {"Responsável na área": a.responsavel}
        if a.status[1] not in STATUS: sys.exit(f"Status deve ser um de: {', '.join(STATUS)}")
        if a.status[1] in ("Concluída", "Cancelada"): extra["Data conclusão"] = date.today()
        mudar(a.status[0], "Status", a.status[1], extra)
    elif a.cobranca: mudar(a.cobranca, "Última cobrança", date.today())
    elif a.concluir: mudar(a.concluir, "Status", "Concluída", {"Data conclusão": date.today()})


if __name__ == "__main__":
    main()
