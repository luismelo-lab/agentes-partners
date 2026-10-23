#!/usr/bin/env python3
"""Briefing determinístico de uma conta Master (célula da Daniela · Central Masters).

Lê o download mais recente do HubSpot (hubspot-*.json, gerado pelo depto-resultados) e,
se existirem, o cruzamento com a Receita (depto-intel-comercial), o scorecard de saúde
(depto-dados) e o controle de pendências da Daniela. Gera um .md com:
pipeline aberto por etapa, ganhos/perdidos recentes, negócios parados (>30 e >90 dias),
qualidade do CRM (campos vazios), principais clientes e, com --negocio, o raio-x do negócio.

Todos os números vêm deste script. Não chama API, não lê .env, não grava no HubSpot.

Uso:
  python3 briefing_master.py --master "D&S"
  python3 briefing_master.py --master agile --negocio 55584269700
  python3 briefing_master.py --negocio 55584269700          # Master deduzido do negócio
Opções: --fonte <hubspot-*.json> · --janela 90 (dias "recentes") · --data-ref AAAA-MM-DD · --saida <pasta>
"""
import argparse, os, statistics, sys
from collections import Counter, defaultdict
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from comum import (CELULA, DADOS_BI, LINK, PENDENCIAS, brl, cnpj_digitos, dt, e_cafe, fonte_hubspot, mi, pct,
                   receita, resolver_master, roster, saude_master, slug, ultima_atividade)
from cnpj import raiz_cnpj  # noqa: E402  (lib no caminho pelo comum; raiz com os zeros à esquerda recompostos)

# campos que a Daniela (ou o time) deveria manter preenchidos, por status
CAMPOS_ABERTO = [("CNPJ", lambda d: bool(cnpj_digitos(d["extras"].get("cnpj")))),
                 ("UF", lambda d: bool(d["extras"].get("uf"))),
                 ("regime tributário", lambda d: bool(d["extras"].get("regime"))),
                 ("valor do negócio", lambda d: d["valor"] > 0),
                 ("faturamento / faixa", lambda d: bool(d["extras"].get("faturamento") or d["extras"].get("faixa_faturamento"))),
                 ("CNAE / segmento", lambda d: bool(d.get("segmento") or d.get("cnae"))),
                 ("e-CAC liberado", lambda d: d["extras"].get("ecac") not in (None, "")),
                 ("atividade registrada", lambda d: bool(d["extras"].get("ultima_atividade")))]
CAMPOS_GANHO = [("valor do negócio", lambda d: d["valor"] > 0),
                ("% honorário contratual", lambda d: d["hContr"] is not None),
                ("% do parceiro", lambda d: d["hParc"] is not None),
                ("tese aplicada", lambda d: bool(d.get("variante"))),
                ("CNAE / segmento", lambda d: bool(d.get("segmento") or d.get("cnae"))),
                ("CNPJ", lambda d: bool(cnpj_digitos(d["extras"].get("cnpj"))))]


def tabela(cab, linhas):
    if not linhas:
        return "_Nada a listar._\n"
    out = "| " + " | ".join(cab) + " |\n|" + "|".join("---" for _ in cab) + "|\n"
    for l in linhas:
        out += "| " + " | ".join(str(x).replace("|", "/") for x in l) + " |\n"
    return out


def link(d):
    return f"[abrir]({LINK.format(d['id'])})"


def nome(d, regras):
    return (d["nome"] or "(sem nome)").strip() + (" ☕" if e_cafe(d, regras) else "")


def pendencias_do_master(master):
    if not os.path.exists(PENDENCIAS):
        return None
    from openpyxl import load_workbook
    ws = load_workbook(PENDENCIAS, data_only=True)["Pendências"]
    cab = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        x = {k: (v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else v) for k, v in zip(cab, r)}
        if x.get("Master") == master and x.get("Status") not in ("Concluída", "Cancelada", None):
            out.append(x)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--master"); ap.add_argument("--negocio"); ap.add_argument("--fonte")
    ap.add_argument("--janela", type=int, default=90); ap.add_argument("--data-ref"); ap.add_argument("--saida")
    a = ap.parse_args()
    if not a.master and not a.negocio:
        ap.error("informe --master e/ou --negocio")
    f, D = fonte_hubspot(a.fonte); N = D["negocios"]
    _, regras = roster()
    por_id = {d["id"]: d for d in N}
    neg = None
    if a.negocio:
        neg = por_id.get(str(a.negocio))
        if not neg:
            sys.exit(f"Negócio {a.negocio} não está no download {os.path.basename(f)} (fora do escopo Área Master ou criado depois).")
    master = resolver_master(a.master) if a.master else neg["parceiro"]
    if neg and neg["parceiro"] != master:
        print(f"Aviso: o negócio {neg['id']} está registrado para '{neg['parceiro']}', não para '{master}'.", file=sys.stderr)
    hoje = date.fromisoformat(a.data_ref) if a.data_ref else date.today()
    J = hoje - timedelta(days=a.janela); A12 = hoje - timedelta(days=365)
    L = [d for d in N if d["parceiro"] == master]
    if not L:
        sys.exit(f"Nenhum negócio de {master} no download.")
    fonte_rf, RF = receita()

    ab = [d for d in L if d["status"] == "aberto"]
    gan = [d for d in L if d["status"] == "ganho"]
    per = [d for d in L if d["status"] == "perdido"]
    fech = lambda d: dt(d["fechado"]) or date.min
    g12 = [d for d in gan if fech(d) >= A12]; p12 = [d for d in per if fech(d) >= A12]
    gJ = sorted([d for d in gan if fech(d) >= J], key=lambda d: d["fechado"], reverse=True)
    pJ = sorted([d for d in per if fech(d) >= J], key=lambda d: d["fechado"], reverse=True)
    criJ = [d for d in L if (dt(d["criado"]) or date.min) >= J]
    conv12 = len(g12) / (len(g12) + len(p12)) * 100 if (g12 or p12) else None
    # ciclo só com datas coerentes: negócios importados (Bitrix) têm criação = data da importação
    ciclos = [(dt(d["fechado"]) - dt(d["criado"])).days for d in g12 if dt(d["fechado"]) and dt(d["criado"]) and dt(d["fechado"]) >= dt(d["criado"])]
    cafe_ab = [d for d in ab if e_cafe(d, regras)]

    # parados: sem atividade registrada (notas) ou, na falta, sem modificação
    def dias_parado(d):
        u = ultima_atividade(d)
        return (hoje - u).days if u else None
    p90 = [d for d in ab if (dias_parado(d) is None or dias_parado(d) > 90)]
    p30 = [d for d in ab if d not in p90 and dias_parado(d) > 30]
    velhos = [d for d in ab if (hoje - (dt(d["criado"]) or hoje)).days > 90]

    num = {"master": master, "negocios": len(L), "abertos": len(ab), "abertos_rs": sum(d["valor"] for d in ab),
           "ganhos_janela": len(gJ), "ganhos_janela_rs": sum(d["valor"] for d in gJ), "perdidos_janela": len(pJ),
           "perdidos_janela_valor_zero": sum(1 for d in pJ if d["valor"] <= 0), "criados_janela": len(criJ),
           "ganhos_12m": len(g12), "ganhos_12m_rs": sum(d["valor"] for d in g12), "perdidos_12m": len(p12),
           "conversao_12m_pct": conv12, "ciclo_mediano_dias": statistics.median(ciclos) if ciclos else None,
           "parados_30_90": len(p30), "parados_mais_90": len(p90), "abertos_criados_ha_mais_90": len(velhos)}

    md = []
    w = md.append
    w(f"# Briefing da conta Master — {master}")
    w(f"*Gerado em {hoje.strftime('%d/%m/%Y')} pelo script `briefing_master.py` (célula da Daniela). "
      f"Fonte: `{os.path.basename(f)}` (HubSpot baixado em {D.get('baixado_em')}). Todos os números vêm do script. "
      f"Uso interno: não enviar ao parceiro sem revisão da Daniela.*\n")
    w(f"Janela \"recente\" = últimos {a.janela} dias (desde {J.strftime('%d/%m/%Y')}). Escopo: negócios da Área Master desde 2024 "
      "atribuídos a este Master pelo roster (`mapeamento-hubspot.json`). ☕ = café (fora da meta desde abr/2026, regra do Luis).\n")
    nomes_brutos = Counter((d["parceiro_bruto"] or "").strip() for d in L)
    w("**Como o parceiro aparece no campo `nome_do_parceiro`:** " + "; ".join(f"{k or '(vazio)'} ({v})" for k, v in nomes_brutos.most_common(8)) + "\n")

    w("## 1. Resumo em números")
    linhas = [("Negócios no histórico (desde 2024)", len(L)),
              ("Pipeline aberto", f"{len(ab)} negócios · {mi(num['abertos_rs'])}" + (f" (dos quais {len(cafe_ab)} de café)" if cafe_ab else "")),
              (f"Criados nos últimos {a.janela} dias", len(criJ)),
              (f"Ganhos nos últimos {a.janela} dias", f"{len(gJ)} · {mi(num['ganhos_janela_rs'])}"),
              (f"Perdidos nos últimos {a.janela} dias", f"{len(pJ)} (valor zero: {num['perdidos_janela_valor_zero']})"),
              ("Ganhos em 12 meses", f"{len(g12)} · {mi(num['ganhos_12m_rs'])}"),
              ("Conversão em 12 meses (ganhos ÷ fechados)", pct(conv12)),
              ("Ciclo mediano dos ganhos em 12 meses (datas coerentes)", f"{num['ciclo_mediano_dias']:.0f} dias ({len(ciclos)} ganhos)" if ciclos else "—"),
              ("Abertos parados 31–90 dias / > 90 dias", f"{len(p30)} / {len(p90)}"),
              ("Abertos criados há mais de 90 dias", len(velhos))]
    w(tabela(["Indicador", "Valor"], linhas))
    arq, s = saude_master(master)
    if s:
        if os.path.getmtime(os.path.join(DADOS_BI, arq)) < os.path.getmtime(f):
            w("_O scorecard abaixo foi gerado antes do download atual do HubSpot: pode divergir um pouco dos números acima (vale o do script)._")
        w(f"**Saúde do parceiro (scorecard do depto-dados, `{arq}`):** {s['sinal']} {s['saude']}/100 · ritmo de entrada vs. média trimestral "
          f"{pct(s.get('ritmo_entrada_vs_media_pct'))} · ritmo de valor fechado vs. média {pct(s.get('ritmo_valor_vs_media_trimestral_pct'))} · "
          f"conversão 12m {pct(s.get('conversao_12m_pct'))} · % honorário preenchido nos ganhos {pct(s.get('honorario_preenchido_pct'))}. "
          "É sinal para conversa, não nota.\n")
    else:
        w("**Saúde do parceiro:** scorecard não encontrado; rode o depto-dados (`dados_bi.py`).\n")

    w("## 2. Pipeline aberto por etapa")
    g = defaultdict(list)
    for d in ab:
        g[(d["pipeline"], d["etapa"] or "(sem etapa)", d["prob"] or 0)].append(d)
    linhas = []
    for (pl, et, pr), X in sorted(g.items(), key=lambda kv: (kv[0][0], kv[0][2], kv[0][1])):
        idades = [(hoje - dt(d["criado"])).days for d in X if dt(d["criado"])]
        linhas.append((pl, et, len(X), brl(sum(d["valor"] for d in X)), f"{statistics.median(idades):.0f}" if idades else "—",
                       sum(1 for d in X if d in p30 or d in p90)))
    w(tabela(["Pipeline", "Etapa", "Negócios", "Valor", "Idade mediana (dias)", "Parados >30d"], linhas))
    top = sorted(ab, key=lambda d: -d["valor"])[:10]
    w("**Maiores negócios abertos**")
    w(tabela(["Negócio", "Etapa", "Valor", "Criado", "Dias sem atividade", "Link"],
             [(nome(d, regras), d["etapa"], brl(d["valor"]), d["criado"], dias_parado(d) if dias_parado(d) is not None else "nunca", link(d)) for d in top]))

    w(f"## 3. Ganhos recentes (últimos {a.janela} dias)")
    w(tabela(["Fechado", "Negócio", "Valor", "Tese", "% honorário", "Link"],
             [(d["fechado"], nome(d, regras), brl(d["valor"]), (d.get("variante") or "—")[:40], pct((d["hContr"] or 0) * 100, 0) if d["hContr"] is not None else "vazio", link(d)) for d in gJ[:15]]))
    if len(gJ) > 15: w(f"_… e mais {len(gJ) - 15}._\n")
    w(f"## 4. Perdidos recentes (últimos {a.janela} dias)")
    w(tabela(["Fechado", "Negócio", "Pipeline", "Valor", "Link"],
             [(d["fechado"], nome(d, regras), d["pipeline"], brl(d["valor"]) if d["valor"] > 0 else "zero", link(d)) for d in pJ[:15]]))
    if len(pJ) > 15: w(f"_… e mais {len(pJ) - 15}._\n")
    if pJ:
        w("Perdidos com valor zero podem ser candidatos à reanálise: quem decide é a Geovana (skill `recuperador`).\n")

    w("## 5. Negócios parados")
    w("Critério: dias desde a última nota/atividade registrada no card; sem nota, usa a última modificação (sinal fraco: qualquer "
      "atualização automática conta). \"nunca\" = sem nenhuma data.\n")
    for titulo, X in ((f"Mais de 90 dias sem atividade ({len(p90)})", p90), (f"De 31 a 90 dias sem atividade ({len(p30)})", p30)):
        w(f"**{titulo}**")
        X = sorted(X, key=lambda d: -d["valor"])
        w(tabela(["Negócio", "Etapa", "Valor", "Dias parado", "Link"],
                 [(nome(d, regras), d["etapa"], brl(d["valor"]), dias_parado(d) if dias_parado(d) is not None else "nunca", link(d)) for d in X[:20]]))
        if len(X) > 20: w(f"_… e mais {len(X) - 20}._\n")

    w("## 6. Qualidade do CRM (campos vazios)")
    linhas = []
    for rot, fn in CAMPOS_ABERTO:
        vaz = [d for d in ab if not fn(d)]
        linhas.append(("Abertos", rot, f"{len(vaz)} de {len(ab)}", pct(len(vaz) / len(ab) * 100 if ab else None, 0)))
    for rot, fn in CAMPOS_GANHO:
        vaz = [d for d in g12 if not fn(d)]
        linhas.append(("Ganhos 12m", rot, f"{len(vaz)} de {len(g12)}", pct(len(vaz) / len(g12) * 100 if g12 else None, 0)))
    w(tabela(["Grupo", "Campo", "Vazios", "%"], linhas))
    imp = sum(1 for d in L if str(d["extras"].get("importado")).lower() == "true")
    w(f"Negócios importados (provável migração do Bitrix; a data de criação é a da importação): {imp} de {len(L)}. "
      "Por isso o histórico usa a data de fechamento, não a de criação, para ganhos e perdidos.\n")
    # checagem com a Receita (dados oficiais), só para abertos
    if RF:
        sit = [(d, RF[d["id"]]) for d in ab if d["id"] in RF and (RF[d["id"]].get("situacao") or "").lower() != "ativa"]
        uf = [(d, RF[d["id"]]) for d in ab if d["id"] in RF and d["extras"].get("uf") and RF[d["id"]].get("uf_oficial") and d["extras"]["uf"] != RF[d["id"]]["uf_oficial"]]
        sem = [d for d in ab if cnpj_digitos(d["extras"].get("cnpj")) and d["id"] not in RF]
        w(f"**Checagem com a Receita** ({fonte_rf}), negócios abertos: situação cadastral diferente de Ativa: {len(sit)} · "
          f"UF do HubSpot diferente da oficial: {len(uf)} · CNPJ preenchido e não localizado na base local: {len(sem)}.")
        det = [(nome(d, regras), "situação", r.get("situacao"), link(d)) for d, r in sit] + \
              [(nome(d, regras), "UF", f"HubSpot {d['extras']['uf']} × Receita {r['uf_oficial']}", link(d)) for d, r in uf]
        if det: w(tabela(["Negócio", "Divergência", "Detalhe", "Link"], det[:15]))
        w("Para gerar a proposta de atualização campo a campo: `crm_checagem.py --negocio <id>`.\n")
    # valor só é cobrado a partir da Apresentação de Diagnóstico (antes disso o valor ainda não existe)
    tem_valor = lambda d: d["valor"] > 0 or (d["prob"] or 0) < 0.4
    corrigir = sorted([d for d in ab if not cnpj_digitos(d["extras"].get("cnpj")) or not tem_valor(d) or not d["extras"].get("regime")], key=lambda d: -d["valor"])
    if corrigir:
        w(f"**Abertos para corrigir primeiro ({len(corrigir)}): sem CNPJ, sem regime ou sem valor já na Apresentação de Diagnóstico ou depois:**")
        w(tabela(["Negócio", "Etapa", "Falta", "Link"],
                 [(nome(d, regras), d["etapa"], ", ".join(x for x, ok in (("CNPJ", cnpj_digitos(d["extras"].get("cnpj"))), ("valor", tem_valor(d)), ("regime", d["extras"].get("regime"))) if not ok), link(d)) for d in corrigir[:15]]))
        if len(corrigir) > 15: w(f"_… e mais {len(corrigir) - 15}._\n")

    w("## 7. Principais clientes (por valor ganho desde 2024)")
    cli = defaultdict(list)
    for d in L:
        c = raiz_cnpj(d["extras"].get("cnpj")) if len(cnpj_digitos(d["extras"].get("cnpj"))) >= 8 else ""
        cli[c or "nome:" + (d["nome"] or "").strip().upper()].append(d)
    rank = []
    for k, X in cli.items():
        G = [d for d in X if d["status"] == "ganho"]
        if not G: continue
        rank.append((sum(d["valor"] for d in G), k, X, G))
    rank.sort(key=lambda r: -r[0])
    w(tabela(["Cliente (grupo pela raiz do CNPJ)", "Ganhos", "Valor ganho", "Último ganho", "Abertos hoje", "UF", "Regime"],
             [(nome(max(G, key=lambda d: d["valor"]), regras), len(G), brl(v), max(d["fechado"] for d in G),
               sum(1 for d in X if d["status"] == "aberto"), G[0]["extras"].get("uf") or "—", G[0]["extras"].get("regime") or "—") for v, k, X, G in rank[:10]]))
    w(f"Clientes distintos com ganho: {len(rank)}. Concentração: os 3 maiores somam "
      f"{pct(sum(r[0] for r in rank[:3]) / sum(r[0] for r in rank) * 100 if rank else None, 0)} do valor ganho.\n")

    if neg:
        d = neg
        w(f"## 8. Raio-x do negócio {d['id']}")
        e = d["extras"]; r = RF.get(d["id"], {})
        linhas = [("Negócio", nome(d, regras)), ("Pipeline / etapa", f"{d['pipeline']} / {d['etapa']} ({d['status']})"),
                  ("Criado / fechado", f"{d['criado']} / {d['fechado'] or '—'}"),
                  ("Idade", f"{(hoje - dt(d['criado'])).days} dias" if dt(d["criado"]) else "—"),
                  ("Dias sem atividade", dias_parado(d) if dias_parado(d) is not None else "nunca"),
                  ("Valor", brl(d["valor"])), ("Parceiro no card", d["parceiro_bruto"] or "—"),
                  ("CNPJ", e.get("cnpj") or "vazio"), ("UF / cidade", f"{e.get('uf') or '—'} / {e.get('cidade') or '—'}"),
                  ("Regime", e.get("regime") or "vazio"), ("Segmento / CNAE", d.get("segmento") or d.get("cnae") or "vazio"),
                  ("Faturamento / faixa", e.get("faturamento") or e.get("faixa_faturamento") or "vazio"),
                  ("e-CAC liberado", e.get("ecac") or "vazio"), ("Tese aplicada", d.get("variante") or "vazio"),
                  ("% honorário / % parceiro", f"{pct(d['hContr']*100,0) if d['hContr'] is not None else 'vazio'} / {pct(d['hParc']*100,0) if d['hParc'] is not None else 'vazio'}"),
                  ("Parecer RT / diagnóstico / parecer jurídico anexados", " / ".join("sim" if e.get(k) else "não" for k in ("parecer_rt", "anexo_diagnostico", "parecer_juridico"))),
                  ("Receita (oficial)", f"{r.get('situacao')} · CNAE {r.get('cnae_oficial')} {r.get('cnae_desc_oficial')} · {r.get('municipio_oficial')}/{r.get('uf_oficial')} · Simples: {r.get('simples') or 'não consta'}" if r else "não cruzado"),
                  ("Link", link(d))]
        w(tabela(["Campo", "Valor"], linhas))
        c = cnpj_digitos(e.get("cnpj"))
        if c:
            hist = [x for x in N if x["id"] != d["id"] and raiz_cnpj(x["extras"].get("cnpj")) == raiz_cnpj(c)]  # raiz com zeros recompostos
            w(f"**Outros negócios do mesmo grupo (raiz do CNPJ), qualquer parceiro: {len(hist)}**")
            w(tabela(["Negócio", "Parceiro", "Status / etapa", "Criado", "Fechado", "Valor", "Link"],
                     [(nome(x, regras), x["parceiro"], f"{x['status']} / {x['etapa']}", x["criado"], x["fechado"] or "—", brl(x["valor"]), link(x)) for x in sorted(hist, key=lambda x: x["criado"])]))
            outros = {x["parceiro"] for x in hist if x["parceiro"] != d["parceiro"] and x["status"] == "aberto"}
            if outros:
                w(f"⚠️ **Possível conflito de canal:** há negócio aberto do mesmo grupo com {', '.join(sorted(outros))}. Decisão do Luis.\n")

    w("## 9. Pendências abertas com este Master (controle da Daniela)")
    P = pendencias_do_master(master)
    if P is None:
        w("_Controle de pendências ainda não criado (`pendencias.py --criar`)._\n")
    else:
        w(tabela(["ID", "Aberta em", "Tipo", "Área", "Descrição", "Prazo", "Situação SLA"],
                 [(x.get("ID"), x.get("Data abertura"), x.get("Tipo de demanda"), x.get("Área de destino"), (x.get("Descrição") or "")[:60], x.get("Prazo solução"), x.get("Situação SLA")) for x in P]))

    w("## 10. Pontos de atenção (regras automáticas)")
    at = []
    if s and s.get("sinal") in ("🔴", "🟡"): at.append(f"Saúde {s['sinal']} ({s['saude']}/100): conversar sobre o ritmo de indicações e o que trava.")
    if not gJ: at.append(f"Nenhum ganho nos últimos {a.janela} dias.")
    if not criJ: at.append(f"Nenhum negócio novo nos últimos {a.janela} dias: risco de esfriamento.")
    if RF:
        susp = [d for d in ab if d["id"] in RF and (RF[d["id"]].get("situacao") or "").lower() != "ativa"]
        if susp: at.append(f"{len(susp)} aberto(s) com CNPJ em situação diferente de Ativa na Receita ({', '.join(sorted({RF[d['id']]['situacao'] for d in susp}))}): confirmar com o parceiro antes de seguir.")
    inv = [d for d in gan if dt(d["fechado"]) and dt(d["criado"]) and dt(d["fechado"]) < dt(d["criado"]) and str(d["extras"].get("importado")).lower() != "true"]
    if inv: at.append(f"{len(inv)} ganho(s) não importado(s) com data de fechamento anterior à de criação: conferir as datas no CRM.")
    if p90: at.append(f"{len(p90)} negócio(s) aberto(s) há mais de 90 dias sem atividade ({mi(sum(d['valor'] for d in p90))}): revisar com o parceiro se seguem vivos.")
    apres = [d for d in ab if "Apresenta" in (d["etapa"] or "") and (hoje - (dt(d["criado"]) or hoje)).days > 60]
    if apres: at.append(f"{len(apres)} negócio(s) em Apresentação de Diagnóstico criados há mais de 60 dias: agendar apresentação ou registrar motivo.")
    contr = [d for d in ab if "Contrato Enviado" in (d["etapa"] or "") or "Pendência Fechamento" in (d["etapa"] or "")]
    if contr: at.append(f"{len(contr)} negócio(s) em Contrato Enviado/Pendência de Fechamento ({mi(sum(d['valor'] for d in contr))}): checar com Contratos o que falta para assinar.")
    semh = [d for d in g12 if d["hContr"] is None]
    if semh: at.append(f"{len(semh)} ganho(s) em 12 meses sem % de honorário: impede calcular margem e repasse (corrigir no CRM).")
    if cafe_ab: at.append(f"{len(cafe_ab)} aberto(s) de café: fora da meta desde abr/2026.")
    if rank and len(rank) >= 3 and sum(r[0] for r in rank[:3]) / sum(r[0] for r in rank) > 0.6: at.append("Carteira concentrada: 3 clientes somam mais de 60% do valor ganho.")
    w("\n".join(f"- {x}" for x in at) if at else "- Nenhum alerta automático.")
    w("\n## 11. Para a Daniela completar antes da reunião")
    w("- Contexto do relacionamento (o que o parceiro reclamou ou elogiou nas últimas conversas).\n- Pendências que estão só no WhatsApp e ainda não entraram no controle.\n"
      "- Objetivo da conversa e o que **não** prometer (valor e êxito nunca são garantidos).")

    dest_dir = os.path.expanduser(a.saida) if a.saida else os.path.join(CELULA, "briefings")
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, f"{hoje.isoformat()}-briefing-{slug(master)}" + (f"-negocio-{neg['id']}" if neg else "") + ".md")
    open(dest, "w").write("\n".join(md) + "\n")
    print(dest)
    print("  ".join(f"{k}={(round(v) if isinstance(v, float) else v)}" for k, v in num.items()))


if __name__ == "__main__":
    main()
