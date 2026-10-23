#!/usr/bin/env python3
"""Conector com os Dados Abertos do CNPJ (Receita Federal) para a Inteligência Comercial.

Fonte oficial: arquivos mensais da Receita Federal (compartilhamento público do Serpro/RFB),
https://arquivos.receitafederal.gov.br/index.php/s/YggdBLfdninEJX9 (≈ 7,7 GB/mês, CSV ';' latin-1, sem cabeçalho).

Privacidade e disco:
- Tudo é processado LOCALMENTE: os CNPJs de clientes nunca saem do Mac.
- Um arquivo por vez: baixa → lê de dentro do zip (sem descompactar) → grava só o filtrado → apaga o zip.
- Sócios (dados pessoais) NÃO são baixados. Telefone e e-mail só entram com "incluir_contatos": true
  no receita-config.json, depois da validação de LGPD pelo Jurídico.
- Consulta externa (Minha Receita, para regime tributário) só com --permitir-externo e só para prospects.

Uso:
  python3 receita_cnpj.py --listar                      # meses disponíveis e tamanho
  python3 receita_cnpj.py --carregar [--mes AAAA-MM]    # baixa e monta a base local filtrada (≈ 30–90 min)
  python3 receita_cnpj.py --bairros [--mes AAAA-MM]     # acrescenta bairro e logradouro à base já carregada (≈ 30–60 min)
  python3 receita_cnpj.py --enriquecer                  # cruza com o HubSpot → cnae/porte/situação oficiais
  python3 receita_cnpj.py --prospectar                  # lista de empresas-alvo (sem clientes atuais)
  python3 receita_cnpj.py --regime --permitir-externo   # regime tributário (Minha Receita) para a lista de prospects
"""
import argparse, csv, glob, io, json, os, re, sqlite3, subprocess, sys, time, zipfile
from datetime import date
sys.path.insert(0, os.path.expanduser("~/agentes-partners/lib"))
from cnpj import normalizar_cnpj  # noqa: E402

SHARE = "YggdBLfdninEJX9"
DAV = "https://arquivos.receitafederal.gov.br/public.php/webdav"
BASE = os.path.expanduser("~/agentes-partners/dados-publicos/receita")
CFG = os.path.expanduser("~/agentes-partners/departamentos/intel-comercial/receita-config.json")
HS = os.path.expanduser("~/agentes-partners/departamentos/resultados/dados")
OUT = os.path.expanduser("~/agentes-partners/departamentos/intel-comercial")
DB = os.path.join(BASE, "cnpj.sqlite")
PORTE = {"00": "Não informado", "01": "Microempresa", "03": "Empresa de Pequeno Porte", "05": "Demais"}
SITUACAO = {"01": "Nula", "02": "Ativa", "03": "Suspensa", "04": "Inapta", "08": "Baixada"}

CFG_PADRAO = {
    "_leia_me": "Filtros da base local. cnaes_alvo aceita prefixos (ex.: '1011' pega todos os CNAEs que começam com 1011). "
                "Empresas dos CNPJs do HubSpot entram sempre, com qualquer CNAE. Sem o dado de contato até validação LGPD.",
    "cnaes_alvo": ["1011", "1012", "1013", "4634", "4722", "4633", "4724", "1031", "1033", "4637", "4639", "4711", "4712", "1091", "4721", "105", "4631", "6920"],
    "ufs_alvo": ["MG", "SP", "ES", "GO", "DF", "PR", "SC", "BA", "RJ", "MT", "MS", "TO"],
    "somente_ativas": True,
    "excluir_simples_na_prospeccao": True,
    "portes_prospeccao": ["Demais", "Empresa de Pequeno Porte"],
    "capital_social_minimo_prospeccao": 100000,
    "capital_social_maximo_prospeccao": 500000000,
    "incluir_contatos": False
}


def cfg():
    os.makedirs(os.path.dirname(CFG), exist_ok=True)
    if not os.path.exists(CFG): json.dump(CFG_PADRAO, open(CFG, "w"), ensure_ascii=False, indent=1)
    c = json.load(open(CFG)); return {**CFG_PADRAO, **c}


def propfind(path=""):
    r = subprocess.run(["curl", "-sS", "-m", "60", "-X", "PROPFIND", "-u", f"{SHARE}:", "-H", "Depth: 1", f"{DAV}/{path}"], capture_output=True, text=True, check=True)
    itens = []
    for it in re.findall(r"<d:response>(.*?)</d:response>", r.stdout, re.S)[1:]:
        h = re.search(r"<d:href>([^<]+)</d:href>", it).group(1).rstrip("/").rsplit("/", 1)[-1]
        s = re.search(r"<d:getcontentlength>(\d+)", it)
        itens.append((h, int(s.group(1)) if s else 0))
    return itens


def ultimo_mes():
    return sorted(h for h, _ in propfind() if re.fullmatch(r"\d{4}-\d{2}", h))[-1]


def baixar(mes, nome, destino):
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    subprocess.run(["curl", "-sS", "-f", "-L", "--retry", "3", "-C", "-", "-u", f"{SHARE}:", "-o", destino, f"{DAV}/{mes}/{nome}"], check=True)


def linhas(zip_path):
    with zipfile.ZipFile(zip_path) as z:
        for n in z.namelist():
            with z.open(n) as f:
                yield from csv.reader(io.TextIOWrapper(f, encoding="latin-1", newline=""), delimiter=";", quotechar='"')


def cnpjs_hubspot():
    fs = sorted(glob.glob(os.path.join(HS, "hubspot-*.json")))
    if not fs: return set()
    return {normalizar_cnpj(d["extras"].get("cnpj")).zfill(14) for d in json.load(open(fs[-1]))["negocios"] if d["extras"].get("cnpj")}


def carregar(mes):
    c = cfg(); mes = mes or ultimo_mes(); tmp = os.path.join(BASE, "tmp"); os.makedirs(tmp, exist_ok=True)
    arquivos = dict(propfind(mes + "/"))
    alvo = tuple(c["cnaes_alvo"]); ufs = set(c["ufs_alvo"]); hs = cnpjs_hubspot(); hs_basico = {x[:8] for x in hs}
    con = sqlite3.connect(DB + ".novo"); cur = con.cursor()
    cur.executescript("""
      CREATE TABLE estab(cnpj TEXT PRIMARY KEY, basico TEXT, matriz TEXT, fantasia TEXT, situacao TEXT, data_situacao TEXT, inicio TEXT,
        cnae TEXT, cnaes_sec TEXT, uf TEXT, municipio TEXT, cep TEXT, telefone TEXT, email TEXT, no_hubspot INTEGER, bairro TEXT, logradouro TEXT);
      CREATE TABLE empresa(basico TEXT PRIMARY KEY, razao TEXT, natureza TEXT, capital REAL, porte TEXT);
      CREATE TABLE simples(basico TEXT PRIMARY KEY, simples TEXT, mei TEXT);
      CREATE TABLE ref(tipo TEXT, codigo TEXT, descricao TEXT, PRIMARY KEY(tipo, codigo));
      CREATE TABLE meta(chave TEXT PRIMARY KEY, valor TEXT);""")
    def etapa(nome):
        z = os.path.join(tmp, nome); print(f"  baixando {nome} ({arquivos.get(nome, 0)/1e6:.0f} MB)…", flush=True); baixar(mes, nome, z); return z
    for tipo, nome in (("cnae", "Cnaes.zip"), ("municipio", "Municipios.zip"), ("natureza", "Naturezas.zip")):
        z = etapa(nome); cur.executemany("INSERT OR REPLACE INTO ref VALUES (?,?,?)", ((tipo, r[0], r[1]) for r in linhas(z) if len(r) >= 2)); os.remove(z)
    basicos = set()
    for i in range(10):
        z = etapa(f"Estabelecimentos{i}.zip"); n = 0
        for r in linhas(z):
            if len(r) < 28: continue
            cnpj = r[0] + r[1] + r[2]; em_hs = cnpj in hs or r[0] in hs_basico
            if not em_hs:
                if c["somente_ativas"] and r[5] != "02": continue
                if r[19] not in ufs or not r[11].startswith(alvo): continue
            tel = (f"({r[21]}) {r[22]}" if r[22] else "") if c["incluir_contatos"] else None
            mail = (r[27] or None) if c["incluir_contatos"] else None
            cur.execute("INSERT OR REPLACE INTO estab VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (cnpj, r[0], "matriz" if r[3] == "1" else "filial", r[4], SITUACAO.get(r[5], r[5]), r[6], r[10], r[11], r[12], r[19], r[20], r[18], tel, mail, int(em_hs),
                         r[17].strip() or None, " ".join(x for x in (r[13], r[14], r[15]) if x).strip() or None))
            basicos.add(r[0]); n += 1
        con.commit(); os.remove(z); print(f"    {n} estabelecimentos mantidos", flush=True)
    for i in range(10):
        z = etapa(f"Empresas{i}.zip")
        cur.executemany("INSERT OR REPLACE INTO empresa VALUES (?,?,?,?,?)",
                        ((r[0], r[1], r[2], float((r[4] or "0").replace(".", "").replace(",", ".") or 0), PORTE.get(r[5], r[5])) for r in linhas(z) if len(r) >= 6 and r[0] in basicos))
        con.commit(); os.remove(z)
    z = etapa("Simples.zip")
    cur.executemany("INSERT OR REPLACE INTO simples VALUES (?,?,?)", ((r[0], r[1], r[4]) for r in linhas(z) if len(r) >= 5 and r[0] in basicos))
    os.remove(z)
    cur.executemany("INSERT OR REPLACE INTO meta VALUES (?,?)", [("mes_referencia", mes), ("carregado_em", time.strftime("%Y-%m-%d %H:%M")), ("fonte", f"Receita Federal · Dados Abertos do CNPJ · {mes}")])
    con.commit(); con.close(); os.replace(DB + ".novo", DB)
    print(f"Base local pronta: {DB} ({os.path.getsize(DB)/1e6:.0f} MB) · referência {mes}")


def bairros(mes=None):
    """Acrescenta bairro e logradouro (endereço da empresa, dado público do CNPJ) aos estabelecimentos já carregados,
    sem refazer a base. Usa o mesmo mês da base, para manter tudo coerente. Um arquivo por vez, apagado depois."""
    con = sqlite3.connect(DB); cur = con.cursor()
    cols = {r[1] for r in cur.execute("PRAGMA table_info(estab)")}
    for c in ("bairro", "logradouro"):
        if c not in cols: cur.execute(f"ALTER TABLE estab ADD COLUMN {c} TEXT")
    mes = mes or dict(cur.execute("SELECT chave, valor FROM meta")).get("mes_referencia") or ultimo_mes()
    tem = {r[0] for r in cur.execute("SELECT cnpj FROM estab")}
    arquivos = dict(propfind(mes + "/")); tmp = os.path.join(BASE, "tmp"); os.makedirs(tmp, exist_ok=True)
    for i in range(10):
        nome = f"Estabelecimentos{i}.zip"; z = os.path.join(tmp, nome)
        print(f"  baixando {nome} ({arquivos.get(nome, 0)/1e6:.0f} MB)…", flush=True); baixar(mes, nome, z); n = 0
        for r in linhas(z):
            if len(r) < 21: continue
            cnpj = r[0] + r[1] + r[2]
            if cnpj in tem:
                cur.execute("UPDATE estab SET bairro=?, logradouro=? WHERE cnpj=?",
                            (r[17].strip() or None, " ".join(x for x in (r[13], r[14], r[15]) if x).strip() or None, cnpj)); n += 1
        con.commit(); os.remove(z); print(f"    {n} endereços atualizados", flush=True)
    cur.execute("INSERT OR REPLACE INTO meta VALUES ('bairros_carregados_em', ?)", (time.strftime("%Y-%m-%d %H:%M"),)); con.commit(); con.close()
    print("Bairros prontos.")


def q(sql, args=()):
    con = sqlite3.connect(DB); con.row_factory = sqlite3.Row
    return [dict(r) for r in con.execute(sql, args)]


VIEW = """SELECT e.cnpj, e.matriz, COALESCE(m.razao,'') razao, e.fantasia, e.situacao, e.inicio, e.cnae, c.descricao cnae_desc, e.cnaes_sec,
          e.uf, mu.descricao municipio, m.capital, m.porte, s.simples, s.mei, e.telefone, e.email, e.no_hubspot
          FROM estab e LEFT JOIN empresa m ON m.basico=e.basico LEFT JOIN simples s ON s.basico=e.basico
          LEFT JOIN ref c ON c.tipo='cnae' AND c.codigo=e.cnae LEFT JOIN ref mu ON mu.tipo='municipio' AND mu.codigo=e.municipio"""


def enriquecer():
    meta = {r["chave"]: r["valor"] for r in q("SELECT * FROM meta")}
    rows = {r["cnpj"]: r for r in q(VIEW + " WHERE e.no_hubspot=1")}
    f = sorted(glob.glob(os.path.join(HS, "hubspot-*.json")))[-1]; N = json.load(open(f))["negocios"]
    out, achados = {}, 0
    for d in N:
        cn = normalizar_cnpj(d["extras"].get("cnpj")).zfill(14)
        r = rows.get(cn)
        if r:
            achados += 1
            out[d["id"]] = {"cnpj": cn, "cnae_oficial": r["cnae"], "cnae_desc_oficial": r["cnae_desc"], "situacao": r["situacao"], "porte_oficial": r["porte"],
                            "capital_social": r["capital"], "simples": r["simples"], "uf_oficial": r["uf"], "municipio_oficial": r["municipio"], "inicio_atividade": r["inicio"]}
    dest = os.path.join(OUT, "enriquecimento-receita.json")
    json.dump({"fonte": meta.get("fonte"), "gerado_em": time.strftime("%Y-%m-%d %H:%M"), "negocios": out}, open(dest, "w"), ensure_ascii=False)
    print(f"{achados} de {len(N)} negócios do HubSpot enriquecidos com dado oficial → {dest}")


def prospectar(cnaes=None):
    """Lista de empresas-alvo, sem clientes nem empresas já no HubSpot, pontuada pelo lead scoring do ICP
    e com sugestão do Master que já ganha naquele segmento × UF."""
    c = cfg(); ex_simples = "AND COALESCE(s.simples,'N')<>'S'" if c["excluir_simples_na_prospeccao"] else ""
    portes = ",".join("?" * len(c["portes_prospeccao"]))
    hs_basicos = {x[:8] for x in cnpjs_hubspot()}
    rows = q(VIEW + f" WHERE e.no_hubspot=0 AND e.situacao='Ativa' AND e.matriz='matriz' {ex_simples} AND m.porte IN ({portes}) AND COALESCE(m.capital,0) >= ?",
             (*c["portes_prospeccao"], c["capital_social_minimo_prospeccao"]))
    rows = [r for r in rows if r["cnpj"][:8] not in hs_basicos and (not cnaes or r["cnae"].startswith(tuple(cnaes)))]
    import importlib.util
    spec = importlib.util.spec_from_file_location("icp", os.path.join(os.path.dirname(os.path.abspath(__file__)), "icp.py"))
    icp = importlib.util.module_from_spec(spec); spec.loader.exec_module(icp)
    fs = sorted(glob.glob(os.path.join(OUT, "icp-*.json"))); I = json.load(open(fs[-1])) if fs else {"lead_scoring_v1": [], "territorio": []}
    pts = {(x["dimensao"], x["valor"]): x["pontos_sugeridos"] for x in I["lead_scoring_v1"]}
    terr = {}
    for t in I["territorio"]:
        k = (t["segmento"], t["uf"])
        if k not in terr or t["valor_ganho_12m_rs"] > terr[k][1]: terr[k] = (t["parceiro"], t["valor_ganho_12m_rs"])
    for r in rows:
        r["segmento"] = icp.grupo(r["cnae_desc"])
        r["score_parcial"] = pts.get(("Segmento", r["segmento"]), 0) + pts.get(("UF", r["uf"]), 0)
        r["master_sugerido"] = terr.get((r["segmento"], r["uf"]), ("a definir (sem histórico no segmento × UF)", 0))[0]
    teto = c.get("capital_social_maximo_prospeccao")
    estrategicas = [r for r in rows if teto and (r["capital"] or 0) > teto]
    rows = [r for r in rows if not (teto and (r["capital"] or 0) > teto)]
    rows.sort(key=lambda r: (-r["score_parcial"], -(r["capital"] or 0)))
    estrategicas.sort(key=lambda r: -(r["capital"] or 0))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Prospects"
    cols = ["score_parcial", "master_sugerido", "segmento", "cnpj", "razao", "fantasia", "cnae", "cnae_desc", "uf", "municipio", "porte", "capital", "simples", "inicio"] + (["telefone", "email"] if c["incluir_contatos"] else [])
    ws.append(cols + ["regime_tributario (Minha Receita)"])
    for j in range(1, len(cols) + 2): ws.cell(row=1, column=j).font = Font(bold=True, color="FFFFFF"); ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="1F1F1F")
    for r in rows: ws.append([r[k] for k in cols] + [""])
    ws.freeze_panes = "A2"
    we = wb.create_sheet("Contas estratégicas"); we.append(cols)
    for r in estrategicas: we.append([r[k] for k in cols])
    resumo = wb.create_sheet("Resumo"); resumo.append(["segmento", "uf", "prospects", "capital social somado (R$)", "master sugerido"])
    from collections import defaultdict
    ag = defaultdict(lambda: [0, 0.0, ""])
    for r in rows: a = ag[(r["segmento"], r["uf"])]; a[0] += 1; a[1] += r["capital"] or 0; a[2] = r["master_sugerido"]
    for (sg, uf), (n, cap, m) in sorted(ag.items(), key=lambda x: -x[1][0]): resumo.append([sg, uf, n, round(cap), m])
    meta = {x["chave"]: x["valor"] for x in q("SELECT * FROM meta")}
    ws2 = wb.create_sheet("Fonte e regras")
    for t in [meta.get("fonte"), "Filtros: " + json.dumps({k: c[k] for k in c if not k.startswith("_")}, ensure_ascii=False) + (f" · CNAEs desta lista: {cnaes}" if cnaes else ""),
              "Score parcial = pontos do segmento + pontos da UF (lead scoring v1, pendente de validação da Jenifer e do gerente de Inteligência). O regime tributário entra depois (Minha Receita, com autorização).",
              "Clientes e empresas já no HubSpot foram excluídos. Uso sujeito à validação de LGPD (legítimo interesse, opt-out, envio pelo parceiro). Sem dados de sócios."]:
        ws2.append([t])
    dest = os.path.join(OUT, f"Prospects-Receita-{date.today().isoformat()}.xlsx"); wb.save(dest); print(dest, len(rows), "prospects +", len(estrategicas), "contas estratégicas")
    return rows


def regime(permitir, somente_com_master=True):
    """Regime tributário (Minha Receita) para os prospects priorizados; soma os pontos de regime ao score e reordena.
    Nunca consulta CNPJ que esteja no HubSpot (clientes/negócios existentes)."""
    if not permitir: sys.exit("Consulta externa bloqueada: use --permitir-externo (envia CNPJs de prospects ao Minha Receita). Nunca use para a lista de clientes.")
    f = sorted(glob.glob(os.path.join(OUT, "Prospects-Receita-*.xlsx")))[-1]
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(f); ws = wb["Prospects"]
    head = [c.value for c in ws[1]]; linhas_ = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    iC, iM, iR, iS = head.index("cnpj"), head.index("master_sugerido"), head.index("regime_tributario (Minha Receita)"), head.index("score_parcial")
    cache_p = os.path.join(BASE, "regime-cache.json"); cache = json.load(open(cache_p)) if os.path.exists(cache_p) else {}
    fs = sorted(glob.glob(os.path.join(OUT, "icp-*.json"))); I = json.load(open(fs[-1]))
    pts = {x["valor"].upper(): x["pontos_sugeridos"] for x in I["lead_scoring_v1"] if x["dimensao"] == "Regime"}
    hs = cnpjs_hubspot(); feitos = 0
    alvo = [l for l in linhas_ if (not somente_com_master or not str(l[iM]).startswith("a definir"))]
    print(f"{len(alvo)} prospects para consultar (cache: {sum(1 for l in alvo if str(l[iC]) in cache)})", flush=True)
    for l in alvo:
        cn = str(l[iC])
        if cn in hs: continue  # proteção: nunca consultar cliente/negócio existente
        if cn not in cache:
            r = subprocess.run(["curl", "-sS", "-m", "20", f"https://minhareceita.org/{cn}"], capture_output=True, text=True)
            try:
                rt = sorted(json.loads(r.stdout).get("regime_tributario") or [], key=lambda x: x.get("ano", 0))
                cache[cn] = f"{rt[-1]['forma_de_tributacao']} ({rt[-1]['ano']})" if rt else "sem registro"
            except Exception: cache[cn] = "erro na consulta"
            feitos += 1; time.sleep(1.0)
            if feitos % 100 == 0: json.dump(cache, open(cache_p, "w")); print(f"  {feitos} consultados", flush=True)
        l[iR] = cache[cn]
    json.dump(cache, open(cache_p, "w"))
    if "score_total" not in head: head.insert(0, "score_total"); linhas_ = [[None] + l for l in linhas_]; iC += 1; iM += 1; iR += 1; iS += 1
    for l in linhas_:
        reg = str(l[iR] or "").upper()
        rp = next((v for k, v in pts.items() if reg.startswith(k)), 0) if reg and not reg.startswith(("SEM", "ERRO")) else 0
        l[0] = (l[iS] or 0) + rp
    linhas_.sort(key=lambda l: (-(l[0] or 0), -(l[head.index("capital")] or 0)))
    del wb["Prospects"]; ws = wb.create_sheet("Prospects", 0); ws.append(head)
    for j in range(1, len(head) + 1): ws.cell(row=1, column=j).font = Font(bold=True, color="FFFFFF"); ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="1F1F1F")
    for l in linhas_: ws.append(l)
    ws.freeze_panes = "A2"
    wf = wb["Fonte e regras"]; wf.append([f"Regime tributário: Minha Receita (código aberto sobre dados da Receita), consultado em {date.today().isoformat()} só para os prospects com Master sugerido. score_total = score_parcial + pontos do regime."])
    wb.save(f)
    from collections import Counter
    print(f, "regime preenchido ·", Counter(str(cache.get(str(l[iC]), "")).split(" (")[0] for l in linhas_ if str(l[iC]) in cache).most_common(6))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); g = ap.add_mutually_exclusive_group(required=True)
    for o in ("listar", "carregar", "bairros", "enriquecer", "prospectar", "regime"): g.add_argument("--" + o, action="store_true")
    ap.add_argument("--mes"); ap.add_argument("--permitir-externo", action="store_true")
    ap.add_argument("--cnaes", help="prefixos separados por vírgula para restringir a lista (ex.: 1011,1012,1013,4634,4633)"); a = ap.parse_args()
    if a.listar:
        for m in sorted(h for h, _ in propfind() if re.fullmatch(r"\d{4}-\d{2}", h))[-3:]:
            print(m, f"{sum(s for _, s in propfind(m + '/'))/1e9:.2f} GB")
    elif a.carregar: carregar(a.mes)
    elif a.bairros: bairros(a.mes)
    elif a.enriquecer: enriquecer()
    elif a.prospectar: prospectar(a.cnaes.split(",") if a.cnaes else None)
    else: regime(a.permitir_externo)
