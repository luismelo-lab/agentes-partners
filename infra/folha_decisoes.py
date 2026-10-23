#!/usr/bin/env python3
"""Folha de decisões do Luis com campos de seleção (Excel, padrão Save) e leitura das respostas.

Fonte: docs/revisao-2026-10-22/decisoes.json (um item = uma decisão). Pedido do Luis (22/10/2026): pendências em campos para aprovar ou reprovar.

  python3 infra/folha_decisoes.py gerar  [--saida <arquivo.xlsx>]   padrão: 04 · Gestão/Central de Agentes/Decisoes-Pendentes-<data>.xlsx
  python3 infra/folha_decisoes.py ler    <arquivo.xlsx>             grava registro/decisoes-respostas.json e mostra o que foi respondido

Ao gerar de novo, as respostas já dadas no arquivo anterior são mantidas (mesmo id).
"""
import glob, json, os, re, sys
from datetime import date
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); AP = os.path.expanduser("~/agentes-partners")
FONTE = os.path.join(REPO, "docs/revisao-2026-10-22/decisoes.json")
OPCOES = ["Aprovar", "Aprovar com ajuste", "Reprovar", "Decidir depois"]
F = "IBM Plex Sans"; OURO = "F3CD60"; OURO_T = "B58A1F"; OURO_S = "FBF1D7"; TX = "0F1115"; TX2 = "4A4F58"; BORDA = "E6E8EC"
COLS = [("Nº", 6), ("Bloco", 15), ("Tema", 18), ("Decisão proposta", 58), ("Por que importa", 34), ("Prazo", 13),
        ("Sua decisão", 20), ("Ajuste ou comentário", 40), ("Detalhe em", 30), ("Quem aplica", 24), ("Se não responder", 26)]


def respostas_anteriores(f):
    if not f or not os.path.exists(f): return {}
    wb = openpyxl.load_workbook(f, data_only=True)
    if "Decisões" not in wb.sheetnames: return {}
    return {str(r[0]).strip(): (r[6], r[7]) for r in wb["Decisões"].iter_rows(min_row=3, values_only=True) if r and len(r) > 7 and r[0] is not None}


def folhas_anteriores(saida):
    """Folhas cujas respostas a folha nova herda, da mais antiga para a mais nova: as Decisoes-Pendentes-AAAA-MM-DD.xlsx da
    pasta com data até a da saída, e a própria saída se já existe. Antes só lia a própria saída, que leva a data de hoje: na
    sexta seguinte a folha nova saía em branco (aconteceu em 23/10: as respostas 8, 11 e 14 da folha de 22/10 sumiram)."""
    pasta, nome = os.path.dirname(os.path.abspath(saida)), os.path.basename(saida)
    rx = re.compile(r"Decisoes-Pendentes-(\d{4}-\d{2}-\d{2})\.xlsx")
    m = rx.fullmatch(nome); teto = m.group(1) if m else "9999-99-99"
    fs = sorted(f for f in glob.glob(os.path.join(pasta, "Decisoes-Pendentes-*.xlsx"))
                if rx.fullmatch(os.path.basename(f)) and rx.fullmatch(os.path.basename(f)).group(1) <= teto)
    if os.path.exists(saida) and os.path.abspath(saida) not in map(os.path.abspath, fs): fs.append(saida)
    return fs


def respostas_herdadas(saida):
    """id → (decisão, comentário); a resposta mais nova vence, mas célula em branco não apaga a resposta de antes."""
    out = {}
    for f in folhas_anteriores(saida):
        for i, (dec, com) in respostas_anteriores(f).items():
            ant = out.get(i, (None, None))
            out[i] = (dec if dec not in (None, "") else ant[0], com if com not in (None, "") else ant[1])
    return out


def gerar(saida):
    D = json.load(open(FONTE))["itens"]; ant = respostas_herdadas(saida)
    print("respostas herdadas de", ", ".join(os.path.basename(f) for f in folhas_anteriores(saida)) or "nenhuma folha anterior")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Decisões"; ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
    c = ws.cell(1, 1, f"Decisões pendentes · escolha na coluna 'Sua decisão' e comente quando precisar · {date.today():%d/%m/%Y}")
    c.font = Font(name=F, bold=True, size=13, color=TX); c.fill = PatternFill("solid", fgColor=OURO); ws.row_dimensions[1].height = 26
    for j, (t, w) in enumerate(COLS, 1):
        c = ws.cell(2, j, t); c.font = Font(name=F, bold=True, size=10, color=TX); c.fill = PatternFill("solid", fgColor=OURO_S)
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=Side(style="medium", color=OURO))
        ws.column_dimensions[get_column_letter(j)].width = w
    fino = Side(style="thin", color=BORDA)
    for i, d in enumerate(D, 3):
        dec, com = ant.get(str(d["id"]).strip(), (None, None))
        vals = [d["id"], d["bloco"], d["tema"], d["decisao"], d["porque"], d["prazo"], dec, com, d["onde"], d["dono_aplicacao"],
                f"Vale a recomendação em {d['prazo']}" if d.get("padrao") else "Continua pendente"]  # infra/alcadas.md
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = Font(name=F, size=10, color=TX2 if j in (9, 10, 11) else TX, bold=(j == 7))
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=fino)
        ws.cell(i, 7).fill = PatternFill("solid", fgColor="FFFFFF"); ws.cell(i, 7).border = Border(left=Side(style="thin", color=OURO), right=Side(style="thin", color=OURO), bottom=fino)
        ws.row_dimensions[i].height = 60
    fim = len(D) + 2
    dv = DataValidation(type="list", formula1='"' + ",".join(OPCOES) + '"', allow_blank=True, showDropDown=False)
    dv.promptTitle = "Sua decisão"; dv.prompt = "Escolha: " + " · ".join(OPCOES); ws.add_data_validation(dv); dv.add(f"G3:G{fim}")
    for txt, cor in (("Aprovar", "DCEFE0"), ("Aprovar com ajuste", "FBF1D7"), ("Reprovar", "FBE3E2"), ("Decidir depois", "F1F2F5")):
        ws.conditional_formatting.add(f"A3:K{fim}", FormulaRule(formula=[f'$G3="{txt}"'], fill=PatternFill("solid", fgColor=cor)))
    ws.conditional_formatting.add(f"B3:B{fim}", FormulaRule(formula=['$B3="Trava a operação"'], font=Font(name=F, bold=True, color=OURO_T)))
    ws.freeze_panes = "D3"; ws.auto_filter.ref = f"A2:K{fim}"
    r = wb.create_sheet("Resumo"); r.sheet_view.showGridLines = False; r.column_dimensions["A"].width = 30; r.column_dimensions["B"].width = 12
    linhas = [("Resumo das respostas", None), ("Total de decisões", f"=COUNTA(Decisões!A3:A{fim})")] + \
             [(o, f'=COUNTIF(Decisões!G3:G{fim},"{o}")') for o in OPCOES] + [("Sem resposta", f'=COUNTBLANK(Decisões!G3:G{fim})'),
             ("Que travam a operação, sem resposta", f'=COUNTIFS(Decisões!B3:B{fim},"Trava a operação",Decisões!G3:G{fim},"")')]
    for i, (a, b) in enumerate(linhas, 1):
        ca = r.cell(i, 1, a); cb = r.cell(i, 2, b); ca.font = Font(name=F, bold=(i == 1), size=12 if i == 1 else 11, color=TX); cb.font = Font(name=F, size=11, bold=True)
    u = wb.create_sheet("Como usar"); u.sheet_view.showGridLines = False; u.column_dimensions["A"].width = 120
    for i, t in enumerate(["Como usar",
        "1. Na aba Decisões, escolha na coluna 'Sua decisão': Aprovar, Aprovar com ajuste, Reprovar ou Decidir depois.",
        "2. Em 'Aprovar com ajuste' e 'Reprovar', escreva o motivo ou o ajuste em 'Ajuste ou comentário'. Para os itens que pedem um dado (ex.: telefone da Canex), escreva o dado no comentário.",
        "3. Salve o arquivo. Os agentes leem as respostas (infra/folha_decisoes.py ler) e aplicam cada uma no material de origem, pela área da coluna 'Quem aplica'.",
        "4. A consolidação de sexta, 16h, gera a folha de novo: mantém as suas respostas, tira o que foi aplicado e acrescenta as decisões novas.",
        "5. Coluna 'Se não responder': item reversível e dentro das regras vale a recomendação no prazo (infra/alcadas.md). Você pode reverter depois marcando Reprovar."], 1):
        c = u.cell(i, 1, t); c.font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11, color=TX); c.alignment = Alignment(wrap_text=True)
    wb.save(saida); print("ok", saida, len(D), "decisões")


def ler(f):
    D = {d["id"]: d for d in json.load(open(FONTE))["itens"]}; resp = []
    for i, (dec, com) in respostas_anteriores(f).items():
        if dec: resp.append({"id": i, "decisao": D.get(i, {}).get("decisao"), "resposta": dec, "comentario": com, "quem_aplica": D.get(i, {}).get("dono_aplicacao")})
    out = os.environ.get("DECISOES_RESPOSTAS") or os.path.join(AP, "registro", "decisoes-respostas.json")  # variável: teste sem tocar o registro
    json.dump({"lido_em": date.today().isoformat(), "arquivo": f, "respostas": resp}, open(out, "w"), ensure_ascii=False, indent=1)
    for r in resp: print(f"{r['id']:>4}  {r['resposta']:<20} {str(r['comentario'] or '')[:60]}")
    print(f"{len(resp)} respondidas de {len(D)} · {out}")


if __name__ == "__main__":
    a = sys.argv[1:] or ["gerar"]
    if a[0] == "gerar":
        sys.path.insert(0, os.path.join(REPO, "lib")); import caminhos
        saida = a[2] if len(a) > 2 and a[1] == "--saida" else os.path.join(caminhos.luis("central"), f"Decisoes-Pendentes-{date.today()}.xlsx")
        gerar(saida)
    elif a[0] == "ler" and len(a) > 1: ler(a[1])
    else: print(__doc__)
