#!/usr/bin/env python3
"""Avaliador mínimo de fórmulas do Excel (para conferir a calculadora sem LibreOffice).

Cobre o subconjunto usado em calculadora_maquila.py: + - * / ^ &, comparações, %, referências
com ou sem aba ('Aba'!$B$3), intervalos, IF, IFERROR, MAX, MIN, SUM, AND, OR, ABS, ROUND,
VLOOKUP (correspondência exata) e COUNTIF simples (">0"). Só lê o arquivo. Não altera nada.
"""
import re
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

TOK = re.compile(r"""\s*(?:
 (?P<str>"(?:[^"]|"")*")|
 (?P<ref>(?:'(?P<sh>[^']+)'!)?\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?)(?![A-Za-z_(])|
 (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)|
 (?P<name>[A-Za-z_][A-Za-z0-9_.]*)|
 (?P<op><>|<=|>=|[-+*/^&=<>(),%])
)""", re.X)


class Erro(Exception):
    pass


class Avaliador:
    def __init__(self, caminho):
        self.wb = openpyxl.load_workbook(caminho)  # fórmulas
        self.cache = {}

    def valor(self, aba, celula):
        chave = (aba, celula)
        if chave in self.cache:
            return self.cache[chave]
        v = self.wb[aba][celula].value
        if isinstance(v, str) and v.startswith("="):
            v = self.formula(v[1:], aba)
        self.cache[chave] = v
        return v

    # ---------- parser
    def formula(self, f, aba):
        # guarda o estado do parser: a avaliação de uma referência chama formula() de novo (reentrância)
        estado = (getattr(self, "toks", None), getattr(self, "p", 0), getattr(self, "aba", None))
        try:
            return self._formula(f, aba)
        finally:
            self.toks, self.p, self.aba = estado

    def _formula(self, f, aba):
        self.toks = [m for m in TOK.finditer(f) if m.group(0).strip()]
        resto = TOK.sub("", f).strip()
        if resto:
            raise Erro(f"não reconhecido: {resto!r} em {f}")
        self.p, self.aba = 0, aba
        v = self.comp()
        if self.p != len(self.toks):
            raise Erro(f"sobra na fórmula {f}")
        return v

    def peek(self):
        return self.toks[self.p] if self.p < len(self.toks) else None

    def op(self, *ops):
        t = self.peek()
        if t is not None and t.group("op") in ops:
            self.p += 1
            return t.group("op")
        return None

    def comp(self):
        a = self.concat()
        while True:
            o = self.op("=", "<>", "<", ">", "<=", ">=")
            if not o:
                return a
            b = self.concat()
            a = {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[o]

    def concat(self):
        a = self.add()
        while self.op("&"):
            b = self.add()
            a = f"{fmt(a)}{fmt(b)}"
        return a

    def add(self):
        a = self.mul()
        while True:
            o = self.op("+", "-")
            if not o:
                return a
            b = self.mul()
            a = num(a) + num(b) if o == "+" else num(a) - num(b)

    def mul(self):
        a = self.pot()
        while True:
            o = self.op("*", "/")
            if not o:
                return a
            b = self.pot()
            if o == "*":
                a = num(a) * num(b)
            else:
                if num(b) == 0:
                    raise ZeroDivisionError
                a = num(a) / num(b)

    def pot(self):
        a = self.un()
        while self.op("^"):
            a = num(a) ** num(self.un())
        return a

    def un(self):
        if self.op("-"):
            return -num(self.un())
        if self.op("+"):
            return num(self.un())
        v = self.prim()
        if self.op("%"):
            v = num(v) / 100
        return v

    def args(self):
        out = []
        if self.op(")"):
            return out
        while True:
            ini = self.p
            out.append(ini)
            self.pular()
            if self.op(")"):
                return out
            if not self.op(","):
                raise Erro("esperava , ou )")

    def pular(self):
        nivel = 0
        while self.p < len(self.toks):
            o = self.toks[self.p].group("op")
            if o == "(":
                nivel += 1
            elif o == ")":
                if nivel == 0:
                    return
                nivel -= 1
            elif o == "," and nivel == 0:
                return
            self.p += 1

    def em(self, pos):
        guarda = self.p
        self.p = pos
        try:
            v = self.comp()
        finally:
            self.p = guarda
        return v

    def prim(self):
        t = self.peek()
        if t is None:
            raise Erro("fim inesperado")
        self.p += 1
        if t.group("str"):
            return t.group("str")[1:-1].replace('""', '"')
        if t.group("num"):
            return float(t.group("num"))
        if t.group("ref"):
            aba = t.group("sh") or self.aba
            if t.group("c2"):
                return Intervalo(self, aba, t.group("c1"), int(t.group("r1")), t.group("c2"), int(t.group("r2")))
            return self.valor(aba, f"{t.group('c1')}{t.group('r1')}")
        if t.group("op") == "(":
            v = self.comp()
            if not self.op(")"):
                raise Erro("falta )")
            return v
        nome = t.group("name")
        if nome is None:
            raise Erro(f"token inesperado {t.group(0)}")
        if nome.upper() in ("TRUE", "FALSE") and not (self.peek() and self.peek().group("op") == "("):
            return nome.upper() == "TRUE"
        if not self.op("("):
            raise Erro(f"nome sem função: {nome}")
        pos = self.args()
        fim = self.p
        r = self.func(nome.upper(), pos)
        self.p = fim
        return r

    def func(self, nome, pos):
        E = self.em
        if nome == "IF":
            c = E(pos[0])
            if c:
                return E(pos[1])
            return E(pos[2]) if len(pos) > 2 else False
        if nome == "IFERROR":
            try:
                return E(pos[0])
            except (ZeroDivisionError, Erro, TypeError):
                return E(pos[1])
        vals = [E(p) for p in pos]
        planos = []
        for v in vals:
            planos.extend(v.lista() if isinstance(v, Intervalo) else [v])
        nums = [x for x in planos if isinstance(x, (int, float)) and not isinstance(x, bool)]
        if nome == "MAX":
            return max(nums) if nums else 0
        if nome == "MIN":
            return min(nums) if nums else 0
        if nome == "SUM":
            return sum(nums)
        if nome == "AND":
            return all(planos)
        if nome == "OR":
            return any(planos)
        if nome == "ABS":
            return abs(num(vals[0]))
        if nome == "ROUND":
            return round(num(vals[0]), int(num(vals[1])))
        if nome == "ISNUMBER":
            return isinstance(vals[0], (int, float)) and not isinstance(vals[0], bool)
        if nome == "VLOOKUP":
            chave, tab, col = vals[0], vals[1], int(num(vals[2]))
            for linha in tab.linhas():
                if linha[0] == chave:
                    return linha[col - 1]
            raise Erro(f"VLOOKUP não achou {chave!r}")
        if nome == "COUNTIF":
            crit = vals[1]
            m = re.match(r"([<>=]+)(.*)", str(crit))
            o, alvo = (m.group(1), float(m.group(2))) if m else ("=", crit)
            f = {">": lambda x: x > alvo, "<": lambda x: x < alvo, ">=": lambda x: x >= alvo,
                 "<=": lambda x: x <= alvo, "=": lambda x: x == alvo}[o]
            return sum(1 for x in vals[0].lista() if isinstance(x, (int, float)) and f(x))
        raise Erro(f"função não suportada: {nome}")


class Intervalo:
    def __init__(self, av, aba, c1, r1, c2, r2):
        self.av, self.aba = av, aba
        self.cols = range(column_index_from_string(c1), column_index_from_string(c2) + 1)
        self.rows = range(r1, r2 + 1)

    def linhas(self):
        return [[self.av.valor(self.aba, f"{get_column_letter(c)}{r}") for c in self.cols] for r in self.rows]

    def lista(self):
        return [x for l in self.linhas() for x in l]


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    raise TypeError(f"não é número: {v!r}")


def fmt(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def avaliar_tudo(caminho):
    """Avalia todas as fórmulas. Devolve ({(aba, célula): valor}, [erros])."""
    av = Avaliador(caminho)
    erros, saida = [], {}
    for ws in av.wb.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith("="):
                    try:
                        saida[(ws.title, c.coordinate)] = av.valor(ws.title, c.coordinate)
                    except Exception as e:  # noqa: BLE001
                        erros.append((ws.title, c.coordinate, c.value, repr(e)))
    return saida, erros
