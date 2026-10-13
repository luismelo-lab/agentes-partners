"""Confere as fórmulas do Scorecard de qualificação recalculando em Python (sem Excel/LibreOffice).

1. Um avaliador mínimo de fórmulas do Excel (IF, AND, IFERROR, COUNTIF(S), COUNTA, SUM, SUMPRODUCT, VLOOKUP, &, comparações)
   lê o TEXTO das fórmulas gravadas na planilha e calcula.
2. Uma implementação independente das regras (direto das listas de critérios do script) calcula o esperado.
3. Compara: exemplos, 400 linhas sintéticas aleatórias no Scorecard e todas as contagens do Resumo contra o montar().
Uso:  python3 conferir_planilha.py
"""
import collections, fnmatch, os, random, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qualificacao_planilha as QP
from qualificacao_dados import PASTA

TOK = re.compile(r"""\s*(?:(?P<num>\d+(?:\.\d+)?)|(?P<str>"(?:[^"]|"")*")|(?P<bool>TRUE|FALSE)(?![\w(!])|(?P<ref>(?:'[^']+'|[A-Za-z_][\w\.]*)?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
                    |(?P<fn>[A-Z][A-Z0-9\.]*)\(|(?P<op><>|>=|<=|[=<>&+\-*/(),]))""", re.X)


def tokens(s):
    out, i = [], 0
    while i < len(s):
        m = TOK.match(s, i)
        if not m or m.end() == i:
            if s[i].isspace(): i += 1; continue
            raise ValueError(f"token em {s[i:i+20]!r}")
        i = m.end()
        k = m.lastgroup; v = m.group(k)
        if k == "str": v = v[1:-1].replace('""', '"')
        out.append((k, v))
    return out


class Avaliador:
    def __init__(self, wb):
        self.wb, self.cache = wb, {}

    def valor(self, sheet, coord):
        key = (sheet, coord)
        if key in self.cache: return self.cache[key]
        v = self.wb[sheet][coord].value
        if isinstance(v, str) and v.startswith("="):
            v = self.formula(sheet, v[1:])
        self.cache[key] = v
        return v

    def ref(self, sheet, txt):
        if "!" in txt:
            sh, txt = txt.rsplit("!", 1); sheet = sh.strip("'")
        txt = txt.replace("$", "")
        if ":" in txt:
            a, b = txt.split(":")
            ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups(); cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
            return [[self.valor(sheet, f"{CL(c)}{r}") for c in range(CI(ca), CI(cb) + 1)] for r in range(int(ra), int(rb) + 1)]
        return self.valor(sheet, txt)

    def formula(self, sheet, f):
        salvo = (getattr(self, "t", None), getattr(self, "p", None), getattr(self, "sheet", None))
        self.t, self.p, self.sheet = tokens(f), 0, sheet
        try:
            v = self.expr(0)
            assert self.p == len(self.t), f"sobrou {self.t[self.p:]} em {f}"
        finally:
            self.t, self.p, self.sheet = salvo
        return v

    PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4}

    def peek(self): return self.t[self.p] if self.p < len(self.t) else (None, None)

    def expr(self, minp):
        left = self.atom()
        while True:
            k, v = self.peek()
            if k == "op" and v in self.PREC and self.PREC[v] >= minp:
                self.p += 1
                left = self.bin(v, left, self.expr(self.PREC[v] + 1))
            else:
                return left

    @staticmethod
    def esc(x):
        if isinstance(x, list): raise ValueError("intervalo em escalar")
        return x

    def bin(self, op, a, b):
        a, b = self.esc(a), self.esc(b)
        if op == "&": return f"{self.txt(a)}{self.txt(b)}"
        if op in "+-*/":
            a, b = self.num(a), self.num(b)
            return a + b if op == "+" else a - b if op == "-" else a * b if op == "*" else a / b
        na, nb = (a if a is not None else ""), (b if b is not None else "")
        if isinstance(na, str) and isinstance(nb, (int, float)) and na == "": na = 0
        if isinstance(nb, str) and isinstance(na, (int, float)) and nb == "": nb = 0
        if isinstance(na, str) and isinstance(nb, str): na, nb = na.lower(), nb.lower()
        if type(na) != type(nb) and not (isinstance(na, (int, float)) and isinstance(nb, (int, float))):
            r = {"=": False, "<>": True}.get(op)
            if r is None: r = (op in (">", ">=")) if isinstance(na, str) else (op in ("<", "<="))   # texto > número no Excel
            return r
        return {"=": na == nb, "<>": na != nb, "<": na < nb, ">": na > nb, "<=": na <= nb, ">=": na >= nb}[op]

    @staticmethod
    def txt(x):
        if x is None: return ""
        if isinstance(x, bool): return "VERDADEIRO" if x else "FALSO"
        if isinstance(x, float) and x.is_integer(): return str(int(x))
        return str(x)

    @staticmethod
    def num(x):
        if x in (None, ""): return 0
        if isinstance(x, bool): return int(x)
        if isinstance(x, str): raise ValueError(f"#VALOR! ({x!r})")
        return x

    def atom(self):
        k, v = self.t[self.p]; self.p += 1
        if k == "num": return float(v)
        if k == "str": return v
        if k == "bool": return v == "TRUE"
        if k == "ref": return self.ref(self.sheet, v)
        if k == "op" and v == "(":
            x = self.expr(0); assert self.t[self.p] == ("op", ")"); self.p += 1; return x
        if k == "op" and v == "-": return -self.num(self.atom())
        if k == "fn": return self.func(v)
        raise ValueError(f"inesperado {k} {v}")

    def args_lazy(self):
        """Devolve as posições dos argumentos (para avaliar IF/IFERROR de forma preguiçosa)."""
        args, depth, start = [], 0, self.p
        while True:
            k, v = self.t[self.p]
            if k == "fn" or (k == "op" and v == "("): depth += 1
            elif k == "op" and v == ")":
                if depth == 0: args.append((start, self.p)); self.p += 1; return args
                depth -= 1
            elif k == "op" and v == "," and depth == 0: args.append((start, self.p)); start = self.p + 1
            self.p += 1

    def ev(self, span):
        t, p = self.t, self.p
        self.t, self.p = t[span[0]:span[1]], 0
        try: v = self.expr(0); assert self.p == len(self.t)
        finally: self.t, self.p = t, p
        return v

    @staticmethod
    def flat(x): return [c for row in x for c in row] if isinstance(x, list) else [x]

    @staticmethod
    def crit(c):
        m = re.match(r"(<>|>=|<=|=|<|>)?(.*)", str(c), re.S); op, val = m.group(1) or "=", m.group(2)
        def f(x):
            if op == "=" and ("*" in val or "?" in val): return isinstance(x, str) and fnmatch.fnmatchcase(x.lower(), val.lower())
            if op == "=": return (x if x is not None else "") != "" and str(Avaliador.txt(x)).lower() == val.lower() if val != "" else x in (None, "")
            if op == "<>": return str(Avaliador.txt(x)).lower() != val.lower()
            try: xv, vv = float(x), float(val)
            except (TypeError, ValueError): return False
            return {"<": xv < vv, ">": xv > vv, "<=": xv <= vv, ">=": xv >= vv}[op]
        return f

    def func(self, name):
        spans = self.args_lazy()
        if name == "IF":
            c = self.esc(self.ev(spans[0]))
            if c is True or (isinstance(c, (int, float)) and c != 0): return self.ev(spans[1])
            return self.ev(spans[2]) if len(spans) > 2 else False
        if name == "IFERROR":
            try: return self.ev(spans[0])
            except (ValueError, KeyError, ZeroDivisionError, LookupError): return self.ev(spans[1])
        a = [self.ev(s) for s in spans]
        if name == "AND": return all(bool(x) for x in a)
        if name == "OR": return any(bool(x) for x in a)
        if name == "SUM": return sum(x for v in a for x in self.flat(v) if isinstance(x, (int, float)) and not isinstance(x, bool))
        if name == "COUNTA": return sum(1 for v in a for x in self.flat(v) if x not in (None, ""))
        if name == "COUNTIF": return sum(1 for x in self.flat(a[0]) if self.crit(a[1])(x))
        if name == "COUNTIFS":
            rs = [(self.flat(a[i]), self.crit(a[i + 1])) for i in range(0, len(a), 2)]
            return sum(1 for i in range(len(rs[0][0])) if all(f(r[i]) for r, f in rs))
        if name == "SUMPRODUCT":
            cols = [self.flat(v) for v in a]
            return sum(__import__("math").prod(self.num(c[i]) for c in cols) for i in range(len(cols[0])))
        if name == "VLOOKUP":
            x, tab, col = a[0], a[1], int(a[2])
            for row in tab:
                if row[0] is not None and str(row[0]).lower() == str(x).lower(): return row[col - 1]
            raise LookupError("#N/D")
        raise ValueError(f"função não suportada: {name}")


# ---------------- regras independentes (esperado) ----------------
PONTOS = {c[0]: dict(c[3]) for c in QP.CRIT}
CHAVES = [c[0] for c in QP.CRIT]
PAR = dict(QP.PARAM)


def esperado(kos, resp):
    tem = [k for k in kos if k]
    if "NÃO" in kos: elim = "Elimina"
    elif kos.count("SIM") == 10: elim = "Passa"
    elif not tem: elim = ""
    else: elim = "Incompleto"
    pts = {ch: PONTOS[ch][r] for ch, r in zip(CHAVES, resp) if r}
    n = len(pts); score = sum(pts.values()) if n else ""
    if elim == "Elimina": faixa = "Não qualificado"
    elif elim != "Passa": faixa = "Incompleto: camada 1"
    elif n < 11: faixa = "Incompleto: camada 2"
    elif (score >= PAR["Corte de Qualificado (pontos)"] and pts["porte"] >= PAR["Mínimo em Porte para Qualificado"]
          and pts["gar"] >= PAR["Mínimo em Garantias para Qualificado"] and pts["ader"] >= PAR["Mínimo em Aderência à entrega para Qualificado"]):
        faixa = "Qualificado"
    elif score >= PAR["Corte de Em desenvolvimento (pontos)"]: faixa = "Em desenvolvimento"
    else: faixa = "Não qualificado"
    return elim, score, n, faixa


def main():
    caminho = os.path.join(PASTA, QP.NOME)
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Scorecard"]
    erros, testadas = [], 0
    random.seed(7)
    # linhas sintéticas: sobrescreve no workbook em memória (não grava)
    for i in range(400):
        r = QP.L0 + 3 + (i % (QP.NLIN_SC - 3))
        ws.cell(r, 1, f"teste {i}")
        modo = random.random()
        kos = [random.choice(["SIM"] * 8 + ["NÃO", None]) if modo < 0.5 else "SIM" for _ in range(10)]
        resp = [random.choice([o for o, _ in c[3]] + ([None] if modo < 0.2 else [])) for c in QP.CRIT]
        for j, v in enumerate(kos): ws.cell(r, 7 + j).value = v
        for j, v in enumerate(resp): ws.cell(r, 18 + j).value = v
        ev = Avaliador(wb)
        got = (ev.valor("Scorecard", f"Q{r}"), ev.valor("Scorecard", f"AC{r}"), ev.valor("Scorecard", f"AD{r}"), ev.valor("Scorecard", f"AE{r}"))
        exp = esperado(kos, resp)
        got = (got[0], got[1] if got[1] == "" else int(got[1]), int(got[2]), got[3])
        if got != exp: erros.append((r, got, exp))
        # lacunas e próximo passo coerentes
        lac, prox = ev.valor("Scorecard", f"AF{r}"), ev.valor("Scorecard", f"AG{r}")
        if exp[3] == "Não qualificado" and exp[0] == "Elimina" and not lac.startswith("Eliminatório:"): erros.append((r, "lacuna", lac))
        if exp[3] == "Qualificado" and not prox.startswith("Camada 3"): erros.append((r, "próximo passo", prox))
        testadas += 1
    # exemplos originais (relê a planilha limpa)
    wb2 = openpyxl.load_workbook(caminho); ev2 = Avaliador(wb2)
    exemplos = []
    for r in range(QP.L0, QP.L0 + 3):
        kos = [wb2["Scorecard"].cell(r, 7 + j).value for j in range(10)]
        resp = [wb2["Scorecard"].cell(r, 18 + j).value for j in range(11)]
        got = (ev2.valor("Scorecard", f"Q{r}"), ev2.valor("Scorecard", f"AC{r}"), ev2.valor("Scorecard", f"AE{r}"), ev2.valor("Scorecard", f"AF{r}"), ev2.valor("Scorecard", f"AG{r}"))
        e = esperado(kos, resp)
        if (got[0], got[1] if got[1] == "" else int(got[1]), got[2]) != (e[0], e[1], e[3]): erros.append(("exemplo", r, got, e))
        exemplos.append((wb2["Scorecard"].cell(r, 1).value, got))
    fila = ev2.valor("Scorecard", f"AE{QP.L0 + 3}")
    # Resumo × montar()
    from qualificacao_dados import montar
    linhas, _ = montar()
    cnt = collections.Counter(l["resultado"] for l in linhas); pr = collections.Counter(l["prior"] for l in linhas if l["resultado"] == "Depende da camada 1")
    R = "Resumo da base"
    chk = {"B5": cnt["Pré-qualificada pelos dados"], "B6": cnt["Depende da camada 1"], "B7": cnt["Descartada por eliminatório já conhecido"],
           "B8": len(linhas), "F5": pr["Alta"], "F6": pr["Média"], "F7": pr["Baixa"]}
    for i, (k, _) in enumerate(QP.STATUS_COLS):
        c = collections.Counter(l[k][0] for l in linhas)
        for j, s in enumerate(["CONFIRMADO: atende", "CONFIRMADO: parcial", "CONFIRMADO: elimina", "ESTIMADO", "DESCONHECIDO"]):
            chk[f"{CL(2 + j)}{20 + i}"] = c[s]
    for fx_i, fx in enumerate("ABC"):
        for j, t in enumerate(["Pré-qualificada pelos dados", "Depende da camada 1", "Descartada por eliminatório já conhecido"]):
            chk[f"{CL(2 + j)}{33 + fx_i}"] = sum(1 for l in linhas if l["faixa_base"] == fx and l["resultado"] == t)
    for cel, v in chk.items():
        g = ev2.valor(R, cel)
        if int(g) != v: erros.append((R, cel, g, v))
    # funil: confere com a conta direta
    grupos = [(chk["B5"], .7, .4), (chk["F5"], .7, .15), (chk["F6"], .5, .05), (chk["F7"], 0, .01)]
    c1 = sum(n * a * b for n, a, b in grupos); q = c1 * 0.5
    for cel, v in (("E44", c1), ("G44", q), ("C48", c1 * .5), ("D48", q * .5), ("C49", c1), ("D49", q), ("C50", c1 * 1.5), ("D50", q * 1.5)):
        if abs(ev2.valor(R, cel) - v) > 1e-9: erros.append((R, cel, ev2.valor(R, cel), v))
    # contagens por linha na Base (CONFIRMADO/ESTIMADO/DESCONHECIDO)
    B = "Base pré-qualificada"; wsb = wb2[B]; nb = 0
    for r in range(QP.L0, QP.L0 + len(linhas)):
        st = [wsb.cell(r, c).value for c in range(13, 23)]
        e = (sum(s.startswith("CONFIRMADO") for s in st), st.count("ESTIMADO"), st.count("DESCONHECIDO"))
        g = tuple(int(ev2.valor(B, f"{CL(c)}{r}")) for c in (33, 34, 35))
        if g != e or sum(g) != 10: erros.append((B, r, g, e))
        nb += 1
    print(f"Scorecard: {testadas} linhas sintéticas + 3 exemplos · Resumo: {len(chk) + 8} células · Base: {nb} linhas × 3 contagens")
    for nome, g in exemplos: print(f"  {nome}: eliminatórios={g[0]!r} pontuação={g[1]!r} faixa={g[2]!r} | lacunas={g[3]!r}")
    print(f"  1ª linha da fila (sem respostas): faixa={fila!r}")
    print("  Resumo:", {k: ev2.valor(R, k) for k in ("B5", "B6", "B7", "B8", "F5", "F6", "F7")}, "| funil base: camada 1 =", round(ev2.valor(R, "E44"), 1), "qualificadas =", round(ev2.valor(R, "G44"), 1),
          "| conservador", round(ev2.valor(R, "D48"), 1), "otimista", round(ev2.valor(R, "D50"), 1))
    print("ERROS:", len(erros))
    for e in erros[:20]: print("  ", e)
    return 1 if erros else 0


if __name__ == "__main__":
    sys.exit(main())
