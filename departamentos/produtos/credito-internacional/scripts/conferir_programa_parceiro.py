"""Confere as fórmulas do simulador do Programa de Crescimento do Parceiro (v2.2) recalculando em Python.

1. Usa o avaliador mínimo de fórmulas de conferir_planilha.py (texto das fórmulas gravadas), com MIN, MAX, INDEX, ROUND e LARGE.
2. Uma implementação independente das regras (abaixo) calcula o esperado.
3. Compara: a jornada gravada e 400 jornadas aleatórias; os 7 perfis da aba Sustentação e 400 perfis aleatórios;
   todas as combinações nível × escopo do simulador; o comparativo, a matriz, a sensibilidade e a aba Canal (gravada e 200 aleatórias).
4. Avalia todas as fórmulas da pasta e acusa qualquer erro (referência quebrada, divisão por zero sem tratamento, função desconhecida).
Uso:  python3 conferir_programa_parceiro.py [--exemplos]
"""
import os, random, sys
import openpyxl
from openpyxl.utils import get_column_letter as CL

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from conferir_planilha import Avaliador as Base
import programa_parceiro_planilha as PP

ARQ = os.path.join(PP.PASTA, PP.NOME)
SUS = "Sustentação"


class Avaliador(Base):
    def func(self, name):
        if name in ("MIN", "MAX", "INDEX", "ROUND", "LARGE"):
            a = [self.ev(s) for s in self.args_lazy()]
            if name in ("MIN", "MAX"):
                xs = [x for v in a for x in self.flat(v) if isinstance(x, (int, float)) and not isinstance(x, bool)]
                return (min if name == "MIN" else max)(xs) if xs else 0
            if name == "INDEX":
                return self.flat(a[0])[int(self.num(a[1])) - 1]
            if name == "LARGE":
                xs = sorted((x for x in self.flat(a[0]) if isinstance(x, (int, float)) and not isinstance(x, bool)), reverse=True)
                return xs[int(self.num(a[1])) - 1]
            return round(self.num(a[0]), int(self.num(a[1])))
        return super().func(name)


# ---------------- parâmetros lidos da planilha ----------------
def parametros(wb):
    ws = wb["Premissas"]; sus = wb[SUS]
    fator = sus[f"B{PP.S_FATOR}"].value
    niv = []
    for i in range(5):
        r = PP.P_N0 + i; s = PP.S_N0 + i
        v = [ws.cell(r, j).value for j in range(1, 17)]
        niv.append(dict(n=v[0], nome=v[1], teto=v[2], trilha=v[3], liq=v[4], lib=v[5], alt=v[6], seg=v[7],
                        lim=v[8], dos=v[9], rec=v[10], oc=v[11], com=v[12], maxdev=v[13], maxrec=v[14], lib24=v[15] or 0,
                        m_liq=round(v[4] * fator, 0), m_lib=sus[f"G{s}"].value, m_reg=sus[f"H{s}"].value,
                        m_lim=sus[f"I{s}"].value, m_tr=sus[f"J{s}"].value, s_tr=sus[f"K{s}"].value))
    esc = [dict(pct=ws.cell(r, 4).value, horas=ws.cell(r, 5).value) for r in range(PP.P_E0, PP.P_E0 + 4)]
    prm = {k: ws[f"B{r}"].value for k, r in
           dict(cambio=PP.P_CAMBIO, fee=PP.P_FEE, save=PP.P_SAVE, op=PP.P_OP, hora=PP.P_HORA, htrim=PP.P_HTRIM,
                amostra=PP.P_AMOSTRA, hb4=PP.P_HB4, hb5=PP.P_HB5, hfix=PP.P_HFIX).items()}
    prm.update({k: sus[f"B{r}"].value for k, r in
                dict(obs=PP.S_OBS, bonus=PP.S_BONUS, minimo=PP.S_MIN, maxext=PP.S_EXT, l24=PP.S_L24, m24=PP.S_M24,
                     form=PP.S_FORM, pos=PP.S_POS).items()})
    return niv, esc, prm


# ---------------- regras independentes ----------------
def janela(t, n, ext):
    return list(range(max(0, t - n + 1), t + 1)) + [t - n - k for k in range(ext) if t - n - k >= 0]


def motor(niv, esc, prm, q, completo):
    """q: lista de dicionários por trimestre. completo=True: aba Jornada (com conhecimento e qualidade); False: perfis."""
    N, Wc, U, stk, acum, acum10 = 1, 0, 99, 0, 0.0, 0.0
    hist = [1]; out = []
    for t, e in enumerate(q):
        p4 = range(max(0, t - 3), t + 1); p8 = range(max(0, t - 7), t + 1)
        ext = min(prm["maxext"], sum(1 for i in p4 if q[i]["X"] == "S"))
        ext24 = min(prm["maxext"], sum(1 for i in p8 if q[i]["X"] == "S"))
        W4, W8 = janela(t, 4, ext), janela(t, 8, ext24)
        s = lambda k, W: sum(q[i][k] for i in W)
        A = lambda i: 1 if (q[i]["Q"] + q[i]["O"] + q[i].get("C", 0)) > 0 or q[i]["Y"] == "S" else 0
        vs, op, q12, reg = s("L", W4), s("O", W4), s("Q", W4), sum(A(i) for i in W4)
        tr = sum(1 for i in W4 if q[i]["L"] > 0)
        vs24, op24, tr24 = s("L", W8), s("O", W8), sum(1 for i in W8 if q[i]["L"] > 0)
        conc = round(max(q[i]["M"] for i in W4) / vs, 4) if vs else 1
        stk = stk + 1 if A(t) else (stk if e["X"] == "S" else 0)
        if completo:
            ind12, rcn, d12 = s("ind", W4), s("recus", W4), s("dos", W4)
            dev = d12 - s("compl", W4)
            rec = round(rcn / ind12, 4) if ind12 else 0
            dos = round((d12 - dev) / d12, 4) if d12 else 0
            oc = sum(q[i]["oc"] for i in p4)
        sobe, sus = {}, {}
        for k in range(2, 6):
            nv = niv[k - 1]
            cap_a, cap_p = nv["lim"] * nv["liq"], nv["m_lim"] * nv["m_liq"]
            va = vs - sum(max(0, q[i]["M"] - cap_a) for i in W4)
            vp = vs - sum(max(0, q[i]["M"] - cap_p) for i in W4)
            v24 = vs24 - sum(max(0, q[i]["M"] - cap_p) for i in W8)
            vol = (op >= nv["lib"] and va >= nv["liq"]) or (nv["alt"] > 0 and q12 >= nv["alt"])
            a = vol and stk >= nv["seg"] and tr >= nv["s_tr"] and op24 >= nv["lib24"]
            m = prm["m24"]
            l12 = vp >= nv["m_liq"] and op >= nv["m_lib"] and tr >= nv["m_tr"]
            l24 = prm["l24"] == "S" and v24 >= m * nv["m_liq"] and op24 >= m * nv["m_lib"] and tr24 >= m * nv["m_tr"]
            b = reg >= nv["m_reg"] and (l12 or l24)
            if completo:
                am = prm["amostra"]
                qual_ok = (e["tri"] >= nv["trilha"] and (dos >= nv["dos"] if d12 >= am else dev <= nv["maxdev"])
                           and (rec <= nv["rec"] if ind12 >= am else rcn <= nv["maxrec"]) and oc <= nv["oc"])
                a = a and qual_ok and (nv["com"] == "N" or e["com"] == "S")
                b = b and qual_ok
            sobe[k], sus[k] = a, b
        E = next((k for k in range(5, 1, -1) if sobe[k]), 1)
        S = next((k for k in range(5, 1, -1) if sus[k]), 1)
        MP = max(hist[-4:])
        lim = prm["obs"] + (prm["bonus"] if any(q[i]["Y"] == "S" for i in p4) else 0)
        prot = (t + 1 <= prm["form"]) or (U + 1 <= prm["pos"])
        vig = N
        if completo and oc > 0:
            N2, W2, sit = 1, 0, "Compliance: ajustado ao N1"
        elif E > N or (S > N and MP > N):
            N2, W2, sit = N + 1, 0, "Promovido" if E > N else "Retomado"
        elif S >= N:
            N2, W2, sit = N, 0, "Sustentado"
        elif e["X"] == "S":
            N2, W2, sit = N, Wc, "Neutro"
        elif Wc < lim:
            N2, W2, sit = N, Wc + 1, "Em observação"
        elif prot:
            N2, W2, sit = N, Wc, "Em observação (período de proteção)"
        else:
            N2, W2 = max(prm["minimo"], N - 1), 0
            sit = f"Ajustado ao N{N - 1}" if N > prm["minimo"] else "No nível de entrada (N1)"
        U = 0 if N2 > N else U + 1
        pct = niv[vig - 1]["teto"] if not completo else min(niv[vig - 1]["teto"], esc[e["esc"] - 1]["pct"])
        ganho = e["L"] * pct; acum += ganho; acum10 += e["L"] * niv[0]["teto"]
        sai = q[t - 3]["L"] if t >= 3 else 0
        nec = max(0, niv[N2 - 1]["m_liq"] - (vs - sai))
        out.append(dict(VS=vs, OP=op, REG=reg, TR=tr, CONC=conc, STK=stk, VS24=vs24, OP24=op24, TR24=tr24, E=E, S=S, MAXP=MP,
                        PROT=prot, N=N2, W=W2, U=U, SIT=sit, VIG=vig, PCT=pct, GANHO=ganho, ACUM=acum, ACUM10=acum10, SAI=sai, NEC=nec))
        N, Wc = N2, W2; hist.append(N2)
    return out


CHAVES = ("VS", "OP", "REG", "TR", "CONC", "STK", "VS24", "OP24", "TR24", "E", "S", "MAXP", "PROT", "N", "W", "U", "SIT", "VIG",
          "PCT", "GANHO", "ACUM", "ACUM10", "SAI", "NEC")


def comparar(wb, sheet, linhas, esp, ncols, rotulo):
    av = Avaliador(wb); erros = 0
    for t in range(ncols):
        c = CL(3 + t)
        for k, v in esp[t].items():
            if k not in linhas: continue
            g = av.valor(sheet, f"{c}{linhas[k]}")
            if isinstance(v, (str, bool)):
                ok = g == v
            else:
                ok = abs(float(g) - float(v)) <= 1e-6
            if not ok:
                erros += 1
                if erros <= 5: print(f"  ✘ {rotulo} trimestre {t + 1} {k}: planilha={g} esperado={v}")
    return erros


# ---------------- Jornada ----------------
def jornada_q(entradas):
    return [dict(zip(PP.CAMPOS_J, e)) for e in entradas]


def escrever_jornada(ws, entradas):
    for t, e in enumerate(entradas):
        for i, v in enumerate(e):
            ws.cell(6 + i, 3 + t).value = v


def linhas_jornada():
    return {k: PP.JR[k] for k in CHAVES}


# ---------------- Perfis da aba Sustentação ----------------
def perfil_q(d):
    return [dict(Q=d["Q"][t], C=d["C"][t], O=d["O"][t], L=d["L"][t], M=d["M"][t] if "M" in d else d["L"][t],
                 X=d["X"][t], Y=d["Y"][t]) for t in range(12)]


def escrever_perfil(ws, r0, d):
    R = {k: r0 + v for k, v in PP.PB.items()}
    for t in range(12):
        c = CL(3 + t)
        for k in ("Q", "C", "O", "L", "X", "Y"):
            ws[f"{c}{R[k]}"].value = d[k][t]
        ws[f"{c}{R['M']}"].value = d.get("M", d["L"])[t]


def linhas_perfil(r0):
    return {k: r0 + PP.PB[k] for k in CHAVES}


def todas_formulas(wb):
    av = Avaliador(wb); erros = 0; n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    try:
                        v = av.valor(ws.title, c.coordinate)
                        if isinstance(v, str) and v.startswith("#"): raise ValueError(v)
                    except Exception as ex:
                        erros += 1
                        if erros <= 10: print(f"  ✘ {ws.title}!{c.coordinate}: {ex}")
    return n, erros


# ---------------- Canal ----------------
def canal_esperado(niv, esc, linhas):
    out = []
    for nome, ops, liq, e, q in linhas:
        lvl = 1
        for k in (5, 4, 3, 2):
            nv = niv[k - 1]
            if (ops >= nv["lib"] and ops * min(liq, nv["lim"] * nv["liq"]) >= nv["liq"]) or (nv["alt"] > 0 and q >= nv["alt"]):
                lvl = k; break
        g = ops * liq; pct = min(niv[lvl - 1]["teto"], esc[e - 1]["pct"])
        out.append(dict(F=lvl, G=g, H=pct, I=g * pct, J=g - g * pct))
    tot = sum(x["G"] for x in out); rep = sum(x["I"] for x in out)
    top3 = sum(sorted((x["G"] for x in out), reverse=True)[:3])
    ss = sum(x["G"] ** 2 for x in out)
    resumo = dict(medio=rep / tot if tot else 0, top3=top3 / tot if tot else 0, efetivos=tot * tot / ss if ss else 0,
                  com_op=sum(1 for l in linhas if l[1] > 0), dist=[sum(1 for x in out if x["F"] == k) for k in range(1, 6)])
    return out, resumo


def conferir_canal(wb, niv, esc, linhas, rotulo):
    ws = wb["Canal"]; n = len(linhas)
    for i, l in enumerate(linhas):
        for j, v in enumerate(l):
            ws.cell(PP.C_L0 + i, 1 + j).value = v
    av = Avaliador(wb); erros = 0
    out, res = canal_esperado(niv, esc, linhas)
    for i, x in enumerate(out):
        for col, v in x.items():
            g = av.valor("Canal", f"{col}{PP.C_L0 + i}")
            if abs(g - v) > 1e-6:
                erros += 1
                if erros <= 5: print(f"  ✘ {rotulo} linha {i + 1} {col}: {g} ≠ {v}")
    t = PP.C_L0 + n
    for r, v in ((t + 2, res["medio"]), (t + 3, res["top3"]), (t + 4, res["efetivos"]), (t + 5, res["com_op"])):
        g = av.valor("Canal", f"B{r}")
        if abs(g - v) > 1e-6:
            erros += 1; print(f"  ✘ {rotulo} resumo B{r}: {g} ≠ {v}")
    for k in range(1, 6):
        if av.valor("Canal", f"B{t + 7 + k}") != res["dist"][k - 1]:
            erros += 1
    return erros, res


def main():
    wb = openpyxl.load_workbook(ARQ)
    niv, esc, prm = parametros(wb)
    total = 0; ver = "--exemplos" in sys.argv

    # 0) todas as fórmulas avaliam sem erro
    n, e = todas_formulas(wb); total += e
    print(f"Fórmulas avaliadas: {n} · {'sem erro' if not e else str(e) + ' com erro'}")

    # 1) jornada gravada
    esp = motor(niv, esc, prm, jornada_q(PP.JORNADA), True)
    e = comparar(wb, "Jornada do parceiro", linhas_jornada(), esp, 8, "jornada"); total += e
    print(f"Jornada de exemplo: {'OK' if not e else str(e) + ' divergências'}")
    if ver:
        for t, x in enumerate(esp):
            print(f"  {PP.TRIMS[t]}: VS={x['VS']:,.0f} E={x['E']} S={x['S']} nível={x['N']} ({x['SIT']}) vigente={x['VIG']} "
                  f"ganho={x['GANHO']:,.0f} acumulado={x['ACUM']:,.0f} (só 10%: {x['ACUM10']:,.0f}) necessário={x['NEC']:,.0f}")

    # 2) jornadas aleatórias
    rnd = random.Random(61); e_rand = 0
    ws = wb["Jornada do parceiro"]
    for k in range(400):
        ent = []
        for t in range(8):
            ind = rnd.randint(0, 4); q = rnd.randint(0, ind); rc = rnd.randint(0, ind - q); d = rnd.randint(0, 3); c = rnd.randint(0, d)
            lib = rnd.choice([0, 0, 1, 1, 2]); liq = sum(rnd.choice([300_000, 400_000, 600_000, 800_000]) for _ in range(lib))
            maior = 0 if not lib else rnd.randint(int(liq / lib), liq)
            ent.append((ind, q, rc, d, c, lib, liq, maior, rnd.randint(1, 4), rnd.randint(1, 5),
                        1 if rnd.random() < 0.05 else 0, rnd.choice("SN"), "S" if rnd.random() < 0.15 else "N",
                        "S" if rnd.random() < 0.2 else "N"))
        escrever_jornada(ws, ent)
        e_rand += comparar(wb, "Jornada do parceiro", linhas_jornada(), motor(niv, esc, prm, jornada_q(ent), True), 8, f"jornada aleatória {k}")
    escrever_jornada(ws, PP.JORNADA)
    total += e_rand
    print(f"400 jornadas aleatórias: {'OK' if not e_rand else str(e_rand) + ' divergências'}")

    # 3) perfis da aba Sustentação
    ws = wb[SUS]; e_p = 0; resumo = []
    for i, (nome, _, d) in enumerate(PP.PERFIS):
        r0 = PP.S_P0 + i * PP.S_BH
        esp = motor(niv, esc, prm, perfil_q(d), False)
        e_p += comparar(wb, SUS, linhas_perfil(r0), esp, 12, f"perfil {nome}")
        resumo.append((nome, esp))
    total += e_p
    print(f"Perfis da aba Sustentação ({len(PP.PERFIS)}): {'OK' if not e_p else str(e_p) + ' divergências'}")
    if ver:
        for nome, esp in resumo:
            print(f"  {nome}: " + " · ".join(f"{PP.TRIMS12[t]} N{x['N']} {x['SIT']}" for t, x in enumerate(esp)))
            print(f"     acumulado {esp[-1]['ACUM']:,.0f} (só 10%: {esp[-1]['ACUM10']:,.0f}); VS final {esp[-1]['VS']:,.0f}; necessário {esp[-1]['NEC']:,.0f}")

    # 4) perfis aleatórios (no bloco do 1º perfil)
    e_r = 0; r0 = PP.S_P0
    for k in range(400):
        d = {"Q": [], "C": [], "O": [], "L": [], "M": [], "X": "", "Y": ""}
        for t in range(12):
            o = rnd.choice([0, 0, 1, 1, 2]); liq = sum(rnd.choice([250_000, 400_000, 600_000, 900_000, 1_800_000]) for _ in range(o))
            d["Q"].append(rnd.randint(0, 3)); d["C"].append(rnd.randint(0, 1)); d["O"].append(o); d["L"].append(liq)
            d["M"].append(0 if not o else rnd.randint(int(liq / o), liq))
            d["X"] += "S" if rnd.random() < 0.12 else "N"; d["Y"] += "S" if rnd.random() < 0.25 else "N"
        escrever_perfil(ws, r0, d)
        e_r += comparar(wb, SUS, linhas_perfil(r0), motor(niv, esc, prm, perfil_q(d), False), 12, f"perfil aleatório {k}")
    escrever_perfil(ws, r0, PP.PERFIS[0][2])
    total += e_r
    print(f"400 perfis aleatórios: {'OK' if not e_r else str(e_r) + ' divergências'}")

    # 5) simulador: todas as combinações nível × escopo, 3 valores de operação
    ws = wb["Simulador de operação"]; e_sim = 0
    for usd in (5_000_000, 7_500_000, 10_000_000):
        for n in range(1, 6):
            for s in range(1, 5):
                ws["B4"].value, ws["B8"].value, ws["B9"].value = usd, n, s
                av = Avaliador(wb)
                liq = usd * prm["cambio"] * prm["fee"] * prm["save"]
                pct = min(niv[n - 1]["teto"], esc[s - 1]["pct"])
                esperado = {"B14": liq, "B19": pct, "B20": liq * pct, "B21": liq * (1 - pct),
                            "B23": liq * min(niv[n - 1]["teto"], esc[0]["pct"]), "D28": liq * 0.3333 * pct,
                            "D32": liq * pct, "C32": liq}
                for ref, v in esperado.items():
                    g = av.valor("Simulador de operação", ref)
                    if abs(g - v) > 1e-6:
                        e_sim += 1; print(f"  ✘ simulador {usd} N{n} E{s} {ref}: {g} ≠ {v}")
                if av.valor("Simulador de operação", "F32") != "OK: as parcelas somam 100%":
                    e_sim += 1
    ws["B4"].value, ws["B8"].value, ws["B9"].value = f"=Premissas!B{PP.P_OP}", 3, 3
    total += e_sim
    print(f"Simulador (60 combinações): {'OK' if not e_sim else str(e_sim) + ' divergências'}")

    # 6) comparativo, matriz, patamares de permanência, unidade e sensibilidade
    av = Avaliador(wb); e_cmp = 0
    liq = prm["op"] * prm["cambio"] * prm["fee"] * prm["save"]
    ops_p, ops_c = wb["Comparativo"]["B5"].value, wb["Comparativo"]["B6"].value
    ret1 = liq * (1 - niv[0]["teto"]); contrib1 = None; contribb1 = None
    linhas_cmp = []
    for i, (n, s) in enumerate([(1, 1), (2, 2), (3, 3), (4, 4), (5, 4), (5, 1)]):
        r = 8 + i; pct = min(niv[n - 1]["teto"], esc[s - 1]["pct"]); ret = liq * (1 - pct)
        h = esc[s - 1]["horas"]
        hb = (prm["hb5"] if n == 5 else prm["hb4"] if n == 4 else 0) * 4 / ops_p
        fixo = prm["hfix"] * 4 * prm["hora"] / ops_c
        contrib = ret - h * prm["hora"]; contribb = contrib - hb * prm["hora"] - fixo
        if contrib1 is None: contrib1, contribb1 = contrib, contribb
        esperado = {"D": pct, "E": liq * pct, "F": ret, "H": liq * pct * ops_p, "J": ret1 / ret - 1,
                    "K": h, "M": prm["htrim"] / h, "N": prm["htrim"] / h * ret, "O": hb, "P": fixo, "Q": contrib,
                    "R": contrib1 / contrib - 1, "S": contribb, "T": contribb1 / contribb - 1}
        linhas_cmp.append((n, s, esperado))
        for col, v in esperado.items():
            g = av.valor("Comparativo", f"{col}{r}")
            if abs(g - v) > 1e-6:
                e_cmp += 1; print(f"  ✘ comparativo linha {r} {col}: {g} ≠ {v}")
    for i in range(5):
        for k in range(4):
            g = av.valor("Premissas", f"{CL(3 + k)}{PP.P_M0 + i}")
            if abs(g - min(niv[i]["teto"], esc[k]["pct"])) > 1e-9:
                e_cmp += 1
        if abs(av.valor(SUS, f"F{PP.S_N0 + i}") - niv[i]["m_liq"]) > 1e-6:
            e_cmp += 1; print(f"  ✘ patamar de permanência N{i + 1}")
    us = 5_000_000 * prm["cambio"] * prm["fee"] * prm["save"]
    if abs(av.valor(SUS, f"B{PP.S_US}") - us) > 1e-6:
        e_cmp += 1
    sens = []
    for i, (fee, sv) in enumerate(PP.SENS):
        r = PP.S_SENS + 1 + i; u = 5_000_000 * prm["cambio"] * fee * sv
        linha = []
        for j, lvl in enumerate((3, 4, 5)):
            v = niv[lvl - 1]["liq"] / u; linha.append(v)
            if abs(av.valor(SUS, f"{CL(4 + j)}{r}") - v) > 1e-9:
                e_cmp += 1; print(f"  ✘ sensibilidade {fee} {sv} N{lvl}")
        sens.append((fee, sv, linha))
    total += e_cmp
    print(f"Comparativo, matriz, patamares e sensibilidade: {'OK' if not e_cmp else str(e_cmp) + ' divergências'}")
    if ver:
        print(f"  Unidade de sustentação: R$ {us:,.2f}; patamares de permanência: " +
              ", ".join(f"N{i + 1} R$ {niv[i]['m_liq']:,.0f}" for i in range(5)))
        for n, s, x in linhas_cmp:
            print(f"  N{n} E{s}: contribuição {x['Q']:,.0f} (empate {x['R']:.1%}); com benefícios e fixo {x['S']:,.0f} (empate {x['T']:.1%})")
        for fee, sv, l in sens:
            print(f"  fee {fee:.0%} · Save {sv:.0%}: N3 {l[0]:.1f} · N4 {l[1]:.1f} · N5 {l[2]:.1f}")

    # 7) aba Canal: cenário gravado e 200 aleatórios
    e_c, res = conferir_canal(wb, niv, esc, PP.CANAL, "canal")
    for k in range(200):
        linhas = [(f"Parceiro {i + 1}", rnd.choice([0, 0, 1, 1, 2, 3, 4]), rnd.choice([0, 300_000, 450_000, 800_000, 2_000_000]),
                   rnd.randint(1, 4), rnd.randint(0, 5)) for i in range(len(PP.CANAL))]
        if not any(l[1] * l[2] for l in linhas):
            linhas[0] = ("Parceiro 1", 1, 450_000, 3, 2)
        e_c += conferir_canal(wb, niv, esc, linhas, f"canal aleatório {k}")[0]
    conferir_canal(wb, niv, esc, PP.CANAL, "canal")
    total += e_c
    print(f"Canal (cenário do plano de 2027 e 200 aleatórios): {'OK' if not e_c else str(e_c) + ' divergências'}")
    if ver:
        print(f"  plano 2027: % médio {res['medio']:.1%}; 3 maiores {res['top3']:.0%}; efetivos {res['efetivos']:.1f}; níveis {res['dist']}")
    print("RESULTADO:", "tudo confere" if total == 0 else f"{total} divergências")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
