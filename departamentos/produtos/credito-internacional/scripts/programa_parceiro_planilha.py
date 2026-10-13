"""Crédito Internacional · Programa de Crescimento do Parceiro: simulador com fórmulas vivas (v2.2, sustentação).

Uso:  python3 programa_parceiro_planilha.py   -> grava Credito-Internacional-Programa-Parceiro-Simulador-v2.xlsx na pasta do produto
Só gera a planilha (openpyxl). Não lê .env, não chama API, não envia nada.
A conferência das fórmulas fica em conferir_programa_parceiro.py.

v2.0 (08/10/2026): patamares para subir e para manter, apuração trimestral em janela móvel, observação antes de cada ajuste,
ajuste de 1 nível por vez até o N1, retorno expresso, trimestres neutros e amortecedor multiproduto.
v2.1 (08/10/2026, revisões Jurídico, Resultados, Partner Success e Português):
- nomenclatura: patamar de acesso (subir), patamar de permanência (manter), "Ajustado ao N_", plano de retomada, retorno expresso;
- regra anti-oscilação: para manter vale a melhor de duas leituras (12 meses ou o dobro do patamar em 24 meses);
- concentração: em vez de reprovar quem traz operação grande, cada operação conta no máximo X% do patamar (N4 e N5);
- período de proteção: sem ajuste nas 4 primeiras apurações após a adesão e nas 2 seguintes a cada subida;
- atividade em outro produto Save conta como atividade; o amortecedor multiproduto olha os últimos 12 meses;
- qualidade com amostra mínima (abaixo de 5 dossiês ou indicações, vale o número de devoluções ou recusas);
- taxa de recusa = recusadas ÷ registradas (entrada própria);
- N4 exige 2 operações liberadas para manter;
- perfil G · Realista 2027 (funil do planejamento 2027) e rótulos dos perfis;
- abas novas: Canal (5 a 10 parceiros) e tabela de sensibilidade dos patamares ao fee e à parcela Save;
- Comparativo com horas de benefícios e custo fixo do programa; "Como usar" na 1ª posição; abas protegidas sem senha.
v2.2 (09/10/2026, parecer do Conselho, R3): para subir ao N4, também 3 operações liberadas em 24 meses (coluna P da aba
Premissas; 0 = não se aplica), para evitar o sobe e desce no N4 de quem tem cerca de 1 operação a cada 6 a 9 meses.
Fonte IBM Plex Sans (Design System Save).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

PASTA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NOME = "Credito-Internacional-Programa-Parceiro-Simulador-v2.xlsx"
VERSAO = "2.2"

FONTE = "IBM Plex Sans"
SUS = "'Sustentação'"   # nome da aba entre aspas para as fórmulas
PRETO, OURO = "111111", "B58A1F"
F_HEAD = PatternFill("solid", fgColor=PRETO); F_SUB = PatternFill("solid", fgColor="FBF1D7")
F_IN = PatternFill("solid", fgColor="FFF2CC"); F_OUT = PatternFill("solid", fgColor="F6F7F9")
BOLD = Font(name=FONTE, bold=True); HFONT = Font(name=FONTE, bold=True, color="F5F5F5")
TITLE = Font(name=FONTE, bold=True, size=14, color="0F1115"); NOTA = Font(name=FONTE, italic=True, size=9, color="4A4F58")
SUBT = Font(name=FONTE, bold=True, size=11, color=OURO)
WRAP = Alignment(wrap_text=True, vertical="top"); CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Border(*(Side(style="thin", color="DEE0E5"),) * 4)
RS = '"R$" #,##0;[Red]-"R$" #,##0'; USD = '"US$" #,##0'; PCT = '0%'; PCT1 = '0.0%'

# ------------------------------------------------------------------ premissas (editáveis)
CAMBIO, FEE, PARCELA_SAVE, OPERACAO = 5.22, 0.03, 0.50, 7_500_000
CUSTO_HORA, HORAS_TRIM = 300, 240
AMOSTRA, HB4, HB5, HFIX = 5, 6, 12, 16
# nº, nome, teto, trilha, patamar de acesso (líquido 12m), liberadas 12m, alternativa qualificadas, trimestres seguidos com
# atividade, limite por operação (% do patamar de acesso), dossiês completos na 1ª entrega mín., recusa máx., ocorrências máx.,
# comitê, devoluções máx. e recusas máx. quando a amostra é menor que a mínima, liberadas em 24 meses para subir (0 = não se aplica)
NIVEIS = [
    (1, "Indicador", 0.10, 1, 0, 0, 0, 0, 1.00, 0.00, 1.00, 0, "N", 99, 99, 0),
    (2, "Qualificador", 0.15, 2, 0, 1, 3, 1, 1.00, 0.00, 0.40, 0, "N", 99, 2, 0),
    (3, "Condutor", 0.20, 3, 350_000, 1, 0, 2, 1.00, 0.70, 0.40, 0, "N", 1, 2, 0),
    (4, "Ponta a Ponta", 0.25, 4, 750_000, 2, 0, 3, 0.50, 0.80, 0.30, 0, "N", 1, 2, 3),
    (5, "Estratégico", 0.30, 5, 1_500_000, 3, 0, 4, 0.50, 0.90, 0.25, 0, "S", 1, 2, 0),
]
# permanência (por nível): operações liberadas 12m, trimestres com atividade (de 4), limite por operação (% do patamar de
# permanência), trimestres com receita (de 4) para manter, trimestres com receita para subir
SUST = [
    (0, 0, 1.00, 0, 0),
    (0, 2, 1.00, 0, 0),
    (1, 2, 1.00, 1, 1),
    (2, 3, 0.80, 1, 2),
    (2, 3, 0.60, 2, 2),
]
FATOR_MANUT, OBS, OBS_BONUS, JAN_REC, NIVEL_MIN, MAX_EXT = 0.75, 1, 1, 4, 1, 2
LEIT24, MULT24, FORMACAO, POS_SUBIDA = "S", 2, 4, 2
# nº, nome, etapas, % do escopo, horas Save por operação [HIPÓTESE]
ESCOPOS = [
    (1, "E1 · Indicação", "Etapa 1 (originação): registra, apresenta o decisor e responde às 3 perguntas de qualificação", 0.10, 60),
    (2, "E2 · Qualificação", "Etapas 1 e 2: ficha de qualificação completa (5 perguntas e sinais de alerta) e participação na conversa inicial", 0.15, 52),
    (3, "E3 · Condução documental", "Etapas 1 a 3, com insumos para a 4: kit documental completo e organizado com o cliente", 0.20, 36),
    (4, "E4 · Ponta a ponta comercial e documental", "Etapas 1 a 9 na parte comercial e documental: acompanha o cliente até o contrato e nas comprovações de cada parcela", 0.30, 24),
]
# Jornada de exemplo (8 trimestres): registradas, qualificadas, recusadas, dossiês, completos, liberadas, líquido, maior operação,
# escopo, trilha, ocorrências, comitê, trimestre neutro, negócio em outro produto Save
TRIMS = ["1T27", "2T27", "3T27", "4T27", "1T28", "2T28", "3T28", "4T28"]
JORNADA = [
    (3, 1, 2, 0, 0, 0, 0, 0, 1, 2, 0, "N", "N", "N"),
    (3, 2, 1, 1, 1, 0, 0, 0, 2, 2, 0, "N", "N", "N"),
    (2, 1, 1, 1, 1, 1, 400_000, 400_000, 3, 3, 0, "N", "N", "N"),
    (3, 2, 1, 2, 2, 0, 0, 0, 3, 3, 0, "N", "N", "N"),
    (2, 2, 0, 2, 2, 1, 450_000, 450_000, 3, 4, 0, "N", "N", "N"),
    (3, 2, 0, 2, 2, 1, 600_000, 600_000, 4, 4, 0, "N", "N", "N"),
    (2, 2, 0, 2, 2, 1, 550_000, 550_000, 4, 5, 0, "S", "N", "N"),
    (3, 2, 0, 2, 2, 1, 650_000, 650_000, 4, 5, 0, "S", "N", "N"),
]
CAMPOS_J = ["ind", "Q", "recus", "dos", "compl", "O", "L", "M", "esc", "tri", "oc", "com", "X", "Y"]
# Perfis da aba Sustentação (12 trimestres): qualificadas, contratadas, liberadas, líquido, maior operação, neutro, outro produto
TRIMS12 = ["1T27", "2T27", "3T27", "4T27", "1T28", "2T28", "3T28", "4T28", "1T29", "2T29", "3T29", "4T29"]
K = 1000
PERFIS = [
    ("A · Constante (4 operações por ano, acima do plano de 2027)",
     "Qualifica todo trimestre e libera uma operação por trimestre a partir do 3T27. É o dobro do que o planejamento prevê para o produto inteiro em 2027: serve para mostrar a escada completa, não o caso típico.",
     dict(Q=[2] * 12, C=[0] + [1] * 11, O=[0, 0] + [1] * 10,
          L=[0, 0, 400 * K, 450 * K, 500 * K, 550 * K, 600 * K, 500 * K, 550 * K, 600 * K, 500 * K, 650 * K], X="N" * 12, Y="N" * 12)),
    ("B · Caso isolado", "Uma operação grande no 3T27 e depois só algumas indicações. O limite por operação impede que o caso compre o N4; quando ele sai da janela, o nível desce um degrau por vez, depois do período de proteção, e para no N1.",
     dict(Q=[1, 1, 1, 1, 1, 1, 0, 0, 0, 0, 0, 0], C=[0, 1] + [0] * 10, O=[0, 0, 1] + [0] * 9,
          L=[0, 0, 2_500 * K] + [0] * 9, X="N" * 12, Y="N" * 12)),
    ("C · Ondas", "Libera em ondas (duas operações a cada dois trimestres). A janela de 12 meses absorve a oscilação.",
     dict(Q=[2, 2, 1, 1] * 3, C=[0, 1, 1, 0] * 3, O=[0, 0, 1, 1] * 3,
          L=[0, 0, 500 * K, 500 * K, 0, 0, 550 * K, 550 * K, 0, 0, 600 * K, 600 * K], X="N" * 12, Y="N" * 12)),
    ("D · Pausa e retomada", "Chega ao N4, fica três trimestres sem operações, passa pela observação, é ajustado ao N3 e volta ao N4 pelo retorno expresso, só com o patamar de permanência.",
     dict(Q=[2, 2, 2, 2, 2, 0, 0, 0, 1, 1, 1, 2], C=[0, 1, 1, 1, 0, 0, 0, 0, 1, 1, 0, 1], O=[0, 0, 1, 1, 1, 0, 0, 0, 0, 1, 1, 1],
          L=[0, 0, 450 * K, 500 * K, 450 * K, 0, 0, 0, 0, 300 * K, 300 * K, 600 * K], X="N" * 12, Y="N" * 12)),
    ("E · Evento não imputável", "A instituição suspende a linha no 3T28 e no 4T28: os trimestres são neutros de pleno direito. O nível fica congelado e a janela é estendida.",
     dict(Q=[2, 2, 2, 2, 2, 2, 1, 1, 2, 2, 2, 2], C=[0, 1, 1, 1, 1, 0, 0, 0, 1, 1, 1, 1], O=[0, 0, 1, 1, 1, 1, 0, 0, 0, 1, 1, 1],
          L=[0, 0, 450 * K, 500 * K, 550 * K, 500 * K, 0, 0, 0, 500 * K, 550 * K, 600 * K], X="NNNNNNSSNNNN", Y="N" * 12)),
    ("F · Multiproduto", "Pausa o Crédito Internacional, mas fecha negócios em outro produto Save: a atividade conta e o amortecedor dá um 2º trimestre de observação. Com a linha Y em N, o resultado muda.",
     dict(Q=[2, 2, 2, 2, 1, 0, 1, 1, 1, 1, 1, 2], C=[0, 1, 1, 0, 0, 0, 0, 1, 0, 1, 0, 1], O=[0, 0, 1, 1, 0, 0, 0, 0, 2, 0, 1, 1],
          L=[0, 0, 450 * K, 500 * K, 0, 0, 0, 0, 600 * K, 0, 600 * K, 600 * K], M=[0, 0, 450 * K, 500 * K, 0, 0, 0, 0, 300 * K, 0, 600 * K, 600 * K],
          X="N" * 12, Y="S" * 12)),
    ("G · Realista 2027", "O funil do planejamento 2027: 1 indicação qualificada por trimestre e cerca de 1 operação liberada por ano (4T27, 3T28 e 4T29). É o caso típico do lançamento: a escada vai do N1 ao N3; como o N4 pede 3 operações em 24 meses, o parceiro fica no N3 e o sustenta, sem sobe e desce.",
     dict(Q=[1] * 12, C=[0, 0, 1, 0, 0, 1, 0, 0, 0, 0, 1, 0], O=[0, 0, 0, 1, 0, 0, 1, 0, 0, 0, 0, 1],
          L=[0, 0, 0, 450 * K, 0, 0, 450 * K, 0, 0, 0, 0, 450 * K], X="N" * 12, Y="N" * 12)),
]
# Canal: parceiro, operações liberadas no ano, líquido médio por operação, escopo típico, indicações qualificadas no ano
CANAL = [("Parceiro 1", 1, 450 * K, 3, 4), ("Parceiro 2", 1, 450 * K, 2, 3), ("Parceiro 3", 0, 0, 2, 3), ("Parceiro 4", 0, 0, 1, 2),
         ("Parceiro 5", 0, 0, 1, 2), ("Parceiro 6", 0, 0, 1, 1), ("Parceiro 7", 0, 0, 1, 1), ("Parceiro 8", 0, 0, 1, 0),
         ("Parceiro 9", 0, 0, 1, 0), ("Parceiro 10", 0, 0, 1, 0)]
SENS = [(0.02, 0.30), (0.02, 0.50), (0.03, 0.30), (0.03, 0.50), (0.04, 0.50), (0.04, 0.70)]

# linhas fixas (usadas pelo conferidor)
P_CAMBIO, P_FEE, P_SAVE, P_OP, P_HORA, P_HTRIM, P_AMOSTRA, P_HB4, P_HB5, P_HFIX = 5, 6, 7, 8, 9, 10, 11, 12, 13, 14
P_N0 = 18   # níveis nas linhas 18 a 22
P_E0 = 27   # escopos nas linhas 27 a 30
P_M0 = 34   # matriz: linhas 34 a 38, colunas C a F
# aba Sustentação
S_FATOR, S_OBS, S_BONUS, S_REC, S_MIN, S_EXT, S_L24, S_M24, S_FORM, S_POS, S_US = 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15
S_N0 = 19   # níveis nas linhas 19 a 23
S_SENS = 27  # sensibilidade: título 26, cabeçalho 27, linhas 28 a 33
S_P0 = 37   # 1º perfil
# chaves das linhas de apuração (comuns às duas abas)
LV = range(2, 6)
APUR = (["A", "EXT", "EXT24", "VS", "OP", "Q12", "REG", "STK", "TR", "CONC", "VS24", "OP24", "TR24"]
        + [f"XA{k}" for k in LV] + [f"XP{k}" for k in LV] + [f"VA{k}" for k in LV] + [f"VP{k}" for k in LV] + [f"V24_{k}" for k in LV]
        + [f"sobe{k}" for k in LV] + ["E"] + [f"sus{k}" for k in LV]
        + ["S", "MAXP", "LIMOBS", "PROT", "N", "W", "U", "SIT", "NOME", "VIG", "PCT", "GANHO", "ACUM", "ACUM10", "RET", "SAI", "NEC"])
EXTRA_J = ["IND12", "RCN", "REC", "DOS12", "DEV", "DOS", "OC"]
# perfis: entradas nas linhas 3 a 9 do bloco
ENT_P = ["Q", "C", "O", "L", "M", "X", "Y"]
PB = dict(tit=0, desc=1, cab=2, **{k: 3 + i for i, k in enumerate(ENT_P)}, lab=10)
PB.update({k: 11 + i for i, k in enumerate(APUR)})
S_BH = 11 + len(APUR) + 3
# jornada: entradas nas linhas 6 a 19
JR = {k: 6 + i for i, k in enumerate(CAMPOS_J)}
JR["lab"] = 6 + len(CAMPOS_J)
JR.update({k: JR["lab"] + 1 + i for i, k in enumerate(EXTRA_J + APUR)})
# Canal
C_L0 = 6   # parceiros nas linhas 6 a 15


def cab(ws, linha, textos, col0=1):
    for j, t in enumerate(textos, col0):
        c = ws.cell(linha, j, t); c.fill = F_HEAD; c.font = HFONT; c.alignment = CENTER; c.border = THIN


def celula(ws, ref, valor, fmt=None, entrada=False, negrito=False):
    c = ws[ref]; c.value = valor; c.border = THIN
    if fmt: c.number_format = fmt
    c.fill = F_IN if entrada else F_OUT
    if entrada: c.protection = Protection(locked=False)
    if negrito: c.font = BOLD
    return c


def validar(ws, ref, tipo, a, b, msg):
    dv = DataValidation(type=tipo, operator="between", formula1=str(a), formula2=str(b), showErrorMessage=True, error=msg)
    ws.add_data_validation(dv); dv.add(ref)


def premissas(wb):
    ws = wb.active; ws.title = "Premissas"
    ws["A1"] = "Premissas do programa · tudo em amarelo é editável"; ws["A1"].font = TITLE
    ws["A2"] = ("Os percentuais intermediários, os patamares de sustentação e o fee são HIPÓTESES para o Luis e a diretoria calibrarem. "
                "O fee de sucesso é definido caso a caso: o valor abaixo é só ilustrativo, não é regra. "
                "Os patamares de permanência e as regras de ajuste ficam na aba Sustentação.")
    ws["A2"].font = NOTA
    ws["A4"] = "Parâmetros da operação e hipóteses de custo"; ws["A4"].font = BOLD
    linhas = [
        (P_CAMBIO, "Câmbio (R$ por US$)", CAMBIO, '0.00', "Referência de 02/10/2026 (21-documento-interno, seção 12)"),
        (P_FEE, "Fee de sucesso ILUSTRATIVO (% do valor da operação)", FEE, '0.00%', "Ilustrativo e editável. O fee real é definido caso a caso (21, seção 10)"),
        (P_SAVE, "Parcela líquida da Save no fee (após a divisão com os demais participantes)", PARCELA_SAVE, PCT, "[HIPÓTESE] Divisão definida caso a caso"),
        (P_OP, "Valor de operação de referência (US$)", OPERACAO, USD, "Faixa do produto: US$ 5 milhões a US$ 10 milhões (21, seções 4 e 5)"),
        (P_HORA, "Custo por hora do time Save (R$)", CUSTO_HORA, RS, "[HIPÓTESE] Para a conta de economia"),
        (P_HTRIM, "Horas do time Save disponíveis para o CI por trimestre", HORAS_TRIM, '0', "[HIPÓTESE] Capacidade do dono do produto e da equipe"),
        (P_AMOSTRA, "Amostra mínima para os critérios de qualidade em percentual (dossiês ou indicações na janela)", AMOSTRA, '0',
         "Abaixo dela, vale o número de devoluções e de recusas (colunas N e O)"),
        (P_HB4, "Horas Save por trimestre com os benefícios de um parceiro N4", HB4, '0', "[HIPÓTESE] Canal direto, plano de conta, listas"),
        (P_HB5, "Horas Save por trimestre com os benefícios de um parceiro N5", HB5, '0', "[HIPÓTESE] Reunião de carteira, plano de conta, evento"),
        (P_HFIX, "Horas fixas do programa por trimestre (apuração, conferência, comitê e extratos)", HFIX, '0', "[HIPÓTESE]"),
    ]
    for r, rot, v, fmt, nota in linhas:
        ws.cell(r, 1, rot).border = THIN
        celula(ws, f"B{r}", v, fmt, entrada=True)
        ws.cell(r, 3, nota).font = NOTA
    validar(ws, f"B{P_FEE}", "decimal", 0, 1, "Informe um percentual entre 0% e 100%")
    validar(ws, f"B{P_SAVE}", "decimal", 0, 1, "Informe um percentual entre 0% e 100%")

    ws.cell(P_N0 - 2, 1, "Níveis e critérios para SUBIR (apuração trimestral, janela móvel de 12 meses = 4 trimestres)").font = BOLD
    cab(ws, P_N0 - 1, ["Nº", "Nível", "% teto (sobre o líquido Save)", "Trilha exigida (nº mín.)",
                       "Patamar de acesso: líquido Save 12m (R$)", "Operações liberadas 12m (mín.)",
                       "Alternativa: qualificadas 12m (0 = não se aplica)", "Trimestres seguidos com atividade",
                       "Limite de contagem por operação (% do patamar de acesso)", "Dossiês completos na 1ª entrega (mín.)",
                       "Taxa de recusa na pré-análise (máx.)", "Ocorrências de compliance ou conflito (máx.)", "Comitê Save exigido (S/N)",
                       "Devoluções máx. com amostra menor que a mínima", "Recusas máx. com amostra menor que a mínima",
                       "Operações liberadas em 24 meses para subir (mín.; 0 = não se aplica)"])
    fmts = [None, None, PCT, '0', RS, '0', '0', '0', PCT, PCT, PCT, '0', None, '0', '0', '0']
    for i, nv in enumerate(NIVEIS):
        r = P_N0 + i
        for j, v in enumerate(nv, 1):
            celula(ws, f"{CL(j)}{r}", v, fmts[j - 1], entrada=j >= 2)
        ws[f"A{r}"].fill = F_OUT
    validar(ws, f"C{P_N0}:C{P_N0 + 4}", "decimal", 0, 1, "Percentual entre 0% e 100%")
    validar(ws, f"I{P_N0}:K{P_N0 + 4}", "decimal", 0, 1, "Percentual entre 0% e 100%")
    ws.cell(P_N0 + 5, 1, ("Limite de contagem: no N4 e no N5, cada operação conta para o patamar no máximo esta fração dele; o excedente "
                          "aparece no extrato, mas não conta. Substitui a antiga trava de concentração: quem traz uma operação grande não é "
                          "reprovado, só não compra o nível com um caso só. Coluna P (v2.2): o N4 pede também 3 operações liberadas em 24 meses, "
                          "para que 2 operações próximas não levem a um N4 que não se sustenta.")).font = NOTA

    ws.cell(P_E0 - 2, 1, "Escopo executado na operação (define quanto o parceiro recebe naquela operação, até o teto do nível)").font = BOLD
    cab(ws, P_E0 - 1, ["Nº", "Escopo", "O que o parceiro faz (etapas da seção 8 do 21)", "% do escopo",
                       "Horas do time Save por operação quando o parceiro executa este escopo [HIPÓTESE]"])
    for i, e in enumerate(ESCOPOS):
        r = P_E0 + i
        for j, v in enumerate(e, 1):
            c = celula(ws, f"{CL(j)}{r}", v, [None, None, None, PCT, '0'][j - 1], entrada=j in (4, 5))
            c.alignment = WRAP

    ws.cell(P_M0 - 2, 1, "Matriz nível × escopo: percentual pago na operação = o MENOR entre o teto do nível e o % do escopo executado").font = BOLD
    cab(ws, P_M0 - 1, ["Nº", "Nível \\ Escopo"] + [e[1] for e in ESCOPOS])
    for i in range(5):
        r = P_M0 + i
        celula(ws, f"A{r}", f"=A{P_N0 + i}")
        celula(ws, f"B{r}", f"=B{P_N0 + i}", negrito=True)
        for k in range(4):
            celula(ws, f"{CL(3 + k)}{r}", f"=MIN($C${P_N0 + i},$D${P_E0 + k})", PCT)
    ws.cell(P_M0 + 6, 1, "Leitura: um parceiro no nível Estratégico que só indicou recebe 10%; um Indicador que executou tudo também recebe 10% (o teto do nível limita).").font = NOTA

    larg = [8, 44, 16, 14, 20, 16, 18, 20, 20, 18, 16, 18, 14, 16, 16, 18]
    for j, w in enumerate(larg, 1):
        ws.column_dimensions[CL(j)].width = w
    ws.column_dimensions["A"].width = 70; ws.column_dimensions["C"].width = 40
    ws.row_dimensions[P_N0 - 1].height = 80; ws.row_dimensions[P_E0 - 1].height = 60
    for i in range(4):
        ws.row_dimensions[P_E0 + i].height = 45
    ws.freeze_panes = "A4"
    return ws


def simulador(wb):
    ws = wb.create_sheet("Simulador de operação")
    ws["A1"] = "Simulador de uma operação"; ws["A1"].font = TITLE
    ws["A2"] = "Amarelo = entrada. As linhas 4 a 7 puxam a aba Premissas; sobrescreva para simular um caso específico."; ws["A2"].font = NOTA
    ent = [
        (4, "Valor da operação (US$)", f"=Premissas!B{P_OP}", USD),
        (5, "Câmbio (R$ por US$)", f"=Premissas!B{P_CAMBIO}", '0.00'),
        (6, "Fee de sucesso ILUSTRATIVO (% do valor)", f"=Premissas!B{P_FEE}", '0.00%'),
        (7, "Parcela líquida da Save no fee", f"=Premissas!B{P_SAVE}", PCT),
        (8, "Nível do parceiro (1 a 5)", 3, '0'),
        (9, "Escopo executado nesta operação (1 a 4)", 3, '0'),
    ]
    for r, rot, v, fmt in ent:
        ws.cell(r, 1, rot).border = THIN
        celula(ws, f"B{r}", v, fmt, entrada=True)
    validar(ws, "B8", "whole", 1, 5, "Informe um nível de 1 a 5")
    validar(ws, "B9", "whole", 1, 4, "Informe um escopo de 1 a 4")
    ws["C8"] = "1 Indicador · 2 Qualificador · 3 Condutor · 4 Ponta a Ponta · 5 Estratégico"; ws["C8"].font = NOTA
    ws["C9"] = "1 Indicação · 2 Qualificação · 3 Condução documental · 4 Ponta a ponta"; ws["C9"].font = NOTA

    ws["A11"] = "Resultado"; ws["A11"].font = BOLD
    res = [
        (12, "Valor da operação (R$)", "=B4*B5", RS),
        (13, "Fee total ILUSTRATIVO (R$)", "=B12*B6", RS),
        (14, "Valor líquido Save (R$): base do repasse", "=B13*B7", RS),
        (15, "Nível", f"=INDEX(Premissas!B{P_N0}:B{P_N0 + 4},B8)", None),
        (16, "Teto do nível", f"=INDEX(Premissas!C{P_N0}:C{P_N0 + 4},B8)", PCT),
        (17, "Escopo executado", f"=INDEX(Premissas!B{P_E0}:B{P_E0 + 3},B9)", None),
        (18, "% do escopo", f"=INDEX(Premissas!D{P_E0}:D{P_E0 + 3},B9)", PCT),
        (19, "% aplicado nesta operação (menor entre teto e escopo)", "=MIN(B16,B18)", PCT),
        (20, "Repasse ao parceiro (R$)", "=B14*B19", RS),
        (21, "Retido pela Save (R$)", "=B14-B20", RS),
        (22, "Margem retida pela Save (% do líquido)", "=IFERROR(B21/B14,0)", PCT1),
        (23, "Repasse se o parceiro só tivesse indicado (E1)", f"=B14*MIN(B16,Premissas!D{P_E0})", RS),
        (24, "Diferença paga pelo escopo executado", "=B20-B23", RS),
    ]
    for r, rot, v, fmt in res:
        ws.cell(r, 1, rot).border = THIN
        celula(ws, f"B{r}", v, fmt, negrito=r in (14, 19, 20, 21))

    ws["A26"] = "Pagamento por parcela (o repasse acompanha cada recebimento efetivo da Save)"; ws["A26"].font = BOLD
    cab(ws, 27, ["Parcela", "% do valor liberado nesta parcela", "Líquido Save da parcela (R$)",
                 "Repasse da parcela (R$)", "Retido pela Save (R$)", "Quando o parceiro recebe"])
    pcts = [0.3333, 0.3333, 0.3334, 0]
    for k in range(4):
        r = 28 + k
        celula(ws, f"A{r}", k + 1)
        celula(ws, f"B{r}", pcts[k], PCT1, entrada=True)
        celula(ws, f"C{r}", f"=$B$14*B{r}", RS)
        celula(ws, f"D{r}", f"=C{r}*$B$19", RS)
        celula(ws, f"E{r}", f"=C{r}-D{r}", RS)
        celula(ws, f"F{r}", "Até 30 dias após o fim do mês em que a Save receber o fee desta parcela, com nota fiscal da PJ")
    validar(ws, "B28:B31", "decimal", 0, 1, "Percentual entre 0% e 100%")
    celula(ws, "A32", "Total", negrito=True)
    celula(ws, "B32", "=SUM(B28:B31)", PCT1, negrito=True)
    for col in "CDE":
        celula(ws, f"{col}32", f"=SUM({col}28:{col}31)", RS, negrito=True)
    celula(ws, "F32", '=IF(ROUND(B32,4)=1,"OK: as parcelas somam 100%","ATENÇÃO: as parcelas precisam somar 100%")', negrito=True)
    ws["A34"] = ("Valores sem retenções (IRRF, CSLL/PIS/COFINS, ISS) e sem custos ordinários do cliente. "
                 "Retenções: [VALIDAR TRIBUTÁRIO]. Se a Save devolver o fee, o repasse da parcela é estornado. "
                 "O percentual fica travado na assinatura do contrato da operação pelo cliente: um ajuste de nível depois disso não muda as parcelas.")
    ws["A34"].font = NOTA
    for col, w in zip("ABCDEF", [52, 22, 24, 22, 22, 60]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[27].height = 45
    return ws


# ------------------------------------------------------------------ janela móvel com extensão por trimestre neutro
def _jan(t, n=4, c0=3):
    """Colunas da janela de n trimestres e das 2 extensões possíveis para o trimestre de índice t (0 = 1ª coluna)."""
    c = CL(c0 + t); ini = CL(c0 + max(0, t - n + 1))
    ext = [CL(c0 + t - n - k) if t - n - k >= 0 else None for k in range(2)]
    return c, ini, ext


def jsum(t, r, rext, n=4):
    c, ini, ext = _jan(t, n)
    f = f"SUM({ini}{r}:{c}{r})"
    for k, e in enumerate(ext):
        if e: f += f"+IF({c}{rext}>={k + 1},{e}{r},0)"
    return f


def jmax(t, r, rext, n=4):
    c, ini, ext = _jan(t, n)
    partes = [f"{ini}{r}:{c}{r}"] + [f"IF({c}{rext}>={k + 1},{e}{r},0)" for k, e in enumerate(ext) if e]
    return f"MAX({','.join(partes)})"


def jcountpos(t, r, rext, n=4):
    c, ini, ext = _jan(t, n)
    f = f'COUNTIF({ini}{r}:{c}{r},">0")'
    for k, e in enumerate(ext):
        if e: f += f"+IF(AND({c}{rext}>={k + 1},{e}{r}>0),1,0)"
    return f


ROTULOS = {
    "A": "Trimestre com atividade (qualificada, contratada, liberada ou negócio em outro produto Save)",
    "EXT": "Trimestres neutros na janela de 12 meses (estendem a janela, máx. 2)",
    "EXT24": "Trimestres neutros na janela de 24 meses (estendem a janela, máx. 2)",
    "VS": "Volume de sustentação: líquido Save 12m (R$)", "OP": "Operações liberadas 12m", "Q12": "Qualificadas 12m",
    "REG": "Regularidade: trimestres com atividade na janela", "STK": "Trimestres seguidos com atividade (neutro não interrompe)",
    "TR": "Regularidade financeira: trimestres com receita na janela",
    "CONC": "Maior operação ÷ volume de sustentação (informativo; o limite por operação substitui a trava)",
    "VS24": "Volume de sustentação em 24 meses (R$)", "OP24": "Operações liberadas em 24 meses", "TR24": "Trimestres com receita em 24 meses",
    "IND12": "Indicações registradas 12m", "RCN": "Indicações recusadas 12m", "REC": "Taxa de recusa na pré-análise 12m (recusadas ÷ registradas)",
    "DOS12": "Dossiês entregues 12m", "DEV": "Dossiês devolvidos 12m", "DOS": "Dossiês completos na 1ª entrega 12m (%)", "OC": "Ocorrências procedentes 12m",
    "E": "Nível elegível para subir", "S": "Nível sustentado (maior nível cujo patamar de permanência é cumprido, em 12 ou 24 meses)",
    "MAXP": "Maior nível nas 4 revisões anteriores (referência do retorno expresso)", "LIMOBS": "Trimestres de observação permitidos",
    "PROT": "Período de proteção (sem ajuste): 4 primeiras apurações ou 2 seguintes a uma subida",
    "N": "Nível apurado na revisão", "W": "Trimestres em observação (contador)", "U": "Apurações desde a última subida",
    "SIT": "Situação na revisão", "NOME": "Nome do nível apurado", "VIG": "Nível vigente no trimestre (para pagar)",
    "PCT": "% aplicado no trimestre", "GANHO": "Ganho do parceiro no trimestre (R$)", "ACUM": "Ganho acumulado do parceiro (R$)",
    "ACUM10": "Comparação: ganho acumulado se ficasse em 10% (R$)", "RET": "Retido pela Save no trimestre (R$)",
    "SAI": "Volume que sai da janela na próxima revisão (R$)",
    "NEC": "Volume necessário na próxima revisão para o patamar de permanência em 12 meses (R$)",
}
for _k in LV:
    ROTULOS[f"XA{_k}"] = f"Excedente acima do limite por operação do patamar de acesso do N{_k} (R$)"
    ROTULOS[f"XP{_k}"] = f"Excedente acima do limite por operação do patamar de permanência do N{_k} (R$)"
    ROTULOS[f"VA{_k}"] = f"Volume que conta para o acesso ao N{_k} (12m, R$)"
    ROTULOS[f"VP{_k}"] = f"Volume que conta para a permanência no N{_k} (12m, R$)"
    ROTULOS[f"V24_{_k}"] = f"Volume que conta para a permanência no N{_k} (24m, R$)"
    ROTULOS[f"sobe{_k}"] = f"Cumpre os critérios para subir ao N{_k}?"
    ROTULOS[f"sus{_k}"] = f"Cumpre o patamar de permanência do N{_k} (12 ou 24 meses)?"


def apurar(ws, R, ncols, completo):
    """Escreve as linhas de apuração (comuns às abas Sustentação e Jornada). R = linhas absolutas por chave."""
    for key in (EXTRA_J if completo else []) + APUR:
        ws.cell(R[key], 1, ROTULOS[key]).border = THIN
    if not completo:
        ws.cell(R["PCT"], 1, "Teto vigente (escopo E4 em todas as operações)").border = THIN
    celula(ws, f"B{R['N']}", 1, '0', entrada=True); celula(ws, f"B{R['W']}", 0, '0'); celula(ws, f"B{R['STK']}", 0, '0')
    celula(ws, f"B{R['U']}", 99, '0'); celula(ws, f"B{R['ACUM']}", 0, RS); celula(ws, f"B{R['ACUM10']}", 0, RS)
    N, SN = P_N0, S_N0
    for t in range(ncols):
        c = CL(3 + t); p = CL(2 + t); ini = CL(max(3, 3 + t - 3)); ini8 = CL(max(3, 3 + t - 7))
        G = lambda k: f"{c}{R[k]}"
        cont = f"+{G('C')}" if "C" in R else ""
        celula(ws, G("A"), f'=IF(OR({G("Q")}{cont}+{G("O")}>0,{G("Y")}="S"),1,0)', '0')
        celula(ws, G("EXT"), f'=MIN({SUS}!$B${S_EXT},COUNTIF({ini}{R["X"]}:{c}{R["X"]},"S"))', '0')
        celula(ws, G("EXT24"), f'=MIN({SUS}!$B${S_EXT},COUNTIF({ini8}{R["X"]}:{c}{R["X"]},"S"))', '0')
        celula(ws, G("VS"), "=" + jsum(t, R["L"], R["EXT"]), RS)
        celula(ws, G("OP"), "=" + jsum(t, R["O"], R["EXT"]), '0')
        celula(ws, G("Q12"), "=" + jsum(t, R["Q"], R["EXT"]), '0')
        celula(ws, G("REG"), "=" + jsum(t, R["A"], R["EXT"]), '0')
        celula(ws, G("STK"), f'=IF({G("A")}>0,{p}{R["STK"]}+1,IF({G("X")}="S",{p}{R["STK"]},0))', '0')
        celula(ws, G("TR"), "=" + jcountpos(t, R["L"], R["EXT"]), '0')
        celula(ws, G("CONC"), f"=IFERROR(ROUND({jmax(t, R['M'], R['EXT'])}/{G('VS')},4),1)", PCT)
        celula(ws, G("VS24"), "=" + jsum(t, R["L"], R["EXT24"], 8), RS)
        celula(ws, G("OP24"), "=" + jsum(t, R["O"], R["EXT24"], 8), '0')
        celula(ws, G("TR24"), "=" + jcountpos(t, R["L"], R["EXT24"], 8), '0')
        if completo:
            celula(ws, G("IND12"), "=" + jsum(t, R["ind"], R["EXT"]), '0')
            celula(ws, G("RCN"), "=" + jsum(t, R["recus"], R["EXT"]), '0')
            celula(ws, G("REC"), f"=IFERROR(ROUND({G('RCN')}/{G('IND12')},4),0)", PCT)
            celula(ws, G("DOS12"), "=" + jsum(t, R["dos"], R["EXT"]), '0')
            celula(ws, G("DEV"), f"={G('DOS12')}-({jsum(t, R['compl'], R['EXT'])})", '0')
            celula(ws, G("DOS"), f"=IFERROR(ROUND(({G('DOS12')}-{G('DEV')})/{G('DOS12')},4),0)", PCT)
            celula(ws, G("OC"), f"=SUM({ini}{R['oc']}:{c}{R['oc']})", '0')
        for k in LV:
            r = N + k - 1; s = SN + k - 1
            celula(ws, G(f"XA{k}"), f"=MAX(0,{G('M')}-Premissas!$I${r}*Premissas!$E${r})", RS)
            celula(ws, G(f"XP{k}"), f"=MAX(0,{G('M')}-{SUS}!$I${s}*{SUS}!$F${s})", RS)
            celula(ws, G(f"VA{k}"), f"={G('VS')}-({jsum(t, R[f'XA{k}'], R['EXT'])})", RS)
            celula(ws, G(f"VP{k}"), f"={G('VS')}-({jsum(t, R[f'XP{k}'], R['EXT'])})", RS)
            celula(ws, G(f"V24_{k}"), f"={G('VS24')}-({jsum(t, R[f'XP{k}'], R['EXT24'], 8)})", RS)
            qual = ""
            if completo:
                am = f"Premissas!$B${P_AMOSTRA}"
                qual = (f",{G('tri')}>=Premissas!$D${r},IF({G('DOS12')}>={am},{G('DOS')}>=Premissas!$J${r},{G('DEV')}<=Premissas!$N${r}),"
                        f"IF({G('IND12')}>={am},{G('REC')}<=Premissas!$K${r},{G('RCN')}<=Premissas!$O${r}),{G('OC')}<=Premissas!$L${r}")
            com = f",OR(Premissas!$M${r}=\"N\",{G('com')}=\"S\")" if completo else ""
            celula(ws, G(f"sobe{k}"),
                   f"=AND(OR(AND({G('OP')}>=Premissas!$F${r},{G(f'VA{k}')}>=Premissas!$E${r}),AND(Premissas!$G${r}>0,{G('Q12')}>=Premissas!$G${r})),"
                   f"{G('STK')}>=Premissas!$H${r},{G('TR')}>={SUS}!$K${s},{G('OP24')}>=Premissas!$P${r}{qual}{com})")
            m24 = f"{SUS}!$B${S_M24}"
            celula(ws, G(f"sus{k}"),
                   f"=AND({G('REG')}>={SUS}!$H${s},OR(AND({G(f'VP{k}')}>={SUS}!$F${s},{G('OP')}>={SUS}!$G${s},{G('TR')}>={SUS}!$J${s}),"
                   f"AND({SUS}!$B${S_L24}=\"S\",{G(f'V24_{k}')}>={m24}*{SUS}!$F${s},{G('OP24')}>={m24}*{SUS}!$G${s},{G('TR24')}>={m24}*{SUS}!$J${s})){qual})")
        celula(ws, G("E"), f"=IF({G('sobe5')},5,IF({G('sobe4')},4,IF({G('sobe3')},3,IF({G('sobe2')},2,1))))", '0')
        celula(ws, G("S"), f"=IF({G('sus5')},5,IF({G('sus4')},4,IF({G('sus3')},3,IF({G('sus2')},2,1))))", '0')
        mp_ini = CL(max(2, 2 + t - 3))
        celula(ws, G("MAXP"), f"=MAX({mp_ini}{R['N']}:{p}{R['N']})", '0')
        celula(ws, G("LIMOBS"), f'={SUS}!$B${S_OBS}+IF(COUNTIF({ini}{R["Y"]}:{c}{R["Y"]},"S")>0,{SUS}!$B${S_BONUS},0)', '0')
        celula(ws, G("PROT"), f"=OR({t + 1}<={SUS}!$B${S_FORM},{p}{R['U']}+1<={SUS}!$B${S_POS})")
        Np, Wp, Up, E, S, MP, X, LO, PR = (f"{p}{R['N']}", f"{p}{R['W']}", f"{p}{R['U']}", G("E"), G("S"), G("MAXP"), G("X"), G("LIMOBS"), G("PROT"))
        sobe = f"OR({E}>{Np},AND({S}>{Np},{MP}>{Np}))"
        oc0, oc1, oc2 = (f"IF({G('OC')}>0,1,", ")", f"IF({G('OC')}>0,0,") if completo else ("", "", "")
        celula(ws, G("N"), f"={oc0}IF({sobe},{Np}+1,IF({S}>={Np},{Np},IF({X}=\"S\",{Np},IF({Wp}<{LO},{Np},IF({PR},{Np},"
                           f"MAX({SUS}!$B${S_MIN},{Np}-1))))))" + (oc1 if completo else ""), '0', negrito=True)
        celula(ws, G("W"), f"={oc2}IF(OR({sobe},{S}>={Np}),0,IF({X}=\"S\",{Wp},IF({Wp}<{LO},{Wp}+1,IF({PR},{Wp},0))))" + (oc1 if completo else ""), '0')
        celula(ws, G("U"), f"=IF({G('N')}>{Np},0,{Up}+1)", '0')
        sit = (f"IF({E}>{Np},\"Promovido\",IF(AND({S}>{Np},{MP}>{Np}),\"Retomado\",IF({S}>={Np},\"Sustentado\","
               f"IF({X}=\"S\",\"Neutro\",IF({Wp}<{LO},\"Em observação\",IF({PR},\"Em observação (período de proteção)\","
               f"IF({Np}>{SUS}!$B${S_MIN},\"Ajustado ao N\"&({Np}-1),\"No nível de entrada (N1)\")))))))")
        celula(ws, G("SIT"), f"=IF({G('OC')}>0,\"Compliance: ajustado ao N1\",{sit})" if completo else "=" + sit)
        celula(ws, G("NOME"), f"=INDEX(Premissas!$B${N}:$B${N + 4},{G('N')})", negrito=True)
        celula(ws, G("VIG"), f"={Np}", '0')
        teto = f"INDEX(Premissas!$C${N}:$C${N + 4},{G('VIG')})"
        celula(ws, G("PCT"), f"=MIN({teto},INDEX(Premissas!$D${P_E0}:$D${P_E0 + 3},{G('esc')}))" if completo else f"={teto}", PCT)
        celula(ws, G("GANHO"), f"={G('L')}*{G('PCT')}", RS)
        celula(ws, G("ACUM"), f"={p}{R['ACUM']}+{G('GANHO')}", RS, negrito=True)
        celula(ws, G("ACUM10"), f"={p}{R['ACUM10']}+{G('L')}*Premissas!$C${N}", RS)
        celula(ws, G("RET"), f"={G('L')}-{G('GANHO')}", RS)
        sai = f"{CL(3 + t - 3)}{R['L']}" if t >= 3 else "0"
        celula(ws, G("SAI"), f"={sai}", RS)
        celula(ws, G("NEC"), f"=MAX(0,INDEX({SUS}!$F${SN}:$F${SN + 4},{G('N')})-({G('VS')}-{G('SAI')}))", RS)


def jornada(wb):
    ws = wb.create_sheet("Jornada do parceiro")
    ws["A1"] = "Jornada do parceiro em 8 trimestres · nível calculado por fórmula em cada revisão (critérios completos)"; ws["A1"].font = TITLE
    ws["A2"] = ("Amarelo = entrada (exemplo hipotético; troque pelos números do parceiro). Regras v2.2: sobe no máximo 1 nível por revisão "
                "(critérios na aba Premissas); mantém o nível enquanto cumprir o patamar de permanência (aba Sustentação) em 12 meses ou, "
                "na leitura de 24 meses, o dobro dele; abaixo dele, 1 trimestre de observação (2 com negócio em outro produto Save) e depois "
                "o nível apurado passa a ser o imediatamente inferior, nunca abaixo do N1; nas 4 primeiras apurações e nas 2 seguintes a "
                "uma subida não há ajuste; volta ao nível anterior em até 4 revisões só com o patamar de permanência (retorno expresso); "
                "trimestre neutro congela o nível e estende a janela; ocorrência julgada procedente ajusta ao N1. "
                "Qualidade: com menos de 5 dossiês ou indicações na janela, vale o número de devoluções e de recusas. "
                "O nível usado para pagar em cada trimestre é o da revisão anterior.")
    ws["A2"].font = NOTA; ws["A2"].alignment = WRAP; ws.merge_cells("A2:J2"); ws.row_dimensions[2].height = 95
    cab(ws, 4, ["Linha", "Início"] + TRIMS)
    rot_in = ["Indicações registradas", "Indicações qualificadas (aprovadas na pré-análise da Save)", "Indicações recusadas na pré-análise",
              "Dossiês entregues", "Dossiês completos na 1ª entrega", "Operações liberadas no trimestre (1ª parcela)",
              "Líquido Save do trimestre (R$)", "Maior operação do trimestre (líquido Save, R$)",
              "Escopo predominante executado (1 a 4)", "Trilha concluída e vigente (nº 1 a 5)",
              "Ocorrências de compliance julgadas procedentes", "Comitê aprovou o nível 5? (S/N)",
              "Trimestre neutro (evento não imputável ou rol do Anexo I)? (S/N)", "Negócio ganho em outro produto Save no trimestre? (S/N)"]
    fm_in = ['0', '0', '0', '0', '0', '0', RS, RS, '0', '0', '0', None, None, None]
    ws.cell(5, 1, "Entradas do trimestre").font = BOLD
    for i, rot in enumerate(rot_in):
        r = 6 + i
        ws.cell(r, 1, rot).border = THIN
        for t in range(8):
            celula(ws, f"{CL(3 + t)}{r}", JORNADA[t][i], fm_in[i], entrada=True)
    ws.cell(JR["lab"], 1, "Apuração na revisão (janela móvel de 12 meses = o trimestre e os 3 anteriores; 24 meses = 8 trimestres)").font = BOLD
    apurar(ws, JR, 8, True)
    ws.column_dimensions["A"].width = 72; ws.column_dimensions["B"].width = 10
    for t in range(8):
        ws.column_dimensions[CL(3 + t)].width = 15
    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=False); ws.add_data_validation(dv)
    for k in ("com", "X", "Y"):
        dv.add(f"C{JR[k]}:J{JR[k]}")
    validar(ws, f"C{JR['esc']}:J{JR['esc']}", "whole", 1, 4, "Escopo de 1 a 4")
    validar(ws, f"C{JR['tri']}:J{JR['tri']}", "whole", 0, 5, "Trilha de 0 a 5")
    ws.freeze_panes = "C5"
    return ws


def sustentacao(wb):
    ws = wb.create_sheet("Sustentação")
    ws["A1"] = "Sustentação de volume financeiro · patamares de acesso e de permanência, observação, ajuste e retorno"; ws["A1"].font = TITLE
    ws["A2"] = ("Amarelo = entrada. Volume de sustentação = líquido Save das operações do parceiro (faturado até a conciliação do recebido; "
                "depois, recebido), em janela móvel de 12 meses, apurado no fechamento de cada trimestre. O patamar de ACESSO vem da aba "
                "Premissas (coluna E); o patamar de PERMANÊNCIA é uma fração dele (faixa de tolerância). Para manter, vale a melhor de duas "
                "leituras: 12 meses ou o dobro em 24 meses. Os perfis usam só volume, regularidade e limite por operação; conhecimento, "
                "qualidade e comitê estão na aba Jornada do parceiro.")
    ws["A2"].font = NOTA; ws["A2"].alignment = WRAP; ws.merge_cells("A2:N2"); ws.row_dimensions[2].height = 70
    ws["A4"] = "Regras de sustentação e de ajuste"; ws["A4"].font = BOLD
    regras = [
        (S_FATOR, "Patamar de permanência = patamar de acesso × este fator (faixa de tolerância)", FATOR_MANUT, PCT,
         "[HIPÓTESE] 75%: quem subiu com folga não é ajustado por uma oscilação pequena"),
        (S_OBS, "Trimestres de observação antes de cada ajuste", OBS, '0', "1 = um trimestre com aviso por escrito antes de qualquer ajuste"),
        (S_BONUS, "Trimestres extras de observação com negócio em outro produto Save nos últimos 12 meses (amortecedor multiproduto)", OBS_BONUS, '0',
         "[HIPÓTESE] Premia o equilíbrio entre produtos sem misturar as bases de cálculo"),
        (S_REC, "Janela do retorno expresso (revisões após o ajuste)", JAN_REC, '0',
         "Dentro dela, o parceiro volta ao nível anterior cumprindo só o patamar de permanência"),
        (S_MIN, "Nível mínimo do programa (o ajuste nunca passa daqui)", NIVEL_MIN, '0',
         "N1 = 10% do líquido Save (D61): o percentual do parceiro nunca fica abaixo de 10%"),
        (S_EXT, "Trimestres neutros que podem estender a janela (máx.)", MAX_EXT, '0', "Evento não imputável ao parceiro: congela o nível e estende a janela"),
        (S_L24, "Leitura de 24 meses para manter (S/N)", LEIT24, None,
         "Regra anti-oscilação: mantém quem cumpre o patamar em 12 meses OU o dobro dele em 24 meses (com o dobro de operações e de trimestres com receita)"),
        (S_M24, "Multiplicador da leitura de 24 meses", MULT24, '0', "2 = o dobro do patamar de permanência, das operações e dos trimestres com receita"),
        (S_FORM, "Apurações sem ajuste depois da adesão (período de formação)", FORMACAO, '0', "Só subida, observação com plano e extrato"),
        (S_POS, "Apurações sem ajuste depois de cada subida", POS_SUBIDA, '0', "Dá tempo de o parceiro se firmar no nível novo"),
    ]
    for r, rot, v, fmt, nota in regras:
        ws.cell(r, 1, rot).border = THIN
        celula(ws, f"B{r}", v, fmt, entrada=True)
        ws.cell(r, 3, nota).font = NOTA
    validar(ws, f"B{S_FATOR}", "decimal", 0.5, 1, "O fator fica entre 50% e 100%")
    validar(ws, f"B{S_OBS}:B{S_REC}", "whole", 0, 4, "Número inteiro de 0 a 4")
    dvl = DataValidation(type="list", formula1='"S,N"', allow_blank=False); ws.add_data_validation(dvl); dvl.add(f"B{S_L24}")
    ws.cell(S_US, 1, "Unidade de sustentação: líquido Save de uma operação mínima (US$ 5 milhões, com as premissas)").border = THIN
    celula(ws, f"B{S_US}", f"=5000000*Premissas!B{P_CAMBIO}*Premissas!B{P_FEE}*Premissas!B{P_SAVE}", RS, negrito=True)
    ws.cell(S_US, 3, "Régua para ler os patamares em número de operações; muda sozinha quando o fee ou a parcela Save mudam").font = NOTA

    ws.cell(S_N0 - 2, 1, "Patamares de sustentação por nível (o de acesso é editado na aba Premissas; os demais, aqui)").font = BOLD
    cab(ws, S_N0 - 1, ["Nº", "Nível", "% teto", "Patamar de acesso: líquido Save 12m (R$)", "Patamar de acesso em operações mínimas",
                       "Patamar de permanência: líquido Save 12m (R$)", "Operações liberadas 12m para manter (mín.)",
                       "Trimestres com atividade na janela para manter (de 4)", "Limite de contagem por operação (% do patamar de permanência)",
                       "Trimestres com receita na janela para manter (de 4)", "Trimestres com receita na janela para subir (de 4)",
                       "Leitura"])
    leitura = ["Entrada (D61). Não há patamar: é o nível mínimo, e o percentual nunca fica abaixo de 10%.",
               "Mantém com atividade em 2 dos 4 trimestres (indicação qualificada, operação ou negócio em outro produto Save).",
               "Mantém com 1 operação liberada em 12 meses (ou 2 em 24) e 75% do patamar de acesso.",
               "Mantém com 2 operações (ou 4 em 24 meses); cada uma conta até 80% do patamar.",
               "Mantém com 2 operações, receita em 2 trimestres e cada operação contando até 60% do patamar."]
    for i in range(5):
        r = S_N0 + i; pr = P_N0 + i; g, h, conc, trm, trs = SUST[i]
        celula(ws, f"A{r}", f"=Premissas!A{pr}", '0')
        celula(ws, f"B{r}", f"=Premissas!B{pr}", negrito=True)
        celula(ws, f"C{r}", f"=Premissas!C{pr}", PCT)
        celula(ws, f"D{r}", f"=Premissas!E{pr}", RS)
        celula(ws, f"E{r}", f"=IFERROR(D{r}/$B${S_US},0)", '0.0')
        celula(ws, f"F{r}", f"=ROUND(D{r}*$B${S_FATOR},0)", RS, entrada=True)
        celula(ws, f"G{r}", g, '0', entrada=True)
        celula(ws, f"H{r}", h, '0', entrada=True)
        celula(ws, f"I{r}", conc, PCT, entrada=True)
        celula(ws, f"J{r}", trm, '0', entrada=True)
        celula(ws, f"K{r}", trs, '0', entrada=True)
        c = celula(ws, f"L{r}", leitura[i]); c.alignment = WRAP
    validar(ws, f"G{S_N0}:H{S_N0 + 4}", "whole", 0, 4, "Número inteiro de 0 a 4")
    validar(ws, f"J{S_N0}:K{S_N0 + 4}", "whole", 0, 4, "Número inteiro de 0 a 4")
    validar(ws, f"I{S_N0}:I{S_N0 + 4}", "decimal", 0, 1, "Percentual entre 0% e 100%")
    ws.cell(S_N0 + 5, 1, ("A coluna F é fórmula (patamar de acesso × fator), mas pode ser sobrescrita. Os patamares em R$ são HIPÓTESES "
                          "calibradas com o fee ilustrativo; são fixados uma vez com o fee real (P15) antes do G2 e só mudam na revisão anual, "
                          "com 90 dias de aviso.")).font = NOTA

    ws.cell(S_SENS - 1, 1, "Sensibilidade: patamares de ACESSO em operações mínimas, conforme o fee e a parcela Save (troque as entradas amarelas)").font = BOLD
    cab(ws, S_SENS, ["Fee (% da operação)", "Parcela Save", "Unidade (R$)", "N3", "N4", "N5"])
    for i, (fee, sv) in enumerate(SENS):
        r = S_SENS + 1 + i
        celula(ws, f"A{r}", fee, '0.0%', entrada=True); celula(ws, f"B{r}", sv, PCT, entrada=True)
        celula(ws, f"C{r}", f"=5000000*Premissas!$B${P_CAMBIO}*A{r}*B{r}", RS)
        for j, lvl in enumerate((3, 4, 5)):
            celula(ws, f"{CL(4 + j)}{r}", f"=IFERROR(Premissas!$E${P_N0 + lvl - 1}/C{r},0)", '0.0')
    validar(ws, f"A{S_SENS + 1}:B{S_SENS + len(SENS)}", "decimal", 0, 1, "Percentual entre 0% e 100%")

    for i, (nome, desc, d) in enumerate(PERFIS):
        perfil(ws, S_P0 + i * S_BH, nome, desc, d)
    ws.column_dimensions["A"].width = 72; ws.column_dimensions["B"].width = 12
    for j in range(3, 15):
        ws.column_dimensions[CL(j)].width = 15
    ws.column_dimensions["L"].width = 40
    ws.row_dimensions[S_N0 - 1].height = 80
    ws.freeze_panes = "B4"
    return ws


def perfil(ws, r0, nome, desc, d):
    R = {k: r0 + v for k, v in PB.items()}
    ws.cell(R["tit"], 1, f"Perfil {nome}").font = SUBT
    ws.cell(R["desc"], 1, desc).font = NOTA
    cab(ws, R["cab"], ["Linha", "Início"] + TRIMS12)
    entradas = [("Q", "Indicações qualificadas no trimestre", '0'), ("C", "Operações contratadas (assinadas pelo cliente) no trimestre", '0'),
                ("O", "Operações liberadas no trimestre (1ª parcela)", '0'), ("L", "Líquido Save do trimestre (R$)", RS),
                ("M", "Maior operação do trimestre (líquido Save, R$)", RS),
                ("X", "Trimestre neutro (evento não imputável ou rol do Anexo I)? (S/N)", None),
                ("Y", "Negócio ganho em outro produto Save no trimestre? (S/N)", None)]
    for key, rot, fmt in entradas:
        ws.cell(R[key], 1, rot).border = THIN
        for t in range(12):
            v = d.get("M", d["L"])[t] if key == "M" else d[key][t]
            celula(ws, f"{CL(3 + t)}{R[key]}", v, fmt, entrada=True)
    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=False); ws.add_data_validation(dv)
    dv.add(f"C{R['X']}:N{R['X']}"); dv.add(f"C{R['Y']}:N{R['Y']}")
    ws.cell(R["lab"], 1, "Apuração trimestral (janela móvel de 12 meses; leitura de 24 meses para manter)").font = BOLD
    apurar(ws, R, 12, False)


def comparativo(wb):
    ws = wb.create_sheet("Comparativo")
    ws["A1"] = "O mesmo parceiro em cada nível · mesma operação e mesmo volume anual"; ws["A1"].font = TITLE
    ws["A2"] = "A operação vem da aba Simulador de operação (líquido Save da linha 14). Horas e custo por hora são hipóteses da aba Premissas."; ws["A2"].font = NOTA
    ws["A4"] = "Líquido Save por operação (R$)"; celula(ws, "B4", "='Simulador de operação'!B14", RS)
    ws["A5"] = "Operações liberadas por ano deste parceiro (premissa)"; celula(ws, "B5", 2, '0', entrada=True)
    ws["A6"] = "Operações liberadas por ano no canal inteiro (plano de 2027: 2)"; celula(ws, "B6", 2, '0', entrada=True)
    validar(ws, "B5:B6", "whole", 1, 100, "Número inteiro de 1 a 100")
    cab(ws, 7, ["Nº", "Nível", "Escopo típico do nível (nº)", "% aplicado", "Repasse por operação (R$)",
                "Retido pela Save por operação (R$)", "% retido pela Save", "Repasse no ano (R$)", "Retido pela Save no ano (R$)",
                "Volume extra que a Save precisa para manter o retido do nível 1", "Horas Save por operação [HIPÓTESE]",
                "Custo das horas Save por operação (R$)", "Operações por trimestre na capacidade do time",
                "Retido pela Save por trimestre na capacidade (R$)", "Horas de benefícios por operação [HIPÓTESE]",
                "Custo fixo do programa por operação (R$) [HIPÓTESE]", "Contribuição por operação: retido − horas Save (R$)",
                "Volume extra para empatar com o N1 pela contribuição", "Contribuição com benefícios e custo fixo (R$)",
                "Volume extra para empatar com o N1 (com benefícios e custo fixo)"])
    linhas = [(i + 1, min(i + 1, 4)) for i in range(5)] + [(5, 1)]
    for i, (nv, esc) in enumerate(linhas):
        r = 8 + i
        celula(ws, f"A{r}", nv, '0')
        nome = f"=INDEX(Premissas!$B${P_N0}:$B${P_N0 + 4},A{r})" if i < 5 else f'=INDEX(Premissas!$B${P_N0}:$B${P_N0 + 4},A{r})&" (só indicou)"'
        celula(ws, f"B{r}", nome, negrito=True)
        celula(ws, f"C{r}", esc, '0', entrada=True)
        celula(ws, f"D{r}", f"=MIN(INDEX(Premissas!$C${P_N0}:$C${P_N0 + 4},A{r}),INDEX(Premissas!$D${P_E0}:$D${P_E0 + 3},C{r}))", PCT)
        celula(ws, f"E{r}", f"=$B$4*D{r}", RS)
        celula(ws, f"F{r}", f"=$B$4-E{r}", RS)
        celula(ws, f"G{r}", f"=IFERROR(F{r}/$B$4,0)", PCT)
        celula(ws, f"H{r}", f"=E{r}*$B$5", RS)
        celula(ws, f"I{r}", f"=F{r}*$B$5", RS)
        celula(ws, f"J{r}", f"=IFERROR($F$8/F{r}-1,0)", PCT1)
        celula(ws, f"K{r}", f"=INDEX(Premissas!$E${P_E0}:$E${P_E0 + 3},C{r})", '0')
        celula(ws, f"L{r}", f"=K{r}*Premissas!$B${P_HORA}", RS)
        celula(ws, f"M{r}", f"=IFERROR(Premissas!$B${P_HTRIM}/K{r},0)", '0.0')
        celula(ws, f"N{r}", f"=M{r}*F{r}", RS)
        celula(ws, f"O{r}", f"=IF(A{r}=5,Premissas!$B${P_HB5},IF(A{r}=4,Premissas!$B${P_HB4},0))*4/$B$5", '0.0')
        celula(ws, f"P{r}", f"=Premissas!$B${P_HFIX}*4*Premissas!$B${P_HORA}/$B$6", RS)
        celula(ws, f"Q{r}", f"=F{r}-L{r}", RS)
        celula(ws, f"R{r}", f"=IFERROR($Q$8/Q{r}-1,0)", PCT1)
        celula(ws, f"S{r}", f"=Q{r}-O{r}*Premissas!$B${P_HORA}-P{r}", RS)
        celula(ws, f"T{r}", f"=IFERROR($S$8/S{r}-1,0)", PCT1)
    ws["A15"] = ("Leitura: subir o percentual reduz o retido por operação. A coluna R mostra quanto volume a mais o parceiro precisa gerar "
                 "para o nível se pagar, contando as horas Save poupadas (incrementalidade); a coluna T inclui as horas dos benefícios do N4 "
                 "e do N5 e o custo fixo do programa. O ganho de capacidade (coluna N) só aparece acima de cerca de 16 operações qualificadas "
                 "por ano; com 2 operações no canal, o programa se justifica pela formação do canal.")
    ws["A15"].font = NOTA; ws["A15"].alignment = WRAP; ws.merge_cells("A15:T15"); ws.row_dimensions[15].height = 60
    for j, w in enumerate([6, 30, 14, 11, 18, 18, 11, 18, 18, 20, 14, 16, 16, 20, 16, 18, 20, 18, 20, 20], 1):
        ws.column_dimensions[CL(j)].width = w
    ws.column_dimensions["A"].width = 44
    ws.row_dimensions[7].height = 90
    return ws


def canal(wb):
    ws = wb.create_sheet("Canal")
    ws["A1"] = "Visão do canal · de 1 a 10 parceiros, nível provável no fim do ano, % médio pago, retido e concentração"; ws["A1"].font = TITLE
    ws["A2"] = ("Amarelo = entrada. Cenário padrão: o plano de 2027 (2 operações liberadas no produto inteiro). O nível provável é uma "
                "aproximação anual (volume, operações e limite por operação; ignora regularidade, qualidade, comitê e a subida de 1 nível "
                "por trimestre). Parceiros identificados por código.")
    ws["A2"].font = NOTA; ws["A2"].alignment = WRAP; ws.merge_cells("A2:K2"); ws.row_dimensions[2].height = 45
    cab(ws, C_L0 - 1, ["Parceiro", "Operações liberadas no ano", "Líquido Save médio por operação (R$)", "Escopo típico (1 a 4)",
                       "Indicações qualificadas no ano", "Nível provável no fim do ano", "Líquido Save no ano (R$)", "% pago",
                       "Repasse (R$)", "Retido pela Save (R$)", "Participação no líquido do canal"])
    n = len(CANAL); fim = C_L0 + n - 1
    for i, (nome, ops, liq, esc, q) in enumerate(CANAL):
        r = C_L0 + i
        celula(ws, f"A{r}", nome, entrada=True)
        celula(ws, f"B{r}", ops, '0', entrada=True); celula(ws, f"C{r}", liq, RS, entrada=True)
        celula(ws, f"D{r}", esc, '0', entrada=True); celula(ws, f"E{r}", q, '0', entrada=True)
        conds = []
        for k in (5, 4, 3, 2):
            pr = P_N0 + k - 1
            conds.append(f"OR(AND(B{r}>=Premissas!$F${pr},B{r}*MIN(C{r},Premissas!$I${pr}*Premissas!$E${pr})>=Premissas!$E${pr}),"
                         f"AND(Premissas!$G${pr}>0,E{r}>=Premissas!$G${pr}))")
        celula(ws, f"F{r}", f"=IF({conds[0]},5,IF({conds[1]},4,IF({conds[2]},3,IF({conds[3]},2,1))))", '0', negrito=True)
        celula(ws, f"G{r}", f"=B{r}*C{r}", RS)
        celula(ws, f"H{r}", f"=MIN(INDEX(Premissas!$C${P_N0}:$C${P_N0 + 4},F{r}),INDEX(Premissas!$D${P_E0}:$D${P_E0 + 3},D{r}))", PCT)
        celula(ws, f"I{r}", f"=G{r}*H{r}", RS)
        celula(ws, f"J{r}", f"=G{r}-I{r}", RS)
        celula(ws, f"K{r}", f"=IFERROR(G{r}/$G${fim + 1},0)", PCT1)
    validar(ws, f"D{C_L0}:D{fim}", "whole", 1, 4, "Escopo de 1 a 4")
    t = fim + 1
    celula(ws, f"A{t}", "Canal", negrito=True)
    for col in "BEGIJ":
        celula(ws, f"{col}{t}", f"=SUM({col}{C_L0}:{col}{fim})", RS if col in "GIJ" else '0', negrito=True)
    celula(ws, f"H{t}", f"=IFERROR(I{t}/G{t},0)", PCT1, negrito=True)
    celula(ws, f"K{t}", f"=SUM(K{C_L0}:K{fim})", PCT1, negrito=True)
    res = [
        (t + 2, "% médio pago pelo canal (repasse ÷ líquido)", f"=H{t}", PCT1),
        (t + 3, "Participação dos 3 maiores no líquido do canal",
         f"=IFERROR((LARGE(G{C_L0}:G{fim},1)+LARGE(G{C_L0}:G{fim},2)+LARGE(G{C_L0}:G{fim},3))/G{t},0)", PCT1),
        (t + 4, "Parceiros efetivos (1 ÷ índice de Herfindahl)", f"=IFERROR(G{t}*G{t}/SUMPRODUCT(G{C_L0}:G{fim},G{C_L0}:G{fim}),0)", '0.0'),
        (t + 5, "Parceiros com operação liberada no ano", f'=COUNTIF(B{C_L0}:B{fim},">0")', '0'),
    ]
    for r, rot, f, fmt in res:
        ws.cell(r, 1, rot).border = THIN
        celula(ws, f"B{r}", f, fmt, negrito=True)
    ws.cell(t + 7, 1, "Distribuição de níveis (nível provável no fim do ano)").font = BOLD
    for k in range(1, 6):
        r = t + 7 + k
        ws.cell(r, 1, f"N{k}").border = THIN
        celula(ws, f"B{r}", f'=COUNTIF(F{C_L0}:F{fim},"={k}")', '0')
    ws.cell(t + 14, 1, ("Leitura: com 2 operações no ano, os 3 maiores fazem 100% do líquido do produto por definição. Por isso o alerta de "
                        "concentração do Crédito Internacional só vale a partir de 10 operações liberadas em 12 meses; antes disso, "
                        "acompanha-se o número de parceiros com indicação qualificada.")).font = NOTA
    ws.column_dimensions["A"].width = 52
    for j in range(2, 12):
        ws.column_dimensions[CL(j)].width = 17
    ws.row_dimensions[C_L0 - 1].height = 60
    return ws


def como_usar(wb):
    ws = wb.create_sheet("Como usar")
    txt = [
        (f"Como usar o simulador do Programa de Crescimento do Parceiro (Crédito Internacional) · versão {VERSAO}", TITLE),
        ("Uso interno. Proposta para aprovação do Luis e da diretoria (decisão D61, 06/10/2026). Nada aqui vai ao parceiro antes do portão G2.", NOTA),
        ("", None),
        ("1. Premissas: câmbio, fee ILUSTRATIVO, parcela líquida da Save, critérios para SUBIR de cada nível (inclui o patamar de acesso em R$, o limite por operação e, no N4, 3 operações em 24 meses) e % de cada escopo.", None),
        ("   O fee de sucesso é definido caso a caso. O 3% da planilha é só para ilustrar a conta e não deve ser citado como regra.", None),
        ("2. Sustentação: fator do patamar de permanência, observação, amortecedor multiproduto, retorno expresso, nível mínimo, trimestres neutros,", None),
        ("   leitura de 24 meses (anti-oscilação) e período de proteção. Sete perfis mostram subida, permanência, observação, ajuste e retorno em 12 trimestres;", None),
        ("   o perfil G é o caso típico do plano de 2027. A tabela de sensibilidade mostra os patamares em operações mínimas para outros fees.", None),
        ("3. Simulador de operação: valor da operação, nível do parceiro (1 a 5) e escopo executado (1 a 4). O % pago é o MENOR entre o teto do nível e o escopo.", None),
        ("4. Jornada do parceiro: um parceiro em 8 trimestres com todos os critérios (conhecimento, volume, regularidade, qualidade com amostra mínima, compliance e comitê).", None),
        ("5. Comparativo: o mesmo parceiro em cada nível; mostra o retido, a contribuição e quanto volume a mais o nível precisa gerar para se pagar.", None),
        ("6. Canal: até 10 parceiros (por código); nível provável, % médio pago, retido, 3 maiores e parceiros efetivos.", None),
        ("", None),
        ("Regras de nível em uma linha", BOLD),
        ("Sobe 1 nível por revisão quando cumpre os critérios · mantém enquanto cumprir o patamar de permanência em 12 meses ou o dobro em 24 meses ·", None),
        ("abaixo dele: 1 trimestre de observação (2 com negócio em outro produto Save) e depois o nível apurado passa a ser o imediatamente inferior ·", None),
        ("sem ajuste nas 4 primeiras apurações e nas 2 seguintes a cada subida · nunca abaixo do N1 (10%) · volta ao nível anterior em até 4 revisões", None),
        ("cumprindo só o patamar de permanência · trimestre neutro congela o nível e estende a janela · ocorrência julgada procedente ajusta ao N1.", None),
        ("", None),
        ("Simplificações da planilha", BOLD),
        ("- Volume e regularidade contam por trimestre; operação liberada = 1ª parcela. A operação conta inteira na apuração da 1ª parcela (estorno do que não for recebido em 18 meses).", None),
        ("- O limite por operação usa a maior operação de cada trimestre (exato quando há até 1 operação acima do limite por trimestre).", None),
        ("- O ganho do trimestre usa o nível vigente (o da revisão anterior). Na regra real, o % fica travado na assinatura do contrato de cada operação:", None),
        ("  um ajuste de nível não muda as parcelas de operações já contratadas (nem as operações protegidas).", None),
        ("- 'Volume necessário na próxima revisão' considera só o patamar em R$ na leitura de 12 meses; os demais critérios aparecem nas linhas acima.", None),
        ("- Não há retenções de tributos nem custos ordinários do cliente. Retenções: [VALIDAR TRIBUTÁRIO].", None),
        ("- Remuneração só a PJ, por indicação comercial e serviços de apoio, nunca como divisão de honorários advocatícios: [VALIDAR JURÍDICO].", None),
        ("- As abas estão protegidas sem senha (só as células amarelas são editáveis). Para mudar uma fórmula: Revisão > Desproteger planilha.", None),
        ("", None),
        ("Conferência: as fórmulas são recalculadas em Python por scripts/conferir_programa_parceiro.py (exemplos e cenários aleatórios contra uma implementação independente das regras).", NOTA),
        ("Documento de referência: 24-programa-crescimento-parceiro.md (v2.2), 24b-anexos-programa.md (Anexo A) e 24a-diagnostico-concentracao.md (v1.2).", NOTA),
    ]
    for i, (t, f) in enumerate(txt, 1):
        c = ws.cell(i, 1, t)
        if f: c.font = f
    ws.column_dimensions["A"].width = 160
    return ws


def aplicar_fonte(wb):
    """IBM Plex Sans em todas as células (Design System Save), preservando negrito, itálico, tamanho e cor."""
    try:
        wb._named_styles["Normal"].font = Font(name=FONTE, size=10)
    except Exception:
        pass
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                f = c.font
                c.font = Font(name=FONTE, bold=f.bold, italic=f.italic, size=f.size or 10, color=f.color)


def montar():
    wb = Workbook()
    premissas(wb); sustentacao(wb); simulador(wb); jornada(wb); comparativo(wb); canal(wb); como_usar(wb)
    aplicar_fonte(wb)
    wb.move_sheet("Como usar", offset=-(len(wb.sheetnames) - 1))
    wb.active = 0
    for ws in wb.worksheets:
        ws.protection.sheet = True   # sem senha: só as células amarelas ficam editáveis
        ws.protection.formatColumns = False; ws.protection.formatRows = False
    return wb


if __name__ == "__main__":
    wb = montar()
    destino = os.path.join(PASTA, NOME)
    wb.save(destino)
    print("ok", destino)
